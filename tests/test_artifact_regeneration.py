"""Failure preservation and opt-in native PPT rebuild checks."""
from __future__ import annotations
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


regen = load("report_regeneration_test", ROOT / "standalone-skills/html-report/scripts/regenerate.py")
comparison = load("ppt_rebuild_comparison_test", ROOT / "standalone-skills/presentation/scripts/compare.py")


class RegenerationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.json"
        self.source.write_text('{"value": 10}', encoding="utf-8")
        self.work = self.root / "work"

    def build(self, data, run):
        (run / "job.json").write_text(json.dumps(data), encoding="utf-8")
        (run / "report.html").write_text("<html><body>" + str(data["value"]) + "</body></html>", encoding="utf-8")
        (run / "evidence").mkdir()
        (run / "evidence/check.json").write_text('{"ok":true}', encoding="utf-8")

    def normal(self):
        result = regen.regenerate(self.source, self.work, self.build)
        self.assertTrue(result["ok"])
        run = Path(result["outputPath"]).parent
        hashes = {str(p.relative_to(run)): hashlib.sha256(p.read_bytes()).hexdigest()
                  for p in run.rglob("*") if p.is_file()}
        return result, run, hashes, (self.work / regen.STATE).read_bytes()

    def assert_previous(self, previous):
        result, run, hashes, state = previous
        self.assertEqual(result, regen.current(self.work))
        self.assertEqual(state, (self.work / regen.STATE).read_bytes())
        self.assertEqual(hashes, {str(p.relative_to(run)): hashlib.sha256(p.read_bytes()).hexdigest()
                                 for p in run.rglob("*") if p.is_file()})

    def test_bad_json_analysis_and_html_failure_keep_entire_previous_run(self):
        previous = self.normal()
        self.source.write_text("{broken", encoding="utf-8")
        with self.assertRaises(json.JSONDecodeError):
            regen.regenerate(self.source, self.work, self.build)
        self.assert_previous(previous)
        self.source.write_text('{"value": 20}', encoding="utf-8")

        def analysis_failure(data, run):
            (run / "analysis.tmp").write_text("partial", encoding="utf-8")
            raise RuntimeError("analysis failed")

        def html_failure(data, run):
            (run / "job.json").write_text(json.dumps(data), encoding="utf-8")
            (run / "report.html").write_text("partial", encoding="utf-8")
            raise OSError("HTML write failed")

        for builder, error in ((analysis_failure, RuntimeError), (html_failure, OSError)):
            with self.subTest(error=error), self.assertRaises(error):
                regen.regenerate(self.source, self.work, builder)
            self.assert_previous(previous)
        recovered = regen.regenerate(self.source, self.work, self.build)
        self.assertNotEqual(previous[0]["outputPath"], recovered["outputPath"])
        self.assertIn("20", Path(recovered["outputPath"]).read_text(encoding="utf-8"))
        self.assertTrue(previous[1].exists())

    def test_incomplete_bundle_or_failed_pointer_switch_keeps_previous(self):
        previous = self.normal()

        def incomplete(data, run):
            self.build(data, run)
            (run / "evidence/check.json").unlink()

        with self.assertRaisesRegex(ValueError, "근거"):
            regen.regenerate(self.source, self.work, incomplete)
        self.assert_previous(previous)
        with patch.object(regen.os, "replace", side_effect=OSError("pointer write failed")):
            with self.assertRaisesRegex(OSError, "pointer"):
                regen.regenerate(self.source, self.work, self.build)
        self.assert_previous(previous)

    def test_changed_normal_bundle_and_foreign_pointer_are_not_overwritten(self):
        _, run, _, state = self.normal()
        (run / "job.json").write_text('{"user":"edit"}', encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "변경"):
            regen.regenerate(self.source, self.work, self.build)
        self.assertEqual(state, (self.work / regen.STATE).read_bytes())
        self.assertEqual('{"user":"edit"}', (run / "job.json").read_text(encoding="utf-8"))
        (self.work / regen.STATE).write_text('{"user":"owned"}', encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "기존 파일"):
            regen.regenerate(self.source, self.work, self.build)
        self.assertEqual('{"user":"owned"}', (self.work / regen.STATE).read_text(encoding="utf-8"))

    def test_concurrent_run_is_blocked_without_removing_its_lock(self):
        previous = self.normal()
        lock = self.work / ".html-report-regeneration.lock"
        lock.write_text("running", encoding="utf-8")
        with self.assertRaises(FileExistsError):
            regen.regenerate(self.source, self.work, self.build)
        self.assertEqual("running", lock.read_text(encoding="utf-8"))
        self.assert_previous(previous)

    def test_external_pointer_change_during_build_is_preserved(self):
        previous = self.normal()

        def external_change(data, run):
            self.build(data, run)
            (self.work / regen.STATE).write_text('{"external":"edit"}', encoding="utf-8")

        with self.assertRaisesRegex(ValueError, "외부"):
            regen.regenerate(self.source, self.work, external_change)
        self.assertEqual('{"external":"edit"}', (self.work / regen.STATE).read_text(encoding="utf-8"))
        self.assertEqual(previous[2], {str(p.relative_to(previous[1])): hashlib.sha256(p.read_bytes()).hexdigest()
                                     for p in previous[1].rglob("*") if p.is_file()})
        # Absence is part of the starting state too: do not replace a new foreign file.
        self.work = self.root / "fresh-work"
        with self.assertRaisesRegex(ValueError, "외부"):
            regen.regenerate(self.source, self.work, external_change)
        self.assertEqual('{"external":"edit"}', (self.work / regen.STATE).read_text(encoding="utf-8"))

    def test_shared_helpers_and_references_are_identical(self):
        for slug, paths in {
            "html-report": ["scripts/regenerate.py", "references/regeneration.md"],
            "presentation": ["scripts/compare.py", "references/rebuild-check.md"],
        }.items():
            for path in paths:
                self.assertEqual((ROOT / "standalone-skills" / slug / path).read_bytes(),
                                 (ROOT / "company-agent-plugin/skills" / slug / path).read_bytes())


class PortableRebuildTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.env = dict(os.environ, TEMP=str(self.root), TMP=str(self.root))

    def stage(self, slug):
        target = self.root / slug
        shutil.copytree(ROOT / "standalone-skills" / slug, target)
        metadata = json.loads((ROOT / "scripts/platform_skills/catalog-metadata.json").read_text(encoding="utf-8-sig"))
        for name in metadata[slug]["packaging"]["runtimeModules"]:
            shutil.copyfile(ROOT / "company-agent-plugin/scripts/company_agent" / name,
                            target / "scripts/visual_engine" / name)
        return target

    def cli(self, script, *args):
        process = subprocess.run([sys.executable, "-X", "utf8", "-B", str(script), *map(str, args)],
                                 env=self.env, capture_output=True, encoding="utf-8", timeout=30)
        self.assertEqual(0, process.returncode, (process.stdout, process.stderr))
        return json.loads(process.stdout)

    def test_html_example_runs_from_clean_copy_and_reports_one_current_output(self):
        skill = self.stage("html-report")
        job = self.root / "job.json"
        job.write_text(json.dumps({"sections": [{"title": "Synthetic", "body": "Only synthetic input"}]}), encoding="utf-8")
        source_hash = hashlib.sha256(job.read_bytes()).hexdigest()
        script = skill / "scripts/regenerate.py"
        first = self.cli(script, "--spec", job, "--work", self.root / "work")
        self.assertEqual({"ok", "outputPath"}, first.keys())
        self.assertTrue(Path(first["outputPath"]).is_file())
        self.assertEqual(first, self.cli(script, "--work", self.root / "work"))
        second = self.cli(script, "--spec", job, "--work", self.root / "work")
        self.assertNotEqual(first["outputPath"], second["outputPath"])
        self.assertEqual(Path(first["outputPath"]).read_bytes(), Path(second["outputPath"]).read_bytes())
        self.assertEqual(source_hash, hashlib.sha256(job.read_bytes()).hexdigest())

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx required")
    def test_same_approved_job_rebuild_matches_and_direct_ppt_patch_is_detected(self):
        from pptx import Presentation
        from pptx.util import Pt
        skill = self.stage("presentation")
        script = skill / "scripts/presentation.py"
        spec = {"title": "Synthetic", "creationMode": "new", "purpose": "test", "audience": "test",
                "designPreset": "business", "slideCount": 1, "slides": [{"title": "Synthetic",
                    "table": {"headers": ["item", "value"], "rows": [["A", 1], ["B", 2]]},
                    "chart": {"type": "column", "categories": ["A", "B"], "series": [{"name": "Value", "values": [1, 2]}]}}]}
        job = self.root / "job.json"
        job.write_text(json.dumps(spec), encoding="utf-8")
        draft = self.cli(script, "draft", "--spec", job, "--output", self.root / "draft.html")
        spec["designReview"] = {**draft["designReview"], "confirmed": True}
        job.write_text(json.dumps(spec), encoding="utf-8")
        first, second = self.root / "first.pptx", self.root / "second.pptx"
        for output in (first, second):
            self.cli(script, "build", "--spec", job, "--output", output)
        result = comparison.compare(first, second)
        self.assertTrue(result["ok"], result)
        self.assertEqual("not_performed", result["visualReview"])
        deck = Presentation(second)
        table = next(s for s in deck.slides[0].shapes if s.has_table)
        self.assertTrue(any(s.has_chart for s in deck.slides[0].shapes))
        table.top += Pt(76)
        deck.save(second)
        result = comparison.compare(first, second)
        self.assertFalse(result["ok"])
        self.assertTrue(any("geometry" in p for p in result["differences"]), result)

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx required")
    def test_package_dates_ignored_but_text_table_and_chart_changes_detected(self):
        from pptx import Presentation
        from pptx.chart.data import ChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        first, second = self.root / "first.pptx", self.root / "second.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Title"
        slide.shapes.add_table(1, 1, 0, Inches(1), Inches(2), Inches(1)).table.cell(0, 0).text = "Cell"
        data = ChartData(); data.categories = ["A"]; data.add_series("Series", [1])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(2), Inches(3), Inches(2), data)
        deck.save(first)
        with ZipFile(first) as source, ZipFile(second, "w") as target:
            for item in source.infolist():
                raw = source.read(item.filename)
                item.date_time = (2020, 1, 1, 0, 0, 0)
                if item.filename == "docProps/core.xml":
                    raw = raw.replace(b"2013-01-27", b"2020-01-27")
                target.writestr(item, raw)
        self.assertTrue(comparison.compare(first, second)["ok"])
        for kind in ("text", "table", "chart", "background"):
            with self.subTest(kind=kind):
                changed = Presentation(first)
                if kind == "text": changed.slides[0].shapes[0].text = "Changed"
                elif kind == "table": changed.slides[0].shapes[1].table.cell(0, 0).text = "Changed"
                elif kind == "chart":
                    data = ChartData(); data.categories = ["A"]; data.add_series("Series", [999])
                    changed.slides[0].shapes[2].chart.replace_data(data)
                else:
                    changed.slides[0].background.fill.solid()
                    changed.slides[0].background.fill.fore_color.rgb = RGBColor(255, 0, 0)
                changed.save(second)
                self.assertFalse(comparison.compare(first, second)["ok"])


if __name__ == "__main__":
    unittest.main()
