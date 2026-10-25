"""Synthetic round-trip checks, with no Office, profile, network or model calls."""
import copy
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'company-agent-plugin/scripts'))
from company_agent import ppt_html, ppt_workflow, presentation_design, ppt_html_import, content_layout


class EditableDeckTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.spec = {'creationMode':'new','designPreset':'business','purpose':'시험','audience':'검토자','slideCount':1,
                     'title':'편집 시험', 'slides':[{'title':'구성', 'elements':[
            {'kind':'text','text':'제목','x':40,'y':20,'w':600,'h':40,'size':28,'bold':True},
            {'kind':'shape','group':'실적카드','x':40,'y':100,'w':380,'h':220,'fill':'EEF3F8'},
            {'kind':'text','group':'실적카드','text':'신청 현황\n기준 12명','x':60,'y':120,'w':330,'h':90,
             'size':24,'bold':False,'color':'173B5E','font':'Malgun Gothic'},
            {'kind':'text','text':'출처: 가상 자료','x':40,'y':480,'w':600,'h':25,'size':12}]}]}

    def render(self, spec=None):
        data = ppt_html.prepare(spec or self.spec)[0]
        target = self.root/'deck.pptx'
        counts = presentation_design.render_python(data, target, None)
        return target, counts, data

    def test_edit_color_weight_font_save_and_reopen(self):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_AUTO_SIZE
        from pptx.util import Pt
        target, counts, _ = self.render()
        deck = Presentation(target)
        group = deck.slides[0].shapes[1]
        self.assertEqual(group.name, '실적카드')
        frame = group.shapes[1].text_frame
        self.assertEqual(frame.auto_size, MSO_AUTO_SIZE.NONE)
        for p in frame.paragraphs:
            for run in p.runs:
                self.assertEqual(str(run.font.color.rgb), '173B5E')
                self.assertIs(run.font.bold, False)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string('C00000')
                run.font.size = Pt(26)
        deck.save(target)
        reopened = Presentation(target).slides[0].shapes[1].shapes[1]
        for p in reopened.text_frame.paragraphs:
            for run in p.runs:
                self.assertEqual((str(run.font.color.rgb), run.font.bold, run.font.size.pt), ('C00000',True,26))
        self.assertEqual(counts['groups'],1)
        self.assertFalse(group._element.xpath('.//a:spLocks[@noGrp="1"] | .//a:spLocks[@noTextEdit="1"] | .//a:grpSpLocks[@noUngrp="1"]'))
        self.assertEqual(ppt_workflow.quality(target)['nativeObjects'], {'groups':1,'textFrames':4})

    def test_group_preserves_geometry_order_and_can_be_ungrouped(self):
        from pptx import Presentation
        target, _, _ = self.render()
        deck = Presentation(target)
        slide = deck.slides[0]
        self.assertEqual([s.name for s in slide.shapes][1], '실적카드')
        group = slide.shapes[1]
        before = [(s.left,s.top,s.width,s.height) for s in group.shapes]
        tree = slide.shapes._spTree
        index = tree.index(group._element)
        # Identity transform: ungroup by replacing the container with its children.
        for offset, shape in enumerate(list(group.shapes)):
            tree.insert(index+offset, shape._element)
        tree.remove(group._element)
        deck.save(target)
        shapes = Presentation(target).slides[0].shapes
        self.assertEqual(len(shapes),4)
        self.assertEqual([(s.left,s.top,s.width,s.height) for s in list(shapes)[1:3]],before)
        self.assertEqual(shapes[2].text,'신청 현황\n기준 12명')

    def test_invalid_group_does_not_silently_reorder_or_rasterize(self):
        for variant in ('noncontiguous','table','badname'):
            spec = copy.deepcopy(self.spec)
            rows = spec['slides'][0]['elements']
            if variant == 'noncontiguous':
                rows[0]['group']='실적카드'; rows[1].pop('group')
            elif variant == 'table':
                rows[1]={'kind':'table','group':'실적카드','x':40,'y':100,'w':380,'h':120,
                         'table':{'headers':['항목'],'rows':[['값']]}}
            else:
                rows[1]['group']='../카드'
            result = ppt_html.create_draft(spec,self.root/'invalid.html')
            self.assertEqual(result['code'],'invalid_ppt_group',result)
            self.assertFalse((self.root/'invalid.html').exists())

    def test_groups_are_checked_for_tiny_text_and_transformed_bounds(self):
        from pptx import Presentation
        from pptx.util import Pt
        target, _, _ = self.render()
        deck = Presentation(target)
        group = deck.slides[0].shapes[1]
        group.shapes[1].text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        group.left = deck.slide_width - Pt(40)
        group.width *= 2
        deck.save(target)
        issues = ppt_workflow.quality(target)['issues']
        self.assertIn('tiny_text',[i['code'] for i in issues])
        self.assertIn('out_of_bounds',[i['code'] for i in issues])

    def test_reference_style_samples_are_bounded_and_keep_group_text(self):
        target, _, _ = self.render()
        result = ppt_workflow.analyze(target, selected_slides=[1])
        self.assertTrue(result['ok'],result)
        page = result['slides'][0]
        self.assertEqual(page['groupCount'],1)
        self.assertLessEqual(len(page['declaredTextStyles']),8)
        self.assertTrue(any(s.get('color')=={'srgbClr':'173B5E'} for s in page['declaredTextStyles']))
        self.assertTrue(any('신청 현황' in s['textPreview'] for s in page['shapes']))
        self.assertFalse(result['savedToMemory'])

    def test_kpi_group_does_not_break_content_order_validation(self):
        spec={'title':'검증','facts':[{'id':'total','label':'신청','op':'value','inputs':[12]}],
              'slides':[{'title':'신청 12명','kpis':[{'fact':'total'}],'body':'가상 자료입니다.',
                         'table':{'headers':['구분','인원'],'rows':[['신청','12']]},
                         'contentOrder':['body','table']}]}
        target, counts, data = self.render(spec)
        self.assertEqual(counts['groups'],1)
        self.assertEqual(content_layout.check_ppt(target,data)['basis'],'pptx-native-coordinates')

    def test_html_sanitizer_keeps_only_declarative_group_attribute(self):
        source = self.root/'source.html'
        source.write_text('<section class="slide"><div data-ppt-group="카드" onclick="bad()">안내</div></section>',encoding='utf-8')
        parser = ppt_html_import.StaticHTML(source)
        parser.feed(source.read_text(encoding='utf-8'))
        output = ''.join(parser.parts)
        self.assertIn('data-ppt-group="카드"',output)
        self.assertNotIn('onclick',output)

    def test_group_metadata_survives_review_and_changes_require_new_approval(self):
        from company_agent import business_artifacts
        preview = ppt_html.create_draft(self.spec,self.root/'draft.html')
        self.assertTrue(preview['ok'],preview)
        approved = copy.deepcopy(self.spec)
        approved['designReview']={**preview['designReview'],'confirmed':True}
        ppt_workflow.verify_design_review(approved,None)
        approved['slides'][0]['elements'][1]['group']='다른카드'
        with self.assertRaises(business_artifacts.ArtifactError):
            ppt_workflow.verify_design_review(approved,None)

    def test_empty_kpi_label_does_not_group_unrelated_title(self):
        spec={'title':'검증','facts':[{'id':'total','label':'','op':'value','inputs':[12]}],
              'slides':[{'title':'신청 현황','kpis':[{'fact':'total'}],'body':'가상 자료'}]}
        _, counts, data = self.render(spec)
        self.assertEqual(counts['groups'],0)
        grouped=[e for e in data['presentationPlan']['pages'][0]['elements'] if e.get('group')]
        self.assertEqual([e['text'] for e in grouped],['12'])


if __name__ == '__main__':
    unittest.main()
