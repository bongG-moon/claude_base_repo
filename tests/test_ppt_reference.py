"""Reference routing uses synthetic files; never opens the user's Office session."""
import contextlib
import io
import json
from pathlib import Path
import re
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "company-agent-plugin/scripts"))
from company_agent import ppt_reference
from company_agent.cli import main

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


class PptReferenceTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def deck(self, count=5):
        if Presentation is None:
            self.skipTest("Synthetic native PPTX requires the existing python-pptx dependency")
        presentation = Presentation()
        for index in range(count):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = f"Synthetic slide {index + 1}"
            slide.placeholders[1].text = f"Verified example {index + 1}"
        file = self.root / "synthetic.pptx"
        presentation.save(file)
        return file

    def prepare(self, spec):
        with patch("company_agent.ppt_workflow.installed_fonts", return_value=set()):
            return ppt_reference.prepare_reference(spec)

    def test_direct_pptx_uses_first_three_without_office_or_source_changes(self):
        file = self.deck()
        before = file.read_bytes()
        with patch("company_agent.business_artifacts._office") as office:
            result = self.prepare({"file": str(file)})
        self.assertTrue(result["ok"], result)
        self.assertEqual([1, 2, 3], result["selectedSlides"])
        self.assertEqual([[1, 2, 3]], result["batches"])
        self.assertEqual(5, result["slideCount"])
        self.assertEqual(3, len(result["analysis"]["slides"]))
        self.assertEqual("not_performed", result["visualReview"])
        office.assert_not_called()
        self.assertEqual(before, file.read_bytes())
        self.assertEqual([file], list(self.root.iterdir()))

    def test_requested_content_is_complete_in_small_batches(self):
        result = self.prepare({"file": str(self.deck()), "purpose": "content"})
        self.assertEqual([1, 2, 3, 4, 5], result["selectedSlides"])
        self.assertEqual([[1, 2, 3], [4, 5]], result["batches"])
        self.assertEqual("read_requested_content", result["nextAction"])
        self.assertEqual("bounded_shape_previews_not_full_content", result["contentCoverage"])

    def test_requested_slide_order_and_structural_content_are_preserved(self):
        result = self.prepare({"file": str(self.deck()), "slides": [5, 2]})
        self.assertEqual([5, 2], result["selectedSlides"])
        self.assertEqual([4, 1], [page["sourceSlide"] for page in result["analysis"]["slides"]])
        self.assertIn("Synthetic slide 5", json.dumps(result["analysis"]["slides"][0]))
        self.assertNotIn("Synthetic slide 1", json.dumps(result))

    def test_one_slide_reference_does_not_request_nonexistent_default_slides(self):
        result = self.prepare({"file": str(self.deck(1))})
        self.assertTrue(result["ok"], result)
        self.assertEqual([1], result["selectedSlides"])

    def test_explicit_slides_after_sixty_are_supported_without_returning_other_content(self):
        result = self.prepare({"file": str(self.deck(63)), "slides": [63, 61]})
        self.assertTrue(result["ok"], result)
        self.assertEqual([63, 61], result["selectedSlides"])
        self.assertEqual([62, 60], [page["sourceSlide"] for page in result["analysis"]["slides"]])
        self.assertNotIn("Synthetic slide 1", json.dumps(result))

    def test_invalid_actual_range_and_large_unselected_content_are_explicit(self):
        file = self.deck(63)
        result = self.prepare({"file": str(file), "purpose": "content"})
        self.assertEqual("reference_range_required", result["code"])
        result = self.prepare({"file": str(file), "slides": [64]})
        self.assertEqual("invalid_slide_selection", result["code"])

    def test_direct_analysis_rejects_malformed_or_unbounded_selection(self):
        from company_agent.ppt_workflow import analyze
        for selection in ([True], [0], [1, 1], list(range(1, 62)), "1,2"):
            result = analyze(self.root / "not-opened.pptx", selected_slides=selection)
            self.assertEqual("invalid_slide_selection", result["code"])

    def test_corrupt_or_compound_header_is_not_a_confirmed_drm_denial(self):
        file = self.root / "unreadable.pptx"
        for body in (b"ordinary bad format", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"synthetic"):
            with self.subTest(body=body):
                file.write_bytes(body)
                result = self.prepare({"file": str(file)})
                self.assertEqual("reference_format_unreadable", result["code"])
                self.assertFalse(result["restrictionConfirmed"])
                self.assertFalse(result["automaticCapture"])
                self.assertEqual("check_known_source_context", result["nextAction"])
                self.assertEqual(body, file.read_bytes())

    def test_real_permission_denial_does_not_offer_capture(self):
        file = self.root / "denied.pptx"
        file.write_bytes(b"synthetic")
        for failure in (PermissionError("access denied"), None):
            mocked = patch.object(ppt_reference, "analyze", side_effect=failure) if failure else patch.object(
                ppt_reference, "analyze", return_value={"ok": False, "status": "blocked", "code": "permission_denied"})
            with mocked:
                result = self.prepare({"file": str(file)})
            self.assertEqual("permission_denied", result["code"])
            self.assertTrue(result["restrictionConfirmed"])
            self.assertFalse(result["retryAllowed"])
            self.assertFalse(result["captured"])
            self.assertNotEqual("use_connected_approved_screen_tool", result.get("nextAction"))

    def test_unexpected_read_failure_is_redacted_and_does_not_claim_drm(self):
        file = self.root / "unstable.pptx"
        file.write_bytes(b"synthetic")
        with patch.object(ppt_reference, "analyze", side_effect=RuntimeError("PRIVATE_SOURCE_CONTENT")):
            result = self.prepare({"file": str(file)})
        self.assertFalse(result["ok"])
        self.assertFalse(result["restrictionConfirmed"])
        self.assertNotIn("PRIVATE_SOURCE_CONTENT", json.dumps(result))

    def test_empty_content_reference_does_not_claim_ready(self):
        result = self.prepare({"file": str(self.deck(0)), "purpose": "content"})
        self.assertEqual("reference_has_no_slides", result["code"])

    def test_screen_plan_neither_reads_document_nor_claims_capture_or_authorization(self):
        file = self.root / "protected-example.pptx"
        file.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
        with patch.object(ppt_reference, "analyze") as analyze, patch.object(Path, "open", side_effect=AssertionError("must not read")):
            result = self.prepare({"file": str(file), "inputMode": "powerpoint-screen", "slides": [1, 7, 9, 12]})
        analyze.assert_not_called()
        self.assertEqual("capture_plan", result["status"])
        self.assertFalse(result["captured"])
        self.assertFalse(result["authorizationGranted"])
        self.assertEqual([[1, 7, 9], [12]], result["batches"])
        self.assertIn("사용자 캡처", result["ifUnavailable"])
        self.assertEqual([file], list(self.root.iterdir()))

    def test_capture_or_ai_denial_stops_before_any_file_access(self):
        for code in ("capture_denied", "ai_processing_denied", "permission_denied"):
            with self.subTest(code=code), patch.object(ppt_reference, "safe_path") as paths:
                result = self.prepare({"file": str(self.root / "not-opened.pptx"), "inputMode": "powerpoint-screen", "accessStatus": code})
            self.assertEqual("blocked", result["status"])
            self.assertFalse(result["retryAllowed"])
            paths.assert_not_called()

    def test_screen_defaults_representative_but_explicit_content_keeps_scope(self):
        file = self.root / "selected.pptx"
        file.write_bytes(b"synthetic")
        for purpose, selection in (("style", "representative"), ("content", "requested_content")):
            result = self.prepare({"file": str(file), "inputMode": "powerpoint-screen", "purpose": purpose})
            self.assertEqual(selection, result["selection"]["mode"])
            self.assertEqual(3, result["batchSize"])

    def test_html_and_image_reference_keep_existing_workflows(self):
        for suffix, action in ((".html", "use_existing_html_workflow"), (".png", "review_selected_reference_image")):
            file = self.root / ("source" + suffix)
            file.write_bytes(b"synthetic")
            with patch.object(ppt_reference, "analyze") as analyze:
                result = self.prepare({"file": str(file)})
            self.assertEqual(action, result["nextAction"])
            self.assertFalse(result["captured"])
            analyze.assert_not_called()

    def test_invalid_inputs_do_not_become_approval_or_unbounded_selection(self):
        path = str(self.root / "source.pptx")
        invalid = [[], {"file": path, "approved": True}, {"file": path, "inputMode": []},
                   {"file": path, "purpose": {}}, {"file": path, "accessStatus": {}},
                   {"file": path, "accessStatus": "DRM"}, {"file": path, "slides": [True]},
                   {"file": path, "slides": [1, 1]}, {"file": path, "slides": list(range(1, 62))}]
        for spec in invalid:
            with self.subTest(spec=spec):
                result = self.prepare(spec)
                self.assertFalse(result["ok"])
                self.assertEqual("invalid_request", result["status"])

    def test_cli_uses_reference_branch_without_personal_state_or_office(self):
        file = self.deck()
        spec = self.root / "reference.json"
        spec.write_text(json.dumps({"file": str(file), "slides": [2, 4]}), encoding="utf-8")
        output = io.StringIO()
        with contextlib.redirect_stdout(output), patch("company_agent.ppt_workflow.installed_fonts", return_value=set()):
            status = main(["business", "ppt-reference", "--spec", str(spec), "--state-root", str(self.root / "isolated-state")])
        self.assertEqual(0, status)
        self.assertEqual([2, 4], json.loads(output.getvalue())["selectedSlides"])
        self.assertFalse((self.root / "isolated-state").exists())

    def test_reference_classification_is_readonly_without_new_permission_grant(self):
        from company_agent.execution_contract import classify_command, safe_permission
        scripts = ROOT / "company-agent-plugin/scripts"
        prefix = f'"{sys.executable}" -B "{scripts / "harness_cli.py"}"'
        command = f'{prefix} business ppt-reference --spec "{self.root / "reference.json"}" --state-root "{self.root}"'
        self.assertEqual("read_only", classify_command(command))
        self.assertIsNone(safe_permission({"hook_event_name": "PermissionRequest", "tool_name": "Bash",
                                          "tool_input": {"command": command}}, self.root))
        for extra in (' --output output.png', '; echo extra', ' > output.json'):
            self.assertEqual("unknown", classify_command(command + extra))

    def test_html_input_support_does_not_make_presentation_an_html_report_competitor(self):
        from company_agent.skill_task_context import task_candidates
        items = []
        for name in ("html-report", "presentation"):
            file = ROOT / "company-agent-plugin/skills" / name / "SKILL.md"
            description = re.search(r"^description: (.+)$", file.read_text(encoding="utf-8"), re.M).group(1)
            items.append({"id": name, "name": name, "description": description,
                          "source": "company", "path": str(file), "invocation": "company-agent:" + name})
        inventory = {"complete": True, "skills": items}
        for prompt, expected in (("HTML 보고서 만들어줘", "html-report"), ("PPT 발표자료 만들어줘", "presentation")):
            result = task_candidates(inventory, prompt)
            names = [row["name"] for group in result["groups"] for row in group["candidates"]]
            self.assertEqual([expected], names)
            self.assertNotIn("competingIds", result)


if __name__ == "__main__":
    unittest.main()
