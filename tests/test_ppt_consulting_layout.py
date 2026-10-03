"""Observable layout, content and approval contracts for template-free reporting."""
import copy
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'company-agent-plugin/scripts'))
from company_agent import business_artifacts as artifacts, ppt_html, presentation_design as design


class ConsultingLayoutTests(unittest.TestCase):
    def chart(self):
        return {'type':'column', 'title':'월별 실적 · 건', 'categories':['7월','8월'],
                'series':[{'name':'실적','values':[110,130]}]}

    def table(self):
        return {'headers':['항목','확인할 내용','상태'],
                'rows':[['품질','실제 개선 효과 확인','진행 중'],['재작업','비용 산정 근거 확인','미정']]}

    def prepared(self, row, width=960, height=540):
        data, *_ = ppt_html.prepare({'slides':[row]})
        return data, design.plan(data, width, height)['pages'][0]['elements']

    def test_layouts_change_evidence_width_and_reading_order(self):
        widths = {}
        for layout in ('summary','evidence','comparison','actions'):
            row={'layout':layout,'title':'8월 실적은 130건입니다', 'chart':self.chart(),
                 'body':'증가 원인은 추가 확인이 필요합니다.'}
            _, elements=self.prepared(row)
            chart=next(e for e in elements if e['kind']=='chart')
            body=next(e for e in elements if e.get('text')==row['body'])
            widths[layout]=chart['w']
            if layout=='actions':
                self.assertGreater(body['y'],chart['y']+chart['h'])
            else:
                self.assertGreater(body['x'],chart['x']+chart['w'])
            self.assertTrue(all(e['kind'] in ('text','table','chart','image') for e in elements))
        self.assertLess(widths['comparison'],widths['summary'])
        self.assertLess(widths['summary'],widths['evidence'])
        self.assertLess(widths['evidence'],widths['actions'])

    def test_summary_separates_key_points_without_inventing_numbers(self):
        row={'layout':'summary','title':'검증 후 적용 범위를 결정합니다',
             'body':'현재는 가상 자료로 구성만 확인했습니다.',
             'bullets':['실제 자료 확인','효과 산정 검증','적용 대상 협의']}
        _, elements=self.prepared(row)
        points=[next(e for e in elements if e.get('text')==value) for value in row['bullets']]
        self.assertEqual(3,len({e['x'] for e in points}))
        self.assertEqual(1,len({e['y'] for e in points}))
        self.assertFalse(any(e['kind']=='chart' or e.get('size')==38 for e in elements))

    def test_comparison_keeps_two_exhibits_and_common_explanation(self):
        row={'layout':'comparison','title':'월별 실적과 확인 사항을 함께 검토합니다',
             'chart':self.chart(),'table':self.table(),'body':'실적 증가와 개선 효과는 구분합니다.'}
        _, elements=self.prepared(row)
        chart=next(e for e in elements if e['kind']=='chart')
        table=next(e for e in elements if e['kind']=='table')
        prose=next(e for e in elements if e.get('text')==row['body'])
        self.assertAlmostEqual(chart['w'],table['w'])
        self.assertGreater(table['x'],chart['x']+chart['w'])
        self.assertGreater(prose['y'],max(chart['y']+chart['h'],table['y']+table['h']))
        with self.assertRaises(artifacts.ArtifactError):
            self.prepared({**row,'layout':'auto'})

    def test_actions_keep_full_width_table_body_and_conclusion_in_order(self):
        row={'layout':'actions','title':'효과 검증 후 적용 범위를 협의합니다','table':self.table(),
             'body':'담당자와 완료일은 아직 정하지 않았습니다.', 'takeaway':'검증 결과를 확인한 뒤 결정합니다.'}
        _, elements=self.prepared(row)
        table=next(e for e in elements if e['kind']=='table')
        body=next(e for e in elements if e.get('text')==row['body'])
        takeaway=next(e for e in elements if e.get('text')==row['takeaway'])
        self.assertEqual(864,table['w'])
        self.assertGreater(body['y'],table['y']+table['h'])
        self.assertGreater(takeaway['y'],body['y'])

    def test_korean_two_line_title_and_both_page_ratios_keep_readable_bounds(self):
        title='실제 업무 자료로 개선 효과를 확인한 뒤 적용 범위와 실행 일정을 결정합니다'
        for layout in ('summary','evidence','comparison','actions'):
            for width,height in ((960,540),(720,540)):
                row={'layout':layout,'title':title,'body':'미확인 효과는 추정하지 않습니다.'}
                _, elements=self.prepared(row,width,height)
                heading=next(e for e in elements if e.get('text')==title)
                self.assertEqual(30,heading['size'])
                self.assertGreater(heading['h'],48)
                for e in elements:
                    self.assertGreaterEqual(min(e['x'],e['y']),0)
                    self.assertLessEqual(e['x']+e['w'],width+.01)
                    self.assertLessEqual(e['y']+e['h'],height+.01)
        with self.assertRaises(artifacts.ArtifactError):
            self.prepared({'layout':'summary','title':title*3,'body':'보존해야 하는 본문'})

    def test_dense_summary_and_table_fail_without_truncation(self):
        for row in (
                {'layout':'summary','title':'확인','body':'긴 설명입니다. '*120},
                {'layout':'evidence','title':'확인','table':{'headers':['항목','설명'],
                    'rows':[['대상','매우 긴 한글 설명을 읽기 좋은 크기로 유지합니다. '*5] for _ in range(7)]}}):
            with self.assertRaises(artifacts.ArtifactError):
                self.prepared(row)

    def test_explicit_content_order_and_auto_keep_existing_geometry(self):
        row={'title':'확인','chart':self.chart(),'body':'검증 내용'}
        _, implicit=self.prepared(row)
        _, explicit=self.prepared({**row,'layout':'auto'})
        self.assertEqual(implicit,explicit)
        chart=next(e for e in implicit if e['kind']=='chart')
        self.assertEqual(48,chart['x'])
        self.assertEqual(170,chart['y'])
        self.assertAlmostEqual(864*.64,chart['w'])
        ordered={'title':'확인','table':self.table(),'body':'표 아래 설명', 'contentOrder':['table','body']}
        _, before=self.prepared(ordered)
        _, after=self.prepared({**ordered,'layout':'evidence'})
        self.assertEqual(before,after)

    @unittest.skipUnless(importlib.util.find_spec('pptx'), 'native library is not installed')
    def test_all_structured_layouts_roundtrip_through_saved_html_without_source_content(self):
        from pptx import Presentation
        job={'purpose':'보고','audience':'팀장','slideCount':4,'slides':[
            {'layout':layout,'title':'PRIVATE_TITLE','body':'PRIVATE_BODY',
             'takeaway':'PRIVATE_TAKEAWAY','source':'PRIVATE_SOURCE','table':{
                 'headers':['PRIVATE_HEADER','상태'], 'rows':[['PRIVATE_CELL','미정']]}}
            for layout in ('summary','evidence','comparison','actions')]}
        job['slides'][2]['chart']={'type':'column','title':'PRIVATE_CHART_TITLE',
            'categories':['PRIVATE_CATEGORY'], 'series':[{'name':'PRIVATE_SERIES','values':[987654321]}]}

        def geometry(elements):
            return [{k:e[k] for k in ('kind','x','y','w','h','size','bold','columnWidths','rowHeights') if k in e}
                    for e in elements]

        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            with patch.object(artifacts,'_office',side_effect=AssertionError('HTML must not start Office')):
                draft=ppt_html.create_draft(job,folder/'source-draft.html')
                self.assertTrue(draft['ok'],draft)
                job['designReview']={**draft['designReview'],'confirmed':True}
                original=ppt_html.prepare(job)[0]['presentationPlan']
                saved=ppt_html.save_template(job,folder/'template.html')
                self.assertTrue(saved['ok'],saved)
                self.assertEqual(4,saved['layoutCount'])
                serialized=(folder/'template.html').read_text(encoding='utf-8')
                self.assertNotIn('PRIVATE_',serialized)
                self.assertNotIn('987654321',serialized)
                template=ppt_html.read_template(folder/'template.html')
                reused={'creationMode':'saved','purpose':'새 보고','audience':'팀장','slideCount':4,'slides':[]}
                for index,(page,layout) in enumerate(zip(original['pages'],template['layouts']),1):
                    self.assertEqual(geometry(page['elements']),geometry(layout['elements']))
                    row={'title':f'새 보고 {index}','layoutIndex':index,
                         'texts':[f'새 업무 내용 {index}-{slot}' for slot in range(layout['textSlots'])],
                         'tables':[{'headers':['새 항목','새 상태'],'rows':[['검증 대상','확인 중']]}]}
                    if index==3:
                        row['charts']=[{'type':'column','categories':['새 항목'],
                                        'series':[{'name':'새 실적','values':[250]}]}]
                    reused['slides'].append(row)
                preview=ppt_html.create_draft(reused,folder/'reused-draft.html',folder/'template.html')
                self.assertTrue(preview['ok'],preview)
                reused['designReview']={**preview['designReview'],'confirmed':True}
                restored=ppt_html.prepare(reused,folder/'template.html')[0]['presentationPlan']
                for before,after in zip(original['pages'],restored['pages']):
                    self.assertEqual(geometry(before['elements']),geometry(after['elements']))
            with patch.object(artifacts,'_office',return_value={'ok':False,'code':'render_refused'}):
                result=artifacts.create_ppt(reused,folder/'reused.pptx',folder/'template.html',require_choices=True)
            self.assertTrue(result['ok'],result)
            deck=Presentation(folder/'reused.pptx')
            for slide,page in zip(deck.slides,original['pages']):
                self.assertEqual(len(page['elements']),len(slide.shapes))
                for shape,element in zip(slide.shapes,page['elements']):
                    for value,key in ((shape.left.pt,'x'),(shape.top.pt,'y'),(shape.width.pt,'w'),(shape.height.pt,'h')):
                        self.assertAlmostEqual(element[key],value,places=3)
                table=next(s.table for s in slide.shapes if s.has_table)
                self.assertEqual('검증 대상',table.cell(1,0).text)
            chart=next(s.chart for slide in deck.slides for s in slide.shapes if s.has_chart)
            self.assertEqual((250,),chart.series[0].values)

    @unittest.skipUnless(importlib.util.find_spec('pptx'), 'native library is not installed')
    def test_default_draft_approval_native_editability_and_changed_layout(self):
        from pptx import Presentation
        job={'purpose':'현황 보고','audience':'부서장','slideCount':4,'facts':[
            {'id':'actual','label':'8월 실적','unit':'건','op':'value','inputs':[130]}],
            'slides':[{'layout':layout,'title':'8월 실적은 {{fact:actual}}건입니다',
                       'table':self.table(),'body':'개선 효과는 추가 확인이 필요합니다.'}
                      for layout in ('summary','evidence','comparison','actions')]}
        job['slides'][2]['chart']=self.chart()
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            with patch.object(artifacts,'_office',side_effect=AssertionError('HTML must not start Office')):
                preview=ppt_html.create_draft(job,folder/'draft.html')
            self.assertTrue(preview['ok'],preview)
            self.assertFalse(preview['officeStarted'])
            self.assertIn('.ppt-element[data-kind="table"]{padding:0}',
                          (folder/'draft.html').read_text(encoding='utf-8'))
            self.assertNotIn('creationMode',job)
            job['designReview']=preview['designReview']
            self.assertEqual('design_confirm',artifacts.create_ppt(job,folder/'not-created.pptx',require_choices=True)['stage'])
            self.assertFalse((folder/'not-created.pptx').exists())
            job['designReview']['confirmed']=True
            with patch.object(artifacts,'_office',return_value={'ok':False,'code':'render_refused'}):
                result=artifacts.create_ppt(job,folder/'native.pptx',require_choices=True)
            self.assertTrue(result['ok'],result)
            deck=Presentation(folder/'native.pptx')
            self.assertEqual(4,len(deck.slides))
            self.assertEqual(4,sum(s.has_table for slide in deck.slides for s in slide.shapes))
            self.assertEqual(1,sum(s.has_chart for slide in deck.slides for s in slide.shapes))
            data, *_=ppt_html.prepare(job)
            for slide,page in zip(deck.slides,data['presentationPlan']['pages']):
                self.assertEqual(len(page['elements']),len(slide.shapes))
                for shape,element in zip(slide.shapes,page['elements']):
                    for coordinate,key in ((shape.left.pt,'x'),(shape.top.pt,'y'),
                                           (shape.width.pt,'w'),(shape.height.pt,'h')):
                        self.assertAlmostEqual(element[key],coordinate,places=3)
                table=next(s.table for s in slide.shapes if s.has_table)
                self.assertEqual('실제 개선 효과 확인',table.cell(1,1).text)
                self.assertEqual('미정',table.cell(2,2).text)
                text='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
                self.assertIn('8월 실적은 130건입니다',text)
                self.assertIn('개선 효과는 추가 확인이 필요합니다.',text)
            chart=next(s.chart for slide in deck.slides for s in slide.shapes if s.has_chart)
            self.assertEqual((110,130),chart.series[0].values)
            changed=copy.deepcopy(job)
            changed['slides'][0]['layout']='actions'
            result=artifacts.create_ppt(changed,folder/'changed.pptx',require_choices=True)
            self.assertEqual('design_review_changed',result['code'])
            self.assertFalse((folder/'changed.pptx').exists())


if __name__ == '__main__':
    unittest.main()
