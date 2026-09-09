from __future__ import annotations

import hashlib
import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "company-agent-plugin" / "scripts"))
from company_agent import business_artifacts as artifacts


class BusinessArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    @staticmethod
    def spec():
        return {"title": "월간 실적 — 검토", "subtitle": "한글 & 안전", "style": "minimalism", "mode": "both",
                "sections": [{"title": "생산 현황", "body": "실적을 확인했습니다.", "bullets": ["이상 항목 확인", "다음 조치"]}]}

    def fixture(self, filename="template.pptx", extras=None):
        result = self.root / filename
        parts = {"[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
                 "ppt/presentation.xml": f'<p:presentation xmlns:p="{artifacts.P}"><p:sldSz cx="12192000" cy="6858000"/></p:presentation>',
                 "ppt/slides/slide1.xml": f'<p:sld xmlns:p="{artifacts.P}"><p:cSld><p:spTree><p:sp/></p:spTree></p:cSld></p:sld>',
                 "ppt/slideLayouts/slideLayout1.xml": f'<p:sldLayout xmlns:p="{artifacts.P}"><p:cSld name="기본 양식"/></p:sldLayout>'}
        parts.update(extras or {})
        with zipfile.ZipFile(result, "w") as archive:
            for name, value in parts.items():
                archive.writestr(name, value)
        return result

    def test_html_all_styles_and_navigation_modes(self):
        for style in artifacts.STYLES:
            for mode in artifacts.MODES:
                with self.subTest(style=style, mode=mode):
                    spec = {**self.spec(), "style": style, "mode": mode}
                    output = self.root / f"{style}-{mode}.html"
                    result = artifacts.create_html(spec, output)
                    self.assertTrue(result["ok"], result)
                    text = output.read_text(encoding="utf-8")
                    self.assertIn(f'data-style="{style}"', text)
                    self.assertIn("Content-Security-Policy", text)
                    self.assertIn("script-src &#x27;sha256-", text)
                    self.assertNotIn('src="https:', text)
                    self.assertNotIn('href="https:', text)
                    self.assertIn("월간 실적 — 검토", text)
                    self.assertEqual(mode == "both", 'id="toggle-view"' in text)

    def test_user_content_cannot_inject_html_css_or_script(self):
        evil = '<img src=x onerror=alert(1)></script><script>alert("x")</script>'
        spec = self.spec()
        spec["title"] = evil
        spec["sections"] = [{"title": evil, "body": evil, "bullets": [evil], "table": {"headers": [evil], "rows": [[evil]]}}]
        result = artifacts.create_html(spec, self.root / "escape.html")
        self.assertTrue(result["ok"], result)
        text = (self.root / "escape.html").read_text(encoding="utf-8")
        self.assertNotIn(evil, text)
        self.assertEqual(1, text.count("<script>"))
        self.assertIn("&lt;img", text)

    def test_html_does_not_overwrite_or_create_unspecified_parent(self):
        target = self.root / "keep.html"
        target.write_text("existing", encoding="utf-8")
        self.assertEqual("output_exists", artifacts.create_html(self.spec(), target)["code"])
        self.assertEqual("existing", target.read_text())
        result = artifacts.create_html(self.spec(), self.root / "new" / "report.html")
        self.assertEqual("output_folder_missing", result["code"])
        self.assertFalse((self.root / "new").exists())

    def test_network_device_and_ads_paths_rejected_before_filesystem_access(self):
        if os.name != "nt":
            # The shared path validator has platform-specific Windows rules.
            return
        for path in (r"\\example.invalid\share\image.png", r"\\?\C:\image.png", r"C:\report.txt:image.png"):
            with self.subTest(path=path), patch.object(Path, "lstat", side_effect=AssertionError("No remote stat allowed")):
                with self.assertRaises(artifacts.ArtifactError) as caught:
                    artifacts._source(Path(path), (".png",))
                self.assertEqual("unsupported_path", caught.exception.code)
                with self.assertRaises(artifacts.ArtifactError):
                    artifacts._target(Path(path), ".html")

    def test_long_titles_have_explicit_mobile_wrapping(self):
        spec = self.spec()
        spec["title"] = "LongTitleWithoutSpaces" * 10
        output = self.root / "long.html"
        self.assertTrue(artifacts.create_html(spec, output)["ok"])
        self.assertIn("h1,h2{overflow-wrap:anywhere}", output.read_text(encoding="utf-8"))

    def test_preview_permission_failure_keeps_completed_ppt_in_partial_result(self):
        fixture = self.fixture()
        output = self.root / "completed.pptx"
        publish = artifacts._publish

        def generate(data, draft, template):
            shutil.copyfile(fixture, draft)
            return {"text": 1}

        def render(data, draft, template, work, render_only):
            preview = work / "preview"
            preview.mkdir()
            (preview / "slide-1.png").write_bytes(b"synthetic preview, not for rendering")
            return {"ok": True}

        def publish_with_denial(source, target):
            if target.suffix == ".png":
                raise PermissionError("PRIVATE-PREVIEW-PATH")
            publish(source, target)

        with patch.object(artifacts.importlib.util, "find_spec", return_value=object()), \
                patch.object(artifacts, "_python_ppt", side_effect=generate), \
                patch.object(artifacts, "_office", side_effect=render), \
                patch.object(artifacts, "_publish", side_effect=publish_with_denial):
            result = artifacts.create_ppt(self.spec(), output)
        self.assertTrue(result["ok"], result)
        self.assertEqual("partial", result["status"])
        self.assertEqual(str(output), result["outputPath"])
        self.assertTrue(output.is_file())
        self.assertEqual("permission_denied", result["validation"]["render"]["code"])
        self.assertNotIn("PRIVATE-PREVIEW-PATH", json.dumps(result))

    def test_protected_input_and_permission_errors_do_not_leak_details(self):
        for flag in ({"drmRestricted": True}, {"protected": True}, {"permissionGranted": False}):
            self.assertEqual("protected_input", artifacts.create_html({**self.spec(), **flag}, self.root / "blocked.html")["code"])
        with patch.object(Path, "write_text", side_effect=PermissionError("PRIVATE-DOCUMENT-SENTINEL")):
            result = artifacts.create_html(self.spec(), self.root / "blocked.html")
        self.assertEqual("permission_denied", result["code"])
        self.assertNotIn("PRIVATE-DOCUMENT-SENTINEL", json.dumps(result))
        self.assertFalse((self.root / "blocked.html").exists())

    def test_invalid_specs_are_blocked_without_output(self):
        for spec in ({}, {**self.spec(), "style": "<script>"}, {**self.spec(), "sections": [{"table": {"headers": ["a"], "rows": [[1, 2]]}}]},
                     {**self.spec(), "sections": [{"chart": {"categories": ["a"], "series": [{"values": [float("inf")]}]}}]}):
            result = artifacts.create_html(spec, self.root / "invalid.html")
            self.assertFalse(result["ok"], result)
            self.assertFalse((self.root / "invalid.html").exists())

    def test_template_inspection_reports_shapes_layout_size_and_hash(self):
        path = self.fixture()
        result = artifacts.inspect_template(path)
        self.assertTrue(result["ok"], result)
        self.assertEqual(1, result["slideCount"])
        self.assertEqual(1, result["slides"][0]["textShapes"])
        self.assertEqual("기본 양식", result["layouts"][0]["name"])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), result["sha256"])
        self.assertFalse(result["hasExternalRelationships"])

    def test_unsafe_template_package_and_entities_are_blocked(self):
        for number, extra in enumerate(({"../escape.xml": "bad"}, {"ppt/vbaProject.bin": "bad"},
                                         {"ppt/embeddings/oleObject1.bin": "bad"},
                                         {"ppt/slides/slide1.xml": '<!DOCTYPE x [<!ENTITY y "value">]><x>&y;</x>'})):
            result = artifacts.inspect_template(self.fixture(f"unsafe-{number}.pptx", extra))
            self.assertFalse(result["ok"], result)
            self.assertEqual("blocked", result["status"])

    def test_encrypted_and_invalid_files_have_no_alternative_reader(self):
        path = self.root / "encrypted.pptx"
        path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1PRIVATE")
        self.assertEqual("protected_or_unsupported", artifacts.inspect_template(path)["code"])
        path.write_bytes(b"not a presentation")
        self.assertEqual("protected_or_invalid", artifacts.inspect_template(path)["code"])

    def test_external_template_relationships_refuse_generation(self):
        path = self.fixture(extras={"ppt/_rels/presentation.xml.rels": f'<Relationships xmlns="{artifacts.R}"><Relationship Id="rId1" TargetMode="External" Target="https://example.invalid/private"/></Relationships>'})
        self.assertTrue(artifacts.inspect_template(path)["hasExternalRelationships"])
        result = artifacts.create_ppt(self.spec(), self.root / "output.pptx", path)
        self.assertEqual("external_template_links", result["code"])
        self.assertFalse((self.root / "output.pptx").exists())

    def test_embedded_chart_workbook_cannot_hide_external_links(self):
        data = BytesIO()
        with zipfile.ZipFile(data, "w") as book:
            book.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{artifacts.R}"><Relationship TargetMode="External" Target="file:///private"/></Relationships>')
        path = self.fixture(extras={"ppt/embeddings/Microsoft_Excel_Sheet1.xlsx": data.getvalue()})
        self.assertEqual("unsafe_workbook", artifacts.inspect_template(path)["code"])

    def test_office_timeout_never_leaks_stdout_or_terminates_user_office(self):
        with patch.object(artifacts, "_windows_powershell", return_value="C:/Windows/System32/WindowsPowerShell/v1.0/powershell.exe"), \
                patch.object(artifacts.subprocess, "run", side_effect=subprocess.TimeoutExpired("powershell", 90, output=b"PRIVATE")):
            result = artifacts._office(self.spec(), self.root / "draft.pptx", None, self.root, True)
        self.assertEqual("office_timeout", result["code"])
        self.assertNotIn("PRIVATE", json.dumps(result))
        script = (ROOT / "company-agent-plugin/scripts/Invoke-BusinessPowerPoint.ps1").read_text(encoding="utf-8")
        self.assertNotIn(".Quit(", script)
        self.assertNotIn("Stop-Process", script)
        self.assertNotIn("ExecutionPolicy", script)
        self.assertIn("AutomationSecurity = 3", script)
        self.assertIn("Permission.Enabled", script)

    def test_office_uses_fixed_absolute_system_executable(self):
        executable = "C:/Windows/System32/WindowsPowerShell/v1.0/powershell.exe"
        response = subprocess.CompletedProcess([], 0, b'{"ok":false,"code":"protected_input","message":"Blocked."}', b"PRIVATE-ERROR")
        with patch.object(artifacts, "_windows_powershell", return_value=executable), patch.object(artifacts.subprocess, "run", return_value=response) as run:
            result = artifacts._office(self.spec(), self.root / "draft.pptx", None, self.root, True)
        self.assertEqual(executable, run.call_args.args[0][0])
        self.assertNotIn("PRIVATE-ERROR", json.dumps(result))
        with patch.object(artifacts.shutil, "which", side_effect=AssertionError("PATH resolution forbidden")):
            detected = artifacts._windows_powershell()
            self.assertTrue(detected is None or Path(detected).is_absolute())

    def test_missing_optional_engines_returns_actionable_result_without_download(self):
        with patch.object(artifacts.importlib.util, "find_spec", return_value=None), patch.object(artifacts, "_office", return_value={
                "ok": False, "status": "unavailable", "code": "powerpoint_unavailable", "message": "Install approved Office separately."}):
            result = artifacts.create_ppt(self.spec(), self.root / "missing.pptx")
        self.assertFalse(result["ok"])
        self.assertEqual("powerpoint_unavailable", result["code"])
        self.assertFalse((self.root / "missing.pptx").exists())

    def test_real_editable_ppt_objects_and_template_original_preserved_when_optional_library_present(self):
        if importlib.util.find_spec("pptx") is None:
            self.assertFalse(artifacts.capabilities()["ppt"]["pythonPptxAvailable"])
            return
        from pptx import Presentation
        template = self.root / "real-template.pptx"
        original = Presentation()
        slide = original.slides.add_slide(original.slide_layouts[0])
        slide.shapes.title.text = "REMOVE-ORIGINAL-EXAMPLE-SENTINEL"
        original.save(str(template))
        before = template.read_bytes()
        spec = self.spec()
        spec["sections"] = [{"title": "표 보고", "table": {"headers": ["월", "실적"], "rows": [["8월", 132]]}},
                            {"title": "차트 보고", "chart": {"type": "column", "categories": ["7월", "8월"], "series": [{"name": "실적", "values": [120, 132]}]}}]
        no_render = {"ok": False, "status": "unknown", "code": "render_refused", "message": "No rendering permission. No fallback."}
        with patch.object(artifacts, "_office", return_value=no_render) as office:
            result = artifacts.create_ppt(spec, self.root / "result.pptx", template)
        self.assertTrue(result["ok"], result)
        self.assertEqual("partial", result["status"])
        self.assertEqual(1, office.call_count)
        self.assertEqual(before, template.read_bytes())
        loaded = Presentation(result["outputPath"])
        self.assertEqual(2, len(loaded.slides))
        self.assertEqual(1, sum(s.has_table for slide in loaded.slides for s in slide.shapes))
        self.assertEqual(1, sum(s.has_chart for slide in loaded.slides for s in slide.shapes))
        self.assertEqual(2, len(loaded.slides[1].shapes[-1].chart.series[0].values))
        with zipfile.ZipFile(result["outputPath"]) as package:
            combined = b"".join(package.read(n) for n in package.namelist() if n.endswith(".xml"))
            self.assertNotIn(b"REMOVE-ORIGINAL-EXAMPLE-SENTINEL", combined)
            self.assertTrue(any(n.startswith("ppt/theme/") for n in package.namelist()))

    def test_dense_slides_are_not_silently_truncated(self):
        spec = self.spec()
        spec["sections"][0]["body"] = "가" * 1000
        result = artifacts.create_ppt(spec, self.root / "dense.pptx")
        self.assertEqual("slide_too_dense", result["code"])
        self.assertFalse((self.root / "dense.pptx").exists())


if __name__ == "__main__":
    unittest.main()
