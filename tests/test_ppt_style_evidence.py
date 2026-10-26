"""Bounded reference evidence and editable Korean text, using synthetic PPTX only."""
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'company-agent-plugin/scripts'))
from company_agent import ppt_workflow, ppt_html, presentation_design


class PptStyleEvidenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'reference.pptx'

    def deck(self, count=1):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        from pptx.enum.text import PP_ALIGN
        from pptx.dml.color import RGBColor
        from pptx.util import Pt
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for i in range(count):
            box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                        Pt(20+i*4), Pt(80), Pt(200), Pt(100))
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor.from_string(f'{0x123450+i:06X}')
            box.line.color.rgb = RGBColor.from_string('445566')
            box.line.width = Pt(1.5)
            box.text = '납기 편차 감소\n원인별 재발 방지 검토'
            paragraph = box.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if i % 2 == 0 else PP_ALIGN.CENTER
            paragraph.space_after = Pt(6+i)
            paragraph.line_spacing = 1.2
            paragraph.font.name = 'Malgun Gothic'
        deck.save(self.path)
        return deck

    def analyze(self):
        with patch.object(ppt_workflow, 'installed_fonts', return_value=set()):
            result = ppt_workflow.analyze(self.path, first_slides=3)
        self.assertTrue(result['ok'], result)
        return result

    def test_shape_and_paragraph_evidence_is_linked_to_original_objects(self):
        self.deck()
        before = self.path.read_bytes()
        result = self.analyze()
        page = result['slides'][0]
        style = page['declaredShapeStyles'][0]
        self.assertEqual(style['preset'], 'roundRect')
        self.assertEqual(style['fill']['color'], {'srgbClr':'123450'})
        self.assertEqual(style['line']['widthEmu'], '19050')
        self.assertEqual(style['line']['paint']['color'], {'srgbClr':'445566'})
        paragraph = page['declaredParagraphStyles'][0]
        self.assertEqual(paragraph['algn'], 'l')
        self.assertEqual(paragraph['spcAft'], {'spcPts':'600'})
        self.assertEqual(paragraph['lnSpc'], {'spcPct':'120000'})
        self.assertEqual(style['shapeId'], paragraph['shapeId'])
        self.assertIn(style['shapeId'], [shape['shapeId'] for shape in page['shapes']])
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(result['visualReview'], 'not_performed')
        self.assertFalse(result['savedToMemory'])

    def test_reference_output_remains_bounded_when_every_object_has_different_style(self):
        self.deck(90)
        page = self.analyze()['slides'][0]
        self.assertEqual(len(page['declaredShapeStyles']), 12)
        self.assertEqual(len(page['declaredParagraphStyles']), 8)
        self.assertEqual(len(page['shapes']), 80)
        self.assertTrue(page['shapesTruncated'])

    def test_inherited_colors_are_not_reported_as_resolved_rgb(self):
        from pptx.enum.dml import MSO_THEME_COLOR
        deck = self.deck()
        shape = deck.slides[0].shapes[0]
        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shape.fill.fore_color.brightness = 0.4
        shape.line.fill.background()
        deck.save(self.path)
        style = self.analyze()['slides'][0]['declaredShapeStyles'][0]
        self.assertEqual(style['fill']['color'], {'schemeClr':'accent1'})
        self.assertEqual(style['fill']['adjustments'], {'lumMod':'60000','lumOff':'40000'})
        self.assertEqual(style['line']['paint'], {'type':'none'})

    def test_missing_declarations_do_not_invent_a_shape_fill_or_paragraph_rule(self):
        from pptx import Presentation
        from pptx.util import Pt
        deck = Presentation()
        box = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(Pt(40), Pt(40), Pt(400), Pt(100))
        box.text = '확인된 사실만 보고'
        deck.save(self.path)
        page = self.analyze()['slides'][0]
        self.assertFalse(page['declaredParagraphStyles'])
        self.assertTrue(all(style.get('fill') in (None, {'type':'none'}) for style in page['declaredShapeStyles']))

    def test_native_korean_font_and_user_chosen_style_survive_roundtrip(self):
        from pptx import Presentation
        spec = {'title':'납기 검토', 'presentationFont':'Malgun Gothic', 'slides':[
            {'title':'납기 편차 검토', 'elements':[
                {'kind':'shape','group':'검토사항','shape':'roundRect','x':40,'y':100,'w':800,'h':220,'fill':'EEF4F7'},
                {'kind':'text','group':'검토사항','text':'납기 편차 감소 검토','x':65,'y':120,'w':740,'h':60,
                 'size':28,'bold':True,'color':'17324D','font':'Malgun Gothic'},
                {'kind':'text','group':'검토사항','text':'표본 20건 · 추가 검증 필요','x':65,'y':210,'w':740,'h':60,
                 'size':20,'bold':False,'color':'30465B','font':'Malgun Gothic'}]}]}
        prepared = ppt_html.prepare(spec)[0]
        presentation_design.render_python(prepared, self.path, None)
        group = Presentation(self.path).slides[0].shapes[0]
        for shape in list(group.shapes)[1:]:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    self.assertEqual(run.font.name, 'Malgun Gothic')
                    self.assertEqual(run._r.rPr.get('lang'), 'ko-KR')
                    self.assertEqual(run._r.rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea').get('typeface'), 'Malgun Gothic')
        self.assertTrue(group.shapes[1].text_frame.paragraphs[0].runs[0].font.bold)
        self.assertFalse(group.shapes[2].text_frame.paragraphs[0].runs[0].font.bold)
        self.assertFalse(group._element.xpath('.//a:grpSpLocks[@noUngrp="1"]'))

    def test_grouped_overflow_warns_without_changing_or_blocking_editable_output(self):
        from pptx import Presentation
        from pptx.util import Pt
        deck = self.deck()
        slide = deck.slides[0]
        shape = slide.shapes[0]
        shape.width, shape.height = Pt(40), Pt(20)
        shape.text = '매우 긴 한글 설명을 작은 영역에 배치한 예시 ' * 20
        for run in shape.text_frame.paragraphs[0].runs:
            run.font.size = Pt(18)
        slide.shapes.add_group_shape([shape])
        deck.save(self.path)
        before = self.path.read_bytes()
        result = ppt_workflow.quality(self.path)
        warning = next(row for row in result['issues'] if row['code'] == 'estimated_text_overflow')
        self.assertEqual(warning['basis'], 'approximate_text_metrics')
        self.assertEqual(warning['severity'], 'warning')
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(len(Presentation(self.path).slides[0].shapes[0].shapes), 1)

    def test_normal_korean_text_does_not_warn_about_overflow(self):
        self.deck()
        self.assertNotIn('estimated_text_overflow', [row['code'] for row in ppt_workflow.quality(self.path)['issues']])


if __name__ == '__main__':
    unittest.main()
