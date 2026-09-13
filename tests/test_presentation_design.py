import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'company-agent-plugin/scripts'))
from company_agent import business_artifacts as artifacts, presentation_design as design


class PresentationDesignTests(unittest.TestCase):
    def spec(self):
        return {'title': '업무 보고', 'facts': [
            {'id': 'actual', 'label': '누적 실적', 'unit': '백만원', 'inputs': [{'chart': [0, 0, 0]}, {'chart': [0, 0, 1]}, {'chart': [0, 0, 2]}]}],
            'slides': [{'title': '실적 {{fact:actual}}백만원', 'kpis': [{'fact': 'actual'}],
                        'chart': {'type': 'column', 'title': '실적 · 백만원', 'categories': ['6월', '7월', '8월'],
                                  'series': [{'name': '실적', 'values': [145, 175, 210]}]}, 'body': '월별 실적을 비교합니다.'}]}

    def prepared(self, spec=None):
        spec = spec or self.spec()
        data = artifacts._normalize(spec)
        validation = design.prepare(spec, data)
        data['presentationPlan'] = design.plan(data)
        return data, validation

    def test_numbers_are_bound_once_to_title_and_metrics(self):
        data, validation = self.prepared()
        self.assertEqual('실적 530백만원', data['sections'][0]['title'])
        self.assertEqual('530', data['sections'][0]['kpis'][0]['value'])
        self.assertEqual('checked', validation['arithmetic'])
        json.dumps(data)  # COM must receive JSON-compatible data, never Decimal.

    def test_wrong_510_is_rejected_before_any_engine_or_file(self):
        spec = self.spec()
        spec['facts'][0]['expected'] = 510
        with tempfile.TemporaryDirectory() as t, patch.object(artifacts, '_office', side_effect=AssertionError('not reached')):
            result = artifacts.create_ppt(spec, Path(t)/'bad.pptx')
            self.assertFalse(result['ok'])
            self.assertEqual('ppt_design_invalid', result['code'])
            self.assertFalse(list(Path(t).iterdir()))

    def test_prose_sits_beside_chart_not_in_equal_vertical_blocks(self):
        data, _ = self.prepared()
        e = data['presentationPlan']['pages'][0]['elements']
        chart = next(x for x in e if x['kind'] == 'chart')
        prose = next(x for x in e if x.get('text') == '월별 실적을 비교합니다.')
        self.assertGreater(prose['x'], chart['x']+chart['w'])
        self.assertLess(abs(prose['y']-chart['y']), 40)
        self.assertTrue(any(x.get('size') == 38 for x in e))

    def test_all_bounds_fit_16_9_and_4_3(self):
        spec = {'title': '표 보고', 'slides': [{'title': '현황', 'table': {'headers': ['항목', '수량'], 'rows': [['자료', '5']]}}]}
        data, _ = self.prepared(spec)
        for w, h in ((960,540), (720,540)):
            plan = design.plan(data, w,h)
            for page in plan['pages']:
                for e in page['elements']:
                    self.assertGreaterEqual(e['x'], 0)
                    self.assertLessEqual(e['x']+e['w'], w)
                    self.assertLessEqual(e['y']+e['h'], h)

    def test_korean_dense_cells_rejected_without_small_fonts(self):
        spec = {'title': '표', 'slides': [{'title': '상세', 'table': {'headers': ['대상', '상세 설명'],
                'rows': [['항목', '설명이매우길어서여러줄이됩니다'*5] for _ in range(7)]}}]}
        with self.assertRaises(design.DesignError):
            self.prepared(spec)

    def test_too_long_title_rejected_without_truncation(self):
        spec = self.spec()
        spec['slides'][0]['title'] = '긴제목'*28
        with self.assertRaises(design.DesignError):
            self.prepared(spec)

    def test_unknown_layout_or_unbound_kpi_rejected(self):
        for extra in ({'layout': 'neumorphic-magic'}, {'kpis': [{'value': 510}]}, {'kpis': [{'fact': 'missing'}]}):
            spec = self.spec()
            spec['slides'][0].update(extra)
            with self.assertRaises(design.DesignError):
                self.prepared(spec)

    def test_no_visual_silently_dropped(self):
        spec = self.spec()
        spec['slides'][0]['table'] = {'headers': ['실적'], 'rows': [[530]]}
        with self.assertRaises(design.DesignError):
            self.prepared(spec)  # chart + table + prose cannot fit this layout

    def test_native_table_fonts_and_column_widths(self):
        spec = {'title': '확인', 'slides': [{'title': '후속 과제', 'table': {
            'headers': ['항목', '확인할 내용', '상태'], 'rows': [['재작업', '품질과 비용의 개선 효과 확인', '미정']]}}]}
        data, _ = self.prepared(spec)
        element = next(e for e in data['presentationPlan']['pages'][0]['elements'] if e['kind'] == 'table')
        self.assertEqual(16, element['size'])
        self.assertGreater(element['columnWidths'][1], element['columnWidths'][0])
        self.assertAlmostEqual(element['h'], sum(element['rowHeights']))

    def test_reference_example_is_valid(self):
        doc = (ROOT/'company-agent-plugin/skills/presentation/references/design-and-quality.md').read_text(encoding='utf-8')
        spec = json.loads(doc.split('```json\n')[1].split('```')[0])
        data, _ = self.prepared(spec)
        self.assertEqual('목표 대비 128.33% 달성', data['sections'][0]['title'])

    def test_real_native_chart_values_and_all_text_are_preserved(self):
        if importlib.util.find_spec('pptx') is None:
            self.assertFalse(artifacts.capabilities()['ppt']['pythonPptxAvailable'])
            return
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as t:
            data, _ = self.prepared()
            path = Path(t)/'native.pptx'
            native = design.render_python(data, path, None)
            prs = Presentation(path)
            chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
            self.assertEqual((145.,175.,210.), chart.series[0].values)
            for axis_id in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
                self.assertGreaterEqual(int(axis_id.get('val')), 0)
            self.assertEqual(1, native['charts'])
            texts = '\n'.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
            self.assertIn('실적 530백만원', texts)
            self.assertNotIn('{{fact:', texts)
            for shape in prs.slides[0].shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        self.assertIn('typeface="Malgun Gothic"', para.font._rPr.xml)
                        self.assertIn('<a:ea', para.font._rPr.xml)

    def test_four_chart_types_keep_editable_series(self):
        if importlib.util.find_spec('pptx') is None:
            self.assertFalse(artifacts.capabilities()['ppt']['pythonPptxAvailable'])
            return
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as t:
            for kind in ('column','bar','line','pie'):
                spec = self.spec()
                spec['slides'][0]['chart']['type'] = kind
                data,_ = self.prepared(spec)
                path = Path(t)/(kind+'.pptx')
                design.render_python(data,path,None)
                self.assertEqual(1, sum(s.has_chart for s in Presentation(path).slides[0].shapes))

    def test_com_consumes_same_geometry_and_keeps_protection(self):
        script = (ROOT/'company-agent-plugin/scripts/Invoke-BusinessPowerPoint.ps1').read_text(encoding='utf-8')
        for token in ('presentationPlan', '$element.columnWidths', '$element.rowHeights', 'missing_presentation_plan', 'office_chart_unavailable'):
            self.assertIn(token,script)
        for token in ('$blockHeight', '.Quit(', 'Stop-Process', 'ExecutionPolicy'):
            self.assertNotIn(token,script)

    def test_unsafe_theme_values_rejected(self):
        for theme in ({'accent': 'red'}, {'font': 'anything'}, {'text': 'FFFFFF'}):
            spec = self.spec()
            spec['presentationTheme'] = theme
            with self.assertRaises(design.DesignError):
                self.prepared(spec)


if __name__ == '__main__':
    unittest.main()
