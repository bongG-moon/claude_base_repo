from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT/'company-agent-plugin/scripts'
sys.path.insert(0,str(SCRIPTS))
from company_agent import artifact_delivery as delivery, business_artifacts as artifacts


class ArtifactDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.state = self.root/'state'
        self.project = self.root/'공통 하네스 테스트'
        self.project.mkdir()
        self.source = self.project/'원본.txt'
        self.source.write_text('사용자 원본',encoding='utf-8')
        self.html = {'title':'임원 보고','style':'minimal','length':'standard','mode':'scroll',
                     'sections':[{'title':'요약','body':'처음 내용'}]}
        self.ppt = {'title':'월간 보고','creationMode':'new','purpose':'보고','audience':'부서장',
                    'slideCount':2,'designPreset':'warm', 'slides':[
                        {'title':'요약','body':'6월부터 8월까지 실적'},
                        {'title':'월별 실적','chart':{'type':'column','categories':['6월','7월','8월'],
                            'series':[{'name':'실적','values':[145,175,210]}]}}]}

    def start(self,name='보고서.html'):
        result = delivery.start(self.state,self.project/name)
        self.assertTrue(result['ok'],result)
        self.assertFalse(result['finalCreated'])
        return result['workFile']

    def build(self,work,operation='html',spec=None):
        result = delivery.build(self.state,work,operation,spec or self.html)
        self.assertTrue(result['ok'],result)
        self.assertEqual([],result['deliverables'])
        self.assertFalse(Path(result['outputPath']).is_relative_to(self.project))
        return result

    def assert_only_source(self):
        self.assertEqual([self.source],list(self.project.iterdir()))

    def test_revisions_publish_one_final_only_and_retry_is_idempotent(self):
        work = self.start()
        self.assert_only_source()
        for text in ('초안','검사 후 수정','확인한 최종본'):
            self.html['sections'][0]['body'] = text
            result = self.build(work)
            self.assert_only_source()
        final = delivery.publish(self.state,work)
        self.assertTrue(final['ok'],final)
        self.assertEqual([str(self.project/'보고서.html')],final['deliverables'])
        self.assertIn('확인한 최종본',Path(final['outputPath']).read_text(encoding='utf-8'))
        self.assertEqual(Path(result['outputPath']).read_bytes(),Path(final['outputPath']).read_bytes())
        self.assertEqual({self.source,self.project/'보고서.html'},set(self.project.iterdir()))
        stamp = (self.project/'보고서.html').stat().st_mtime_ns
        with patch.object(artifacts,'create_html',side_effect=AssertionError('must not regenerate')):
            self.assertTrue(delivery.build(self.state,work,'html',self.html)['alreadyDelivered'])
        self.assertTrue(delivery.publish(self.state,work)['alreadyDelivered'])
        self.assertEqual(stamp,(self.project/'보고서.html').stat().st_mtime_ns)
        self.assertEqual('사용자 원본',self.source.read_text(encoding='utf-8'))

    def test_later_request_may_create_a_second_result_without_changing_first(self):
        first = self.start()
        self.build(first)
        self.assertTrue(delivery.publish(self.state,first)['ok'])
        before = (self.project/'보고서.html').read_bytes()
        second = self.start('추가 요청.html')
        self.build(second)
        self.assertTrue(delivery.publish(self.state,second)['ok'])
        self.assertEqual(before,(self.project/'보고서.html').read_bytes())
        self.assertEqual(2,len(list(self.project.glob('*.html'))))

    def test_preexisting_or_racing_output_is_preserved_without_auto_version(self):
        work = self.start()
        self.build(work)
        final = self.project/'보고서.html'
        final.write_text('기존 사용자 파일',encoding='utf-8')
        self.assertEqual('output_exists',delivery.publish(self.state,work)['code'])
        self.assertEqual('output_exists',delivery.start(self.state,final)['code'])
        self.assertEqual('기존 사용자 파일',final.read_text(encoding='utf-8'))
        self.assertEqual([final],list(self.project.glob('*.html')))

    def test_changed_delivered_file_is_not_overwritten(self):
        work = self.start()
        self.build(work)
        delivery.publish(self.state,work)
        final = self.project/'보고서.html'
        final.write_text('사용자가 직접 수정',encoding='utf-8')
        for result in (delivery.publish(self.state,work),delivery.build(self.state,work,'html',self.html)):
            self.assertEqual('delivered_file_changed',result['code'])
        self.assertEqual('사용자가 직접 수정',final.read_text(encoding='utf-8'))

    def test_failed_or_protected_latest_revision_invalidates_old_candidate(self):
        for bad in ({**self.html,'style':'unknown'}, {**self.html,'protection':'blocked'}):
            with self.subTest(bad=bad):
                work = self.start()
                self.build(work)
                self.assertFalse(delivery.build(self.state,work,'html',bad)['ok'])
                self.assertEqual('artifact_not_ready',delivery.publish(self.state,work)['code'])
                self.assert_only_source()

    def test_candidate_changed_outside_builder_must_be_rebuilt(self):
        work = self.start()
        result = self.build(work)
        Path(result['outputPath']).write_text('unverified edit',encoding='utf-8')
        self.assertEqual('artifact_candidate_changed',delivery.publish(self.state,work)['code'])
        self.assert_only_source()

    def test_concurrent_publish_creates_exactly_one_result(self):
        work = self.start()
        self.build(work)
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: delivery.publish(self.state,work),range(2)))
        self.assertTrue(all(r['ok'] for r in results),results)
        self.assertEqual(1,sum(r.get('alreadyDelivered',False) for r in results))
        self.assertEqual([self.project/'보고서.html'],list(self.project.glob('*.html')))

    def test_no_workspace_for_invalid_final_format_or_missing_folder(self):
        for output in (self.project/'report.pdf',self.project/'missing/report.html'):
            self.assertFalse(delivery.start(self.state,output)['ok'])
        self.assertFalse(self.state.exists())

    def test_draft_is_not_final_and_old_or_foreign_draft_cannot_be_used(self):
        work = self.start('발표.pptx')
        first = self.build(work,'ppt-design-preview',self.ppt)
        self.assertEqual('artifact_not_ready',delivery.publish(self.state,work)['code'])
        second = self.build(work,'ppt-design-preview',self.ppt)
        self.ppt['designReview'] = {**first['designReview'],'confirmed':True}
        self.assertEqual('artifact_preview_required',delivery.build(self.state,work,'ppt',self.ppt)['code'])
        foreign = self.start('다른 작업.pptx')
        self.ppt['designReview'] = {**second['designReview'],'confirmed':True}
        self.assertEqual('artifact_preview_required',delivery.build(self.state,foreign,'ppt',self.ppt)['code'])
        self.assert_only_source()

    def test_ppt_keeps_drafts_and_review_images_internal_and_native_objects_editable(self):
        work = self.start('월간 보고.pptx')
        for text in ('초안','수정','최종 내용'):
            self.ppt['slides'][0]['body'] = text
            draft = self.build(work,'ppt-design-preview',self.ppt)
        self.ppt['designReview'] = {**draft['designReview'],'confirmed':True}
        def render(_data,_draft,_template,folder,_render_only):
            preview = folder/'preview'
            preview.mkdir()
            for index in (1,2):
                (preview/f'slide-{index}.png').write_bytes(b'synthetic test image')
            return {'ok':True}
        with patch.object(artifacts,'_office',side_effect=render):
            built = self.build(work,'ppt',self.ppt)
        self.assertEqual(2,len(built['previews']))
        for preview in built['previews']:
            self.assertTrue(Path(preview).is_relative_to(Path(work).parent))
        self.assert_only_source()
        final = delivery.publish(self.state,work)
        self.assertTrue(final['ok'],final)
        self.assertEqual({self.source,self.project/'월간 보고.pptx'},set(self.project.iterdir()))
        from pptx import Presentation
        deck = Presentation(final['outputPath'])
        self.assertEqual(2,len(deck.slides))
        self.assertEqual(1,sum(s.has_chart for slide in deck.slides for s in slide.shapes))
        self.assertIn('최종 내용',' '.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame))

    def test_partial_render_status_is_preserved_on_publish_and_repeat(self):
        work = self.start('부분 확인.pptx')
        draft = self.build(work,'ppt-design-preview',self.ppt)
        self.ppt['designReview'] = {**draft['designReview'],'confirmed':True}
        with patch.object(artifacts,'_office',return_value={'ok':False,'code':'render_refused'}):
            built = self.build(work,'ppt',self.ppt)
        self.assertEqual('partial',built['status'])
        for _ in range(2):
            final = delivery.publish(self.state,work)
            self.assertEqual('partial',final['status'])
            self.assertEqual(built['warnings'],final['warnings'])

    def test_template_is_separate_only_when_requested_and_remains_one_html(self):
        work = self.start('발표.pptx')
        draft = self.build(work,'ppt-design-preview',self.ppt)
        self.ppt['designReview'] = {**draft['designReview'],'confirmed':True}
        saved = self.start('대표 양식.html')
        self.build(saved,'ppt-template',self.ppt)
        self.assert_only_source()
        final = delivery.publish(self.state,saved)
        self.assertTrue(final['ok'],final)
        self.assertNotIn('6월부터',Path(final['outputPath']).read_text(encoding='utf-8'))
        self.assertEqual({self.source,self.project/'대표 양식.html'},set(self.project.iterdir()))

    def test_lost_publish_receipt_reuses_identical_final_without_rewriting(self):
        work = self.start()
        built = self.build(work)
        final = self.project/'보고서.html'
        final.write_bytes(Path(built['outputPath']).read_bytes())
        stamp = final.stat().st_mtime_ns
        self.assertTrue(delivery.publish(self.state,work)['ok'])
        self.assertEqual(stamp,final.stat().st_mtime_ns)

    def test_wrong_scope_and_arbitrary_manifest_path_fail_without_writes(self):
        work = self.start()
        self.assertFalse(delivery.build(self.root/'other-state',work,'html',self.html)['ok'])
        fake = self.project/'work.json'
        fake.write_text(Path(work).read_text(encoding='utf-8'),encoding='utf-8')
        self.assertFalse(delivery.publish(self.state,fake)['ok'])
        self.assertFalse((self.project/'보고서.html').exists())

    def test_foreign_candidate_path_and_wrong_output_kind_are_rejected(self):
        work = self.start()
        built = self.build(work)
        data = json.loads(Path(work).read_text(encoding='utf-8'))
        data['candidate']['path'] = str(self.source)
        Path(work).write_text(json.dumps(data),encoding='utf-8')
        self.assertEqual('invalid_artifact_work',delivery.publish(self.state,work)['code'])
        self.assertEqual('artifact_kind_mismatch',delivery.build(self.state,work,'ppt',self.ppt)['code'])
        self.assertEqual('artifact_not_ready',delivery.publish(self.state,work)['code'])
        self.assert_only_source()

    def test_linked_state_is_rejected(self):
        real = self.root/'real'
        real.mkdir()
        linked = self.root/'linked'
        try:
            linked.symlink_to(real,target_is_directory=True)
        except OSError:
            self.skipTest('This Windows account cannot create symbolic links.')
        self.assertFalse(delivery.start(linked,self.project/'보고서.html')['ok'])
        self.assertEqual([],list(real.iterdir()))

    def test_cli_korean_paths_and_output_work_exclusion(self):
        def run(*args):
            return subprocess.run([sys.executable,'-X','utf8','-B',str(SCRIPTS/'harness_cli.py'),
                'business',*args,'--state-root',str(self.state)],capture_output=True,text=True,encoding='utf-8',timeout=20)
        started = run('artifact-start','--output',str(self.project/'보고서.html'))
        self.assertEqual(0,started.returncode,started.stderr)
        record = json.loads(started.stdout)
        spec = Path(record['jobPath'])
        spec.write_text(json.dumps(self.html,ensure_ascii=False),encoding='utf-8')
        call = run('html','--spec',str(spec),'--work',record['workFile'])
        self.assertEqual(0,call.returncode,call.stdout+call.stderr)
        self.assertTrue(json.loads(call.stdout)['ok'])
        self.assert_only_source()
        self.assertNotEqual(0,run('html','--spec',str(spec),'--work',record['workFile'],
                                 '--output',str(self.project/'extra.html')).returncode)
        denied = {**self.html,'protection':'blocked'}
        spec.write_text(json.dumps(denied,ensure_ascii=False),encoding='utf-8')
        self.assertEqual('protection_blocked',json.loads(run('html','--spec',str(spec),'--work',record['workFile']).stdout)['code'])
        self.assertEqual('artifact_not_ready',json.loads(run('artifact-publish','--work',record['workFile']).stdout)['code'])
        spec.write_text(json.dumps(self.html,ensure_ascii=False),encoding='utf-8')
        run('html','--spec',str(spec),'--work',record['workFile'])
        spec.write_text('{invalid json',encoding='utf-8')
        self.assertFalse(json.loads(run('html','--spec',str(spec),'--work',record['workFile']).stdout)['ok'])
        self.assertEqual('artifact_not_ready',json.loads(run('artifact-publish','--work',record['workFile']).stdout)['code'])
        spec.write_text(json.dumps(self.html,ensure_ascii=False),encoding='utf-8')
        run('html','--spec',str(spec),'--work',record['workFile'])
        final = json.loads(run('artifact-publish','--work',record['workFile']).stdout)
        self.assertTrue(final['ok'],final)
        self.assertEqual([str(self.project/'보고서.html')],final['deliverables'])


if __name__ == '__main__':
    unittest.main()
