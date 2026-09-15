import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'company-agent-plugin/scripts'))
from company_agent import ppt_workflow as flow, business_artifacts as artifacts


class PptWorkflowTests(unittest.TestCase):
    def test_questions_are_sequential_and_keep_known_brief(self):
        spec={'purpose':'보고','audience':'부서장','slideCount':2}
        self.assertEqual('method',flow.choices(spec)['stage'])
        spec['creationMode']='reference'
        self.assertEqual('reference_file',flow.choices(spec)['stage'])
        self.assertEqual('reference_scope',flow.choices(spec,'confirmed.pptx')['stage'])
        spec['referenceMode']='preserve'
        self.assertEqual('design_preview',flow.choices(spec,'confirmed.pptx')['stage'])
        self.assertEqual(2,flow.choices(spec,'confirmed.pptx')['preservedChoices']['slideCount'])

    def test_new_and_saved_choices(self):
        self.assertEqual(['purpose','audience','slideCount'],flow.choices({'creationMode':'new'})['missing'])
        self.assertEqual('reference_file',flow.choices({'creationMode':'saved'})['stage'])
        for spec in ({'creationMode':'bad'},{'slideCount':True},{'slideCount':61},{'audience':''}):
            self.assertEqual('invalid_choice',flow.choices(spec)['code'])

    def test_cli_does_not_create_before_choices(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            spec=folder/'spec.json'
            spec.write_text(json.dumps({'slides':[{'title':'내용'}]}),encoding='utf-8')
            p=subprocess.run([sys.executable,'-B',str(ROOT/'company-agent-plugin/scripts/harness_cli.py'),
                              'business','ppt','--spec',str(spec),'--output',str(folder/'not-created.pptx'),
                              '--state-root',str(folder/'state')],capture_output=True,text=True,timeout=15)
            self.assertEqual(0,p.returncode,p.stderr)
            self.assertEqual('method',json.loads(p.stdout)['stage'])
            self.assertFalse((folder/'not-created.pptx').exists())
            self.assertFalse((folder/'state').exists())

    def fixture(self,folder):
        from pptx import Presentation
        from pptx.util import Inches,Pt
        from pptx.dml.color import RGBColor
        deck=Presentation()
        for title in ('첫 장','둘째 장'):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1.5))
            run=box.text_frame.paragraphs[0].add_run()
            run.text=title
            run.font.name='Malgun Gothic'
            run.font.size=Pt(30)
            run.font.color.rgb=RGBColor.from_string('224466')
        path=folder/'reference.pptx'
        deck.save(str(path))
        return path

    def test_analysis_is_read_only_and_returns_real_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=self.fixture(Path(tmp)); before=path.read_bytes()
            result=flow.analyze(path)
            self.assertTrue(result['ok'],result)
            self.assertEqual([0,1],[s['sourceSlide'] for s in result['slides']])
            self.assertEqual(2,result['slides'][0]['shapes'][0]['shapeId'])
            self.assertIn('224466',result['declaredColors'])
            self.assertFalse(result['savedToMemory'])
            self.assertEqual(before,path.read_bytes())

    def test_fill_preserves_geometry_font_color_and_source(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); path=self.fixture(folder); before=path.read_bytes()
            spec={'creationMode':'reference','referenceMode':'preserve','purpose':'보고','audience':'부서장','slideCount':1,
                  'slides':[{'title':'새 결과'}], 'templateSlides':[{'sourceSlide':1,'title':2}]}
            with patch.object(artifacts,'_office',return_value={'ok':False,'code':'render_refused','message':'미리보기 제한'}):
                preview=artifacts.create_ppt(spec,folder/'draft.pptx',path,require_choices=True,preview_only=True)
                self.assertTrue(preview['ok'],preview)
                spec['designReview']={**preview['designReview'],'confirmed':True}
                result=artifacts.create_ppt(spec,folder/'out.pptx',path,require_choices=True)
            self.assertTrue(result['ok'],result)
            self.assertEqual('partial',result['status'])
            deck=Presentation(str(folder/'out.pptx'))
            old=Presentation(str(path)).slides[1].shapes[0]; new=deck.slides[0].shapes[0]
            self.assertEqual('새 결과',new.text)
            self.assertEqual((old.left,old.top,old.width,old.height),(new.left,new.top,new.width,new.height))
            self.assertEqual(old.text_frame.paragraphs[0].runs[0].font.color.rgb,new.text_frame.paragraphs[0].runs[0].font.color.rgb)
            self.assertEqual(old.text_frame.paragraphs[0].runs[0].font.name,new.text_frame.paragraphs[0].runs[0].font.name)
            self.assertEqual(before,path.read_bytes())
            self.assertEqual(1,len(deck.slides))

    def test_missing_mapping_duplicate_and_overflow_do_not_publish(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); template=self.fixture(folder)
            for mapping,text,code in (({},'새 내용','template_mapping_invalid'),
                                      ({'sourceSlide':0,'title':999},'새 내용','template_mapping_invalid'),
                                      ({'sourceSlide':0,'title':2},'긴 내용 '*150,'invalid_spec')):
                result=artifacts.create_ppt({'referenceMode':'preserve','slides':[{'title':text}],
                                            'templateSlides':[mapping]},folder/'bad.pptx',template)
                self.assertFalse(result['ok'],result)
                self.assertEqual(code,result['code'])
                self.assertFalse((folder/'bad.pptx').exists())

    def test_unmapped_old_content_blocks(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); template=self.fixture(folder)
            result=artifacts.create_ppt({'referenceMode':'preserve','slides':[{'body':'새 내용'}],
                'templateSlides':[{'sourceSlide':0,'body':2,'title':999}]},folder/'out.pptx',template)
            # Provided unused IDs must not be accepted as though they were applied.
            self.assertFalse(result['ok'],result)

    def test_font_must_be_available(self):
        from company_agent import presentation_design as design
        spec={'slides':[{'title':'보고'}],'presentationFont':'Missing Font'}
        with patch.object(flow,'installed_fonts',return_value=set()):
            with self.assertRaises(design.DesignError):
                design.prepare(spec,artifacts._normalize(spec))

    def test_optional_checker_does_not_install(self):
        import importlib.metadata
        with patch('importlib.metadata.version',side_effect=importlib.metadata.PackageNotFoundError), patch('subprocess.run') as run:
            self.assertEqual('not_installed',flow.optional_archforge(Path('absent.pptx'))['status'])
            run.assert_not_called()

    def test_optional_checker_reads_warning_summary_and_ignores_nearby_config(self):
        response={'summary':{'pass':True,'error_count':0,'warn_count':2},'private':'private slide'}
        proc=subprocess.CompletedProcess([],0,json.dumps(response).encode(),b'')
        with patch('importlib.metadata.version',return_value='0.11.0'),patch('subprocess.run',return_value=proc) as run:
            result=flow.optional_archforge(Path('out.pptx'))
        self.assertEqual('issues_found',result['status'])
        self.assertEqual(2,result['warnings'])
        self.assertIn('--no-config',run.call_args.args[0])
        self.assertNotIn('private',json.dumps(result))

    def test_native_table_and_chart_remain_editable(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); source=folder/'native.pptx'; output=folder/'out.pptx'
            deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
            shape=slide.shapes.add_table(2,2,Inches(0.5),Inches(1),Inches(4),Inches(3))
            for row in shape.table.rows:
                for cell in row.cells: cell.text='old'
            chart_data=CategoryChartData(); chart_data.categories=['A']; chart_data.add_series('old',[1])
            chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(5),Inches(1),Inches(4),Inches(3),chart_data)
            deck.save(str(source)); before=source.read_bytes()
            row={'table':{'headers':['항목','값'],'rows':[['실적','9']]},
                 'chart':{'type':'column','categories':['실적'],'series':[{'name':'당월','values':[9]}]}}
            native=flow.fill_template({'sections':[row]},output,source,[{'sourceSlide':0,'table':shape.shape_id,'chart':chart.shape_id}])
            new=Presentation(str(output)).slides[0]
            self.assertEqual('9',new.shapes[0].table.cell(1,1).text)
            self.assertEqual((9.0,),new.shapes[1].chart.series[0].values)
            self.assertEqual({'text':0,'tables':1,'charts':1,'images':0},native)
            self.assertEqual(before,source.read_bytes())

    def test_preview_refusal_no_capture_or_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); template=self.fixture(folder)
            with patch.object(artifacts,'_office',return_value={'ok':False,'status':'blocked','code':'render_refused'}) as call:
                result=artifacts.preview_template(template,folder/'preview')
            self.assertEqual('render_refused',result['code'])
            call.assert_called_once()
            self.assertFalse((folder/'preview').exists())

    def test_quality_detects_small_text_and_outside_shape(self):
        from pptx import Presentation
        from pptx.util import Inches,Pt
        with tempfile.TemporaryDirectory() as tmp:
            path=self.fixture(Path(tmp)); deck=Presentation(str(path))
            shape=deck.slides[0].shapes[0]; shape.left=Inches(-1)
            shape.text_frame.paragraphs[0].runs[0].font.size=Pt(6)
            deck.save(str(path))
            result=flow.quality(path)
            self.assertEqual({'out_of_bounds','tiny_text'},{i['code'] for i in result['issues']})
            self.assertEqual('not_performed',result['visualReview'])


if __name__=='__main__': unittest.main()
