"""Declared metric metadata reuse, not model or free-prose meaning validation."""
import copy
import csv
import hashlib
import html
import importlib.util
from pathlib import Path
import re
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'company-agent-plugin/scripts'))
from company_agent import business_artifacts as artifacts, presentation_design as design
from company_agent.report_facts import FactError, bind, resolve


class MetricSemanticsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.output = Path(self.temp.name) / 'report.html'

    def tearDown(self):
        self.temp.cleanup()

    def spec(self, period='9월', actual='262', target='280'):
        return {'title': '{{fact:actual.period}} {{fact:actual.label}} {{fact:actual.valueWithUnit}}, '
                         '{{fact:rate.label}} {{fact:rate.valueWithUnit}}',
                'facts': [
                    {'id': 'actual', 'op': 'value', 'inputs': [{'table': [0, 0, 1]}],
                     'label': '순매출', 'unit': '백만원', 'period': period, 'decimals': 1},
                    {'id': 'target', 'op': 'value', 'inputs': [{'table': [0, 0, 2]}],
                     'label': '목표', 'unit': '백만원', 'period': period},
                    {'id': 'rate', 'op': 'ratio', 'inputs': [{'fact': 'actual'}, {'fact': 'target'}],
                     'label': '달성률', 'unit': '%', 'period': period, 'decimals': 1}],
                'sections': [
                    {'title': '{{fact:actual.period}} {{fact:actual.label}}와 {{fact:rate.label}}',
                     'body': '{{fact:actual.valueWithUnit}} / {{fact:rate.valueWithUnit}}',
                     'table': {'headers': ['기간', '{{fact:actual.label}}({{fact:actual.unit}})',
                                          '{{fact:target.label}}({{fact:target.unit}})',
                                          '{{fact:rate.label}}({{fact:rate.unit}})'],
                               'rows': [[period, actual, target, '{{fact:rate}}']]},
                     'kpis': [{'fact': 'actual'}, {'fact': 'rate', 'label': '{{fact:rate.label}}'}]},
                    {'chart': {'type': 'bar',
                               'title': '{{fact:actual.period}} {{fact:actual.label}}({{fact:actual.unit}})',
                               'categories': ['{{fact:actual.period}}'],
                               'series': [{'name': '{{fact:actual.label}}', 'values': [float(actual)]}]}}]}

    def generate(self, spec):
        result = artifacts.create_html(spec, self.output)
        self.assertTrue(result['ok'], result)
        return self.output.read_text(encoding='utf-8'), result

    def test_correction_reuses_label_period_unit_in_title_headers_chart_and_kpi(self):
        text, result = self.generate(self.spec())
        self.assertIn('9월 순매출 262.0백만원, 달성률 93.6%', text)
        self.assertIn('<th scope="col">순매출(백만원)</th>', text)
        self.assertIn('<th scope="col">달성률(%)</th>', text)
        self.assertIn('9월 순매출(백만원)', text)
        self.assertNotIn('{{fact:', text)
        self.assertEqual(result['validation']['arithmetic'], 'checked')
        self.assertEqual(result['validation']['metricMeaning'], 'not_verified')

    def test_source_csv_and_spec_unchanged_with_overall_and_september_totals(self):
        source = Path(self.temp.name) / 'source.csv'
        source.write_text('월,순매출,목표\n7월,210,220\n8월,218,220\n9월,262,280\n', encoding='utf-8')
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        with source.open(encoding='utf-8', newline='') as handle:
            rows = list(csv.reader(handle))[1:]
        spec = self.spec()
        spec['sections'][0]['table']['rows'] = [row + [''] for row in rows]
        spec['facts'][0]['inputs'] = [{'table': [0, 2, 1]}]
        spec['facts'][1]['inputs'] = [{'table': [0, 2, 2]}]
        spec['facts'] += [
            {'id': 'total', 'label': '순매출', 'unit': '백만원', 'period': '7~9월',
             'inputs': [{'table': [0, i, 1]} for i in range(3)], 'decimals': 1},
            {'id': 'totalTarget', 'inputs': [{'table': [0, i, 2]} for i in range(3)]},
            {'id': 'totalRate', 'op': 'ratio', 'inputs': [{'fact': 'total'}, {'fact': 'totalTarget'}],
             'label': '달성률', 'unit': '%', 'decimals': 1}]
        spec['subtitle'] = ('{{fact:total.period}} {{fact:total.label}} {{fact:total.valueWithUnit}}, '
                            '{{fact:totalRate.label}} {{fact:totalRate.valueWithUnit}}')
        snapshot = copy.deepcopy(spec)
        text, _ = self.generate(spec)
        self.assertIn('7~9월 순매출 690.0백만원, 달성률 95.8%', text)
        self.assertIn('9월 순매출 262.0백만원, 달성률 93.6%', text)
        self.assertEqual(snapshot, spec)
        self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())

    def test_varied_period_and_values_not_hardcoded(self):
        cases = [('10월', '305', '320', '95.3'), ('2027년 1분기', '1543.25', '2000', '77.2')]
        for n, (period, actual, target, rate) in enumerate(cases):
            with self.subTest(period=period):
                self.output = Path(self.temp.name) / f'{n}.html'
                text, _ = self.generate(self.spec(period, actual, target))
                self.assertIn(f'{period} 순매출', text)
                self.assertIn(f'달성률 {rate}%', text)
                self.assertNotIn('9월 순매출', text)

    def test_legacy_numeric_binding_and_default_metadata_are_unchanged(self):
        facts, _ = resolve({'facts': [{'id': 'x', 'inputs': [1234.5], 'decimals': 1}]}, [])
        self.assertEqual('1,234.5%', bind('{{fact:x}}%', facts))
        self.assertEqual('x /  / 1,234.5', bind('{{fact:x.label}} / {{fact:x.unit}} / {{fact:x.valueWithUnit}}', facts))

    def test_unit_join_preserves_declared_spacing_and_does_not_duplicate_percent(self):
        facts, _ = resolve({'facts': [{'id': 'x', 'inputs': [5], 'unit': ' kg'}]}, [])
        self.assertEqual('5 kg', bind('{{fact:x.valueWithUnit}}', facts))
        text, _ = self.generate(self.spec())
        self.assertNotIn('93.6%%', text)

    def test_unknown_fields_and_ids_fail_before_publication(self):
        for token in ('{{fact:actual.value}}', '{{fact:actual.currency}}', '{{fact:actual.period.start}}',
                      '{{fact:actual.__class__}}', '{{fact:actual.valueWithUnitx}}', '{{fact:missing.label}}'):
            with self.subTest(token=token):
                spec = self.spec(); spec['title'] = token
                result = artifacts.create_html(spec, self.output)
                self.assertEqual(result['code'], 'numeric_validation_failed', result)
                self.assertFalse(self.output.exists())

    def test_explicit_period_reference_requires_declared_period(self):
        spec = self.spec(); del spec['facts'][0]['period']
        result = artifacts.create_html(spec, self.output)
        self.assertEqual(result['code'], 'numeric_validation_failed')
        self.assertIn('period', result['message'])
        self.assertFalse(self.output.exists())

    def test_invalid_period_values_are_rejected(self):
        for period in ('', ' ', None, 2026, ['9월'], 'x' * 101, '9월\n10월', '9월\x00'):
            with self.subTest(period=period):
                spec = self.spec(); spec['facts'][0]['period'] = period
                result = artifacts.create_html(spec, self.output)
                self.assertEqual(result['code'], 'numeric_validation_failed')
                self.assertFalse(self.output.exists())

    def test_metadata_is_escaped_at_render_and_not_interpreted_as_markup(self):
        spec = self.spec(); label = '<script>alert(1)</script>'
        spec['facts'][0].update(label=label, unit='<img src=x>', period='<9&월>')
        text, _ = self.generate(spec)
        self.assertNotIn(label, text)
        self.assertIn(html.escape(label), text)
        self.assertIn('&lt;9&amp;월&gt;', text)
        self.assertEqual(text.count('<script>'), 1)
        ET.fromstring(re.search(r'<svg\b.*?</svg>', text, re.S).group())

    def test_metadata_binding_is_not_recursive(self):
        spec = self.spec(); spec['facts'][0]['label'] = '{{fact:rate}}'
        result = artifacts.create_html(spec, self.output)
        self.assertEqual(result['code'], 'numeric_validation_failed')
        self.assertFalse(self.output.exists())

    def test_default_kpi_label_remains_literal_metadata(self):
        spec = {'facts': [{'id': 'actual', 'inputs': [262], 'label': '{{fact:rate}}'},
                          {'id': 'rate', 'inputs': ['93.6'], 'decimals': 1, 'label': '달성률'}],
                'sections': [{'kpis': [{'fact': 'actual'},
                                       {'fact': 'rate', 'label': '{{fact:rate.label}}'}]}]}
        text, _ = self.generate(spec)
        self.assertIn('<div class="kpi-label">{{fact:rate}}</div>', text)
        self.assertNotIn('<div class="kpi-label">93.6</div>', text)
        self.assertIn('<div class="kpi-label">달성률</div>', text)

    def test_header_bindings_do_not_make_cells_valid_numeric_inputs(self):
        spec = self.spec(); spec['sections'][0]['table']['rows'][0][1] = '{{fact:target}}'
        result = artifacts.create_html(spec, self.output)
        self.assertEqual(result['code'], 'numeric_validation_failed')
        self.assertFalse(self.output.exists())

    def test_zero_denominator_still_rejected_and_zero_actual_works(self):
        bad = artifacts.create_html(self.spec(target='0'), self.output)
        self.assertEqual(bad['code'], 'numeric_validation_failed')
        self.assertFalse(self.output.exists())
        text, _ = self.generate(self.spec(actual='0'))
        self.assertIn('순매출 0.0백만원, 달성률 0.0%', text)

    def test_legacy_unit_and_free_prose_are_not_new_semantic_claims(self):
        spec = self.spec(); spec['facts'][2]['unit'] = 'legacy-unit'
        spec['sections'][0]['body'] = '순매출 93.6백만원'  # Deliberately wrong free prose, not bound.
        text, result = self.generate(spec)
        self.assertIn('순매출 93.6백만원', text)
        self.assertIn('93.6legacy-unit', text)
        self.assertEqual(result['validation']['metricMeaning'], 'not_verified')
        self.assertEqual(result['validation']['freeTextClaims'], 'not_verified')

    def test_no_facts_keeps_old_generation_and_reports_semantic_limit(self):
        text, result = self.generate({'sections': [{'title': '일반 안내', 'body': '숫자가 없는 설명'}]})
        self.assertIn('숫자가 없는 설명', text)
        self.assertEqual(result['validation']['arithmetic'], 'not_provided')
        self.assertEqual(result['validation']['metricMeaning'], 'not_verified')

    def test_ppt_preparation_binds_headers_chart_labels_and_keeps_numeric_series(self):
        spec = self.spec()
        spec['sections'][0].pop('kpis')
        data = artifacts._normalize(spec)
        original_values = copy.deepcopy(data['sections'][1]['chart']['series'][0]['values'])
        validation = design.prepare(spec, data)
        table, chart = data['sections'][0]['table'], data['sections'][1]['chart']
        self.assertEqual(table['headers'][1], '순매출(백만원)')
        self.assertEqual(chart['categories'], ['9월'])
        self.assertEqual(chart['series'][0]['name'], '순매출')
        self.assertEqual(chart['series'][0]['values'], original_values)
        self.assertEqual(validation['metricMeaning'], 'not_verified')

    def test_ppt_bound_chart_labels_still_enforce_length_limits(self):
        for key in ('label', 'period'):
            with self.subTest(field=key):
                spec = self.spec(); spec['sections'][0].pop('kpis')
                spec['facts'][0][key] = '가' * 25
                data = artifacts._normalize(spec)
                with self.assertRaises(design.DesignError):
                    design.prepare(spec, data)

    @unittest.skipUnless(importlib.util.find_spec('pptx'), 'python-pptx is not available')
    def test_native_ppt_preserves_bound_chart_metadata_as_text(self):
        from pptx import Presentation
        spec = {'title': '{{fact:actual.period}} {{fact:actual.label}}',
                'facts': [{'id': 'actual', 'inputs': [262], 'label': '<매출&>',
                           'period': '9&월', 'unit': '백만원'}],
                'slides': [{'title': '{{fact:actual.period}} {{fact:actual.label}}',
                            'chart': {'type': 'bar', 'title': '{{fact:actual.label}}({{fact:actual.unit}})',
                                      'categories': ['{{fact:actual.period}}'],
                                      'series': [{'name': '{{fact:actual.label}}', 'values': [262]}]}}]}
        data = artifacts._normalize(spec)
        design.prepare(spec, data)
        data['presentationPlan'] = design.plan(data)
        output = Path(self.temp.name) / 'metadata.pptx'
        design.render_python(data, output, None)
        deck = Presentation(output)
        chart = next(shape.chart for shape in deck.slides[0].shapes if shape.has_chart)
        self.assertEqual(chart.series[0].name, '<매출&>')
        self.assertEqual(chart.plots[0].categories[0].label, '9&월')
        self.assertEqual(chart.series[0].values, (262.0,))
        text = '\n'.join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
        self.assertIn('9&월 <매출&>', text)
        self.assertNotIn('{{fact:', text)


if __name__ == '__main__':
    unittest.main()
