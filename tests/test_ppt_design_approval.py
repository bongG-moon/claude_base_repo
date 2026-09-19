import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'company-agent-plugin/scripts'))
from company_agent import ppt_workflow as flow, business_artifacts as artifacts
from company_agent.state import begin_turn,record_activity,stop_decision,load_session
from company_agent.paths import ensure_user_layout


class PptDesignApprovalTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.job={'creationMode':'new','purpose':'월간 실적','audience':'부서장','slideCount':5,
                  'slides':[{'title':title,'body':'테스트용 가상 자료'} for title in
                            ('월간 실적','핵심 요약','월별 비교','상세표','확인할 사항')]}

    def preview(self,job=None):
        with patch.object(artifacts,'_office',side_effect=AssertionError('HTML draft must not start Office')):
            return artifacts.create_ppt(job or self.job,self.root/'representative.html',require_choices=True,preview_only=True)

    def test_new_mode_does_not_skip_design_or_confirmation(self):
        self.assertEqual('design',flow.choices(self.job)['stage'])
        self.job['designPreset']='monochrome'
        self.assertEqual('design_preview',flow.choices(self.job)['stage'])
        preview=self.preview()
        self.assertTrue(preview['ok'],preview)
        self.assertEqual(5,preview['slides'])
        self.assertEqual(5,len(preview['outline']))
        self.job['designReview']=preview['designReview']
        self.assertFalse(self.job['designReview']['confirmed'])
        self.assertEqual('design_confirm',flow.choices(self.job)['stage'])
        target=self.root/'final.pptx'
        self.assertEqual('design_confirm',artifacts.create_ppt(self.job,target,require_choices=True)['stage'])
        self.assertFalse(target.exists())
        self.job['designReview']['confirmed']=True  # Simulated explicit user confirmation.
        with patch.object(artifacts,'_office',return_value={'ok':False,'code':'render_refused'}):
            final=artifacts.create_ppt(self.job,target,require_choices=True)
        self.assertTrue(final['ok'],final)
        self.assertEqual(5,final['slides'])
        self.assertEqual('partial',final['status'])

    def test_changed_job_or_preview_needs_new_confirmation(self):
        self.job['designPreset']='business'
        self.job['designReview']=self.preview()['designReview']
        self.job['designReview']['confirmed']=True
        for field,value in [('designPreset','warm'),('title','수정된 제목'),('slideCount',4)]:
            changed=copy.deepcopy(self.job);changed[field]=value
            target=self.root/(field+'.pptx')
            result=artifacts.create_ppt(changed,target,require_choices=True)
            self.assertEqual('design_review_changed',result['code'])
            self.assertFalse(target.exists())
        (self.root/'representative.html').write_bytes(b'changed')
        self.assertEqual('design_preview_changed',artifacts.create_ppt(self.job,self.root/'bad.pptx',require_choices=True)['code'])

    def test_metadata_write_is_not_a_business_change_but_job_is(self):
        state=self.root/'state';ensure_user_layout(state)
        begin_turn('ppt-choice','MEDIUM',True,(),state)
        choices={k:v for k,v in self.job.items() if k!='slides'}
        def record(spec,path=None):
            record_activity({'session_id':'ppt-choice','tool_name':'Write','tool_input':{
                'file_path':str(path or state/'tmp/ppt-choices-test.json'),'content':json.dumps(spec)}},state)
        record(choices)
        record({**choices,'designPreset':'monochrome'})
        record({**choices,'creationMode':'reference','referenceImages':[str(self.root/'reference.png')]})
        self.assertEqual(0,load_session('ppt-choice',state)['mutationCount'])
        self.assertEqual({},stop_decision({'session_id':'ppt-choice'},state))
        record(self.job)
        self.assertEqual(1,load_session('ppt-choice',state)['mutationCount'])
        record(choices)
        self.assertEqual('block',stop_decision({'session_id':'ppt-choice'},state)['decision'])

    def test_other_path_and_invalid_metadata_do_not_get_exemption(self):
        state=self.root/'state';ensure_user_layout(state)
        begin_turn('ppt-choice','MEDIUM',True,(),state)
        for content,path in [({'creationMode':'new'},self.root/'ppt-choices-test.json'),
                             ({'slideCount':True},state/'tmp/ppt-choices-test.json'),
                             ({'creationMode':'reference','referenceImages':['https://example.invalid/a.png']},state/'tmp/ppt-choices-test.json'),
                             ({'creationMode':'reference','referenceImages':[str(self.root/'a.png')]*4},state/'tmp/ppt-choices-test.json')]:
            record_activity({'session_id':'ppt-choice','tool_name':'Write','tool_input':{
                'file_path':str(path),'content':json.dumps(content)}},state)
        self.assertEqual(4,load_session('ppt-choice',state)['mutationCount'])

    def test_palettes_are_actually_applied(self):
        from company_agent import presentation_design
        for key,(_,_,palette) in flow.DESIGNS.items():
            job={**self.job,'designPreset':key}
            data=artifacts._normalize(job)
            presentation_design.prepare(job,data)
            self.assertEqual(palette,data['presentationTheme'])

    def test_preview_resolves_full_job_numbers_and_final_preserves_native_evidence(self):
        from pptx import Presentation
        self.job['designPreset']='business'
        self.job['facts']=[{'id':'total','op':'sum','label':'누적 실적','unit':'백만원',
                           'inputs':[{'chart':[2,1,i]} for i in range(3)]}]
        self.job['slides'][1]={'title':'누적 실적 {{fact:total}}백만원','kpis':[{'fact':'total'}]}
        self.job['slides'][2]={'title':'월별 비교','chart':{'type':'column','categories':['6월','7월','8월'],
            'series':[{'name':'목표','values':[150,150,150]},{'name':'실적','values':[145,175,210]}]}}
        self.job['slides'][3]={'title':'상세표','table':{'headers':['월','목표','실적'],
            'rows':[['6월','150','145'],['7월','150','175'],['8월','150','210']]}}
        result=self.preview()
        self.assertTrue(result['ok'],result)
        draft=(self.root/'representative.html').read_text(encoding='utf-8')
        self.assertEqual(5,draft.count('class="ppt-page"'))
        self.assertEqual('누적 실적 530백만원',result['outline'][1])
        self.assertIn('530',draft)
        self.job['designReview']={**result['designReview'],'confirmed':True}
        with patch.object(artifacts,'_office',return_value={'ok':False,'code':'render_refused'}):
            final=artifacts.create_ppt(self.job,self.root/'final.pptx',require_choices=True)
        self.assertTrue(final['ok'],final)
        deck=Presentation(str(self.root/'final.pptx'))
        self.assertEqual(5,len(deck.slides))
        chart=next(s.chart for s in deck.slides[2].shapes if s.has_chart)
        self.assertEqual((145.0,175.0,210.0),chart.series[1].values)
        table=next(s.table for s in deck.slides[3].shapes if s.has_table)
        self.assertEqual('210',table.cell(3,2).text)

    def test_parser_exposes_draft_and_final_as_separate_commands(self):
        import argparse
        from company_agent import business
        parser=argparse.ArgumentParser();business.register(parser.add_subparsers())
        args=parser.parse_args(['business','ppt-design-preview','--spec','job.json','--output','draft.html'])
        self.assertEqual('ppt-design-preview',args.business_action)

    def test_digest_tracks_image_bytes_at_unchanged_path(self):
        import base64
        image=self.root/'logo.png'
        raw=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aMycAAAAASUVORK5CYII=')
        image.write_bytes(raw)
        spec={**self.job,'slides':[{'title':'예시','image':{'path':str(image)}}]}
        first=flow.design_digest(spec)
        image.write_bytes(raw+b'changed')
        self.assertNotEqual(first,flow.design_digest(spec))


if __name__=='__main__':unittest.main()
