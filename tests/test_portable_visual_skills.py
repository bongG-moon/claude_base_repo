"""Exercise bundled HTML/PPT skills without a harness/profile/Office process."""
from __future__ import annotations
import ast
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
MODULES = "business_artifacts business_safety content_layout explanation_diagram explanation_export html_reference ppt_html ppt_html_import ppt_scene ppt_workflow presentation_design report_design report_facts report_styles report_typography".split()
FILES = [n + ".py" for n in MODULES] + ["ppt_dom_capture.js"]
ISOLATED_PPTX = subprocess.run([sys.executable, "-I", "-c", "import pptx"],
                               capture_output=True, timeout=20).returncode == 0


class PortableVisualTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.home = self.root / "empty-profile"
        self.home.mkdir()
        self.env = {key: value for key, value in os.environ.items()
                    if not key.upper().startswith(("COMPANY", "ANTHROPIC", "CLAUDE", "PYTHONPATH"))}
        self.env.update(HOME=str(self.home), USERPROFILE=str(self.home), CLAUDE_CONFIG_DIR=str(self.home / ".claude"))
        self.spec = {"title": "가상 교육 현황", "creationMode": "new", "designPreset": "business",
                     "purpose": "진행 확인", "audience": "담당자", "slideCount": 2, "slides": [
            {"title": "신청", "elements": [
                {"kind": "text", "text": "교육 운영 현황", "x": 40, "y": 20, "w": 800, "h": 50, "size": 30, "bold": True},
                {"kind": "shape", "group": "신청카드", "shape": "roundRect", "x": 40, "y": 100, "w": 380, "h": 220, "fill": "EEF4F7"},
                {"kind": "text", "group": "신청카드", "text": "신청 20명", "x": 60, "y": 120, "w": 330, "h": 80,
                 "size": 24, "bold": False, "color": "173B5E", "font": "Malgun Gothic"}]},
            {"title": "부서별 인원", "table": {"headers": ["부서", "인원"], "rows": [["생산", "12"], ["품질", "8"]]},
             "chart": {"type": "column", "categories": ["생산", "품질"], "series": [{"name": "신청", "values": [12, 8]}]}}
        ]}

    def stage(self, slug):
        target = self.root / slug
        if not target.exists():
            shutil.copytree(ROOT / "standalone-skills" / slug, target)
            for name in FILES:
                shutil.copyfile(ROOT / "company-agent-plugin/scripts/company_agent" / name,
                                target / "scripts/visual_engine" / name)
        return target

    def run_cli(self, slug, *args, expected=0):
        skill = self.stage(slug)
        script = "report.py" if slug == "html-report" else "presentation.py"
        proc = subprocess.run([sys.executable, "-I", "-X", "utf8", str(skill / "scripts" / script), *map(str, args)],
                              cwd=self.root, env=self.env, text=True, encoding="utf-8", capture_output=True, timeout=30)
        self.assertEqual(proc.returncode, expected, (proc.stdout, proc.stderr))
        self.assertEqual(proc.stderr, "")
        self.assertEqual(list(self.home.iterdir()), [], "A portable run must not install or alter a user profile")
        return json.loads(proc.stdout)

    def write_spec(self, spec):
        path = self.root / "job.json"
        path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
        return path

    def approved_spec(self):
        path = self.write_spec(self.spec)
        result = self.run_cli("presentation", "draft", "--spec", path, "--output", self.root / "draft.html")
        self.assertTrue(result["ok"], result)
        self.spec["designReview"] = {**result["designReview"], "confirmed": True}
        return self.write_spec(self.spec)

    def test_runtime_dependency_closure_has_no_harness_imports(self):
        for name in MODULES:
            tree = ast.parse((ROOT / "company-agent-plugin/scripts/company_agent" / (name + ".py")).read_text(encoding="utf-8-sig"))
            for node in ast.walk(tree):
                if isinstance(node, ast.ImportFrom) and node.level:
                    imports = [node.module.split(".")[0]] if node.module else [n.name for n in node.names]
                    self.assertTrue(set(imports) <= set(MODULES), (name, imports))
            self.assertNotIn(name, {"state", "cli", "memory", "session", "vision_routing", "hook_router"})

    def test_local_document_references_resolve(self):
        for slug in ("html-report", "presentation"):
            skill = ROOT / "standalone-skills" / slug
            for file in skill.rglob("*.md"):
                for target in re.findall(r"`(references/[^`\n]+\.md)`", file.read_text(encoding="utf-8")):
                    self.assertTrue((skill / target).is_file(), (file, target))

    def test_html_offline_font_and_calculation(self):
        job = self.write_spec({"title": "교육 신청 현황", "fontChoice": "noto-sans-kr", "fontSource": "system",
            "facts": [{"id": "total", "op": "sum", "inputs": [{"table": [0, 0, 1]}, {"table": [0, 1, 1]}]}],
            "sections": [{"title": "신청 {{fact:total}}명", "table": {"headers": ["부서", "인원"], "rows": [["생산", 12], ["품질", 8]]},
                          "body": "가상 자료", "contentOrder": ["table", "body"]}]})
        result = self.run_cli("html-report", "build", "--spec", job, "--output", self.root / "report.html")
        self.assertTrue(result["ok"], result)
        self.assertTrue(result["offline"])
        text = (self.root / "report.html").read_text(encoding="utf-8")
        self.assertIn("신청 20명", text)
        self.assertIn("Noto Sans KR", text)
        self.assertNotIn("fonts.googleapis.com", text)

    @unittest.skipUnless(ISOLATED_PPTX, "This Python's isolated environment needs python-pptx; run with a prepared runtime")
    def test_ppt_native_groups_edit_and_ungroup(self):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        job = self.approved_spec()
        path = self.root / "result.pptx"
        result = self.run_cli("presentation", "build", "--spec", job, "--output", path)
        self.assertTrue(result["ok"], result)
        self.assertFalse(result["officeStarted"])
        self.assertEqual(result["visualReview"], "not_verified")
        deck = Presentation(path)
        group = deck.slides[0].shapes[1]
        self.assertEqual(group.name, "신청카드")
        self.assertFalse(group._element.xpath('.//a:grpSpLocks[@noUngrp="1"] | .//a:spLocks[@noTextEdit="1"]'))
        run = group.shapes[1].text_frame.paragraphs[0].runs[0]
        self.assertEqual((str(run.font.color.rgb), run.font.bold), ("173B5E", False))
        run.font.color.rgb = RGBColor.from_string("C00000")
        run.font.bold = True
        run.font.name = "Arial"
        tree = deck.slides[0].shapes._spTree
        index = tree.index(group._element)
        before = [(s.left, s.top, s.width, s.height) for s in group.shapes]
        for offset, shape in enumerate(list(group.shapes)):
            tree.insert(index + offset, shape._element)
        tree.remove(group._element)
        deck.save(path)
        deck = Presentation(path)
        self.assertEqual([(s.left, s.top, s.width, s.height) for s in list(deck.slides[0].shapes)[1:]], before)
        run = deck.slides[0].shapes[2].text_frame.paragraphs[0].runs[0]
        self.assertEqual((str(run.font.color.rgb), run.font.bold, run.font.name), ("C00000", True, "Arial"))
        self.assertTrue(any(s.has_table for s in deck.slides[1].shapes))
        self.assertTrue(any(s.has_chart for s in deck.slides[1].shapes))

    def test_unapproved_draft_and_changed_content_are_not_exported(self):
        job = self.write_spec(self.spec)
        result = self.run_cli("presentation", "build", "--spec", job, "--output", self.root / "no.pptx", expected=1)
        self.assertFalse(result["ok"])
        self.assertFalse((self.root / "no.pptx").exists())
        self.approved_spec()
        self.spec["slides"][0]["elements"][2]["text"] = "수정 99명"
        job = self.write_spec(self.spec)
        result = self.run_cli("presentation", "build", "--spec", job, "--output", self.root / "no.pptx", expected=1)
        self.assertFalse(result["ok"])
        self.assertFalse((self.root / "no.pptx").exists())

    def test_changed_html_is_not_approved(self):
        job = self.approved_spec()
        (self.root / "draft.html").write_text("changed", encoding="utf-8")
        result = self.run_cli("presentation", "build", "--spec", job, "--output", self.root / "no.pptx", expected=1)
        self.assertFalse(result["ok"])

    def test_reusable_html_template_contains_synthetic_content(self):
        job = self.approved_spec()
        target = self.root / "template.html"
        result = self.run_cli("presentation", "template", "--spec", job, "--output", target)
        self.assertTrue(result["ok"], result)
        self.assertFalse(result["businessContentStored"])
        self.assertNotIn("신청 20명", target.read_text(encoding="utf-8"))
        self.assertIn("presentation.py", result["reuse"])
        reference = self.run_cli("presentation", "reference", "--file", target)
        self.assertTrue(reference["ok"], reference)
        self.assertEqual(reference["format"], "html")

    @unittest.skipUnless(ISOLATED_PPTX, "This Python's isolated environment needs python-pptx; run with a prepared runtime")
    def test_reference_inspection_is_bounded(self):
        job = self.approved_spec()
        path = self.root / "deck.pptx"
        self.run_cli("presentation", "build", "--spec", job, "--output", path)
        result = self.run_cli("presentation", "reference", "--file", path)
        self.assertTrue(result["ok"], result)
        self.assertLessEqual(len(result["slides"]), 3)
        self.assertEqual(result["slides"][0]["groupCount"], 1)
        result = self.run_cli("presentation", "reference", "--file", path, "--slides", "1,2,3,4", expected=1)
        self.assertFalse(result["ok"])

    def test_existing_output_is_preserved(self):
        job = self.approved_spec()
        target = self.root / "existing.pptx"
        target.write_bytes(b"user data")
        result = self.run_cli("presentation", "build", "--spec", job, "--output", target, expected=1)
        self.assertFalse(result["ok"])
        self.assertEqual(target.read_bytes(), b"user data")

    def test_incomplete_skill_copy_has_actionable_error(self):
        skill = self.stage("html-report")
        (skill / "scripts/visual_engine/business_artifacts.py").unlink()
        job = self.write_spec({"sections": [{"body": "example"}]})
        result = self.run_cli("html-report", "build", "--spec", job, "--output", self.root / "no.html", expected=2)
        self.assertEqual(result["code"], "incomplete_package")

    def test_schema_examples_execute(self):
        for slug in ("html-report", "presentation"):
            text = (ROOT / "standalone-skills" / slug / "references/schema.md").read_text(encoding="utf-8")
            examples = re.findall(r"```json\n(.*?)\n```", text, re.S)
            for index, source in enumerate(examples):
                spec = json.loads(source)
                if slug == "presentation" and "elements" in spec:
                    spec = {"title": "예시", "slides": [spec]}
                path = self.write_spec(spec)
                result = self.run_cli(slug, "build" if slug == "html-report" else "draft", "--spec", path,
                                      "--output", self.root / f"{slug}-{index}.html")
                self.assertTrue(result["ok"], result)


if __name__ == "__main__":
    unittest.main()
