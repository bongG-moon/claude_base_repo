import base64
import copy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'company-agent-plugin/scripts'))
from company_agent import business_artifacts as artifacts, ppt_html, ppt_workflow


class PptHtmlTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.spec = {'creationMode':'new','purpose':'보고','audience':'부서장','slideCount':3,'designPreset':'warm',
                     'title':'비공개 고객 A 보고','slides':[
                         {'title':'비공개 요약','body':'고객 비공개 본문 481'},
                         {'title':'실적','chart':{'type':'column','categories':['6월','7월','8월'],
                             'series':[{'name':'비공개 실적','values':[145,175,210]}]}},
                         {'title':'상세','table':{'headers':['월','실적'],'rows':[['6월',145],['7월',175],['8월',210]]}}]}

    def draft(self, spec=None, name='draft.html', template=None):
        result = ppt_html.create_draft(spec or self.spec,self.root/name,template)
        self.assertTrue(result['ok'], result)
        return result

    def approve(self):
        result = self.draft()
        self.spec['designReview'] = {**result['designReview'],'confirmed':True}
        return result

    def test_draft_is_offline_without_office_or_ppt_generation(self):
        with patch.object(artifacts,'_office',side_effect=AssertionError()), patch.object(artifacts,'_python_ppt',side_effect=AssertionError()):
            result = self.draft()
        self.assertFalse(result['officeStarted'])
        self.assertEqual('offline-html', result['engine'])
        self.assertEqual([], list(self.root.glob('*.pptx')))
        self.assertEqual(3, result['slides'])
        document = (self.root/'draft.html').read_text(encoding='utf-8')
        self.assertIn('비공개 고객 A 보고', document)
        self.assertIn('script-src &#x27;none&#x27;', document)
        self.assertNotIn('<script', document)
        self.assertNotIn('https://', document)
        self.assertNotIn('fetch(', document)

    def test_draft_does_not_hide_errors_on_later_slides(self):
        self.spec['slides'][2]['title'] = '긴제목'*50
        result = ppt_html.create_draft(self.spec,self.root/'bad.html')
        self.assertFalse(result['ok'])
        self.assertFalse((self.root/'bad.html').exists())

    def test_content_is_escaped_not_executed(self):
        self.spec['slides'][0] = {'title':'<script>danger()</script>', 'body':'<img src=x onerror=alert(1)>'}
        result = self.draft()
        document = Path(result['outputPath']).read_text(encoding='utf-8')
        self.assertIn('&lt;script&gt;',document)
        self.assertNotIn('<script>',document)
        self.assertNotIn('<img src=x',document)

    def test_saved_representative_is_one_content_free_reusable_html(self):
        self.approve()
        saved = self.root/'대표 양식.html'
        result = ppt_html.save_template(self.spec,saved)
        self.assertTrue(result['ok'],result)
        document = saved.read_text(encoding='utf-8')
        for private in ('비공개','481','145','175','210','draft.html',str(self.root)):
            self.assertNotIn(private, document)
        self.assertFalse(result['businessContentStored'])
        self.assertFalse(result['referenceImagesStored'])
        meta = ppt_html.read_template(saved)
        new = {'creationMode':'saved','purpose':'다른 보고','audience':'팀원','slideCount':1,
               'slides':[{'title':'새 업무','body':'새 본문'}]}
        self.assertEqual('design_preview', ppt_workflow.choices(new,saved)['stage'])
        preview = self.draft(new,'reused.html',saved)
        new['designReview'] = {**preview['designReview'],'confirmed':True}
        with patch.object(artifacts,'_office',return_value={'ok':False,'code':'render_refused'}):
            final = artifacts.create_ppt(new,self.root/'final.pptx',saved,require_choices=True)
        self.assertTrue(final['ok'],final)
        from pptx import Presentation
        deck = Presentation(self.root/'final.pptx')
        self.assertEqual(1,len(deck.slides))
        self.assertEqual(meta['presentationTheme']['background'],str(deck.slides[0].background.fill.fore_color.rgb))
        self.assertNotIn('비공개',' '.join(s.text for s in deck.slides[0].shapes if s.has_text_frame))

    def test_template_save_requires_approval_and_never_overwrites(self):
        self.assertEqual('design_preview',ppt_html.save_template(self.spec,self.root/'saved.html')['stage'])
        self.spec['designReview'] = self.draft()['designReview']
        self.assertEqual('design_confirm',ppt_html.save_template(self.spec,self.root/'saved.html')['stage'])
        self.spec['designReview']['confirmed'] = True
        self.assertTrue(ppt_html.save_template(self.spec,self.root/'saved.html')['ok'])
        before = (self.root/'saved.html').read_bytes()
        self.assertEqual('output_exists',ppt_html.save_template(self.spec,self.root/'saved.html')['code'])
        self.assertEqual(before,(self.root/'saved.html').read_bytes())
        self.assertEqual('output_exists',ppt_html.create_draft(self.spec,self.root/'saved.html')['code'])

    def test_saved_frame_is_reused_by_html_and_native_ppt(self):
        self.spec['presentationFrame'] = {'width':720,'height':540}
        self.approve()
        saved = self.root/'four-three.html'
        self.assertTrue(ppt_html.save_template(self.spec,saved)['ok'])
        new = {'creationMode':'saved','purpose':'보고','audience':'팀원','slideCount':1,
               'slides':[{'title':'새 보고','body':'다른 내용'}]}
        draft = self.draft(new,'four-three-draft.html',saved)
        self.assertIn('--pw:720',Path(draft['outputPath']).read_text(encoding='utf-8'))
        new['designReview'] = {**draft['designReview'],'confirmed':True}
        with patch.object(artifacts,'_office',return_value={'ok':False}):
            result = artifacts.create_ppt(new,self.root/'four-three.pptx',saved,require_choices=True)
        self.assertTrue(result['ok'],result)
        from pptx import Presentation
        deck = Presentation(self.root/'four-three.pptx')
        self.assertEqual((720,540),(deck.slide_width/12700,deck.slide_height/12700))

    def test_invalid_template_metadata_is_never_executed_or_accepted(self):
        self.approve()
        saved = self.root/'saved.html'
        self.assertTrue(ppt_html.save_template(self.spec,saved)['ok'])
        original = saved.read_text(encoding='utf-8')
        meta = ppt_html.read_template(saved)
        for value in ({**meta,'schema':ppt_html.SCHEMA,'layoutVersion':2},
                      {**meta,'schema':ppt_html.SCHEMA,'layoutVersion':1,'command':'anything'},
                      {**meta,'schema':ppt_html.SCHEMA,'layoutVersion':1,'presentationFrame':{'width':True,'height':540}}):
            with self.subTest(value=value):
                source = self.root/'invalid.html'
                source.write_text(f'<script id="{ppt_html.TEMPLATE_ID}" type="application/json">{json.dumps(value)}</script>',encoding='utf-8')
                with self.assertRaises(artifacts.ArtifactError):
                    ppt_html.read_template(source)
        duplicate = self.root/'duplicate.html'
        duplicate.write_text(original+original,encoding='utf-8')
        with self.assertRaises(artifacts.ArtifactError):
            ppt_html.read_template(duplicate)

    def test_template_or_reference_changes_invalidate_review(self):
        self.approve()
        saved = self.root/'saved.html'
        self.assertTrue(ppt_html.save_template(self.spec,saved)['ok'])
        new = {**self.spec,'creationMode':'saved'}
        new.pop('designReview')
        preview = self.draft(new,'new.html',saved)
        new['designReview'] = {**preview['designReview'],'confirmed':True}
        saved.write_text(saved.read_text(encoding='utf-8')+' ',encoding='utf-8')
        result = artifacts.create_ppt(new,self.root/'denied.pptx',saved,require_choices=True)
        self.assertEqual('design_review_changed',result['code'])
        self.assertFalse((self.root/'denied.pptx').exists())

    def test_screenshots_are_optional_style_inputs_not_slide_bitmaps(self):
        question = ppt_workflow.choices({'creationMode':'reference'})
        self.assertEqual('reference_file',question['stage'])
        self.assertIn('2~3장',question['question'])
        path = self.root/'reference.png'
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aMycAAAAASUVORK5CYII=')
        path.write_bytes(png)
        spec = {**self.spec,'creationMode':'reference','referenceImages':[str(path)],
                'presentationTheme':{'accent':'80452C'}}
        self.assertEqual('design_preview',ppt_workflow.choices(spec)['stage'])
        preview = self.draft(spec)
        self.assertNotIn(base64.b64encode(png).decode(),(self.root/'draft.html').read_text(encoding='utf-8'))
        spec['designReview'] = {**preview['designReview'],'confirmed':True}
        with patch.object(artifacts,'_office',return_value={'ok':False}):
            final = artifacts.create_ppt(spec,self.root/'native.pptx',require_choices=True)
        self.assertTrue(final['ok'],final)
        self.assertEqual(0,final['editability']['images'])
        self.assertEqual(1,final['editability']['tables'])
        self.assertEqual(1,final['editability']['charts'])
        path.write_bytes(png+b'changed')
        self.assertEqual('design_review_changed',artifacts.create_ppt(spec,self.root/'changed.pptx',require_choices=True)['code'])
        spec['referenceMode'] = 'preserve'
        self.assertEqual('invalid_choice',ppt_workflow.choices(spec)['code'])

    def test_no_screenshot_requirement_for_new_design_and_missing_reference_not_ignored(self):
        self.assertEqual('design_preview',ppt_workflow.choices(self.spec)['stage'])
        spec = {**self.spec,'creationMode':'reference','referenceImages':[str(self.root/'missing.png')],
                'presentationTheme':{'accent':'80452C'}}
        result = ppt_html.create_draft(spec,self.root/'out.html')
        self.assertEqual('input_missing',result['code'])
        self.assertFalse((self.root/'out.html').exists())
        spec['referenceImages'] *= 4
        self.assertEqual('invalid_choice',ppt_workflow.choices(spec)['code'])

    def test_arbitrary_html_is_not_treated_as_editable_template(self):
        source = self.root/'foreign.html'
        source.write_text('<html><script>alert(1)</script></html>',encoding='utf-8')
        result = ppt_html.create_draft({**self.spec,'creationMode':'saved'},self.root/'out.html',source)
        self.assertEqual('invalid_ppt_template',result['code'])
        self.assertFalse((self.root/'out.html').exists())

    def test_changed_html_cannot_be_silently_exported_from_old_spec(self):
        self.approve()
        draft = self.root/'draft.html'
        draft.write_text(draft.read_text(encoding='utf-8').replace('비공개 요약','수정된 요약'),encoding='utf-8')
        self.assertEqual('design_preview_changed',artifacts.create_ppt(self.spec,self.root/'out.pptx',require_choices=True)['code'])
        self.assertEqual('design_preview_changed',ppt_html.save_template(self.spec,self.root/'saved.html')['code'])

    def test_cli_draft_and_template_end_to_end(self):
        specfile = self.root/'job.json'
        def run(action,output):
            specfile.write_text(json.dumps(self.spec,ensure_ascii=False),encoding='utf-8')
            call = subprocess.run([sys.executable,'-X','utf8','-B',str(ROOT/'company-agent-plugin/scripts/harness_cli.py'),
                'business',action,'--spec',str(specfile),'--output',str(output),'--state-root',str(self.root/'state')],
                capture_output=True,text=True,encoding='utf-8',timeout=15)
            self.assertEqual(0,call.returncode,call.stderr)
            return json.loads(call.stdout)
        draft = run('ppt-design-preview',self.root/'초안.html')
        self.assertTrue(draft['ok'],draft)
        self.spec['designReview'] = {**draft['designReview'],'confirmed':True}
        saved = run('ppt-template',self.root/'양식.html')
        self.assertTrue(saved['ok'],saved)
        self.assertFalse((self.root/'state').exists())

    def test_old_ppt_draft_requires_new_html_but_not_permission_error(self):
        self.spec['designReview'] = {'previewPath':str(self.root/'old.pptx')}
        result = ppt_workflow.choices(self.spec)
        self.assertEqual('design_preview',result['stage'])
        self.assertIn('ppt-design-preview',result['previewCommand'])
        self.assertIn('--work WORK_FILE',result['previewCommand'])
        self.assertNotIn('--output',result['previewCommand'])


if __name__ == '__main__':
    unittest.main()
