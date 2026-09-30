from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'company-agent-plugin/scripts'))
from company_agent import business_artifacts as artifacts, content_layout, ppt_html, presentation_design, ppt_workflow


class ContentLayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.row = {'title': '현황', 'table': {'headers': ['항목', '값'], 'rows': [['가', 5], ['나', 3]]},
                    'body': '가 항목이 많습니다. 나 항목은 확인이 필요합니다.', 'contentOrder': ['table', 'body']}
        self.spec = {'title': '검토', 'slides': [self.row]}

    def test_html_all_styles_order_in_dom_not_css_reversal(self):
        for style in artifacts.STYLES:
            result = artifacts.create_html({**self.spec, 'style': style}, self.root/(style+'.html'))
            self.assertTrue(result['ok'], result)
            doc = Path(result['outputPath']).read_text(encoding='utf-8')
            self.assertLess(doc.index('data-content-block="table"'), doc.index('data-content-block="body"'))
            self.assertIn('flex-direction:column!important', doc)
            self.assertEqual('checked', result['contentLayoutValidation']['status'])
            self.assertEqual('not-performed', result['contentLayoutValidation']['visualReview'])
            self.assertEqual('not-checked', result['contentLayoutValidation']['sentenceCount'])

    def test_default_output_keeps_copy_then_visual_structure(self):
        del self.row['contentOrder']
        result = artifacts.create_html(self.spec, self.root/'default.html')
        doc = Path(result['outputPath']).read_text(encoding='utf-8')
        self.assertIn('<div class="section-content"><div class="copy-block">', doc)
        self.assertNotIn('data-content-block=', doc)
        self.assertEqual('not-requested', result['contentLayoutValidation']['status'])
        data = ppt_html.prepare(self.spec)[0]
        self.assertNotIn('contentBlocks', data['presentationPlan']['pages'][0])
        content = data['presentationPlan']['pages'][0]['elements']
        table = next(e for e in content if e['kind'] == 'table')
        body = next(e for e in content if e.get('text') == self.row['body'])
        self.assertGreater(body['x'], table['x']+table['w'])

    def test_invalid_orders_do_not_save_or_ignore_content(self):
        cases = [None, '', [], ['table'], ['body','table','image'], ['table','body','body'],
                 ['table','summary'], ['table',{}], ['table','body','takeaway']]
        for n, order in enumerate(cases):
            with self.subTest(order=order):
                self.row['contentOrder'] = order
                for suffix, create in (('html', artifacts.create_html), ('pptx', artifacts.create_ppt)):
                    output = self.root/f'bad-{n}.{suffix}'
                    result = create(self.spec, output)
                    self.assertFalse(result['ok'], result)
                    self.assertEqual('invalid_content_layout',result['code'])
                    self.assertFalse(output.exists())
        self.row['contentOrder'] = ['table','body']
        self.row['contentLayout'] = 'horizontal'
        self.assertEqual('invalid_content_layout', artifacts.create_html(self.spec,self.root/'horizontal.html')['code'])

    def test_layout_without_order_is_rejected(self):
        del self.row['contentOrder']
        self.row['contentLayout'] = 'vertical'
        self.assertEqual('invalid_content_layout',artifacts.create_html(self.spec,self.root/'no-order.html')['code'])

    def test_document_level_order_is_not_silently_ignored(self):
        spec = {**self.spec, 'contentOrder': ['table','body']}
        self.assertEqual('invalid_content_layout',artifacts.create_html(spec,self.root/'top-level.html')['code'])

    def test_document_level_fields_are_rejected_before_import_or_preserve(self):
        for field, value in [('contentOrder',['table','body']),('contentLayout','vertical')]:
            for route in ({'htmlSource':{'path':'unopened.html'}},{'referenceMode':'preserve'}):
                with self.subTest(field=field,route=route), \
                     patch.object(ppt_html,'prepare_import',side_effect=AssertionError('IMPORT_REACHED')), \
                     patch.object(artifacts,'inspect_template',side_effect=AssertionError('TEMPLATE_REACHED')):
                    with self.assertRaises(artifacts.ArtifactError) as caught:
                        ppt_html.prepare({**self.spec,**route,field:value},self.root/'unopened.pptx')
                    self.assertEqual('invalid_content_layout',caught.exception.code)

    def test_changed_order_invalidates_an_actual_draft_approval(self):
        spec = {**self.spec,'creationMode':'new','designPreset':'warm','purpose':'현황','audience':'동료','slideCount':1}
        result = ppt_html.create_draft(spec,self.root/'approved.html')
        self.assertTrue(result['ok'],result)
        spec['designReview'] = {**result['designReview'],'confirmed':True}
        self.row['contentOrder'] = ['body','table']
        with patch.object(artifacts,'_python_ppt',side_effect=AssertionError('must not export')):
            result = artifacts.create_ppt(spec,self.root/'stale.pptx',require_choices=True)
        self.assertFalse(result['ok'],result)
        self.assertEqual('design_review_changed',result['code'])
        self.assertFalse((self.root/'stale.pptx').exists())

    def test_html_saved_dom_check_detects_renderer_regression(self):
        from company_agent import report_design
        real = report_design.render
        def broken(*args):
            return real(*args).replace('data-content-block="table"','data-content-block="body"')
        with patch.object(report_design, 'render', side_effect=broken):
            result = artifacts.create_html(self.spec,self.root/'regression.html')
        self.assertEqual('layout_validation_failed', result['code'])
        self.assertFalse((self.root/'regression.html').exists())

    def test_ppt_native_roundtrip_matches_shared_preview_scene(self):
        from pptx import Presentation
        data = ppt_html.prepare(self.spec)[0]
        page = data['presentationPlan']['pages'][0]
        table = next(e for e in page['elements'] if e['kind']=='table')
        body = next(e for e in page['elements'] if e.get('text')==self.row['body'])
        self.assertLessEqual(table['y']+table['h'], body['y'])
        draft = ppt_html.create_draft(self.spec,self.root/'draft.html',require_choices=False)
        self.assertTrue(draft['ok'],draft)
        self.assertEqual('shared-scene-coordinates',draft['contentLayoutValidation']['basis'])
        with patch.object(artifacts,'_office',return_value={'ok':False,'message':'테스트에서는 Office 렌더 생략'}), \
             patch.object(ppt_workflow,'optional_archforge',return_value={'status':'unavailable'}):
            result = artifacts.create_ppt(self.spec,self.root/'result.pptx')
        self.assertTrue(result['ok'],result)
        self.assertEqual('pptx-native-coordinates',result['contentLayoutValidation']['basis'])
        prs = Presentation(result['outputPath'])
        shapes = list(prs.slides[0].shapes)
        actual_table = next(s for s in shapes if s.has_table)
        actual_body = next(s for s in shapes if s.has_text_frame and s.text==self.row['body'])
        self.assertLessEqual(actual_table.top+actual_table.height,actual_body.top)
        self.assertEqual(0,result['editability']['images'])
        self.assertEqual(1,result['editability']['tables'])
        actual_body.text = '수정 가능한 설명'
        actual_table.table.cell(1,1).text = '6'
        prs.save(self.root/'edited.pptx')
        reread = Presentation(self.root/'edited.pptx')
        self.assertIn('수정 가능한 설명',[s.text for s in reread.slides[0].shapes if s.has_text_frame])

    def test_saved_ppt_coordinate_mismatch_is_detected(self):
        from pptx import Presentation
        data = ppt_html.prepare(self.spec)[0]
        draft = self.root/'moved.pptx'
        presentation_design.render_python(data,draft,None)
        prs = Presentation(draft)
        body = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text==self.row['body'])
        body.top = 0
        prs.save(draft)
        with self.assertRaises(content_layout.LayoutError):
            content_layout.check_ppt(draft,data)

    def test_requested_order_can_be_changed_only_with_new_approval(self):
        first = ppt_workflow.design_digest(self.spec)
        self.row['contentOrder'] = ['body','table']
        self.assertNotEqual(first, ppt_workflow.design_digest(self.spec))
        data = ppt_html.prepare(self.spec)[0]
        self.assertEqual(['body','table'], [b['name'] for b in data['presentationPlan']['pages'][0]['contentBlocks']])

    def test_saved_template_preserves_vertical_geometry_without_original_content(self):
        spec = {**self.spec,'creationMode':'new','designPreset':'warm','purpose':'현황','audience':'동료','slideCount':1}
        result = ppt_html.create_draft(spec,self.root/'draft.html')
        spec['designReview'] = {**result['designReview'],'confirmed':True}
        saved = ppt_html.save_template(spec,self.root/'template.html')
        self.assertTrue(saved['ok'],saved)
        self.assertEqual(1,saved['layoutCount'])
        template = ppt_html.read_template(Path(saved['outputPath']))
        elements = template['layouts'][0]['elements']
        table = next(e for e in elements if e['kind']=='table')
        self.assertTrue(any(e['kind']=='text' and e['y']>=table['y']+table['h'] for e in elements))
        self.assertNotIn(self.row['body'],Path(saved['outputPath']).read_text(encoding='utf-8'))

    def test_ordered_template_entire_serialized_html_has_no_business_chart_content(self):
        spec = {'title':'PRIVATE_REPORT','subtitle':'PRIVATE_SUBTITLE',
                'creationMode':'new','designPreset':'warm','purpose':'PRIVATE_PURPOSE','audience':'PRIVATE_AUDIENCE','slideCount':2,
                'slides':[{'title':'PRIVATE_SLIDE','eyebrow':'PRIVATE_EYEBROW','source':'PRIVATE_SOURCE','takeaway':'PRIVATE_TAKEAWAY',
                           'body':'PRIVATE_BODY','contentOrder':['chart','body'],
                           'chart':{'title':'CONFIDENTIAL_CUSTOMER_X_REVENUE','type':'column','categories':['PRIVATE_CATEGORY'],
                                    'metadata':'PRIVATE_CHART_METADATA',
                                    'series':[{'name':'PRIVATE_SERIES','values':[999777321],'metadata':'PRIVATE_SERIES_METADATA'}]}},
                          {'title':'PRIVATE_TABLE_TITLE','body':'PRIVATE_TABLE_BODY','contentOrder':['table','body'],
                           'table':{'headers':['PRIVATE_HEADER'],'rows':[['PRIVATE_CELL']],'metadata':'PRIVATE_TABLE_METADATA'}}]}
        preview = ppt_html.create_draft(spec,self.root/'private-draft.html')
        self.assertTrue(preview['ok'],preview)
        self.assertIn('CONFIDENTIAL_CUSTOMER_X_REVENUE',Path(preview['outputPath']).read_text(encoding='utf-8'))
        spec['designReview'] = {**preview['designReview'],'confirmed':True}
        saved = ppt_html.save_template(spec,self.root/'content-free-template.html')
        self.assertTrue(saved['ok'],saved)
        self.assertFalse(saved['businessContentStored'])
        serialized = Path(saved['outputPath']).read_text(encoding='utf-8')
        for marker in ('PRIVATE_', 'CONFIDENTIAL_CUSTOMER_X_REVENUE', '999777321'):
            self.assertNotIn(marker,serialized)
        self.assertEqual(2,len(ppt_html.read_template(Path(saved['outputPath']))['layouts']))

    def test_template_data_payload_is_rebuilt_not_retained_with_unknown_fields(self):
        import json
        from company_agent.ppt_scene import template_layouts
        original = {'width':960,'height':540,'pages':[{'elements':[
            {'kind':'chart','x':50,'y':50,'w':300,'h':200,'chart':{
                'type':'column','title':'PRIVATE_TITLE','categories':['PRIVATE_CATEGORY'],
                'metadata':{'owner':'PRIVATE_OWNER'},
                'series':[{'name':'PRIVATE_SERIES','values':[999777321],'source':'PRIVATE_SERIES_SOURCE'}]}},
            {'kind':'table','x':50,'y':280,'w':300,'h':100,'table':{
                'headers':['PRIVATE_HEADER'],'rows':[['PRIVATE_CELL']], 'metadata':'PRIVATE_TABLE_SOURCE'}}]}]}
        cleaned = template_layouts(original)
        for marker in ('PRIVATE_', '999777321'):
            self.assertNotIn(marker,json.dumps(cleaned))
        self.assertEqual('PRIVATE_TITLE',original['pages'][0]['elements'][0]['chart']['title'])

    def test_dense_order_fails_before_any_export_no_silent_side_by_side(self):
        self.row['table']['rows'] = [['항목',str(n)] for n in range(10)]
        with patch.object(artifacts,'_office',side_effect=AssertionError('must not run')), \
             patch.object(artifacts,'_python_ppt',side_effect=AssertionError('must not run')):
            result = artifacts.create_ppt(self.spec,self.root/'dense.pptx')
        self.assertFalse(result['ok'])
        self.assertEqual('ppt_design_invalid',result['code'])
        self.assertIn('본문 순서',result['message'])
        self.assertFalse((self.root/'dense.pptx').exists())

    def test_explicit_order_not_silently_ignored_in_import_or_preserve(self):
        for extra in ({'htmlSource': {'path':'unopened.html'}},{'referenceMode':'preserve'}):
            result = artifacts.create_ppt({**self.spec,**extra},self.root/'unsupported.pptx')
            self.assertEqual('content_layout_unsupported',result['code'])
        self.row['elements'] = [{'kind':'text'}]
        result = artifacts.create_ppt(self.spec,self.root/'elements.pptx')
        self.assertEqual('ppt_design_invalid',result['code'])

    def test_body_bullets_table_are_not_merged_or_lost(self):
        self.row['bullets'] = ['다음 확인 사항']
        self.row['contentOrder'] = ['table','bullets','body']
        data = ppt_html.prepare(self.spec)[0]
        self.assertEqual(self.row['contentOrder'],[b['name'] for b in data['presentationPlan']['pages'][0]['contentBlocks']])
        self.assertEqual('checked',content_layout.check_plan(data)['status'])

    def test_chart_with_title_and_contained_image_roundtrip(self):
        from PIL import Image
        picture = self.root/'wide.png'
        Image.new('RGB',(240,30),'white').save(picture)
        for name, payload in [('chart',{'title':'항목 비교','categories':['A','B'],'series':[{'name':'값','values':[1,2]}]}),
                              ('image',{'path':str(picture),'alt':'가상 그림'})]:
            spec = {'slides':[{'title':'먼저 읽는 기본 장','body':'이 장에는 명시 순서 없음'},
                              {'title':'근거','body':'확인된 내용입니다.',name:payload,'contentOrder':[name,'body']}]}
            data = ppt_html.prepare(spec)[0]
            output = self.root/(name+'.pptx')
            presentation_design.render_python(data,output,None)
            result = content_layout.check_ppt(output,data)
            self.assertEqual('pptx-native-coordinates',result['basis'])
            self.assertEqual(2,result['items'][0]['slide'])


if __name__ == '__main__':
    unittest.main()
