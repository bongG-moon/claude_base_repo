"""Bounded, read-only Excel observation through xlwings on Windows.

Only the worker owns Excel. Cell data travels in memory; no workbook copy,
Save/SaveAs, macro execution, refresh, decryption, or security-setting downgrade.
Range.Value(11) is Microsoft's XML Spreadsheet range observation, not a file
export. Formula values are observations, not proof of fresh recalculation.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import threading
import time
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

MAX_CELLS = 250_000
MAX_SHEETS = 100
MAX_NAMES = 1000
MAX_TABLES = 500
MAX_MERGES = 5000
MAX_XML = 48 * 1024 * 1024
MAX_OUTPUT = 64 * 1024 * 1024
MAX_MODULES = 64
MAX_MODULE_SOURCE = 2 * 1024 * 1024
MAX_TOTAL_SOURCE = 8 * 1024 * 1024
SS = "{urn:schemas-microsoft-com:office:spreadsheet}"


class NativeReadError(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


def _limit(condition, message):
    if not condition:
        raise NativeReadError("native_limit_exceeded", message)


def _bounds(row, col, rows, cols):
    _limit(all(isinstance(v, int) and v > 0 for v in (row, col, rows, cols)), "Excel 범위가 올바르지 않습니다.")
    _limit(row + rows - 1 <= 100_000 and col + cols - 1 <= 512,
           "Excel 읽기는 최대 100,000행·512열 범위까지 지원합니다.")
    _limit(rows * cols <= MAX_CELLS, "Excel 사용 범위가 250,000셀을 넘습니다. 분석 범위를 나누어 주세요.")


def _matrix(value, rows, cols):
    if rows == cols == 1:
        return [[value]]
    _limit(isinstance(value, (tuple, list)) and len(value) == rows,
           "Excel 수식 배열의 행 수가 맞지 않습니다.")
    result = [list(row) if isinstance(row, (tuple, list)) else [row] for row in value]
    _limit(all(len(row) == cols for row in result), "Excel 수식 배열의 열 수가 맞지 않습니다.")
    return result


def _number(text):
    value = float(text)
    _limit(math.isfinite(value), "유한하지 않은 Excel 숫자는 처리하지 않습니다.")
    return int(value) if value.is_integer() else value


def parse_range_xml(xml, formulas, origin, shape, formula_cells=None, values=None):
    """Normalize a single rectangular COM range; preserve actual cell types.

    Excel XML range indexes are relative to the exported range. Some exports
    include TopCell/LeftCell; these are validated instead of guessing offsets.
    """
    _limit(isinstance(xml, str) and len(xml.encode("utf-8")) <= MAX_XML,
           "Excel 범위 XML이 크기 한도를 넘습니다.")
    normalized = xml.replace("\x00", "").upper()
    _limit("<!DOCTYPE" not in normalized and "<!ENTITY" not in normalized,
           "DTD/entity가 포함된 Excel 범위 XML은 처리하지 않습니다.")
    root = ET.fromstring(xml)
    rows, cols = shape
    first_row, first_col = origin
    _bounds(first_row, first_col, rows, cols)
    matrix = _matrix(formulas, rows, cols)
    observed_values = _matrix(values, rows, cols) if values is not None else None
    styles = {}
    for style in root.findall(f"{SS}Styles/{SS}Style"):
        _limit(len(styles) < 4096, "Excel 셀 서식 수가 한도를 넘습니다.")
        number = style.find(SS + "NumberFormat")
        font, fill = style.find(SS + "Font"), style.find(SS + "Interior")
        styles[style.get(SS + "ID")] = (style.get(SS + "Parent"),
                                      number.get(SS + "Format") if number is not None else None,
                                      font.get(SS + "Bold") == "1" if font is not None else None,
                                      fill.get(SS + "Pattern", "None") != "None" if fill is not None else None)

    def number_format(style_id):
        seen = set()
        while style_id in styles:
            _limit(style_id not in seen, "Excel 서식 상속에 순환이 있습니다.")
            seen.add(style_id)
            parent, fmt, _, _ = styles[style_id]
            if fmt is not None:
                # SpreadsheetML's named formats are not OOXML format codes.
                return {"General Number": "General", "Short Date": "yyyy-mm-dd",
                        "Long Date": "yyyy-mm-dd", "General Date": "yyyy-mm-dd hh:mm:ss"}.get(fmt, fmt)
            style_id = parent or ("Default" if style_id != "Default" else None)
        return "General"

    def header_style(style_id):
        seen, bold, fill = set(), None, None
        while style_id in styles:
            _limit(style_id not in seen, "Excel 서식 상속에 순환이 있습니다.")
            seen.add(style_id)
            parent, _, current_bold, current_fill = styles[style_id]
            if bold is None:
                bold = current_bold
            if fill is None:
                fill = current_fill
            style_id = parent or ("Default" if style_id != "Default" else None)
        return {"bold": bool(bold), "fill": bool(fill)}

    tables = root.findall(f"{SS}Worksheet/{SS}Table")
    _limit(len(tables) == 1, "Excel 범위 XML에서 표를 하나로 확인하지 못했습니다.")
    table = tables[0]
    # Range XML is a new range-local table; TopCell/LeftCell, if present, declare
    # the original origin rather than changing Row/Cell Index relativity.
    for key, expected in (("TopCell", first_row), ("LeftCell", first_col)):
        if table.get(SS + key) is not None:
            _limit(int(table.get(SS + key)) in {1, expected}, "Excel XML 범위 시작점이 일치하지 않습니다.")
    cells, merges = [], []
    row_index = 0
    for row in table.findall(SS + "Row"):
        row_index = int(row.get(SS + "Index", row_index + 1))
        _limit(1 <= row_index <= rows, "Excel XML 행 범위가 맞지 않습니다.")
        col_index = 0
        for cell in row.findall(SS + "Cell"):
            col_index = int(cell.get(SS + "Index", col_index + 1))
            _limit(1 <= col_index <= cols, "Excel XML 열 범위가 맞지 않습니다.")
            across, down = int(cell.get(SS + "MergeAcross", 0)), int(cell.get(SS + "MergeDown", 0))
            _limit(across >= 0 and down >= 0 and col_index + across <= cols and row_index + down <= rows,
                   "Excel XML 병합 범위가 맞지 않습니다.")
            r, c = first_row + row_index - 1, first_col + col_index - 1
            if across or down:
                merges.append(f"{get_column_letter(c)}{r}:{get_column_letter(c+across)}{r+down}")
                _limit(len(merges) <= MAX_MERGES, "병합 범위 수가 한도를 넘습니다.")
            data = cell.find(SS + "Data")
            value, kind = None, "n"
            if data is not None:
                text = "".join(data.itertext())
                xml_type = data.get(SS + "Type", "String")
                if xml_type == "Number":
                    value = _number(text)
                elif xml_type == "Boolean":
                    _limit(text in {"0", "1"}, "Excel 논리값이 올바르지 않습니다.")
                    value, kind = text == "1", "b"
                elif xml_type == "DateTime":
                    value, kind = dt.datetime.fromisoformat(text).replace(tzinfo=None).isoformat(), "d"
                elif xml_type == "Error":
                    value, kind = text, "e"
                elif xml_type == "String":
                    # Excel's XML range export may trim leading/trailing spaces
                    # even with xml:space='preserve'. Value2 is the lossless bulk
                    # observation for strings, including formula-cached strings.
                    value = observed_values[row_index-1][col_index-1] if observed_values is not None else text
                    _limit(value is None or isinstance(value, str), "Excel 문자열 값과 XML 셀 형식이 일치하지 않습니다.")
                    kind = "s"
                else:
                    raise NativeReadError("native_type_unsupported", "지원하지 않는 Excel XML 셀 형식입니다.")
            formula = None
            if ((r, c) in formula_cells if formula_cells is not None else cell.get(SS + "Formula") is not None):
                formula = matrix[row_index - 1][col_index - 1]
                _limit(isinstance(formula, str) and formula.startswith("="), "Excel A1 수식을 확인하지 못했습니다.")
            style_id = cell.get(SS + "StyleID", row.get(SS + "StyleID", "Default"))
            fmt = number_format(style_id)
            if data is not None or formula is not None:
                cells.append([r, c, value, kind, formula, fmt, header_style(style_id)])
            col_index += across
    if formula_cells is not None:
        actual = {(cell[0], cell[1]) for cell in cells if cell[4] is not None}
        _limit(actual == formula_cells, "Excel XML에 모든 수식 셀이 포함되지 않았습니다.")
    return cells, merges


def _formula_cells(used, origin, shape):
    """XML range exports omit formulas referring outside the range/workbook.

    Read the true formula mask as bounded native Areas, never infer it from an
    '=' string (which may be literal text), never perform a per-cell COM read.
    """
    has_formula = used.HasFormula
    if has_formula is False or has_formula == 0:
        return set()
    row, col = origin
    rows, cols = shape
    if has_formula is True or has_formula == -1:
        return {(r, c) for r in range(row, row+rows) for c in range(col, col+cols)}
    areas = used.SpecialCells(-4123).Areas  # xlCellTypeFormulas
    _limit(int(areas.Count) <= 5000, "분산된 수식 영역 수가 5,000개를 넘습니다.")
    result = set()
    for index in range(1, int(areas.Count)+1):
        area = areas.Item(index)
        ar, ac = int(area.Row), int(area.Column)
        nr, nc = int(area.Rows.Count), int(area.Columns.Count)
        _bounds(ar, ac, nr, nc)
        _limit(row <= ar and col <= ac and ar+nr <= row+rows and ac+nc <= col+cols,
               "수식 영역이 읽기 범위를 벗어납니다.")
        result.update((r, c) for r in range(ar, ar+nr) for c in range(ac, ac+nc))
    return result


def _vba(book):
    from vba_extract import outcome, module_path
    try:
        if not bool(book.HasVBProject):
            return outcome("no_vba", "내장 VBA 프로젝트가 없습니다.")
    except Exception:
        return outcome("access_denied", "VBA 존재 여부를 확인할 수 없습니다. 셀·수식 분석만 계속합니다.")
    try:
        project = book.VBProject  # One authorized access; never lower VBOM protection.
        if int(project.Protection) != 0:
            return outcome("access_denied", "VBA 프로젝트가 잠겨 있습니다. 잠금을 풀지 않고 셀·수식만 처리합니다.")
        components = project.VBComponents
        count = int(components.Count)
        _limit(0 < count <= MAX_MODULES, "VBA 모듈 수가 추출 한도를 넘습니다.")
        modules, total = [], 0
        for index in range(1, count + 1):
            component = components.Item(index)
            code = component.CodeModule
            lines = int(code.CountOfLines)
            _limit(0 <= lines <= 100_000, "VBA 모듈 줄 수가 추출 한도를 넘습니다.")
            chunks, module_bytes = [], 0
            for first in range(1, lines+1, 1000):
                chunk = str(code.Lines(first, min(1000, lines-first+1)))
                module_bytes += len(chunk.encode("utf-8"))
                _limit(module_bytes <= MAX_MODULE_SOURCE and total + module_bytes <= MAX_TOTAL_SOURCE,
                       "VBA 원문 크기가 추출 한도를 넘습니다.")
                chunks.append(chunk)
                # CodeModule.Lines can omit the separator after each chunk.
                if first + 1000 <= lines and not chunk.endswith(("\n", "\r")):
                    chunks.append("\r\n")
            source = "".join(chunks)
            encoded = source.encode("utf-8")
            total += len(encoded)
            _limit(len(encoded) <= MAX_MODULE_SOURCE and total <= MAX_TOTAL_SOURCE,
                   "VBA 원문 크기가 추출 한도를 넘습니다.")
            name = str(component.Name)
            _limit(len(name) <= 255, "VBA 모듈 이름이 한도를 넘습니다.")
            extension = "bas" if int(component.Type) == 1 else "cls"
            modules.append({"name": name, "relative_path": module_path(index, name, extension),
                            "module_type": extension, "stream_path": "VBProject/" + name,
                            "source": source, "source_sha256": hashlib.sha256(encoded).hexdigest(),
                            "source_encoding": "unicode_from_excel", "decoded_with_replacement": False,
                            "source_bytes": len(encoded), "line_count": lines})
        result = outcome("extracted", "Excel에서 허용한 VBA 원문을 실행 없이 읽었습니다.")
        result.update(modules=modules, module_count=count, source_bytes=total,
                      limitations=["VBA 속성·UserForm/FRX·p-code는 추출하지 않습니다.",
                                   "매크로 실행 및 Python과의 동등성은 이 읽기로 검증되지 않습니다."])
        return result
    except NativeReadError:
        return outcome("limit_exceeded", "VBA 추출 한도를 넘었습니다. 셀·수식 분석만 계속합니다.")
    except Exception:
        return outcome("access_denied", "VBA 원문 접근이 거절되었거나 사용할 수 없습니다. 설정을 바꾸지 않고 셀·수식만 처리합니다.")


def _address(native_range):
    address = str(native_range.Address).replace("$", "")
    c1, r1, c2, r2 = range_boundaries(address)
    _bounds(r1, c1, r2-r1+1, c2-c1+1)
    return address


def _native_snapshot(source, xw, report_owner):
    """Injection point for safety tests; production called only by child."""
    app = book = blank = None
    evidence = {"reader": "xlwings", "drm_status": "not_determined", "read_only": True,
                "macro_executed": False, "source_saved": False, "security_policy_changed": False,
                "automation_security": 3, "events_enabled": False, "update_links": 0,
                "calculation": "manual", "values_origin": "excel_observed_not_recalculated",
                "fresh_recalculation_verified": False, "refresh_requested": False,
                "refresh_on_open_prevention": "not_guaranteed_by_excel_com",
                "incomplete_metadata": [], "non_worksheet_sheets_not_read": 0}
    try:
        app = xw.App(visible=False, add_book=False)
        report_owner(int(app.pid))
        api = app.api
        api.AutomationSecurity = 3
        api.EnableEvents = False
        api.DisplayAlerts = False
        api.AskToUpdateLinks = False
        # Excel requires a workbook before Calculation can be set reliably.
        blank = api.Workbooks.Add(-4167)  # Built-in empty worksheet, not a template.
        api.Calculation = -4135
        _limit(int(api.AutomationSecurity) == 3 and not bool(api.EnableEvents)
               and int(api.Calculation) == -4135, "Excel 안전 읽기 설정을 확인하지 못했습니다.")
        book = api.Workbooks.Open(str(source), UpdateLinks=0, ReadOnly=True,
                                 Password="", WriteResPassword="", IgnoreReadOnlyRecommended=True,
                                 Notify=False, AddToMru=False, CorruptLoad=0)
        _limit(bool(book.ReadOnly), "읽기 전용 Excel 열기를 확인하지 못했습니다.")
        if int(book.FileFormat) not in {51, 52, 61}:
            raise NativeReadError("native_format_unsupported", "실제 Excel 형식이 xlsx/xlsm이 아닙니다. 확장자만 바꾼 파일은 처리하지 않습니다.")
        evidence["file_format"] = int(book.FileFormat)
        _limit(int(book.Sheets.Count) <= MAX_SHEETS + 20, "Excel 시트 수가 안전 읽기 한도를 넘습니다.")
        for position in range(1, int(book.Sheets.Count)+1):
            if int(book.Sheets.Item(position).Type) in {3, 4, -4116}:
                raise NativeReadError("native_legacy_sheet_unsupported", "Excel 4.0 매크로·대화상자 시트는 안전한 읽기 대상에서 제외합니다.")
        count = int(book.Worksheets.Count)
        _limit(0 < count <= MAX_SHEETS, "워크시트 수가 100개 한도를 넘거나 워크시트가 없습니다.")
        result = {"sheets": [], "names": [], "features": [], "date1904": bool(book.Date1904),
                  "evidence": evidence}
        evidence["excel_version"] = str(api.Version)[:30]
        evidence["non_worksheet_sheets_not_read"] = max(0, int(book.Sheets.Count)-count)
        scanned = merge_area = table_area = table_count = 0
        for index in range(1, count+1):
            sheet = book.Worksheets.Item(index)
            used = sheet.UsedRange
            row, col = int(used.Row), int(used.Column)
            rows, cols = int(used.Rows.Count), int(used.Columns.Count)
            _bounds(row, col, rows, cols)
            scanned += rows * cols
            _limit(scanned <= MAX_CELLS, "워크북의 누적 사용 범위가 250,000셀을 넘습니다.")
            xml = used._oleobj_.Invoke(6, 0, 2, True, 11)
            formula_cells = _formula_cells(used, (row, col), (rows, cols))
            try:
                native_formulas = used.Formula2
            except Exception:
                native_formulas = used.Formula
                evidence["incomplete_metadata"].append(f"sheet_{index}_formula2_unavailable_used_legacy_formula")
            cells, merges = parse_range_xml(xml, native_formulas, (row, col), (rows, cols),
                                             formula_cells, used.Value2)
            for merged in merges:
                c1, r1, c2, r2 = range_boundaries(merged)
                merge_area += (c2-c1+1)*(r2-r1+1)
                _limit(merge_area <= MAX_CELLS, "누적 병합 셀 수가 한도를 넘습니다.")
            record = {"name": str(sheet.Name), "state": {-1: "visible", 0: "hidden", 2: "veryHidden"}[int(sheet.Visible)],
                      "cells": cells, "merges": merges, "tables": []}
            try:
                native_tables = sheet.ListObjects
                size = int(native_tables.Count)
                table_count += size
                _limit(table_count <= MAX_TABLES, "Excel Table 수가 한도를 넘습니다.")
                for position in range(1, size+1):
                    table = native_tables.Item(position)
                    reference = _address(table.Range)
                    c1, r1, c2, r2 = range_boundaries(reference)
                    table_area += (c2-c1+1)*(r2-r1+1)
                    _limit(table_area <= 1_000_000, "Excel Table 누적 면적이 한도를 넘습니다.")
                    record["tables"].append({"name": str(table.Name), "ref": reference,
                                             "headerRowCount": 1 if bool(table.ShowHeaders) else 0,
                                             "totalsRowCount": 1 if bool(table.ShowTotals) else 0})
            except NativeReadError:
                raise
            except Exception:
                record["tables"] = []
                evidence["incomplete_metadata"].append(f"sheet_{index}_tables")
            result["sheets"].append(record)
            for label, getter in (("query_tables", lambda: sheet.QueryTables.Count),
                                  ("pivot_tables", lambda: sheet.PivotTables().Count)):
                try:
                    if int(getter()) > 0 and label not in result["features"]:
                        result["features"].append(label)
                except Exception:
                    evidence["incomplete_metadata"].append(f"sheet_{index}_{label}")
        try:
            names = book.Names
            _limit(int(names.Count) <= MAX_NAMES, "정의된 이름 수가 한도를 넘습니다.")
            for index in range(1, int(names.Count)+1):
                name = names.Item(index)
                label, ref = str(name.Name), str(name.RefersTo)
                _limit(len(label) <= 512 and len(ref) <= 8192, "정의된 이름 길이가 한도를 넘습니다.")
                result["names"].append({"name": label, "ref": ref, "hidden": not bool(name.Visible)})
        except NativeReadError:
            raise
        except Exception:
            result["names"] = []
            evidence["incomplete_metadata"].append("defined_names")
        for label, getter in (("connections", lambda: book.Connections.Count),
                              ("external_links", lambda: len(book.LinkSources(1) or []))):
            try:
                if int(getter()) > 0:
                    result["features"].append(label)
            except Exception:
                evidence["incomplete_metadata"].append(label)
        result["extraction"] = _vba(book)
        if result["extraction"]["status"] != "no_vba":
            result["features"].append("macros")
        evidence["scanned_cells"] = scanned
        evidence["vba_status"] = result["extraction"]["status"]
        return result
    finally:
        # Never touch xw.apps.active or attach to an existing user's workbook.
        for owned_book in (book, blank):
            if owned_book is not None:
                try:
                    owned_book.Close(SaveChanges=False)
                except Exception:
                    evidence["incomplete_metadata"].append("owned_book_close_failed")
        if app is not None:
            try:
                app.quit()
            except Exception:
                evidence["incomplete_metadata"].append("owned_excel_quit_failed")


def snapshot_to_workbooks(snapshot):
    formula_wb, observed_wb = openpyxl.Workbook(), openpyxl.Workbook()
    for workbook in (formula_wb, observed_wb):
        workbook.remove(workbook.active)
        workbook.epoch = CALENDAR_MAC_1904 if snapshot["date1904"] else CALENDAR_WINDOWS_1900
    for item in snapshot["sheets"]:
        ws, observed = (wb.create_sheet(item["name"]) for wb in (formula_wb, observed_wb))
        ws.sheet_state = observed.sheet_state = item["state"]
        for record in item["cells"]:
            row, col, value, kind, formula, fmt = record[:6]
            style = record[6] if len(record) > 6 else {}
            if kind == "d" and value is not None:
                value = dt.datetime.fromisoformat(value)
            for target, is_formula in ((ws, bool(formula)), (observed, False)):
                cell = target.cell(row, col)
                cell.value = formula if is_formula else value
                cell.data_type = "f" if is_formula else kind
                cell.number_format = fmt
                if style.get("bold"):
                    cell.font = Font(bold=True)
                if style.get("fill"):
                    cell.fill = PatternFill(patternType="solid")
        for merged in item["merges"]:
            ws.merge_cells(merged)
            observed.merge_cells(merged)
        for table in item["tables"]:
            ws.add_table(Table(displayName=table["name"], ref=table["ref"],
                               headerRowCount=table["headerRowCount"], totalsRowCount=table["totalsRowCount"]))
    for item in snapshot["names"]:
        label, ref = item["name"], item["ref"]
        ref = ref[1:] if ref.startswith("=") else ref
        scope = None
        if "!" in label:
            prefix, local = label.rsplit("!", 1)
            sheet = prefix.strip("'").replace("''", "'")
            if sheet in formula_wb.sheetnames:
                scope, label = sheet, local
        defined = DefinedName(label, attr_text=ref, hidden=item["hidden"])
        if scope is None:
            formula_wb.defined_names.add(defined)
        else:
            formula_wb[scope].defined_names.add(defined)
    return formula_wb, observed_wb, snapshot["features"], snapshot["extraction"], snapshot["evidence"]


def _process_ids():
    import ctypes
    from ctypes import wintypes
    ids = (wintypes.DWORD * 65536)()
    needed = wintypes.DWORD()
    if not ctypes.windll.psapi.EnumProcesses(ids, ctypes.sizeof(ids), ctypes.byref(needed)):
        raise NativeReadError("native_ownership_unavailable", "Excel 프로세스 소유 범위를 확인할 수 없습니다.")
    _limit(needed.value < ctypes.sizeof(ids), "프로세스 목록 한도를 넘었습니다.")
    return set(ids[:needed.value // ctypes.sizeof(wintypes.DWORD)])


class _OwnedProcess:
    """Hold a Windows handle so PID reuse cannot terminate somebody else's app."""
    def __init__(self, pid, previous_ids, earliest):
        import ctypes
        from ctypes import wintypes
        self.handle = None
        self.kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        self.kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
        self.kernel.OpenProcess.restype = wintypes.HANDLE
        self.kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        self.kernel.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
        self.kernel.TerminateProcess.argtypes = [wintypes.HANDLE, wintypes.UINT]
        self.kernel.GetProcessTimes.argtypes = [wintypes.HANDLE] + [ctypes.POINTER(wintypes.FILETIME)]*4
        self.kernel.QueryFullProcessImageNameW.argtypes = [wintypes.HANDLE, wintypes.DWORD, wintypes.LPWSTR, ctypes.POINTER(wintypes.DWORD)]
        if pid in previous_ids or pid <= 0:
            raise NativeReadError("native_ownership_unavailable", "기존 Excel 인스턴스는 사용하거나 종료하지 않습니다.")
        handle = self.kernel.OpenProcess(0x1000 | 0x00100000 | 0x0001, False, pid)
        if not handle:
            raise NativeReadError("native_ownership_unavailable", "새 Excel의 프로세스 핸들을 확인하지 못했습니다.")
        try:
            times = [wintypes.FILETIME() for _ in range(4)]
            if not self.kernel.GetProcessTimes(handle, *(ctypes.byref(t) for t in times)):
                raise ValueError()
            created = ((times[0].dwHighDateTime << 32) | times[0].dwLowDateTime)/10_000_000 - 11644473600
            path, size = ctypes.create_unicode_buffer(32768), wintypes.DWORD(32768)
            if not self.kernel.QueryFullProcessImageNameW(handle, 0, path, ctypes.byref(size)):
                raise ValueError()
            if Path(path.value).name.casefold() != "excel.exe" or created < earliest:
                raise ValueError()
            self.handle = handle
        except Exception:
            self.kernel.CloseHandle(handle)
            raise NativeReadError("native_ownership_unavailable", "새로 만든 Excel임을 확인하지 못했습니다.")

    def close(self):
        terminated = False
        if self.handle:
            if self.kernel.WaitForSingleObject(self.handle, 1000) == 258:
                terminated = bool(self.kernel.TerminateProcess(self.handle, 1))
                self.kernel.WaitForSingleObject(self.handle, 3000)
            self.kernel.CloseHandle(self.handle)
            self.handle = None
        return terminated


def load_via_excel(source, timeout=90):
    source = Path(source).resolve()
    if os.name != "nt":
        raise NativeReadError("native_windows_required", "xlwings 읽기는 Windows와 설치된 Excel이 필요합니다.")
    _limit(isinstance(timeout, (int, float)) and math.isfinite(timeout) and 1 <= timeout <= 300,
           "Excel 읽기 제한 시간은 1~300초입니다.")
    if not source.is_file() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise NativeReadError("native_invalid_source", "읽을 xlsx/xlsm 파일을 확인해 주세요.")
    _limit(source.stat().st_size <= 64*1024*1024, "시험판 입력 한도는 64 MiB입니다.")
    previous_ids, earliest = _process_ids(), time.time()
    proc = subprocess.Popen([sys.executable, "-X", "utf8", "-B", str(Path(__file__).resolve()),
                             "--worker", str(source)], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                            stderr=subprocess.DEVNULL, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    output, state = bytearray(), {"owner": None, "error": None}
    finished = threading.Event()

    def collect():
        prefix, owner_seen = bytearray(), False
        try:
            while True:
                chunk = proc.stdout.read1(65536)
                if not chunk:
                    break
                if not owner_seen:
                    prefix.extend(chunk)
                    if b"\n" not in prefix:
                        _limit(len(prefix) <= 8192, "Excel 작업자 시작 응답이 너무 큽니다.")
                        continue
                    first, rest = prefix.split(b"\n", 1)
                    message = json.loads(first)
                    if message.get("kind") == "owner":
                        state["owner"] = _OwnedProcess(int(message["pid"]), previous_ids, earliest)
                        proc.stdin.write(b"owned\n")
                        proc.stdin.flush()
                    elif message.get("kind") == "error":
                        output.extend(first)
                    else:
                        raise NativeReadError("native_protocol_error", "Excel 작업자 시작 응답이 올바르지 않습니다.")
                    owner_seen = True
                    chunk = rest
                output.extend(chunk)
                _limit(len(output) <= MAX_OUTPUT, "Excel 작업자 결과가 메모리 한도를 넘습니다.")
        except Exception as error:
            state["error"] = error
        finally:
            finished.set()

    thread = threading.Thread(target=collect, daemon=True)
    thread.start()
    forced_cleanup = False
    try:
        if not finished.wait(timeout):
            raise NativeReadError("native_timeout", "Excel 읽기 제한 시간이 지났습니다. 승인 창·파일 접근 상태를 확인해 주세요.")
        if state["error"]:
            raise state["error"]
        result = json.loads(output)
        if result.get("kind") == "error":
            raise NativeReadError(result.get("code", "native_failed"), result.get("message", "Excel 읽기를 마치지 못했습니다."))
        if state["owner"] is None or result.get("kind") != "result":
            raise NativeReadError("native_protocol_error", "소유가 확인된 Excel 결과가 아닙니다.")
        return_value = snapshot_to_workbooks(result["snapshot"])
    finally:
        if proc.poll() is None:
            proc.kill()  # Our Python worker only; never process-name based cleanup.
        proc.wait(timeout=5)
        thread.join(timeout=2)
        if state["owner"] is not None:
            forced_cleanup = state["owner"].close()
        for pipe in (proc.stdin, proc.stdout):
            if pipe:
                pipe.close()
    return_value[-1]["owned_excel_forced_cleanup"] = forced_cleanup
    return return_value


def _worker(source):
    def report_owner(pid):
        print(json.dumps({"kind": "owner", "pid": pid}), flush=True)
        if sys.stdin.readline(20).strip() != "owned":
            raise NativeReadError("native_ownership_unavailable", "Excel 소유 확인이 되지 않았습니다.")
    try:
        import xlwings as xw
        snapshot = _native_snapshot(Path(source), xw, report_owner)
        result = {"kind": "result", "snapshot": snapshot}
        encoded = json.dumps(result, ensure_ascii=False, allow_nan=False)
        _limit(len(encoded.encode("utf-8")) <= MAX_OUTPUT, "Excel 결과가 크기 한도를 넘습니다.")
        print(encoded, flush=True)
    except ImportError:
        print(json.dumps({"kind": "error", "code": "xlwings_missing",
                          "message": "xlwings 의존성이 없습니다. 별도 스킬의 requirements-excel.txt를 준비해 주세요."}), flush=True)
    except NativeReadError as error:
        print(json.dumps({"kind": "error", "code": error.code, "message": str(error)}, ensure_ascii=False), flush=True)
    except Exception as error:
        print(json.dumps({"kind": "error", "code": "native_excel_access_failed",
                          "message": "Excel에서 읽기를 완료하지 못했습니다. 파일을 직접 열 수 있는지와 회사의 자동화 접근 허용 여부를 확인해 주세요.",
                          "error_type": type(error).__name__}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--worker":
        _worker(sys.argv[2])
    else:
        raise SystemExit("Import load_via_excel; this script is an internal bounded worker.")
