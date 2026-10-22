"""Offline, metadata-only Excel lineage report. No model calls or row values."""
from __future__ import annotations

import base64
from collections import Counter, defaultdict
from html import escape
from pathlib import Path
import re

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from formula_engine import FormulaFailure, Parser, FUNCTIONS
from compact_project import expression_code, FUNCTION_NOTES, OPERATOR_NOTES

MAX_NODES = 48
MAX_EDGES = 160
MAX_DETAIL_COLUMNS = 40
MAX_FORMULA_GROUPS = 16
MAX_EXPLAIN_FORMULAS = 20_000
MAX_FORMULA_REFERENCES = 10
MAX_FORMULA_STEPS = 10


def short(value, limit=100):
    text = " ".join(str(value or "").replace("\x00", "").split())
    return text if len(text) <= limit else text[:limit - 1] + "…"


def table_name(table):
    """Names are cautious interpretations of labels, never sampled row values."""
    if table.get("display_name"):
        return short(table["display_name"], 80), "확인한 명세의 이름"
    sheet = short(table.get("sheet"), 35)
    columns = [short(c, 24) for c in table.get("columns", [])
               if c and not re.fullmatch(r"(?:column|열|col)[ _]?(?:\d+|[A-Z]{1,3})", str(c), re.I)]
    generic_sheet = not sheet or bool(re.fullmatch(r"(?:(?:sheet|시트|data)\s*\d*|기준표?|입력|데이터)", sheet, re.I))
    # A meaningful sheet title is stronger evidence than a single header keyword.
    if not generic_sheet:
        return sheet + " 테이블", "시트 이름 기준 · 업무 의미 확인 필요"
    text = " ".join(columns)
    combinations = ((('과정코드', '과정명'), '교육 과정 기준'),
                    (('부서', '예산'), '부서 예산'),
                    (('회차', '시작일', '정원'), '교육 회차'),
                    (('신청', '수료'), '신청·수료 실적'))
    normalized = text.replace(" ", "")
    for keywords, meaning in combinations:
        if all(keyword in normalized for keyword in keywords):
            return meaning + " 테이블", "열 조합 기반 추정"
    rules = (("신청", "신청 현황"), ("강사", "강사 정보"), ("정산", "정산 내역"),
             ("교육", "교육 과정"), ("재고", "재고 현황"), ("매출", "매출 내역"),
             ("급여", "급여 내역"), ("발주", "발주 내역"), ("고객", "고객 정보"))
    for keyword, meaning in rules:
        if keyword in text:
            return meaning + " 테이블", "시트·열 이름 기반 추정"
    if columns:
        return " · ".join(columns[:2]) + " 테이블", "열 이름 기준 · 업무 의미 확인 필요"
    return "데이터 테이블", "이름 확인 필요"


def _safe_formula(formula):
    """Hide string constants before shortening; never export cached/row values."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return "원본 수식 형식을 확인해야 합니다."
    # Also hides a malformed unterminated string, so failure is not a data leak.
    masked = re.sub(r'"(?:""|[^"])*(?:"|$)', '"[문자 상수 가림]"', formula)
    if "[" in masked.replace("[문자 상수 가림]", "") or re.search(r"(?:https?://|[A-Za-z]:[\\/]|\\\\)", masked):
        return "외부 경로·구조화 참조가 포함된 수식 — 원문은 상세 진단에서 확인하세요."
    return short(masked, 360)


def _table_bounds(table):
    try:
        bounds = range_boundaries(str(table.get("range", "")))
        return bounds if all(isinstance(n, int) and n > 0 for n in bounds) else None
    except (TypeError, ValueError):
        return None


def _reference_label(sheet, reference, tables):
    """Map observed references to headers without inspecting any data cells."""
    labels = []
    try:
        left, top, right, bottom = range_boundaries(str(reference).replace("$", ""))
        if None in (left, top, right, bottom):
            raise ValueError()
        for table in tables:
            bounds = _table_bounds(table)
            if table.get("sheet") != sheet or not bounds:
                continue
            c1, r1, c2, r2 = bounds
            if right < c1 or left > c2 or bottom < r1 or top > r2:
                continue
            for col in range(max(c1, left), min(c2, right) + 1):
                index = col - c1
                if index < len(table.get("columns", [])):
                    labels.append(short(table["columns"][index], 32))
                if len(labels) >= 4:
                    break
            if len(labels) >= 4:
                break
    except (TypeError, ValueError):
        pass
    location = short(sheet, 60) + "!" + short(reference, 60)
    labels = list(dict.fromkeys(labels))
    return location + (" · " + ", ".join(labels[:3]) + (" 등" if len(labels) > 3 else "") if labels else "")


def _reference_entry(sheet, reference, tables):
    identity = (str(sheet).casefold(), str(reference).replace("$", "").upper())
    return identity, _reference_label(sheet, reference, tables)


def _argument_label(node, sheet, tables, depth=0):
    """Describe an argument's structure, never resolve its data value."""
    if depth > 2:
        return "안쪽 계산식"
    kind = node[0]
    if kind == "ref":
        reference = node[2] if node[2] == node[3] else node[2] + ":" + node[3]
        return _reference_label(node[1] or sheet, reference, tables)
    if kind == "name":
        return "이름 정의 " + short(node[2], 50)
    if kind == "literal":
        return "문자 상수(내용 가림)" if isinstance(node[1], str) else str(node[1])
    if kind == "binary":
        return short(f"({_argument_label(node[2], sheet, tables, depth + 1)}) {node[1]} ({_argument_label(node[3], sheet, tables, depth + 1)})", 200)
    if kind == "unary":
        return short(node[1] + " " + _argument_label(node[2], sheet, tables, depth + 1), 200)
    if kind == "call":
        return node[1] + " 계산 결과"
    return "원본 오류 또는 해석 확인이 필요한 값"


def _argument_roles(node, sheet, tables):
    """Frequent formula roles only; unfamiliar forms keep the existing fallback."""
    if node[0] != "call":
        return None
    name, args = node[1], node[2]
    label = lambda index: _argument_label(args[index], sheet, tables)
    if name == "SUMIFS" and len(args) >= 3 and len(args) % 2 == 1:
        conditions = [f"조건 범위 {label(i)}에 조건 {label(i + 1)} 적용" for i in range(1, min(len(args), 7), 2)]
        extra = f" 외 조건 {(len(args) - 1) // 2 - 3}개" if len(args) > 7 else ""
        return f"SUMIFS — 합계 범위: {label(0)}. " + " / ".join(conditions) + extra + ". 모든 조건을 만족하는 행의 합계 범위 값만 더합니다."
    if name == "SUMIF" and len(args) in (2, 3):
        return f"SUMIF — 조건 범위 {label(0)}에 조건 {label(1)}을 적용하고, 맞는 행의 {label(2) if len(args) == 3 else label(0)} 값을 더합니다."
    if name == "SUM" and args:
        inputs = " / ".join(label(i) for i in range(min(4, len(args))))
        return "SUM — " + inputs + (f" 외 입력 {len(args) - 4}개" if len(args) > 4 else "") + "의 숫자를 합산합니다."
    if name == "IF" and len(args) == 3:
        return f"IF — 조건 {label(0)}을 확인합니다. 참이면 {label(1)}, 거짓이면 {label(2)}를 결과로 사용합니다."
    if name == "IFERROR" and len(args) == 2:
        return f"IFERROR — {label(0)} 계산이 오류인 경우에만 {label(1)}을 대체 결과로 사용합니다."
    if name == "ROUND" and len(args) == 2:
        return f"ROUND — 반올림 대상은 {label(0)}, 자릿수는 {label(1)}입니다. 자릿수가 음수이면 십·백 단위로 반올림합니다."
    return None


def _formula_structure(ast, sheet, tables):
    """Reuse the generator's Korean operation notes, with bounded actual inputs."""
    refs, steps, functions = [], [], []

    def visit(node):
        kind = node[0]
        if kind == "ref":
            reference = node[2] if node[2] == node[3] else node[2] + ":" + node[3]
            refs.append(_reference_entry(node[1] or sheet, reference, tables))
        elif kind == "name":
            refs.append((("name", str(node[1] or sheet).casefold(), str(node[2]).casefold()),
                         "이름 정의 " + short(node[2], 60) + " · 실제 범위는 분석된 참조와 함께 확인"))
        elif kind == "call":
            for child in node[2]:
                visit(child)
            functions.append(node[1])
            steps.append(_argument_roles(node, sheet, tables) or node[1] + " — " + FUNCTION_NOTES.get(node[1], "지원하지 않는 함수입니다. 해당 분기의 계산은 검증되지 않았습니다."))
        elif kind == "binary":
            visit(node[2])
            visit(node[3])
            inputs = (f"입력 {_argument_label(node[2], sheet, tables)} / {_argument_label(node[3], sheet, tables)}. "
                      if node[1] in ("+", "-", "*", "/") else "")
            steps.append(node[1] + " — " + inputs + OPERATOR_NOTES[node[1]])
        elif kind == "unary":
            visit(node[2])
            steps.append({"%": "퍼센트 — 값을 100으로 나눕니다.",
                          "-": "음수 — 값의 부호를 바꿉니다.", "+": "양수 — 숫자 값으로 처리합니다."}[node[1]])
        elif kind == "error":
            steps.append("원본 오류 " + short(node[1], 30) + "가 포함되어 있습니다. 정상 계산으로 처리하지 않습니다.")

    visit(ast)
    if ast[0] in ("ref", "name"):
        steps.append("참조한 값을 결과 위치에 가져옵니다.")
    elif ast[0] == "literal":
        steps.append("수식에 지정된 상수를 결과 위치에 사용합니다. 문자 상수의 내용은 가렸습니다.")
    return list(dict.fromkeys(refs)), list(dict.fromkeys(steps)), sorted(set(functions))


def _row_ranges(column, rows):
    """Do not imply contiguous application when excluded/blank rows make gaps."""
    runs = []
    for row in sorted(set(rows)):
        if runs and row == runs[-1][1] + 1:
            runs[-1][1] = row
        else:
            runs.append([row, row])
    shown = [f"{column}{start}" + (f":{column}{end}" if start != end else "") for start, end in runs[:6]]
    return ", ".join(shown), max(0, len(runs) - 6)


def table_explanations(result, tables):
    """Group existing formula records only; no workbook load, evaluation or LLM."""
    stages = {str(stage["table"]): stage for stage in result.get("pipeline_plan", {}).get("stages", [])}
    by_sheet = defaultdict(list)
    for record in result.get("formulas", []):
        try:
            row, col = coordinate_to_tuple(record["cell"])
        except (TypeError, ValueError, KeyError):
            continue
        by_sheet[record.get("sheet")].append((row, col, record))
    budget = MAX_EXPLAIN_FORMULAS
    output = {}
    for table in tables:
        bounds = _table_bounds(table)
        if not bounds:
            output[str(table["id"])] = {"columns": [], "groups": [], "omitted_columns": 0,
                "omitted_groups": 0, "omitted_formulas": 0, "group_count": 0,
                "formula_count": 0, "row_count": None, "scope_note": "표 범위를 확인하지 못해 열·계산 설명은 미확인입니다."}
            continue
        c1, r1, c2, r2 = bounds
        own = sorted(((row, col, record) for row, col, record in by_sheet[table.get("sheet")]
                      if r1 <= row <= r2 and c1 <= col <= c2), key=lambda entry: (entry[1], entry[0]))
        rows_by_column = defaultdict(set)
        for row, col, _ in own:
            rows_by_column[col].add(row)
        source_rows = set(table.get("source_rows", []))
        planned_columns = stages.get(str(table["id"]), {}).get("formula_columns", [])
        column_details = []
        for index, name in enumerate(table.get("columns", [])[:MAX_DETAIL_COLUMNS]):
            col = c1 + index
            formula_rows = rows_by_column[col]
            data_formula_rows = formula_rows & source_rows
            mode = ("수식 계산 + 원본 입력/빈칸" if data_formula_rows and source_rows - formula_rows
                    else "수식 계산" if data_formula_rows else "원본 입력/빈칸")
            if "source_rows" not in table and formula_rows:
                mode = "수식 있음 · 입력 혼합 여부 미확인"
            elif formula_rows - source_rows:
                mode += " · 데이터행 밖 수식 있음"
            if sum(rule.get("formula_count", 0) for rule in planned_columns if rule.get("column") == name) > len(formula_rows):
                mode = "수식 기록 불완전 · 입력 혼합 여부 미확인"
            column_details.append({"name": short(name, 80), "column": get_column_letter(col),
                                   "mode": mode, "formulas": len(formula_rows)})
        selected = own[:budget]
        budget -= len(selected)
        groups = {}
        for row, col, record in selected:
            formula = record.get("formula", "")
            ast = None
            try:
                if not isinstance(formula, str) or not formula.startswith("=") or len(formula) > 8192:
                    raise FormulaFailure("formula_length_or_format")
                ast = Parser(formula[1:], 2048, 64).parse()
                signature = expression_code(ast, row)
            except (FormulaFailure, RecursionError, ValueError, KeyError, OverflowError):
                ast = None
                # Keep distinct unknown formulas distinct; never group on redacted text.
                signature = ("unparsed", str(formula))
            identity = (col, signature)
            if identity not in groups:
                shown = len(groups) < MAX_FORMULA_GROUPS
                refs, steps, functions = _formula_structure(ast, table.get("sheet", ""), result.get("tables", [])) if ast is not None and shown else ([], [], [])
                for dep in record.get("dependencies", []) if shown else []:
                    if isinstance(dep, dict) and dep.get("sheet") and dep.get("range"):
                        refs.append(_reference_entry(dep["sheet"], dep["range"], result.get("tables", [])))
                # AST keeps $ anchors; analyzed dependencies may omit them. One
                # physical range is one input, retaining the first display form.
                distinct_refs = {}
                for identity_key, ref_label in refs:
                    distinct_refs.setdefault(identity_key, ref_label)
                refs = list(distinct_refs.values())
                index = col - c1
                name = table.get("columns", [])[index] if index < len(table.get("columns", [])) else get_column_letter(col) + "열"
                groups[identity] = {"column": get_column_letter(col), "name": short(name, 80),
                    "representative_cell": short(record.get("cell"), 20), "formula": _safe_formula(formula),
                    "refs": refs[:MAX_FORMULA_REFERENCES], "omitted_refs": max(0, len(refs) - MAX_FORMULA_REFERENCES),
                    "steps": steps[:MAX_FORMULA_STEPS], "omitted_steps": max(0, len(steps) - MAX_FORMULA_STEPS),
                    "unsupported_functions": [f for f in functions if f not in FUNCTIONS],
                    "conditional": bool({"IF", "IFERROR"} & set(functions)),
                    "parsed": ast is not None, "rows": [], "unresolved": 0,
                    "string_constants_hidden": isinstance(formula, str) and '"' in formula}
            group = groups[identity]
            group["rows"].append(row)
            group["unresolved"] += record.get("status") != "computed"
        ordered = sorted(groups.values(), key=lambda g: (coordinate_to_tuple(g["representative_cell"])[1], min(g["rows"])))
        displayed = []
        for group in ordered[:MAX_FORMULA_GROUPS]:
            locations, omitted_ranges = _row_ranges(group["column"], group["rows"])
            displayed.append({key: value for key, value in group.items() if key != "rows"} |
                             {"range": locations, "omitted_ranges": omitted_ranges, "count": len(group["rows"])})
        output[str(table["id"])] = {"columns": column_details, "groups": displayed,
            "omitted_columns": max(0, len(table.get("columns", [])) - MAX_DETAIL_COLUMNS),
            "omitted_groups": max(0, len(ordered) - MAX_FORMULA_GROUPS),
            "omitted_formulas": len(own) - len(selected), "group_count": len(ordered),
            "formula_count": len(own), "row_count": len(source_rows) if "source_rows" in table else None,
            "scope_note": "수식 원문에서 확인한 동작이며 실제 데이터 값과 업무 의미를 확정한 설명은 아닙니다."}
    return output


def _components(nodes, edges):
    """Small bounded graph: Tarjan SCC, then dependency layers on its DAG."""
    adjacency = {n["key"]: [] for n in nodes}
    for source, target in edges:
        adjacency[source].append(target)
    index, stack, on_stack, indices, low, groups = 0, [], set(), {}, {}, []

    def visit(v):
        nonlocal index
        indices[v] = low[v] = index
        index += 1
        stack.append(v)
        on_stack.add(v)
        for w in adjacency[v]:
            if w not in indices:
                visit(w)
                low[v] = min(low[v], low[w])
            elif w in on_stack:
                low[v] = min(low[v], indices[w])
        if low[v] == indices[v]:
            group = []
            while True:
                w = stack.pop()
                on_stack.remove(w)
                group.append(w)
                if w == v:
                    break
            groups.append(group)

    for v in adjacency:
        if v not in indices:
            visit(v)
    owner = {v: i for i, group in enumerate(groups) for v in group}
    predecessors = {i: set() for i in range(len(groups))}
    for source, target in edges:
        if owner[source] != owner[target]:
            predecessors[owner[target]].add(owner[source])
    ranks = {}

    def rank(i):
        if i not in ranks:
            ranks[i] = max((rank(p) + 1 for p in predecessors[i]), default=0)
        return ranks[i]

    cyclic = {v for group in groups if len(group) > 1 for v in group}
    cyclic.update(s for s, t in edges if s == t)
    return {v: rank(owner[v]) for v in adjacency}, cyclic


def build_workflow(result):
    """Use only recognized table references; missing evidence stays visible."""
    tables = result.get("tables", [])
    explanations = table_explanations(result, tables[:MAX_NODES])
    stages = {s["table"]: s for s in result.get("pipeline_plan", {}).get("stages", [])}
    nodes, lookup, external, edges, warnings = [], {}, {}, set(), []
    represented_formulas = sum(r.get("formula_count", 0) for stage in stages.values()
                               for r in stage.get("formula_columns", []))
    unrepresented_formulas = max(0, len(result.get("formulas", [])) - represented_formulas)
    if unrepresented_formulas:
        warnings.append(f"표 처리 기록에 포함되지 않은 수식 {unrepresented_formulas}개가 있습니다. 그 계산 흐름은 이 표 연결도로 확인하지 않았습니다.")
    for table in tables[:MAX_NODES]:
        sheet = table.get("sheet", "")
        original = str(table["id"])
        # Reuse the decomposed table's identifier; the diagram adds no ID system.
        alias = original
        stage = stages.get(original, {})
        name, basis = table_name(table)
        rules = stage.get("formula_columns", [])
        formulas = sum(r.get("formula_count", 0) for r in rules)
        unresolved = sum(r.get("unresolved", 0) for r in rules)
        node = {"key": f"n{len(nodes)}", "alias": alias, "id": original, "name": name,
                "basis": basis, "sheet": short(sheet, 80), "range": short(table.get("range"), 60),
                "columns": [short(c, 28) for c in table.get("columns", [])[:6]],
                "omitted_columns": max(0, len(table.get("columns", [])) - 6),
                "formulas": formulas, "unresolved": unresolved,
                "functions": sorted({str(f) for r in rules for f in r.get("functions", [])})[:8],
                "role": "unknown" if not stage else "calculation" if formulas else "input", "outside": False,
                "missing_plan": not stage, "explanation": explanations[original]}
        nodes.append(node)
        lookup[original] = node["key"]
        if unresolved:
            warnings.append(f"{alias}: 미지원·오류 수식 {unresolved}개가 있습니다.")
        if not stage:
            warnings.append(f"{alias}: 참조 분석 기록이 없어 연결을 확정하지 않았습니다.")
    duplicate_names = Counter((n["sheet"], n["name"]) for n in nodes)
    for node in nodes:
        if duplicate_names[(node["sheet"], node["name"])] > 1:
            meaningful_columns = [c for c in node["columns"]
                                  if not re.fullmatch(r"(?:column|열|col)[ _]?(?:\d+|[A-Z]{1,3})", c, re.I)]
            qualifier = " · ".join(meaningful_columns[:2]) or node["range"]
            node["name"] = node["name"].removesuffix(" 테이블") + " · " + short(qualifier, 42) + " 테이블"
            node["basis"] += " · 같은 시트의 표는 열 이름/범위로 구분"
    omitted_references = 0
    for table in tables[:MAX_NODES]:
        target = lookup[str(table["id"])]
        stage = stages.get(str(table["id"]), {})
        for dep in stage.get("depends_on_tables", []):
            if str(dep) in lookup:
                edges.add((lookup[str(dep)], target))
            else:
                omitted_references += 1
        for dep in stage.get("references_outside_detected_tables", []):
            if not isinstance(dep, dict) or not dep.get("sheet") or not dep.get("range"):
                omitted_references += 1
                continue
            identity = (str(dep["sheet"]), str(dep["range"]))
            if identity not in external:
                if len(nodes) >= MAX_NODES:
                    omitted_references += 1
                    continue
                node = {"key": f"n{len(nodes)}", "alias": f"표 밖 참조 {len(external) + 1}",
                        "id": "", "name": "표 밖의 참조 범위", "basis": "수식 참조에서 확인",
                        "sheet": short(identity[0], 80), "range": short(identity[1], 60),
                        "columns": [], "omitted_columns": 0, "formulas": 0, "unresolved": 0,
                        "functions": [], "role": "reference", "outside": True, "missing_plan": False,
                        "explanation": None}
                external[identity] = node["key"]
                nodes.append(node)
            edges.add((external[identity], target))
    all_edges = sorted(edges)
    edges = all_edges[:MAX_EDGES]
    # The display edge cap must not turn a processing table into a false endpoint.
    ranks, cyclic = _components(nodes, all_edges)
    outgoing = {source for source, _ in all_edges}
    outgoing_table_ids = {str(dep) for stage in stages.values()
                          for dep in stage.get("depends_on_tables", [])}
    for node in nodes:
        node["rank"] = ranks[node["key"]]
        node["intertwined"] = node["key"] in cyclic
        if (node["role"] == "calculation" and node["key"] not in outgoing
                and node["id"] not in outgoing_table_ids and not omitted_references):
            node["role"] = "output"
    if cyclic:
        warnings.append("표 간 상호 참조가 있습니다. 표 수준 연결만으로 셀 순환 오류나 실행 순서를 단정하지 않습니다.")
    vba = result.get("summary", {}).get("vba", {})
    if vba and vba.get("status") != "no_vba":
        warnings.append("VBA 동작은 이 연결도에 포함하지 않았습니다. 추출 여부와 무관하게 매크로 실행 흐름은 별도 검토가 필요합니다.")
    if result.get("summary", {}).get("non_cell_features"):
        warnings.append("외부 연결·비셀 기능의 동작은 이 연결도에서 확인하지 않았습니다.")
    omitted_tables = max(0, len(tables) - min(len(tables), MAX_NODES))
    if omitted_tables or omitted_references or len(all_edges) > MAX_EDGES:
        warnings.append(f"표시 한도: 표 {omitted_tables}개, 참조 {omitted_references}건, 연결 {max(0, len(all_edges) - MAX_EDGES)}건을 생략했습니다. 전체 연결도라고 해석하지 마세요.")
    return {"nodes": nodes, "edges": edges, "warnings": warnings,
            "omitted_tables": omitted_tables, "omitted_edges": max(0, len(all_edges) - MAX_EDGES),
            "source": short(result.get("summary", {}).get("source_name", "Excel 분석"), 100),
            "selection": [short(s, 90) for s in result.get("selection", {}).get("targets", [])[:12]],
            "formula_count": len(result.get("formulas", [])), "table_count": len(tables)}


CSS = """
:root{color-scheme:light;--ink:#172a3a;--muted:#647384;--line:#d9e2e8;--blue:#176b9d;--teal:#217a70;--warm:#b06a17}
*{box-sizing:border-box}body{margin:0;background:#f3f6f8;color:var(--ink);font-family:'Noto Sans KR',sans-serif;font-size:15px;line-height:1.75}
main{max-width:1440px;margin:auto;padding:48px 36px 64px}a{color:#176b9d;text-decoration-thickness:1px;text-underline-offset:3px}
.eyebrow{color:var(--teal);font-size:12px;font-weight:700;letter-spacing:.14em}h1{font-size:34px;line-height:1.35;margin:10px 0 15px;letter-spacing:-.04em}h2{font-size:22px;margin:0 0 8px;letter-spacing:-.03em}p{margin:8px 0}header{max-width:950px;margin-bottom:28px}.muted{color:var(--muted)}
.summary{display:flex;gap:12px;flex-wrap:wrap;margin:24px 0}.stat{min-width:150px;border:1px solid var(--line);border-radius:14px;background:white;padding:15px 22px}.stat strong{display:block;font-size:28px;line-height:1.4}.stat span{color:var(--muted);font-size:13px}
.panel{background:white;border:1px solid var(--line);border-radius:18px;margin-top:22px;padding:26px;box-shadow:0 4px 18px #24425b05}.legend{display:flex;gap:20px;flex-wrap:wrap;font-size:13px;margin:16px 0}.dot{display:inline-block;width:9px;height:9px;border-radius:50%;margin-right:6px;background:var(--teal)}.dot.calc{background:var(--blue)}.dot.out{background:#7654a4}.dot.ref{background:#7b838b}
.scroll{overflow:auto;max-height:780px;border:1px solid #e8edf1;border-radius:12px;background:linear-gradient(#edf2f6 1px,transparent 1px),linear-gradient(90deg,#edf2f6 1px,transparent 1px);background-size:24px 24px}.canvas{position:relative}.canvas svg{position:absolute;inset:0;overflow:visible}.lane{position:absolute;top:16px;font-size:12px;color:var(--muted);font-weight:600}.empty-graph{padding:28px;min-height:150px;background:#f5f8fa;border-radius:12px}.scroll-hint{font-size:12px;color:var(--muted);margin:9px 0}
.node{position:absolute;width:248px;height:182px;padding:15px 17px;background:white;border:1px solid #cad8e0;border-top:4px solid var(--teal);border-radius:11px;box-shadow:0 4px 10px #1639510b;overflow:hidden}.node.calculation{border-top-color:var(--blue)}.node.output{border-top-color:#7654a4}.node.reference{border-top-color:#7b838b;border-style:dashed}.node.cycle{outline:2px dashed #b8842b;outline-offset:3px}.node small{font-size:11px;color:var(--muted)}.node h3{font-size:16px;line-height:1.45;margin:6px 0;display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden;word-break:break-word}.node p{font-size:12px;margin:3px 0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.alias{font-family:inherit;letter-spacing:.04em;font-size:11px;font-weight:700;color:var(--teal)}
.note{background:#f3f7fa;border-left:3px solid #96afc2;padding:12px 16px;font-size:13px;margin:16px 0}.warning{background:#fff8ed;border-left-color:#c18c37}.list{display:grid;grid-template-columns:repeat(auto-fit,minmax(300px,1fr));gap:14px;margin-top:20px}.detail{border:1px solid var(--line);border-radius:12px;padding:18px;overflow-wrap:anywhere}.detail h3{font-size:17px;line-height:1.5;margin:4px 0 10px}.detail p{font-size:13px}.detail strong{font-weight:600}.chips{display:flex;gap:6px;flex-wrap:wrap}.chip{font-size:11px;border-radius:5px;background:#eef4f7;padding:2px 7px}.edge-table{border-collapse:collapse;width:100%;font-size:13px;min-width:480px}.edge-table th,.edge-table td{text-align:left;border-bottom:1px solid var(--line);padding:10px 12px}.edge-table th{color:var(--muted);font-weight:600}.foot{font-size:12px;color:var(--muted);margin-top:24px}details summary{cursor:pointer;font-size:14px;font-weight:600}details[open] summary{margin-bottom:12px}code{background:#eef3f6;padding:1px 4px;border-radius:3px;font-size:12px}
@media(max-width:640px){main{padding:28px 16px}h1{font-size:28px}.panel{padding:18px}.stat{min-width:120px;flex:1}.list{grid-template-columns:1fr}}@media print{body{background:white}main{padding:0}.panel{box-shadow:none;break-inside:avoid}.scroll{overflow:visible}.graph-panel{display:none}.list{grid-template-columns:1fr 1fr}.detail{break-inside:avoid}details{display:block}}
"""

CSS += """
.list{grid-template-columns:repeat(2,minmax(0,1fr));align-items:start}.detail{min-width:0}.table-explanation{border-top:1px solid var(--line);margin-top:16px;padding-top:13px}.table-explanation h4{font-size:14px;margin:8px 0}.table-explanation details{margin:12px 0}.column-table{width:100%;table-layout:fixed;border-collapse:collapse;font-size:12px;line-height:1.65}.column-table th,.column-table td{text-align:left;padding:7px 5px;border-bottom:1px solid #e4ebf0;vertical-align:top;overflow-wrap:anywhere}.column-table th:first-child{width:46%}.column-table th:nth-child(2){width:12%}.column-table thead{color:var(--muted)}.formula-details{background:#f6f9fb;border:1px solid #e0e8ee;border-radius:9px;padding:12px}.formula-details summary{font-size:13px;line-height:1.65}.formula-code{white-space:pre-wrap;overflow-wrap:anywhere;word-break:break-word;margin:10px 0;background:white;border:1px solid var(--line);border-radius:6px;padding:11px;font-family:'Noto Sans KR',sans-serif;font-size:12px;line-height:1.7}.formula-steps{padding-left:20px;margin:8px 0;font-size:12px}.formula-steps li{margin:5px 0}.formula-refs{font-size:12px;line-height:1.8;padding-left:18px}.explanation-limit{font-size:12px!important;color:#8a5b16;background:#fff7e8;padding:7px 10px;border-radius:6px}.calc-description{font-size:13px;color:#30576f;background:#eff6fa;border-radius:6px;padding:10px 12px}.table-explanation .muted{font-size:12px}
@media(max-width:900px){.list{grid-template-columns:1fr}}@media print{.list{grid-template-columns:1fr}.table-explanation{break-inside:avoid}}
"""


def _render_explanation(node):
    info = node.get("explanation")
    if info is None:
        return '<section class="table-explanation"><p>표 밖의 참조 범위입니다. 열 정의와 내부 계산 구조는 이 표 목록에서 확인하지 않았습니다.</p></section>'
    e = lambda value: escape(str(value), quote=True)
    row_note = f'분해된 데이터 {info["row_count"]:,}행 · ' if info["row_count"] is not None else ''
    contents = ['<section class="table-explanation"><h4>어떤 데이터를 다루나요?</h4>',
                f'<p>{row_note}실제 데이터 값은 싣지 않고 열 이름·계산 위치만 표시합니다.</p>']
    if info["columns"]:
        entries = ''.join(f'<tr><td>{e(c["name"])}</td><td>{e(c["column"])}열</td><td>{e(c["mode"])}' +
                          (f' · 수식 {c["formulas"]:,}개' if c["formulas"] else '') + '</td></tr>' for c in info["columns"])
        contents.append(f'<details class="column-details"><summary>데이터 항목과 원본 열 보기 ({len(info["columns"]) + info["omitted_columns"]}개)</summary><table class="column-table"><thead><tr><th>데이터 항목</th><th>열</th><th>값의 구성</th></tr></thead><tbody>{entries}</tbody></table><p class="muted">입력/빈칸 여부는 셀 값을 읽어 구분하지 않았습니다. 수식 위치는 헤더·합계·제외 행을 포함할 수 있으므로 아래 실제 범위를 확인하세요.</p></details>')
    else:
        contents.append('<p class="muted">확인된 열 목록이 없습니다.</p>')
    if info["omitted_columns"]:
        contents.append(f'<p class="explanation-limit">열 표시 한도: {info["omitted_columns"]}개를 생략했습니다.</p>')
    if info["formula_count"]:
        contents.append(f'<h4>어떤 계산으로 결과를 만드나요?</h4><p class="calc-description">수식 {info["formula_count"]:,}개에서 반복 계산 {info["group_count"]:,}묶음을 확인했습니다. 각 묶음은 아래 결과 열·범위에 값을 만듭니다.</p>')
        for index, group in enumerate(info["groups"]):
            location = node["sheet"] + '!' + group["range"]
            contents.append(f'<details class="formula-details"{" open" if index == 0 else ""}><summary>{e(group["name"])} · {e(group["range"])} · {group["count"]:,}개 셀</summary><p><strong>만드는 결과</strong> {e(group["name"])} ({e(group["column"])}열), {e(location)}</p><p class="muted">대표식은 {e(group["representative_cell"])} 기준입니다. 상대 행은 적용 행에 맞게 이동하고 $로 고정한 행은 유지합니다.</p><pre class="formula-code">{e(group["formula"])}</pre>')
            if group["string_constants_hidden"]:
                contents.append('<p class="muted">조건·반환값의 문자 상수는 가렸습니다. 서로 다른 조건은 별도 계산 묶음으로 유지합니다.</p>')
            if group["refs"]:
                contents.append('<p><strong>사용하는 입력 · 대표식 기준</strong></p><ul class="formula-refs">' + ''.join(f'<li>{e(ref)}</li>' for ref in group["refs"]) + '</ul>')
            else:
                contents.append('<p class="muted">셀 입력 참조 없음 또는 해석 미확인 — 상수식·미지원 참조 여부를 확인하세요.</p>')
            if group["steps"]:
                contents.append('<p><strong>계산 동작 읽기</strong></p><ol class="formula-steps">' + ''.join(f'<li>{e(step)}</li>' for step in group["steps"]) + '</ol>')
            if not group["parsed"] or group["unresolved"] or group["unsupported_functions"]:
                notes = []
                if not group["parsed"]:
                    notes.append('대표식을 자동 해석하지 못했습니다')
                if group["unresolved"]:
                    notes.append(f'미지원·오류·미확인 {group["unresolved"]}개 셀')
                if group["unsupported_functions"]:
                    notes.append('미지원 함수: ' + ', '.join(group["unsupported_functions"]))
                contents.append('<p class="explanation-limit">검토 필요 — ' + e(' · '.join(notes)) + '. 이 설명만으로 계산 성공·VBA 변환 완료를 판단하지 마세요.</p>')
            if group["omitted_refs"] or group["omitted_steps"] or group["omitted_ranges"]:
                contents.append(f'<p class="explanation-limit">설명 한도: 입력 참조 {group["omitted_refs"]}개, 계산 설명 {group["omitted_steps"]}개, 적용 구간 {group["omitted_ranges"]}개를 생략했습니다.</p>')
            contents.append('</details>')
        if any(group["conditional"] for group in info["groups"]):
            contents.append('<p class="muted">IF·IFERROR는 조건에 따라 필요한 분기만 계산하므로 모든 분기가 실행된다는 뜻은 아닙니다.</p>')
        contents.append('<p class="muted">계산 동작은 수식의 구조 설명입니다. 분석기의 계산 기록은 실제 Excel 실행 결과와 동일하다는 검증을 대신하지 않습니다.</p>')
    else:
        message = ('계산 분석 기록이 없어 표의 처리 방식은 미확인입니다.' if node["missing_plan"]
                   else '처리 계획에는 수식이 있지만 원문 기록이 없어 계산 설명은 미확인입니다.' if node["formulas"]
                   else '이 표 안에서 확인된 셀 수식은 없습니다. 원본 입력을 다른 표에서 참조할 수 있으며, VBA나 외부 처리까지 없다는 뜻은 아닙니다.')
        contents.append('<p class="calc-description">' + message + '</p>')
    if info["omitted_groups"] or info["omitted_formulas"]:
        contents.append(f'<p class="explanation-limit">설명 한도: 계산 묶음 {info["omitted_groups"]}개, 분석 예산 밖 수식 {info["omitted_formulas"]}개를 생략했습니다. 표시된 묶음만으로 전체 계산을 설명하지 않습니다.</p>')
    if node["formulas"] > info["formula_count"]:
        contents.append('<p class="explanation-limit">처리 계획에만 있고 원본 수식 기록에는 없는 계산이 있습니다. 해당 계산 설명은 미확인입니다.</p>')
    contents.append('<p class="muted">' + e(info["scope_note"]) + '</p></section>')
    return ''.join(contents)


def render_workflow(result, font_path=None):
    model = build_workflow(result)
    font_path = Path(font_path) if font_path else Path(__file__).resolve().parent.parent / "assets/NotoSansKR-Variable.woff"
    # A missing bundled font is a packaging error, not a reason to contact a CDN.
    font = font_path.read_bytes()
    if not font.startswith(b"wOFF") or len(font) > 12_000_000:
        raise ValueError("승인된 Noto Sans KR WOFF 자산을 확인하세요.")
    font_css = "@font-face{font-family:'Noto Sans KR';font-style:normal;font-weight:100 900;font-display:swap;src:url(data:font/woff;base64," + base64.b64encode(font).decode("ascii") + ") format('woff')}"
    license_text = font_path.with_name("OFL.txt").read_text(encoding="utf-8")
    e = lambda value: escape(str(value), quote=True)
    nodes = model["nodes"]
    lookup = {n["key"]: n for n in nodes}
    connected_keys = {key for edge in model["edges"] for key in edge}
    graph_nodes = [node for node in nodes if node["key"] in connected_keys]
    disconnected_count = len(nodes) - len(graph_nodes)
    lanes = defaultdict(list)
    for node in graph_nodes:
        lanes[node["rank"]].append(node)
    width = max(1, len(lanes)) * 340 + 24
    height = max([1, *(len(group) for group in lanes.values())]) * 216 + 64
    positions = {n["key"]: (24 + rank * 340, 52 + i * 216)
                 for rank, group in lanes.items() for i, n in enumerate(group)}
    paths = []
    for source, target in model["edges"]:
        x1, y1 = positions[source]
        x2, y2 = positions[target]
        if x1 == x2:
            # Intra-component links use the free right gutter; do not imply a linear order.
            gutter = x1 + 278
            path = f"M{x1 + 248},{y1 + 100} C{gutter},{y1 + 100} {gutter},{y2 + 110} {x2 + 248},{y2 + 110}"
        else:
            middle = (x1 + 248 + x2) / 2
            path = f"M{x1 + 248},{y1 + 91} C{middle},{y1 + 91} {middle},{y2 + 91} {x2},{y2 + 91}"
        paths.append(f'<path d="{path}" fill="none" stroke="#89a5b8" stroke-width="1.6" marker-end="url(#arrow)"/>')
    cards, details, connections = [], [], []
    role_labels = {"input": "입력 표", "calculation": "수식 처리", "output": "연결의 마지막 계산 표", "reference": "표 밖의 참조", "unknown": "분석 기록 없음"}
    for n in nodes:
        title = f'{n["name"]} ({n["alias"]})'
        location = f'{n["sheet"]}!{n["range"]}'
        if n["key"] in connected_keys:
            x, y = positions[n["key"]]
            cards.append(f'<a class="node {n["role"]}{" cycle" if n["intertwined"] else ""}" href="#{n["key"]}" style="left:{x}px;top:{y}px" title="{e(title)}"><span class="alias">{e(n["alias"])} · {role_labels[n["role"]]}</span><h3>{e(n["name"])}</h3><p>{e(location)}</p><p>수식 {n["formulas"]:,}개 · 검토 {n["unresolved"]:,}개</p><small>상세 정보 보기</small></a>')
        inputs = [lookup[s]["alias"] for s, t in model["edges"] if t == n["key"]]
        output_aliases = [lookup[t]["alias"] for s, t in model["edges"] if s == n["key"]]
        chips = ''.join(f'<span class="chip">{e(c)}</span>' for c in n["columns"])
        extra = f' 외 {n["omitted_columns"]}개 열' if n["omitted_columns"] else ''
        details.append(f'<article class="detail" id="{n["key"]}"><span class="alias">{e(n["alias"])} · {role_labels[n["role"]]}</span><h3>{e(title)}</h3><p><strong>원본 위치</strong> {e(location)}</p><p><strong>기존 표 ID</strong> {e(n["id"] or "표 범위 밖")}</p><p><strong>이름 근거</strong> {e(n["basis"])}</p><div class="chips">{chips}</div><p>{e(extra)}</p><p><strong>먼저 참고</strong> {e(", ".join(inputs) or "확인된 표 간 참조 없음")}</p><p><strong>다음으로 연결</strong> {e(", ".join(output_aliases) or "확인된 후속 표 없음")}</p><p><strong>주요 함수</strong> {e(", ".join(n["functions"]) or "확인된 함수 없음")}</p><p>수식 {n["formulas"]:,}개 · 미지원/오류 {n["unresolved"]:,}개</p>{"<p>상호 참조 그룹: 단일 실행 순서로 해석하지 마세요.</p>" if n["intertwined"] else ""}{_render_explanation(n)}</article>')
    for source, target in model["edges"]:
        a, b = lookup[source], lookup[target]
        connections.append(f'<tr><td><a href="#{source}">{e(a["alias"])} · {e(a["name"])}</a></td><td><a href="#{target}">{e(b["alias"])} · {e(b["name"])}</a></td><td>원본 수식의 참조 범위</td></tr>')
    lane_headers = ''.join(f'<div class="lane" style="left:{24 + rank * 340}px">참조 깊이 {rank + 1}{" · 시작 범위" if rank == 0 else ""}</div>' for rank in sorted(lanes))
    graph = (f'<div class="scroll" tabindex="0" aria-label="표 참조 연결도. 넓은 경우 좌우, 긴 경우 위아래로 이동하세요."><div class="canvas" style="width:{width}px;height:{height}px"><svg width="{width}" height="{height}" aria-hidden="true"><defs><marker id="arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#89a5b8"/></marker></defs>{"".join(paths)}</svg>{lane_headers}{"".join(cards)}</div></div><p class="scroll-hint">연결도가 넓거나 길면 이 영역 안에서 좌우·위아래로 이동하세요. 같은 연결은 아래 연결 근거 목록에서도 확인할 수 있습니다.</p>'
             if graph_nodes else '<div class="empty-graph"><strong>표 간 연결을 확인하지 못했습니다.</strong><p>독립된 입력이거나 참조 분석이 불완전한 경우일 수 있습니다. 임의의 순서를 그리지 않았습니다. 아래 표 목록과 확인 사항을 살펴보세요.</p></div>')
    if disconnected_count:
        graph += f'<p class="note">표시할 연결이 없는 표·참조 {disconnected_count}개는 아래 표 목록에 별도로 정리했습니다. 연결되지 않았다는 사실만으로 불필요한 데이터라고 판단하지 않습니다.</p>'
    warning_html = ''.join(f'<p class="note warning">{e(w)}</p>' for w in model["warnings"][:4])
    if len(model["warnings"]) > 4:
        warning_html += '<details class="note warning"><summary>추가 확인 사항 ' + str(len(model["warnings"]) - 4) + '건</summary>' + ''.join(f'<p>{e(w)}</p>' for w in model["warnings"][4:]) + '</details>'
    selection = f'<p class="note">선택한 결과: {e(", ".join(model["selection"]))}. 현재 남겨진 표와 선행 참조만 표시합니다.</p>' if model["selection"] else ''
    csp = "default-src 'none'; style-src 'unsafe-inline'; font-src data:; img-src 'none'; script-src 'none'; connect-src 'none'; base-uri 'none'; form-action 'none'"
    return f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta http-equiv="Content-Security-Policy" content="{csp}"><title>엑셀 업무 흐름 · {e(model['source'])}</title><style>{font_css}{CSS}</style></head><body><main>
<header><div class="eyebrow">EXCEL WORKFLOW</div><h1>데이터가 결과로 이어지는 흐름</h1><p>{e(model['source'])}</p><p class="muted">어떤 표를 먼저 읽고, 어떤 계산을 거쳐 결과를 만드는지 원본 수식의 연결로 살펴봅니다. 원본 파일과 기존 표 ID는 바꾸지 않습니다.</p></header>
<div class="summary"><div class="stat"><strong>{model['table_count']:,}</strong><span>분석한 표</span></div><div class="stat"><strong>{model['formula_count']:,}</strong><span>분석한 수식</span></div><div class="stat"><strong>{len(model['edges']):,}</strong><span>표시한 참조 연결</span></div></div>{selection}{warning_html}
<section class="panel graph-panel"><h2>입력에서 계산 결과까지</h2><p class="muted">화살표는 앞의 데이터를 뒤의 표가 참고한다는 뜻입니다. 카드를 누르면 이름 근거와 참조 범위를 볼 수 있습니다.</p><div class="legend"><span><i class="dot"></i>입력 표</span><span><i class="dot calc"></i>계산</span><span><i class="dot out"></i>마지막 계산 표</span><span><i class="dot ref"></i>표 밖 참조</span></div>{graph}<p class="note">이 그림은 수식에서 확인한 데이터 의존성입니다. 실제 업무 순서·VBA 실행 순서와 같다는 뜻은 아닙니다. 표 안의 셀 참조는 요약되어 있으며 상세 계산은 PROCESS의 한국어 주석을 확인하세요.</p></section>
<section class="panel"><h2>표별 데이터와 계산 설명</h2><p class="muted">각 표가 담는 항목 → 사용하는 입력 → 계산 동작 → 결과 열을 순서대로 읽어 보세요. 이름은 시트·열 이름에서 해석한 안내용 이름입니다. 원본 데이터 행과 계산된 값은 싣지 않습니다. 시트·열 이름에도 민감 정보가 있을 수 있으므로 원본과 같은 범위에서 보관하세요.</p><div class="list">{''.join(details) or '<p>표로 식별된 영역이 없습니다. 범위와 헤더를 먼저 확인하세요.</p>'}</div></section>
<section class="panel"><details><summary>연결 근거 목록 ({len(connections)}건)</summary><div style="overflow:auto"><table class="edge-table"><thead><tr><th>참고하는 원본</th><th>이를 사용하는 표</th><th>근거</th></tr></thead><tbody>{''.join(connections) or '<tr><td colspan="3">확인된 표 간 연결이 없습니다. 서로 독립된 입력일 수도, 참조를 분석하지 못한 경우일 수도 있습니다.</td></tr>'}</tbody></table></div></details></section>
<p class="foot">생성 시점의 분석 구조입니다. Python 계산을 수정했다면 원본 수식 기반 연결과 달라질 수 있습니다. DATA는 입력, PROCESS는 계산, main.py는 실행 시작점입니다. 이 HTML은 네트워크·스크립트 없이 열립니다. Noto Sans KR 글꼴은 SIL Open Font License에 따라 내장했습니다.</p>
<details class="foot"><summary>글꼴 사용 허가</summary><pre style="white-space:pre-wrap">{e(license_text)}</pre></details></main></body></html>'''
