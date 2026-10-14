"""Change detection for an explicitly reviewed migration, not a trust signature.

Create a contract only while authoring/reviewing a Python translation. Never
rebuild it automatically to make a mismatching input pass at execution time.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import openpyxl

from excel_pipeline import InputError, preflight
from vba_extract import extract_vba


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def module_identity(extraction):
    if extraction.get("status") != "extracted" or not extraction.get("modules"):
        raise InputError("내장 VBA 원문을 확보하지 못해 검토된 Python 처리를 실행하지 않았습니다: " +
                         extraction.get("status", "unknown"))
    # Include every module, even an empty document class: additions/deletions
    # and event source changes must not be hidden by an entrypoint allowlist.
    rows = []
    for module in extraction["modules"]:
        source = module["source"].replace("\r\n", "\n").replace("\r", "\n")
        rows.append({"name": module["name"], "path": module["relative_path"],
                     "module_type": module.get("module_type", "unknown"),
                     "source_sha256": hashlib.sha256(source.encode("utf-8")).hexdigest()})
    return sorted(rows, key=lambda m: (m["name"], m["path"]))


def observe_contract(source, header_ranges, extraction=None):
    source = Path(source).resolve()
    preflight(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    extraction = extraction if extraction is not None else extract_vba(source)
    if extraction.get("container_sha256") != before:
        raise InputError("VBA 추출 파일과 현재 입력이 달라 처리를 중단했습니다.")
    modules = module_identity(extraction)
    workbook = openpyxl.load_workbook(source, data_only=False, keep_links=False)
    try:
        formulas = []
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        if not isinstance(cell.value, str):
                            raise InputError("배열·특수 수식은 계약 작성 시 별도 검토가 필요합니다.")
                        formulas.append([sheet.title, cell.coordinate, cell.value])
        headers = []
        for spec in header_ranges:
            if set(spec) != {"sheet", "range"} or spec["sheet"] not in workbook.sheetnames:
                raise InputError("검토된 입력 표의 헤더 범위가 없습니다.")
            from excel_pipeline import bounded_range, safe_value
            bounded_range(spec["range"], "헤더", max_area=1024)
            cells = workbook[spec["sheet"]][spec["range"]]
            if not isinstance(cells, tuple):
                cells = ((cells,),)
            headers.append({**spec, "values": [[safe_value(c.value) for c in row] for row in cells]})
        names = sorted([["workbook", name, definition.attr_text] for name, definition in workbook.defined_names.items()] +
                       [["sheet:" + sheet.title, name, definition.attr_text]
                        for sheet in workbook for name, definition in sheet.defined_names.items()])
        identity = {"version": 1, "modules": modules, "sheets": workbook.sheetnames,
                    "formula_count": len(formulas), "formula_sha256": digest(sorted(formulas)),
                    "headers": headers, "defined_names": names}
    finally:
        workbook.close()
    if hashlib.sha256(source.read_bytes()).hexdigest() != before:
        raise InputError("확인 중 입력이 변경되어 처리를 중단했습니다.")
    return identity


def verify_contract(source, contract_path):
    contract_path = Path(contract_path)
    if contract_path.stat().st_size > 1024 * 1024:
        raise InputError("이전 계약 크기가 한도를 넘습니다.")
    expected = json.loads(contract_path.read_text(encoding="utf-8"))
    if expected.get("version") != 1 or not isinstance(expected.get("headers"), list) or len(expected["headers"]) > 128:
        raise InputError("검토된 이전 계약 형식이 아닙니다.")
    headers = [{"sheet": h["sheet"], "range": h["range"]} for h in expected["headers"]]
    observed = observe_contract(source, headers)
    differences = [key for key in observed if expected.get(key) != observed[key]]
    if differences:
        raise InputError("VBA·수식·입력 구조가 검토본과 달라 기존 Python 처리를 실행하지 않았습니다. "
                         "변경 부분만 다시 분석해 주세요: " + ", ".join(differences))
    return {"status": "matched", "contract_sha256": hashlib.sha256(contract_path.read_bytes()).hexdigest(),
            "module_count": len(observed["modules"]), "formula_count": observed["formula_count"],
            "note": "원문 변경 감지이며 안전성 서명이나 컴파일 코드와의 일치 검증이 아닙니다."}
