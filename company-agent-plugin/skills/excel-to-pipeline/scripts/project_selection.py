"""Static, conservative output slicing. No source deletion or VBA execution."""
import copy
from collections import deque

from openpyxl.utils.cell import range_boundaries, get_column_letter
from formula_engine import FormulaEngine, FormulaFailure, FUNCTIONS


def addresses(reference):
    left, top, right, bottom = range_boundaries(reference.replace("$", ""))
    if (None in (left, top, right, bottom) or right < left or bottom < top
            or left < 1 or top < 1 or right > 16384 or bottom > 1048576):
        raise ValueError("결과 범위는 A1:D20처럼 유한한 셀 범위로 지정하세요.")
    if (bottom - top + 1) * (right - left + 1) > 10000:
        raise ValueError("결과/참조 범위는 한 번에 10,000셀까지 지원합니다.")
    for row in range(top, bottom + 1):
        for col in range(left, right + 1):
            yield f"{get_column_letter(col)}{row}"


def dependencies(engine, expression, sheet):
    """Both IF branches are retained; unresolved references never mean no deps."""
    result = set()

    def visit(ast):
        kind = ast[0]
        if kind == "call":
            if ast[1] not in FUNCTIONS:
                raise ValueError(f"미지원 함수 {ast[1]}의 의존성을 확정할 수 없습니다.")
            for arg in ast[2]:
                visit(arg)
        elif kind == "name":
            scope = engine._sheet(ast[1] or sheet)
            visit(engine._name_node(ast[2], scope))
        elif kind == "ref":
            scope = engine._sheet(ast[1] or sheet)
            for address in addresses(ast[2] + ":" + ast[3]):
                result.add((scope, address))
        elif kind == "binary":
            visit(ast[2])
            visit(ast[3])
        elif kind == "unary":
            visit(ast[2])
        elif kind not in ("literal", "error"):
            raise ValueError("동적 참조의 의존성을 확정할 수 없습니다.")
    try:
        visit(engine._parse(expression))
    except (FormulaFailure, RecursionError) as error:
        raise ValueError(f"{sheet}: 참조 범위를 확정하지 못해 축소하지 않았습니다 ({error}).") from error
    return result


def workbook_from_result(result):
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in result["sheets"]:
        wb.create_sheet(sheet["name"])
    for cell in result["source_cells"]:
        target = wb[cell["sheet"]][cell["cell"]]
        target.value = cell["formula"] or cell["value"]
        target.data_type = "f" if cell["formula"] else cell["type"]
    for item in result["workbook_structure"]["names"]:
        owner = wb if item["sheet"] is None else wb[item["sheet"]]
        owner.defined_names.add(DefinedName(item["name"], attr_text=item["text"]))
    return wb


def select_result(result, targets):
    """Build a NEW analysis containing only requested outputs + literal/formula deps."""
    from project_builder import make_rules
    from project_runtime import calculate_data
    if not targets:
        return result
    vba = result["summary"].get("vba", {})
    if vba and vba.get("status") != "no_vba":
        raise ValueError("VBA 동작이 결과에 미치는 범위를 확정하지 못했습니다. 매크로 이전 검토 전에는 자동 축소하지 않습니다.")
    if result["summary"].get("non_cell_features"):
        raise ValueError("외부 연결·비셀 기능이 있어 필요한 입력을 확정하지 못했습니다. 자동 축소하지 않습니다.")
    tables = {table["id"]: table for table in result["tables"]}
    rules = make_rules(result)
    wb = workbook_from_result(result)
    try:
        engine = FormulaEngine(wb)
        requested, specs = set(), []
        for index, target in enumerate(dict.fromkeys(targets), 1):
            if target in tables:
                table = tables[target]
                sheet, reference = table["sheet"], table["range"]
                spec = {**rules["tables"][list(tables).index(target)]}
            else:
                try:
                    sheet, reference = target.rsplit("!", 1)
                    sheet = sheet.strip("'").replace("''", "'")
                    sheet = engine._sheet(sheet)
                    list(addresses(reference))
                except (ValueError, FormulaFailure):
                    raise ValueError(f"결과 지정 {target!r}을 찾지 못했습니다. 표 ID 또는 시트!A1:D20을 사용하세요.")
                spec = {"id": f"result_{index:02d}", "sheet": sheet, "range": reference.replace("$", ""),
                        "header_rows": 0, "exclude_rows": [], "origin": "user_output",
                        "flags": [], "columns": [], "input_rows": []}
            specs.append(spec)
            requested.update((sheet, address) for address in addresses(reference))
        # Table extraction displays a merged region's anchor even if the requested
        # cell is a child. Scalar formula refs to a child still remain blank.
        for sheet, ranges in result["workbook_structure"]["merges"].items():
            for reference in ranges:
                region = list(addresses(reference))
                if any((sheet, address) in requested for address in region):
                    requested.add((sheet, region[0]))
        observations = {(c["sheet"], c["cell"]): c for c in result["source_cells"]}
        needed, queue = set(requested), deque(sorted(requested))
        while queue:
            key = queue.popleft()
            cell = observations.get(key)
            if cell and cell["formula"]:
                for dep in dependencies(engine, cell["formula"], key[0]):
                    if dep not in needed:
                        needed.add(dep)
                        queue.append(dep)
            if len(needed) > 250000:
                raise ValueError("필요한 셀 수가 분석 한도를 넘습니다. 축소하지 않았습니다.")
        # Include physically empty references as explicit None inputs.
        for sheet, address in needed:
            key = f"{sheet}!{address}"
            if key not in rules["inputs"] and key not in rules["formulas"]:
                rules["inputs"][key] = {"sheet": sheet, "cell": address, "kind": "any",
                                       "number_format": "General", "cell_type": "n"}
        keys = {f"{s}!{a}" for s, a in needed}
        rules["inputs"] = {k: v for k, v in rules["inputs"].items() if k in keys}
        rules["formulas"] = {k: v for k, v in rules["formulas"].items() if k in keys}
        rules["tables"] = specs
        # Keep names/merges metadata for exact semantics, not unneeded input values.
        values = {k: observations.get((m["sheet"], m["cell"]), {}).get("value")
                  for k, m in rules["inputs"].items()}
        selected = calculate_data({"cells": values, "source_name": result["summary"]["source_name"]}, rules)
        selected["project_source"] = result["project_source"]
        selected["workbook_structure"] = copy.deepcopy(result["workbook_structure"])
        selected["selection"] = {"targets": list(targets), "source_formula_count": len(result["formulas"]),
                                 "kept_formula_count": len(selected["formulas"]),
                                 "kept_input_count": len(values), "source_modified": False}
        selected["summary"]["selection"] = selected["selection"]
        selected["summary"]["source_unchanged"] = result["summary"]["source_unchanged"]
        selected["summary"]["sha256"] = result["summary"]["sha256"]
        selected["input_evidence"] = result.get("input_evidence", {})
        return selected
    finally:
        wb.close()
