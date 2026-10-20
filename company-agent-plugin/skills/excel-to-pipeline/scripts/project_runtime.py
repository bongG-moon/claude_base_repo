"""Shared DATA adapters and pure PROCESS calculation for generated projects.

Never executes source code, writes a workbook, runs VBA, or uses saved answers as
computed values. Input acquisition is separate from calculate_data().
"""
import copy
import datetime as dt
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.defined_name import DefinedName


def _decode(value, meta):
    if value is None:
        return None
    kind = meta["kind"]
    if kind == "date":
        if isinstance(value, (dt.datetime, dt.date, dt.time)):
            return value
        if isinstance(value, str):
            try:
                return dt.datetime.fromisoformat(value)
            except ValueError:
                try:
                    return dt.time.fromisoformat(value)
                except ValueError:
                    pass
        raise ValueError("날짜는 원래 날짜 자료형 또는 ISO 날짜/시간 문자열이어야 합니다.")
    if kind == "text" and not isinstance(value, str):
        raise ValueError("문자열이어야 합니다. 사번의 앞자리 0이나 대소문자를 바꾸지 마세요.")
    if kind == "number" and (isinstance(value, bool) or not isinstance(value, (int, float))):
        raise ValueError("숫자여야 합니다. 빈 값은 None으로 전달하세요.")
    if kind == "boolean" and not isinstance(value, bool):
        raise ValueError("참/거짓 값이어야 합니다.")
    if not isinstance(value, (str, int, float, bool)) or isinstance(value, float) and not math.isfinite(value):
        raise ValueError("지원하지 않는 값 또는 무한대/NaN입니다.")
    return value


def validate_data(data, rules):
    """Validate without case folding, trimming or inferred numeric conversion."""
    if not isinstance(data, dict) or not isinstance(data.get("cells"), dict):
        raise ValueError("DATA/load.py는 cells 사전을 반환해야 합니다. README.md를 확인하세요.")
    actual, expected = set(data["cells"]), set(rules["inputs"])
    missing, extra = expected - actual, actual - expected
    if missing or extra:
        raise ValueError(f"DATA 입력 구조가 다릅니다. 누락 {len(missing)}개 {sorted(missing)[:3]}, "
                         f"추가 {len(extra)}개 {sorted(extra, key=str)[:3]}. 계산 규칙을 자동 변경하지 않았습니다.")
    values = {}
    for key, meta in rules["inputs"].items():
        try:
            value = _decode(data["cells"][key], meta)
            if "header" in meta and value != meta["header"]:
                raise ValueError("열 제목이 바뀌었습니다. 같은 의미의 열인지 확인하세요.")
            values[key] = value
        except ValueError as error:
            raise ValueError(f"입력 {key}: {error}") from error
    return values


def read_excel_data(source, rules, reader="auto", excel_timeout=90):
    """Load current literals only; reject structural drift instead of rewriting PROCESS."""
    from workbook_input import load_workbooks
    from excel_pipeline import safe_value, has_value
    source = Path(source).resolve()
    if not source.is_file():
        raise ValueError("DATA/sources.py에 지정한 입력 파일이 없습니다.")
    if source.stat().st_size > 64 * 1024 * 1024:
        raise ValueError("입력 파일이 64 MiB 한도를 넘습니다.")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    wb, cached, features, vba, evidence = load_workbooks(source, reader, excel_timeout)
    try:
        from excel_pipeline import MAX_RECTANGLE, MAX_CELLS
        if wb.sheetnames != [sheet["name"] for sheet in rules["sheets"]]:
            raise ValueError("시트 구성이 달라졌습니다. DATA 연결 또는 PROCESS 규칙을 확인하세요.")
        if any(ws.max_row * ws.max_column > MAX_RECTANGLE for ws in wb.worksheets):
            raise ValueError("시트 사용 범위가 분석 한도를 넘습니다.")
        current_formulas, populated, count = {}, set(), 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if not has_value(cell):
                        continue
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError("셀 수가 분석 한도를 넘습니다.")
                    key = f"{ws.title}!{cell.coordinate}"
                    populated.add(key)
                    if cell.data_type == "f":
                        current_formulas[key] = cell.value
        allowed = set(rules["inputs"]) | set(rules["formulas"])
        selected = bool(rules.get("selection"))
        compared_formulas = ({k: v for k, v in current_formulas.items() if k in allowed}
                             if selected else current_formulas)
        if (not selected and populated - allowed) or compared_formulas != rules["formulas"]:
            raise ValueError("행·열·수식 구조가 달라졌습니다. 기존 PROCESS를 덮어쓰지 않았습니다. 변경 범위를 검토하세요.")
        structure = {"epoch": wb.epoch.isoformat(),
                     "names": [{"name": n.name, "text": n.attr_text, "sheet": scope}
                               for scope, names in [(None, wb.defined_names)] +
                               [(ws.title, getattr(ws, "defined_names", {})) for ws in wb.worksheets]
                               for n in names.values()],
                     "merges": {ws.title: [str(m) for m in ws.merged_cells.ranges] for ws in wb.worksheets}}
        if structure != rules["structure"]:
            raise ValueError("이름 범위·병합·날짜 기준이 달라졌습니다. 계산 규칙을 먼저 확인하세요.")
        data = {"cells": {key: safe_value(wb[meta["sheet"]][meta["cell"]].value)
                          for key, meta in rules["inputs"].items()},
                "source_name": source.name, "input_evidence": evidence,
                "sha256": digest, "source_unchanged": True}
        validate_data(data, rules)
        if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
            raise ValueError("읽는 동안 원본이 바뀌었습니다. 저장이 끝난 뒤 다시 실행하세요.")
        return data
    finally:
        wb.close()
        cached.close()


def replace_tables(data, replacements, rules):
    """Map Python row dictionaries to existing inputs without touching formulas.

    Rows retain their original positions. This adapter intentionally rejects row
    expansion, missing fields and calculated-column overrides instead of losing data.
    """
    result = copy.deepcopy(data)
    tables = {t["id"]: t for t in rules["tables"]}
    for name, rows in replacements.items():
        if name not in tables:
            raise ValueError(f"알 수 없는 표 ID: {name}")
        mapping = tables[name]["input_rows"]
        if not isinstance(rows, list) or len(rows) != len(mapping):
            raise ValueError(f"{name}: 기존 {len(mapping)}행과 다릅니다. 행 추가·삭제는 PROCESS 범위 검토가 필요합니다.")
        for index, (row, positions) in enumerate(zip(rows, mapping), 1):
            if not isinstance(row, dict) or set(row) != set(positions):
                raise ValueError(f"{name}의 {index}번째 행: 필요한 입력 열과 다릅니다. input_rows의 열만 전달하세요.")
            for column, key in positions.items():
                result["cells"][key] = row[column]
    validate_data(result, rules)
    # The Excel checksum describes the initial observation, not this replaced data.
    result.pop("sha256", None)
    result["source_unchanged"] = None
    result["input_evidence"] = {"reader": "python_adapter", "values_origin": "provided_python_values",
                                "excel_started": False, "incomplete_metadata": []}
    return result


def calculate_data(data, rules, compiled_nodes=None, sheet_order=None):
    """Recompute from values and frozen rules, with no source I/O or Office process."""
    from excel_pipeline import (FormulaEngine, extract_table, merged_lookup, processing_plan,
                                safe_value, VERSION)
    values = validate_data(data, rules)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.epoch = dt.datetime.fromisoformat(rules["structure"]["epoch"])
    for name in sheet_order or [sheet["name"] for sheet in rules["sheets"]]:
        wb.create_sheet(name)
    try:
        for key, value in values.items():
            meta = rules["inputs"][key]
            cell = wb[meta["sheet"]][meta["cell"]]
            cell.value = value
            # A text value starting with '=' is data, never a newly supplied formula.
            if isinstance(value, str):
                cell.data_type = "e" if meta.get("cell_type") == "e" else "s"
            cell.number_format = meta.get("number_format", "General")
        for key, expression in rules["formulas"].items():
            sheet, address = key.rsplit("!", 1)
            wb[sheet][address] = expression
        for sheet, merges in rules["structure"]["merges"].items():
            for reference in merges:
                wb[sheet].merge_cells(reference)
        for item in rules["structure"]["names"]:
            owner = wb if item["sheet"] is None else wb[item["sheet"]]
            owner.defined_names.add(DefinedName(item["name"], attr_text=item["text"]))
        if compiled_nodes is not None:
            expected = {tuple(key.rsplit("!", 1)) for key in rules["formulas"]}
            if set(compiled_nodes) != expected:
                raise ValueError("PROCESS 계산 위치가 입력 계약과 다릅니다. 수식 위치 추가·삭제는 재검토가 필요합니다.")
        engine = FormulaEngine(wb, compiled_nodes=compiled_nodes)
        formulas = engine.analyze()
        for formula in formulas:
            formula.update(value=safe_value(formula.get("value")), cached_value=None,
                           cached_value_origin="not_loaded_for_reusable_processing", cache_comparison="unavailable")
        tables = []
        for table in rules["tables"]:
            spec = {**table, "bounds": range_boundaries(table["range"])}
            tables.append(extract_table(wb[table["sheet"]], spec, merged_lookup(wb[table["sheet"]]), engine))
        flags = sorted({f for t in tables for f in t["flags"]})
        if any(f["status"] != "computed" for f in formulas):
            flags.append("formula_review_required")
        if rules.get("vba") and rules["vba"].get("status") != "no_vba":
            flags.append("vba_semantic_migration_required")
        if rules.get("non_cell_features"):
            flags.append("non_cell_features_not_executed")
        evidence = data.get("input_evidence") or {"reader": "python_adapter", "excel_started": False,
                                                 "values_origin": "provided_python_values", "incomplete_metadata": []}
        observations = [{"sheet": meta["sheet"], "cell": meta["cell"], "value": safe_value(values[key]),
                         "formula": None, "type": meta["cell_type"], "number_format": meta["number_format"], "cached_value": None}
                        for key, meta in rules["inputs"].items()]
        observations += [{"sheet": key.rsplit("!", 1)[0], "cell": key.rsplit("!", 1)[1], "value": None,
                          "formula": text, "type": "f", "number_format": "General", "cached_value": None}
                         for key, text in rules["formulas"].items()]
        summary = {"version": VERSION, "source_name": data.get("source_name", "Python 입력"),
                   "sha256": data.get("sha256"), "source_unchanged": data.get("source_unchanged"),
                   "sheet_count": len(wb.worksheets), "table_count": len(tables),
                   "data_row_count": sum(len(t["rows"]) for t in tables), "formula_count": len(formulas),
                   "formula_status": dict(Counter(f["status"] for f in formulas)),
                   "cache_comparison": dict(Counter(f["cache_comparison"] for f in formulas)),
                   "review_flags": flags, "status": "review_required" if flags else "extracted",
                   "reader": evidence["reader"], "values_origin": evidence["values_origin"], "model_calls": 0,
                   "non_cell_features": rules.get("non_cell_features", []),
                   "scope": "reused_fixed_structure_and_supported_formulas_not_vba_execution"}
        if rules.get("vba"):
            summary["vba"] = {**rules["vba"], "status": "not_reextracted", "macro_executed": False,
                              "translation_generated": False}
        return {"summary": summary, "tables": tables, "formulas": formulas, "source_cells": observations,
                "sheets": rules["sheets"], "pipeline_plan": processing_plan(tables, formulas), "input_evidence": evidence}
    finally:
        wb.close()
