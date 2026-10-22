"""Beginner-facing Python project with one module per sheet and grouped formulas."""
import re
from pathlib import Path
from collections import defaultdict

from formula_engine import FormulaFailure, FUNCTIONS
from project_selection import workbook_from_result, dependencies
from formula_engine import FormulaEngine
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter


FUNCTION_NOTES = {
    "SUM": "입력한 숫자들의 합계를 구합니다.",
    "AVERAGE": "입력한 숫자들의 평균을 구합니다.",
    "MIN": "가장 작은 숫자를 선택합니다.", "MAX": "가장 큰 숫자를 선택합니다.",
    "COUNT": "숫자가 있는 항목의 개수를 셉니다.",
    "COUNTA": "비어 있지 않은 항목의 개수를 셉니다.",
    "IF": "첫 인자의 조건이 참이면 두 번째 값, 거짓이면 세 번째 값을 사용합니다.",
    "IFERROR": "첫 계산이 오류이면 두 번째 대체 값을 사용합니다.",
    "ROUND": "첫 값을 두 번째 인자의 자릿수로 반올림합니다. 음수 자릿수는 십·백 단위입니다.",
    "SUMIF": "조건 범위에서 조건에 맞는 항목의 값을 합산합니다.",
    "SUMIFS": "여러 조건을 모두 만족하는 항목의 값을 합산합니다.",
    "COUNTIF": "지정한 조건에 맞는 항목의 개수를 셉니다.",
    "COUNTIFS": "여러 조건을 모두 만족하는 항목의 개수를 셉니다.",
    "VLOOKUP": "표의 첫 열에서 값을 찾고 지정한 열의 값을 반환합니다. 마지막 인자의 정확/근사 일치 설정을 확인하세요.",
    "AND": "모든 조건이 참인지 확인합니다.", "OR": "하나 이상의 조건이 참인지 확인합니다.",
}
OPERATOR_NOTES = {
    "+": "왼쪽 값에 오른쪽 값을 더합니다.", "-": "왼쪽 값에서 오른쪽 값을 뺍니다.",
    "*": "두 값을 곱합니다.", "/": "왼쪽 값을 오른쪽 값으로 나눕니다. 분모가 0이면 오류입니다.",
    "^": "왼쪽 값을 오른쪽 지수만큼 거듭제곱합니다.", "&": "두 값을 문자열로 이어 붙입니다.",
    "=": "두 값이 같은지 비교합니다.", "<>": "두 값이 다른지 비교합니다.",
    ">": "왼쪽 값이 더 큰지 비교합니다.", "<": "왼쪽 값이 더 작은지 비교합니다.",
    ">=": "왼쪽 값이 크거나 같은지 비교합니다.", "<=": "왼쪽 값이 작거나 같은지 비교합니다.",
}


def formula_comments(formula, ast, sheet, start, end, column):
    """Deterministic documentation: no model calls or guessed business meaning."""
    notes = [
        f"결과 위치: {sheet}!{column}{start}:{column}{end} ({end - start + 1}개 셀)",
        f"원본 Excel 수식 ({column}{start} 기준): {formula}",
        "설명 범위: 수식에서 확인한 계산 동작입니다. 매출·수익 등 업무 의미는 임의로 추정하지 않습니다.",
    ]
    refs, operations = [], []

    def visit(node):
        kind = node[0]
        if kind == "ref":
            location = f"{node[1] or sheet}!{node[2]}"
            if node[2] != node[3]:
                location += ":" + node[3]
            refs.append(location)
        elif kind == "name":
            refs.append(f"이름 정의 {node[1:]!r} (원본 이름 정의의 참조 범위 사용)")
        elif kind == "call":
            for child in node[2]:
                visit(child)
            operations.append(node[1] + ": " + FUNCTION_NOTES.get(
                node[1], "자동 지원 여부를 확인해야 합니다. 실행되는 분기에서 미지원이면 오류로 남깁니다."))
        elif kind == "binary":
            visit(node[2])
            visit(node[3])
            operations.append(node[1] + ": " + OPERATOR_NOTES[node[1]])
        elif kind == "unary":
            visit(node[2])
            operations.append({"%": "퍼센트: 값을 100으로 나눕니다.",
                               "-": "단항 음수: 값의 부호를 반대로 바꿉니다.",
                               "+": "단항 양수: 숫자 값으로 처리합니다."}[node[1]])
        elif kind == "error":
            operations.append(f"원본 오류 {node[1]}를 유지합니다. 정상 값으로 대체하지 않습니다.")

    if ast is not None:
        visit(ast)
        if ast[0] in ("ref", "name"):
            operations.append("참조한 셀 또는 이름 정의의 값을 결과 위치에 사용합니다.")
        elif ast[0] == "literal":
            operations.append("원본 수식의 상수 값을 그대로 사용합니다.")
    notes.append("입력 참조 (첫 행 기준): " + (" / ".join(dict.fromkeys(refs)) or "셀 참조 없음 — 상수 또는 원본 오류"))
    for number, operation in enumerate(dict.fromkeys(operations), 1):
        notes.append(f"계산 동작 {number}: {operation}")
    if ast is None:
        notes.append("자동 해석 불가: 원본 수식의 함수·참조를 확인해야 합니다. 계산했다고 처리하지 않습니다.")
    notes += [
        "행 적용: row는 Excel 행 번호입니다. 상대 참조 행은 row와 함께 이동하고 $로 고정한 행은 유지합니다.",
        "수정 위치: 아래 book.set의 계산식·상수·조건을 바꾸면 이 묶음 전체에 적용됩니다.",
        "주의: 참조 열은 원본 그대로입니다. 행 범위를 늘리거나 참조 구조를 바꾸면 입력 계약도 재검토하세요.",
        "이 주석은 생성 시점의 원본 설명입니다. Python 계산을 수정하면 주석도 함께 갱신하세요.",
    ]
    # Keep workbook line breaks inside comments rather than executable source.
    return ["    # " + line.replace("\r", r"\r").replace("\n", r"\n").replace("\x00", r"\0")
            for line in notes]


def coordinate_code(address, origin_row):
    """Relative row -> loop variable; $ row stays fixed, columns stay exact."""
    match = re.fullmatch(r"\$?([A-Za-z]+)(\$?)([0-9]+)", address)
    column, absolute, row = match.groups()
    if absolute:
        return repr(column.upper() + row)
    offset = int(row) - origin_row
    term = "row" if not offset else f"row {'+' if offset > 0 else '-'} {abs(offset)}"
    return 'f' + repr(column.upper() + "{" + term + "}")


def expression_code(ast, row):
    kind = ast[0]
    if kind == "literal":
        return f"VALUE({ast[1]!r})"
    if kind == "error":
        return f"ERROR({ast[1]!r})"
    if kind == "ref":
        first, last = coordinate_code(ast[2], row), coordinate_code(ast[3], row)
        return f"REF({ast[1]!r}, {first}" + (f", {last}" if ast[2] != ast[3] else "") + ")"
    if kind == "name":
        return f"NAME({ast[1]!r}, {ast[2]!r})"
    if kind == "call":
        if ast[1] not in FUNCTIONS:
            return "UNSUPPORTED_CALL(" + repr(ast[1]) + (", " if ast[2] else "") + ", ".join(expression_code(arg, row) for arg in ast[2]) + ")"
        return ast[1] + "(" + ", ".join(expression_code(arg, row) for arg in ast[2]) + ")"
    if kind == "unary":
        inner = expression_code(ast[2], row)
        return f"PERCENT({inner})" if ast[1] == "%" else f"({ast[1]}{inner})"
    if kind == "binary":
        left, right = expression_code(ast[2], row), expression_code(ast[3], row)
        op = ast[1]
        if op == "&":
            return f"CONCAT({left}, {right})"
        if op in ("=", "<>", "<", ">", "<=", ">="):
            return f"COMPARE({left}, {op!r}, {right})"
        return f"({left} {'**' if op == '^' else op} {right})"
    raise FormulaFailure("unsupported_syntax")


def sheet_schedule(result, engine):
    """Kahn ordering. Mutual sheet refs use the existing cell-level DAG evaluator."""
    names = [s["name"] for s in result["sheets"]]
    graph = {s: set() for s in names}
    uncertain = set()
    for f in result["formulas"]:
        try:
            graph[f["sheet"]].update(s for s, _ in dependencies(engine, f["formula"], f["sheet"]) if s != f["sheet"])
        except ValueError:
            uncertain.add(f["sheet"])
    pending = {s: set(deps) for s, deps in graph.items()}
    order = []
    while pending:
        ready = [s for s in names if s in pending and not pending[s]]
        if not ready:
            break
        order.extend(ready)
        for s in ready:
            del pending[s]
        for deps in pending.values():
            deps.difference_update(ready)
    # A sheet-level cycle need not be a cell-level cycle. Do not invent an order.
    intertwined = [s for s in names if s in pending]
    return order + intertwined, graph, intertwined, sorted(uncertain)


def compile_sheets(result):
    wb = workbook_from_result(result)
    try:
        engine = FormulaEngine(wb)
        order, graph, intertwined, uncertain = sheet_schedule(result, engine)
        groups = defaultdict(lambda: defaultdict(list))
        unsupported = defaultdict(list)
        originals = {(f["sheet"], f["cell"]): f["formula"] for f in result["formulas"]}
        for formula in result["formulas"]:
            row, col = coordinate_to_tuple(formula["cell"])
            try:
                code = expression_code(engine._parse(formula["formula"]), row)
                groups[formula["sheet"]][(col, code)].append(row)
            except (FormulaFailure, RecursionError) as error:
                unsupported[formula["sheet"]].append((formula["cell"], str(error)))
        modules, group_count = {}, 0
        # Formula-free input sheets need no editable PROCESS file.
        for sheet in order:
            if not groups[sheet] and not unsupported[sheet]:
                continue
            index = len(modules) + 1
            safe = re.sub(r"\W+", "_", sheet, flags=re.UNICODE).strip("_")[:40] or "sheet"
            name = f"sheet_{index:02d}_{safe}"
            lines = [repr(f"시트 {sheet} 계산: 아래 반복문 한 곳을 수정하면 같은 행 패턴에 적용됩니다."),
                     "from python_formulas import (VALUE, REF, NAME, ERROR, COMPARE, CONCAT, PERCENT,",
                     "    SUM, AVERAGE, MIN, MAX, COUNT, COUNTA, IF, IFERROR, ROUND,",
                     "    SUMIF, SUMIFS, COUNTIF, COUNTIFS, VLOOKUP, AND, OR, UNSUPPORTED_CALL)", "",
                     f"SHEET = {sheet!r}", "", "def register(book):",
                     "    # REF(None, 'A1')은 이 시트의 셀, REF('기준', 'A1')은 다른 시트입니다.",
                     "    # SUM/IF/ROUND 등은 Excel의 빈 값·조건·반올림 의미를 지키는 Python 함수입니다.",
                     "    # row는 원본 Excel 행 번호입니다. 행 범위는 자동으로 늘리지 않습니다."]
            for (col, code), rows in sorted(groups[sheet].items()):
                rows = sorted(rows)
                runs = []
                for row in rows:
                    if runs and row == runs[-1][-1] + 1:
                        runs[-1].append(row)
                    else:
                        runs.append([row])
                for run in runs:
                    group_count += 1
                    start, end = run[0], run[-1]
                    column = get_column_letter(col)
                    original = originals[(sheet, f"{column}{start}")]
                    lines += formula_comments(original, engine._parse(original), sheet, start, end, column)
                    lines += [f"    # {column}{start}:{column}{end} — 같은 계산을 해당 행 전체에 적용",
                              f"    for row in range({start}, {end + 1}):",
                              f'        book.set(SHEET, f"{column}{{row}}", {code})']
            for address, code in unsupported[sheet]:
                row, col = coordinate_to_tuple(address)
                lines += formula_comments(originals[(sheet, address)], None, sheet, row, row, get_column_letter(col))
                lines += ["    # 자동 변환하지 못한 수식입니다. 0이나 과거 값으로 대신하지 않습니다.",
                          f"    book.unsupported(SHEET, {address!r}, {code!r})"]
            lines += ["", "def process_tables(tables):",
                      "    # 이 시트의 표에 추가 정리·결합·집계가 필요하면 여기에 작성합니다.",
                      "    # 기본은 계산된 표를 그대로 반환합니다. 입력 출처는 DATA에서만 변경합니다.",
                      "    return tables", ""]
            modules[name] = (sheet, "\n".join(lines))
        return modules, {"sheet_order": order, "dependencies": {s: sorted(d) for s, d in graph.items()},
                         "intertwined_sheets": intertwined, "uncertain_sheets": uncertain,
                         "formula_count": len(result["formulas"]), "loop_groups": group_count}
    finally:
        wb.close()


MAIN = '''"""실행 시작점: DATA 읽기 → 시트 순서대로 PROCESS 계산 → 필요한 표만 저장."""
from pathlib import Path
from datetime import datetime
import argparse
import json
import sys

ROOT = Path(__file__).resolve().parent
sys.dont_write_bytecode = True  # 실행할 때 불필요한 __pycache__ 폴더를 만들지 않습니다.
sys.path.insert(0, str(ROOT / "_support"))
from DATA.load import load_data
from PROCESS.pipeline import run
from excel_pipeline import publish, write_json


def main():
    parser = argparse.ArgumentParser(description="자료를 읽고 시트별 Python 계산을 실행합니다.")
    parser.add_argument("--output", help="새 결과 폴더. 생략하면 outputs 아래에 저장합니다.")
    parser.add_argument("--diagnostics", action="store_true", help="셀별 상세 검증 기록도 저장합니다.")
    args = parser.parse_args()
    output = Path(args.output).resolve() if args.output else ROOT / "outputs" / datetime.now().strftime("run-%Y%m%d-%H%M%S-%f")
    try:
        if output.exists() or output.is_symlink():
            raise ValueError("결과 폴더가 이미 있습니다. 다른 이름을 지정하세요.")
        result, tables = run(load_data())
        publish(result, output, emit_project=False, layout="compact", diagnostics=args.diagnostics, sample_rows=SAMPLE_ROWS, context_chars=CONTEXT_CHARS)
        write_json(output / "결과표.json", tables)
        print(json.dumps({"output": str(output), "summary": result["summary"]}, ensure_ascii=False))
        return 0
    except (ValueError, OSError, KeyError, TypeError) as error:
        print("처리하지 못했습니다: " + str(error), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
'''

LOAD = '''"""입력 담당: 다른 파일·Python·DB로 바꾸려면 이 파일을 수정합니다."""
import json
from pathlib import Path
from project_runtime import read_excel_data, replace_tables
from .sources import SOURCE, READER, EXCEL_TIMEOUT

ROOT = Path(__file__).resolve().parents[1]


def load_data():
    # 입력 계약은 자동 생성한 보조 자료입니다. 계산 코드는 PROCESS에 있습니다.
    rules = json.loads((ROOT / "_support" / "contract.json").read_text(encoding="utf-8"))
    source = SOURCE if SOURCE.is_absolute() else ROOT / SOURCE
    data = read_excel_data(source, rules, reader=READER, excel_timeout=EXCEL_TIMEOUT)
    # 외부 Python의 표를 연결할 때: data = replace_tables(data, {"표ID": 행목록}, rules)
    # Excel 없이도 {"cells": {"시트!A1": 값, ...}, "source_name": "자료 이름"}을 반환할 수 있습니다.
    return data
'''


def write_compact_project(result, output, sample_rows, context_chars, compiled=None):
    from project_builder import make_rules
    from excel_pipeline import write_json
    modules, schedule = compiled or compile_sheets(result)
    rules = make_rules(result)
    output = Path(output)
    for directory in ("DATA", "PROCESS", "outputs"):
        (output / directory).mkdir()
    for directory in ("DATA", "PROCESS"):
        (output / directory / "__init__.py").write_text("", encoding="utf-8")
    (output / "main.py").write_text(MAIN.replace("SAMPLE_ROWS", str(sample_rows)).replace("CONTEXT_CHARS", str(context_chars)), encoding="utf-8")
    (output / "DATA" / "sources.py").write_text(
        '"""입력 위치: 다른 파일을 사용하려면 SOURCE만 수정하세요."""\nfrom pathlib import Path\n\n'
        f"SOURCE = Path({result['project_source']!r})\n"
        '# auto는 일반 파일을 직접 읽고 필요한 경우에만 Excel을 사용합니다.\nREADER = "auto"\nEXCEL_TIMEOUT = 90\n', encoding="utf-8")
    (output / "DATA" / "load.py").write_text(LOAD, encoding="utf-8")
    write_json(output / "_support" / "contract.json", rules)
    lines = ['"""전체 처리: 선행 시트부터 실행합니다. 셀 간 참조도 계산기가 확인합니다."""',
             "import copy", "import json", "from pathlib import Path",
             "from python_formulas import FormulaBook", "from project_runtime import calculate_data"]
    for name, (_, source) in modules.items():
        (output / "PROCESS" / (name + ".py")).write_text(source, encoding="utf-8")
        lines.append(f"from . import {name}")
    lines += ["", f"SHEET_ORDER = {schedule['sheet_order']!r}",
              "# 서로 참조하는 시트는 단일 순서를 강제하지 않고 셀 참조 순서로 계산합니다.",
              f"INTERTWINED_SHEETS = {schedule['intertwined_sheets']!r}", "", "def run(data):",
              '    rules = json.loads((Path(__file__).resolve().parents[1] / "_support" / "contract.json").read_text(encoding="utf-8"))',
              "    book = FormulaBook()", "    # 각 시트의 Python 계산을 등록합니다. 실제 값 계산은 아래 순서로 수행합니다."]
    for name in modules:
        lines.append(f"    {name}.register(book)")
    lines += ["    result = calculate_data(data, rules, compiled_nodes=book.nodes, sheet_order=SHEET_ORDER)",
              f"    result['summary']['calculation_order'] = SHEET_ORDER",
              f"    result['summary']['intertwined_sheets'] = INTERTWINED_SHEETS",
              f"    result['summary']['uncertain_dependency_sheets'] = {schedule['uncertain_sheets']!r}",
              f"    result['summary']['selection'] = rules.get('selection')",
              "    tables = copy.deepcopy(result['tables'])"]
    for name, (sheet, _) in modules.items():
        lines += [f"    selected = [t for t in tables if t['sheet'] == {sheet!r}]",
                  f"    updated = {name}.process_tables(selected)",
                  f"    tables = [t for t in tables if t['sheet'] != {sheet!r}] + updated"]
    lines += ["    return result, tables", ""]
    (output / "PROCESS" / "pipeline.py").write_text("\n".join(lines), encoding="utf-8")
    write_json(output / "_details" / "calculation-order.json", schedule)
    selection = result.get("selection")
    note = (f"\n선택한 결과: {', '.join(selection['targets'])}. 원본 {selection['source_formula_count']}개 중 "
            f"{selection['kept_formula_count']}개 수식과 필요한 입력만 포함했습니다. 원본은 삭제하지 않았습니다.\n") if selection else ""
    (output / "README.md").write_text(
        "# 시작하기\n\n`python -X utf8 main.py`를 실행하세요. 새 결과는 `outputs` 안에 저장됩니다.\n\n"
        "계산 구조를 먼저 보려면 `업무흐름.html`을 여세요. 입력 표·계산·결과의 실제 수식 참조를 그림으로 보여줍니다.\n"
        "표 이름은 시트·열 이름에 근거한 제안이며 기존 표 ID는 유지합니다. 그림은 생성 시점 기준이고 재실행 때 추가 생성하지 않습니다.\n\n"
        "## 수정할 곳은 두 군데입니다\n\n"
        "- **DATA**: 입력 파일 위치는 sources.py, Python·JSON·DB 연결은 load.py\n"
        "- **PROCESS**: 시트당 Python 파일 하나. 같은 행 계산은 반복문 한 곳에서 수정합니다.\n"
        "  pipeline.py는 시트 간 참조로 정한 실행 순서입니다.\n\n"
        "## 계산 코드 읽기\n\n"
        "`for row in range(2, 101)`은 Excel 2~100행에 같은 계산을 적용합니다.\n"
        "`REF(None, f'B{row}') * REF(None, f'C{row}')`는 현재 행의 B열과 C열을 곱합니다.\n"
        "`REF('단가', 'B2')`는 다른 시트의 고정 셀입니다. VALUE(0.1)은 숫자 상수입니다.\n"
        "SUM·IF·ROUND 등은 Excel의 빈 값·참조·반올림 의미를 보존하는 Python 보조 함수입니다.\n"
        "일반 pandas 코드로 전부 번역한 것은 아니며 미지원 함수와 VBA는 검토가 필요합니다.\n\n"
        f"계산 순서: {' → '.join(schedule['sheet_order'])}\n"
        f"수식 {schedule['formula_count']}개를 반복 계산 {schedule['loop_groups']}묶음으로 구성했습니다.\n"
        "시트가 서로 참조하면 셀 단위로 선행 값을 계산하며, 실제 순환 참조는 오류로 남깁니다.\n"
        "의존성을 확인하지 못한 함수는 순서가 검증됐다고 처리하지 않습니다.\n" + note +
        "\n## 알아둘 점\n\n"
        "- 값·출처만 바꾸면 DATA만 수정해 PROCESS를 재사용할 수 있습니다. 앞자리 0·대소문자·빈 값을 유지하세요.\n"
        "- 입력 행·열·수식 구조가 달라지면 새 프로젝트로 다시 분석합니다. 반복문의 범위를 무작정 늘리지 마세요.\n"
        "- 필요한 결과만 지정하면 그 결과와 연결된 입력·수식만 새 프로젝트에 남깁니다. 원본이나 기존 결과는 지우지 않습니다.\n"
        "- `_details`는 상세 진단, `_support`는 자동 생성된 계산 지원 코드와 입력 계약입니다. 평소에는 열 필요 없습니다.\n"
        "- 초기 추출 표는 tables, 실행 결과는 outputs입니다. 미지원 계산은 README 요약과 _details에서 확인하세요.\n"
        "- 상세 셀별 검증이 필요할 때만 `python -X utf8 main.py --diagnostics`로 실행하세요.\n"
        "- 기본 준비물: Python 3.11 이상과 openpyxl. Excel 경유 읽기에는 xlwings가 필요합니다. 자동 설치하지 않습니다.\n"
        "- 원본 수정·VBA 실행·추가 모델 호출은 하지 않습니다.\n",
        encoding="utf-8")
