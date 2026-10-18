"""Generate a local DATA/PROCESS project; no workbook data is executed as code."""
from pathlib import Path
import json
from openpyxl.utils.cell import range_boundaries, get_column_letter


MAIN = '''"""실행 시작점: 데이터 읽기 → 계산 → 업무 처리 → 새 결과 저장."""
from pathlib import Path
from datetime import datetime
import argparse
import json
import sys

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "pipeline_lib"))
from DATA.load import load_data
from PROCESS.pipeline import run
from excel_pipeline import publish, write_json


def main():
    parser = argparse.ArgumentParser(description="DATA의 입력으로 PROCESS를 다시 실행합니다.")
    parser.add_argument("--output", help="새 결과 폴더. 생략하면 프로젝트의 outputs 안에 만듭니다.")
    args = parser.parse_args()
    output = Path(args.output).resolve() if args.output else ROOT / "outputs" / datetime.now().strftime("run-%Y%m%d-%H%M%S-%f")
    try:
        # 기존 결과를 덮어쓰지 않습니다. 경로부터 확인해 불필요한 Excel 읽기를 피합니다.
        if output.exists() or output.is_symlink():
            raise ValueError("결과 폴더가 이미 있습니다. 다른 이름을 지정하세요.")
        data = load_data()  # 파일·DB 등 입력 출처는 DATA/load.py에서 바꿉니다.
        result, processed = run(data)  # 계산과 업무 규칙은 PROCESS에서 관리합니다.
        publish(result, output, emit_project=False, sample_rows=SAMPLE_ROWS, context_chars=CONTEXT_CHARS)
        # 수식 분석 원본과 사용자가 추가한 업무 결과를 구분해서 보관합니다.
        write_json(output / "processed_tables.json", processed)
        print(json.dumps({"output": str(output), "summary": result["summary"]}, ensure_ascii=False))
        return 0
    except (ValueError, OSError, KeyError, TypeError) as error:
        print("처리하지 못했습니다: " + str(error), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
'''

LOAD = '''"""입력 담당: 이 파일만 바꾸면 다른 Python 코드나 DB의 자료도 연결할 수 있습니다."""
import json
from pathlib import Path
from .sources import SOURCE, READER, EXCEL_TIMEOUT
from project_runtime import read_excel_data, replace_tables

ROOT = Path(__file__).resolve().parents[1]


def load_data():
    """표의 열 이름·자료형·행 배치를 유지한 입력을 반환합니다."""
    rules = json.loads((ROOT / "PROCESS" / "workbook_rules.json").read_text(encoding="utf-8"))
    # 정상 파일은 직접 읽습니다. 필요한 파일만 기존 xlwings 경로를 사용합니다.
    # 읽기 실패 시 다른 자료로 조용히 대체하지 않습니다.
    # 상대 경로를 적어도 실행 창의 위치가 아니라 이 프로젝트 폴더를 기준으로 찾습니다.
    source = SOURCE if SOURCE.is_absolute() else ROOT / SOURCE
    data = read_excel_data(source, rules, reader=READER, excel_timeout=EXCEL_TIMEOUT)
    # 다른 코드의 표를 연결할 때: data = replace_tables(data, {"표ID": 행목록}, rules)
    # 행목록은 {"열이름": 값}의 목록입니다. 사용 예와 빈 값 규칙은 README.md를 보세요.
    # Excel을 전혀 쓰지 않을 때도 같은 cells 구조를 반환하면 PROCESS는 그대로 작동합니다.
    return data
'''

PIPELINE = '''"""처리 담당: 입력 검사 → 저장된 수식 계산 → 표별 업무 처리 순서입니다."""
import copy
import json
from pathlib import Path
from project_runtime import calculate_data
from .table_steps import process_tables


def run(data):
    rules = json.loads(Path(__file__).with_name("workbook_rules.json").read_text(encoding="utf-8"))
    # 새 파일의 수식으로 몰래 바꾸지 않고, 프로젝트 생성 때 확인한 규칙을 사용합니다.
    # 입력 열·자료형이 맞지 않으면 0으로 채우지 않고 수정할 위치를 알려줍니다.
    result = calculate_data(data, rules)
    # 업무별 추가 처리만 아래 함수에서 수행합니다. 원본 분석 표는 따로 보존합니다.
    processed = process_tables(copy.deepcopy(result["tables"]))
    return result, processed
'''


def value_kind(cell):
    value = cell.get("value")
    if cell.get("type") == "d":
        return "date"
    if value is None:
        return "any"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    return "text"


def make_rules(result):
    cells = {(c["sheet"], c["cell"]): c for c in result["source_cells"]}
    formulas = {f"{c['sheet']}!{c['cell']}": c["formula"] for c in result["source_cells"] if c.get("formula")}
    inputs = {f"{s}!{a}": {"sheet": s, "cell": a, "kind": value_kind(c),
                            "number_format": c.get("number_format", "General"), "cell_type": c.get("type", "n")}
              for (s, a), c in cells.items() if not c.get("formula")}
    tables = []
    merged_children = set()
    for sheet, ranges in result["workbook_structure"]["merges"].items():
        for reference in ranges:
            left, top, right, bottom = range_boundaries(reference)
            for row in range(top, bottom + 1):
                for column in range(left, right + 1):
                    if (column, row) != (left, top):
                        merged_children.add(f"{sheet}!{get_column_letter(column)}{row}")
    for table in result["tables"]:
        left, top, right, bottom = range_boundaries(table["range"])
        addresses = []
        for row in table["source_rows"]:
            mapping = {}
            for column, label in enumerate(table["columns"], left):
                address = f"{get_column_letter(column)}{row}"
                key = f"{table['sheet']}!{address}"
                if key not in formulas and key not in merged_children:
                    inputs.setdefault(key, {"sheet": table["sheet"], "cell": address, "kind": "any", "cell_type": "n", "number_format": "General"})
                    mapping[label] = key
            addresses.append(mapping)
        # Header text fixes the meaning of each column; data values are not frozen.
        for row in range(top, top + table["header_rows"]):
            for column in range(left, right + 1):
                key = f"{table['sheet']}!{get_column_letter(column)}{row}"
                if key in inputs:
                    inputs[key]["header"] = cells[(table["sheet"], f"{get_column_letter(column)}{row}")]["value"]
        tables.append({"id": table["id"], "sheet": table["sheet"], "range": table["range"],
                       "header_rows": table["header_rows"], "exclude_rows": table["excluded_rows"],
                       "columns": table["columns"], "input_rows": addresses,
                       "origin": table["origin"], "flags": [f for f in table["flags"] if f not in {
                           "summary_rows_retained", "unresolved_formula_values", "internal_blank_rows", "repeated_header_rows"}]})
    return {"version": 1, "source_name": result["summary"]["source_name"],
            "sheets": result["sheets"], "structure": result["workbook_structure"],
            "inputs": inputs, "formulas": formulas, "tables": tables,
            "vba": result["summary"].get("vba"),
            "non_cell_features": result["summary"].get("non_cell_features", [])}


def write_project(result, output, sample_rows, context_chars):
    from excel_pipeline import write_json
    output = Path(output)
    rules = make_rules(result)
    for name in ("DATA", "PROCESS", "outputs"):
        (output / name).mkdir()
    for name in ("DATA", "PROCESS"):
        (output / name / "__init__.py").write_text('"""' + ("데이터를 읽는 코드입니다." if name == "DATA" else "입력을 계산하고 정리하는 코드입니다.") + '"""\n', encoding="utf-8")
    (output / "main.py").write_text(MAIN.replace("SAMPLE_ROWS", str(sample_rows)).replace("CONTEXT_CHARS", str(context_chars)), encoding="utf-8")
    sources = ('"""바꿔도 되는 설정: 입력 파일 위치와 읽기 방식."""\nfrom pathlib import Path\n\n'
               '# 다른 파일을 쓰려면 아래 경로만 수정하세요. 원본은 저장하지 않습니다.\n'
               f'SOURCE = Path({result["project_source"]!r})\n'
               '# auto: 일반 파일 직접 읽기, 필요한 파일만 Excel 읽기.\nREADER = "auto"\n'
               '# Excel을 사용할 때만 적용되는 대기 시간(초)입니다.\nEXCEL_TIMEOUT = 90\n')
    (output / "DATA" / "sources.py").write_text(sources, encoding="utf-8")
    (output / "DATA" / "load.py").write_text(LOAD, encoding="utf-8")
    (output / "PROCESS" / "pipeline.py").write_text(PIPELINE, encoding="utf-8")
    write_json(output / "PROCESS" / "workbook_rules.json", rules)
    lines = ['"""표별 추가 업무 규칙: 필요한 정리·결합·집계를 이곳에 작성합니다."""', '']
    stages = {stage["table"]: stage for stage in result["pipeline_plan"]["stages"]}
    meanings = {"SUM": "값 합산", "SUMIF": "조건에 맞는 값 합산", "SUMIFS": "여러 조건에 맞는 값 합산",
                "COUNTIF": "조건에 맞는 건수", "COUNTIFS": "여러 조건에 맞는 건수", "IF": "조건에 따라 값 선택",
                "ROUND": "정해진 자릿수 반올림", "VLOOKUP": "기준표에서 같은 키의 값 찾기",
                "AVERAGE": "평균", "IFERROR": "원래 수식이 정한 오류 처리", "AND": "모든 조건 확인", "OR": "하나 이상의 조건 확인"}
    for i, table in enumerate(result["tables"], 1):
        lines += [f'def step_{i:03d}(table):', f'    # 대상 표: {table["id"]!r}, 시트: {table["sheet"]!r}, 범위: {table["range"]!r}',
                  '    # rows에는 이미 지원 수식을 계산한 값이 들어 있습니다.']
        for rule in stages[table["id"]]["formula_columns"]:
            explanation = ' / '.join(meanings.get(name, name + " 함수(지원 여부는 formulas.json 확인)") for name in rule["functions"]) or "원래 수식의 사칙연산·셀 참조"
            lines.append(f'    # {rule["column"]!r}: {explanation!r}. 원본 예: {rule["examples"][0]["formula"]!r}')
        lines += [
                  '    # columns는 열 이름, source_rows는 원래 Excel 행 번호입니다.',
                  '    # 합계 후보 행과 미지원 수식의 빈 값은 임의로 삭제하거나 0으로 바꾸지 마세요.',
                  '    # 현재는 추가 업무 규칙 없이 표를 그대로 반환합니다. 필요한 처리만 여기에 추가하세요.',
                  '    return table', '']
    dispatch = ', '.join(f'{t["id"]!r}: step_{i:03d}' for i, t in enumerate(result["tables"], 1))
    lines += ['def process_tables(tables):', '    # 각 표를 담당 함수에 연결합니다. 데이터 출처를 바꿀 때는 이 코드를 수정하지 않습니다.',
              f'    steps = {{{dispatch}}}', '    return [steps[table["id"]](table) for table in tables]', '']
    (output / "PROCESS" / "table_steps.py").write_text('\n'.join(lines), encoding="utf-8")
    (output / "README.md").write_text(README, encoding="utf-8")


README = '''# 엑셀 처리 프로젝트 사용법

`python -X utf8 main.py`를 실행하면 DATA에서 자료를 읽고 PROCESS에서 계산한 뒤
새 결과를 `outputs/run-날짜-시간/`에 저장합니다. 프로젝트 폴더 밖에서 실행해도 됩니다.
기본 Python 패키지는 requirements.txt에 있습니다. Excel 경로와 VBA 추출의 추가 준비물은
requirements-excel.txt와 requirements-vba.txt를 참고하세요. 자동 설치하지 않습니다.

## 어디를 고치면 되나요?

| 바꾸려는 내용 | 수정할 파일 |
| --- | --- |
| 같은 양식의 다른 엑셀 파일 | DATA/sources.py의 SOURCE |
| 다른 Python·JSON·CSV·DB에서 자료 가져오기 | DATA/load.py의 load_data |
| 표별 정리·결합·집계 규칙 | PROCESS/table_steps.py |
| 원래 수식·표 범위가 달라짐 | 변경 내용을 확인한 뒤 새 프로젝트 생성 또는 PROCESS 규칙 수정 |

각 Python 파일에 역할과 수정 위치를 한국어 주석으로 설명했습니다. `pipeline_lib/`는
공통 읽기·계산 지원 코드이므로 보통 수정하지 않습니다. 수식은 `PROCESS/workbook_rules.json`에
원본 식으로 보존되며 Python 계산기가 지원하는 부분만 실행합니다. 개별 수식을 모두 자연어
업무 코드로 번역한 것은 아닙니다. VBA·외부 동작은 자동 실행/번역하지 않습니다.

## DATA만 바꿔서 재사용하기

`load_data()`는 `{"cells": {"시트!셀": 값}, "source_name": "자료 이름"}`을 반환합니다.
cells의 필수 주소·자료형은 PROCESS/workbook_rules.json의 inputs에 있습니다.
문자열 ID의 앞자리 0·대소문자, 숫자와 빈 값(None), 날짜를 보존하세요.
날짜는 ISO 문자열을 사용할 수 있습니다. 계산할 수식은 DATA에서 전달하지 않습니다.
사번을 숫자로 바꾸는 등 자료형이 달라지면 어느 입력이 다른지 안내하고 중단합니다.

다른 Python에서 행 목록을 받는 경우 `replace_tables(data, {"표ID": 행목록}, rules)`로
같은 열 이름을 원래 위치에 연결할 수 있습니다. 각 행은 {"열이름": 값} 사전이며
수식 전용 열은 보내지 않습니다. 행별 필요한 열은 rules의 tables/input_rows에 나옵니다.
이 함수는 원본 data를 변경하지 않고 새 자료를 반환합니다. 표 밖의 설정 셀도 입력이므로
빠뜨리지 마세요. 모든 입력을 Python/JSON에서 제공하면 Excel 파일 없이 실행할 수 있습니다.

예시(주소·표 ID는 실제 rules에 맞춰 사용):
```python
# DATA/load.py 안에서 필요한 자료만 연결합니다. PROCESS는 그대로 둡니다.
# data = {"cells": 내가_읽은_입력값, "source_name": "사내 데이터"}
# data = replace_tables(data, {"s02_t01": 내가_읽은_행목록}, rules)
# return data
```

## 그대로 재사용할 수 있는 범위

데이터 값과 출처만 바뀌고 표의 열 의미·행 배치·자료형이 같으면 PROCESS를 재사용합니다.
현재 자동 생성은 고정 셀 수식을 보존하므로 행 추가·삭제, 열 이동, 수식·이름 범위·병합 변경은
자동 확장하지 않습니다. 다른 구조를 기존 규칙에 억지로 맞추지 않고 재검토를 요청합니다.
원본이 이미 빈 값인 칸과 이후 비게 된 값도 실제 빈 값으로 처리하며 0으로 채우지 않습니다.

처음 분석한 표·수식·관찰 근거는 프로젝트 루트의 JSON과 report.md에 있습니다.
재실행 결과는 outputs 아래에 저장합니다. tables/는 수식 계산 결과이고,
processed_tables.json은 table_steps.py의 추가 업무 처리 결과입니다.
미지원 수식·오류·VBA 미이전은 summary.json과 formulas.json에서 확인하세요.
새 입력에는 과거 저장 캐시를 정답으로 재사용하지 않습니다.

`transform.py`는 이전 사용법을 위한 구조 재분석 명령입니다. PROCESS를 재사용하려면
항상 main.py를 쓰세요. 기존 결과 폴더는 덮어쓰지 않으며 원본 Excel도 수정하지 않습니다.
모델용 요약의 행·글자 제한은 유지됩니다. 전체 셀·수식 파일은 필요한 때만 읽으세요.
'''
