"""Opt-in semantic comparison of a delivered PPTX and its rebuilt counterpart.

No Office process, rendering, modification, approval or rebuild is performed.
Package timestamps and core properties are not part of this comparison.
"""
from __future__ import annotations

import argparse
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _xml(raw, relationships=None, seen=()):
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise ValueError("지원하지 않는 XML 선언입니다.")
    root = ET.fromstring(raw)
    for element in root.iter():
        # Object identifiers and generated names do not affect displayed content.
        if element.tag.endswith("}cNvPr"):
            element.attrib.pop("id", None)
            element.attrib.pop("name", None)
        for key, value in list(element.attrib.items()):
            if key.startswith("{" + REL_NS + "}") and relationships is not None:
                rel = relationships[value]
                element.set(key, rel.target_ref if rel.is_external else _part(rel.target_part, seen))
        element.attrib.update(sorted(element.attrib.items()))
    return ET.canonicalize(ET.tostring(root, encoding="unicode"))


def _part(part, seen=()):
    name = str(part.partname)
    if name in seen:
        return "cycle:" + part.content_type
    raw = part.blob
    if raw.startswith(b"PK"):
        # Embedded XLSX: compare cell values/formulas/styles, not ZIP/core dates.
        with ZipFile(BytesIO(raw)) as archive:
            _bounded(archive)
            value = [(n, _xml(archive.read(n)) if n.endswith((".xml", ".rels"))
                      else hashlib.sha256(archive.read(n)).hexdigest())
                     for n in sorted(archive.namelist()) if not n.startswith("docProps/") and not n.endswith("/")]
        raw = json.dumps(value, ensure_ascii=True).encode()
    elif name.endswith(".xml"):
        raw = _xml(raw, part.rels, (*seen, name)).encode()
    return hashlib.sha256(raw).hexdigest()


def _bounded(archive):
    entries = archive.infolist()
    if (len(entries) > 10000 or sum(e.file_size for e in entries) > 256 * 1024 * 1024
            or any(e.flag_bits & 1 for e in entries)):
        raise ValueError("검사 가능한 PPTX/XLSX 크기를 넘었거나 암호화되어 있습니다.")


def snapshot(path):
    from pptx import Presentation
    path = Path(path)
    if path.suffix.lower() != ".pptx" or path.stat().st_size > 40 * 1024 * 1024:
        raise ValueError("40MB 이하의 PPTX 파일을 사용해 주세요.")
    with ZipFile(path) as archive:
        _bounded(archive)
    deck = Presentation(path)

    def shape_record(shape, part):
        item = {"type": int(shape.shape_type),
                "geometry": [shape.left, shape.top, shape.width, shape.height, shape.rotation],
                "text": shape.text if shape.has_text_frame else None,
                "nativeXml": _xml(ET.tostring(shape._element), part.rels)}
        if shape.has_table:
            item["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
        if shape.has_chart:
            item["chart"] = _part(shape.chart.part)
        if hasattr(shape, "shapes"):
            item["children"] = [shape_record(child, part) for child in shape.shapes]
        return item

    return {"size": [deck.slide_width, deck.slide_height], "slides": [
        {"shapes": [shape_record(s, slide.part) for s in slide.shapes],
         "slideXml": _xml(ET.tostring(slide._element), slide.part.rels),
         "layout": _part(slide.slide_layout.part)} for slide in deck.slides]}


def compare(expected, rebuilt):
    left, right = snapshot(expected), snapshot(rebuilt)
    differences = []

    def walk(a, b, path):
        if len(differences) >= 20 or a == b:
            return
        if isinstance(a, dict) and isinstance(b, dict) and a.keys() == b.keys():
            for key in a:
                walk(a[key], b[key], f"{path}.{key}")
        elif isinstance(a, list) and isinstance(b, list) and len(a) == len(b):
            for index, (x, y) in enumerate(zip(a, b)):
                walk(x, y, f"{path}[{index}]")
        else:
            differences.append(path)

    walk(left, right, "deck")
    return {"ok": not differences, "code": "rebuild_matches" if not differences else "rebuild_differs",
            "differences": differences, "differenceLimit": 20, "visualReview": "not_performed",
            "basis": "native objects, text, tables, charts, layout and embedded data; no ZIP timestamps"}


def main(argv=None):
    parser = argparse.ArgumentParser(description="선택형 PPT 재빌드 내용 및 배치 비교")
    parser.add_argument("--expected", type=Path, required=True)
    parser.add_argument("--rebuilt", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        result = compare(args.expected, args.rebuilt)
    except Exception as exc:
        result = {"ok": False, "code": "comparison_failed", "message": str(exc)}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
