"""Shared, bounded slide composition for Python and Office. No generated code.

Coordinates are points on the actual page, not browser pixels. Conservative text
estimates reject dense input instead of truncating it or shrinking body text.
Estimates are not a substitute for rendering and inspecting every slide.
"""
from __future__ import annotations

import math
import re
import unicodedata
from . import report_facts


class DesignError(ValueError):
    pass


def _text(value, limit=200):
    if not isinstance(value, str) or len(value) > limit:
        raise DesignError("슬라이드의 제목·설명 길이와 형식을 확인해 주세요.")
    return value


def prepare(spec, data):
    facts, validation = report_facts.resolve(spec, data['sections'])
    bind = lambda s: report_facts.bind(s, facts)
    for key in ('title', 'subtitle'):
        data[key] = bind(data[key])
    raw_rows = spec.get('sections', spec.get('slides'))
    for raw, row in zip(raw_rows, data['sections']):
        for key in ('title', 'body'):
            row[key] = bind(row[key])
        row['bullets'] = [bind(s) for s in row['bullets']]
        if 'table' in row:
            row['table']['headers'] = [bind(s) for s in row['table']['headers']]
            row['table']['rows'] = [[bind(s) for s in r] for r in row['table']['rows']]
        row['eyebrow'] = bind(_text(raw.get('eyebrow', ''), 70))
        row['takeaway'] = bind(_text(raw.get('takeaway', ''), 160))
        row['source'] = bind(_text(raw.get('source', ''), 180))
        row['layout'] = raw.get('layout', 'auto')
        if row['layout'] not in ('auto', 'summary', 'evidence', 'comparison', 'actions'):
            raise DesignError("PPT 구성은 auto, summary, evidence, comparison, actions 중 선택해 주세요.")
        kpis = raw.get('kpis', [])
        if not isinstance(kpis, list) or len(kpis) > 3:
            raise DesignError("한 장의 핵심 지표는 3개 이하로 나누어 주세요.")
        row['kpis'] = []
        for item in kpis:
            if not isinstance(item, dict) or not isinstance(item.get('fact'), str) or item['fact'] not in facts:
                raise DesignError("핵심 지표는 계산한 공통 수치의 fact ID로 연결해 주세요.")
            fact = facts[item['fact']]
            row['kpis'].append({'label': _text(fact['label'], 32), 'value': fact['display'],
                                'unit': _text(fact['unit'], 12)})
        if 'chart' in row:
            row['chart']['title'] = bind(_text(raw['chart'].get('title', ''), 90))
            if len(row['chart']['categories']) > 12 or len(row['chart']['series']) > 4:
                raise DesignError("PPT 차트는 항목 12개·계열 4개 이하로 나누어 주세요.")
            if any(len(s) > 22 for s in row['chart']['categories']) or any(len(s['name']) > 24 for s in row['chart']['series']):
                raise DesignError("차트의 항목명과 범례를 짧게 정리해 주세요.")
    # Explicit theme tokens allow an inspected template's palette to be followed.
    # They do not claim exact reconstruction of arbitrary masters or layouts.
    tokens = {'title': '17324D', 'accent': '087F8C', 'text': '30465B',
              'muted': '52667C', 'background': 'FFFFFF', 'tint': 'EEF4F7'}
    if 'designPreset' in spec:
        from .ppt_workflow import DESIGNS
        if spec['designPreset'] not in DESIGNS:
            raise DesignError('지원하는 PPT 디자인을 선택해 주세요.')
        tokens.update(DESIGNS[spec['designPreset']][2])
    custom = spec.get('presentationTheme', {})
    if not isinstance(custom, dict) or set(custom) - set(tokens):
        raise DesignError("PPT 색상 설정 항목을 확인해 주세요.")
    for key, value in custom.items():
        if not isinstance(value, str) or not re.fullmatch('[0-9A-Fa-f]{6}', value):
            raise DesignError("PPT 색상은 # 없이 여섯 자리 색상값으로 지정해 주세요.")
        tokens[key] = value.upper()
    def luminance(value):
        channels = [int(value[i:i+2], 16)/255 for i in (0, 2, 4)]
        linear = [v/12.92 if v <= .04045 else ((v+.055)/1.055)**2.4 for v in channels]
        return sum(v*w for v, w in zip(linear, (.2126,.7152,.0722)))
    for front, back in (('text','background'), ('text','tint'), ('title','background'), ('muted','background'), ('accent','background')):
        a, b = sorted((luminance(tokens[front]), luminance(tokens[back])))
        if (b+.05)/(a+.05) < 4.5:
            raise DesignError("글자와 배경 색상이 너무 비슷합니다. 읽기 쉬운 대비의 색상을 선택해 주세요.")
    data['presentationTheme'] = tokens
    font = spec.get('presentationFont', 'Malgun Gothic')
    if not isinstance(font,str) or not font.strip() or len(font)>80:
        raise DesignError('글꼴 이름을 확인해 주세요.')
    if font != 'Malgun Gothic':
        from .ppt_workflow import installed_fonts
        if font not in installed_fonts():
            raise DesignError('양식의 글꼴을 이 PC에서 확인하지 못했습니다. 설치된 글꼴을 선택해 주세요. 자동 다운로드하지 않습니다.')
    data['presentationFont'] = font
    return validation


def _lines(text, width, size):
    # Hangul is roughly one em, Latin half an em; reserve wrapping tolerance.
    capacity = max(1, (width - 12) / (size * 1.08))
    return sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(c) in 'WF' else .56 for c in line) / capacity)) for line in text.split('\n'))


def plan(data, width=960., height=540.):
    if not (600 <= width <= 1800 and 400 <= height <= 1100 and 1.25 <= width / height <= 2.1):
        raise DesignError("이 양식의 화면 비율은 자동 배치를 지원하지 않습니다. 4:3 또는 16:9 양식을 선택해 주세요.")
    margin, right, bottom = width * .05, width * .95, height - 48
    usable = right - margin
    pages = []
    for number, row in enumerate(data['sections'], 1):
        elements = []

        def text(value, x, y, w, h, size=18, color='text', bold=False):
            if not value:
                return
            if _lines(value, w, size) * size * 1.27 + 4 > h:
                raise DesignError(f"{number}장 글자가 배치 공간보다 많습니다. 문장을 줄이거나 장을 나누어 주세요.")
            elements.append(dict(kind='text', text=value, x=x, y=y, w=w, h=h, size=size, color=color, bold=bold))

        text(row.get('eyebrow', ''), margin, 12, usable, 20, 11, 'muted')
        text(row['title'], margin, 37, usable, 48, 30, 'title', True)
        text(row.get('takeaway', ''), margin, 91, usable, 30, 17, 'text')
        text(row.get('source', ''), margin, height-31, usable-45, 22, 10, 'muted')
        text(f'{number:02d}', right-30, height-31, 30, 22, 10, 'muted')
        top = 139.
        if row.get('kpis'):
            kpi_w = usable / len(row['kpis'])
            for index, kpi in enumerate(row['kpis']):
                x = margin + index*kpi_w
                text(kpi['label'] + (' · '+kpi['unit'] if kpi['unit'] else ''), x, top, kpi_w-16, 26, 15, 'muted')
                text(kpi['value'], x, top+28, kpi_w-16, 57, 38, 'accent', True)
            top += 103
        visuals = [k for k in ('chart', 'table', 'image') if row.get(k)]
        prose = '\n\n'.join(([row['body']] if row['body'] else []) + row['bullets'])
        if len(visuals) > 2 or (len(visuals) == 2 and prose):
            raise DesignError(f"{number}장은 표·차트·설명이 너무 많습니다. 근거와 설명을 별도 장으로 나누어 주세요.")
        gap = 28.
        slots = []
        if len(visuals) == 2:
            slots = [(margin, top, (usable-gap)/2, bottom-top), (margin+(usable+gap)/2, top, (usable-gap)/2, bottom-top)]
        elif visuals and prose:
            visual_w = usable*.64
            slots = [(margin, top, visual_w, bottom-top)]
            text(prose, margin+visual_w+gap, top+8, usable-visual_w-gap, bottom-top-8)
        elif visuals:
            slots = [(margin, top, usable, bottom-top)]
        elif prose:
            text(prose, margin, top+10, usable, bottom-top-10, 21)
        for kind, (x, y, w, h) in zip(visuals, slots):
            element = dict(kind=kind, x=x, y=y, w=w, h=h)
            if kind == 'table':
                table = row['table']
                all_rows = [table['headers'], *table['rows']]
                # Give longer text columns more space, within a bounded weight.
                weights = [min(3.5, max(1, max(sum(1 if unicodedata.east_asian_width(c) in 'WF' else .56 for c in r[j]) for r in all_rows)/9)) for j in range(len(table['headers']))]
                widths = [w*v/sum(weights) for v in weights]
                heights = [max(36, max(_lines(cell, widths[j], 16) for j, cell in enumerate(r))*21+16) for r in all_rows]
                if sum(heights) > h:
                    raise DesignError(f"{number}장 표를 읽기 좋은 크기로 담을 수 없습니다. 행·열을 나누어 주세요.")
                element.update(h=sum(heights), columnWidths=widths, rowHeights=heights, size=16)
            elif kind == 'chart':
                if h < 180 or w < 260:
                    raise DesignError(f"{number}장 차트 공간이 작습니다. 지표나 설명을 다른 장으로 옮겨 주세요.")
                title = row['chart'].get('title', '')
                if title:
                    text(title, x, y, w, 27, 15, 'muted', True)
                    element.update(y=y+31, h=h-31)
            elements.append(element)
        if not prose and not visuals and not row.get('kpis'):
            raise DesignError(f"{number}장에 제목 외의 핵심 내용이 없습니다.")
        for e in elements:
            if min(e['x'], e['y']) < 0 or e['x']+e['w'] > width+.01 or e['y']+e['h'] > height+.01:
                raise DesignError("슬라이드 영역을 벗어나는 배치를 저장하지 않았습니다.")
        pages.append({'elements': elements})
    return {'width': width, 'height': height, 'pages': pages, 'font': data.get('presentationFont','Malgun Gothic'), 'theme': data['presentationTheme']}


def render_python(data, draft, template):
    import base64
    from io import BytesIO
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.enum.text import MSO_ANCHOR
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from pptx.oxml.xmlchemy import OxmlElement
    prs = Presentation(str(template)) if template else Presentation()
    if not template:
        prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    for relation in list(prs.slides._sldIdLst):
        prs.part.drop_rel(relation.rId)
        prs.slides._sldIdLst.remove(relation)
    design = data['presentationPlan']
    palette = design['theme']
    color = lambda name: RGBColor.from_string(palette.get(name, name))
    native = {'text': 0, 'tables': 0, 'charts': 0, 'images': 0}
    blank = min(prs.slide_layouts, key=lambda layout: len(layout.placeholders))

    def typeface(font):
        font.name = design['font']
        # Latin name alone leaves Korean runs dependent on the template's EA font.
        for tag in ('a:ea', 'a:cs'):
            child = font._rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
            if child is None:
                child = OxmlElement(tag)
                font._rPr.append(child)
            child.set('typeface', design['font'])
        font._rPr.set('lang', 'ko-KR')

    def format_text(frame, size, ink='text', bold=False):
        frame.word_wrap = True
        frame.margin_top = frame.margin_bottom = Pt(2)
        frame.margin_left = frame.margin_right = Pt(0)
        for p in frame.paragraphs:
            typeface(p.font)
            p.font.size = Pt(size)
            for run in p.runs:
                typeface(run.font)
            p.font.bold, p.font.color.rgb = bold, color(ink)
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = 1.15

    for row, page in zip(data['sections'], design['pages']):
        slide = prs.slides.add_slide(blank)
        for shape in list(slide.placeholders):
            shape._element.getparent().remove(shape._element)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color('background')
        for e in page['elements']:
            box = tuple(Pt(e[k]) for k in ('x', 'y', 'w', 'h'))
            kind = e['kind']
            if kind == 'text':
                shape = slide.shapes.add_textbox(*box)
                shape.text_frame.text = e['text']
                format_text(shape.text_frame, e['size'], e['color'], e['bold'])
                native['text'] += 1
            elif kind == 'table':
                content = row[kind]
                table = slide.shapes.add_table(len(content['rows'])+1, len(content['headers']), *box).table
                for j, width in enumerate(e['columnWidths']):
                    table.columns[j].width = Pt(width)
                for i, values in enumerate([content['headers'], *content['rows']]):
                    table.rows[i].height = Pt(e['rowHeights'][i])
                    for j, value in enumerate(values):
                        cell = table.cell(i, j)
                        cell.text = value
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = color('title' if i == 0 else ('tint' if i % 2 else 'background'))
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        format_text(cell.text_frame, 16, 'background' if i == 0 else 'text', i == 0)
                        cell.margin_left = cell.margin_right = Pt(6)
                native['tables'] += 1
            elif kind == 'chart':
                content = row[kind]
                chart_data = CategoryChartData()
                chart_data.categories = content['categories']
                for s in content['series']:
                    chart_data.add_series(s['name'], s['values'])
                kinds = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED,
                         'line': XL_CHART_TYPE.LINE, 'pie': XL_CHART_TYPE.PIE}
                chart = slide.shapes.add_chart(kinds[content['type']], *box, chart_data).chart
                # Some python-pptx chart templates use signed axis IDs. OOXML
                # readers expect UInt32; update IDs and their references together.
                axis_nodes = chart._chartSpace.xpath('.//c:axId | .//c:crossAx')
                axis_map = {value: str(100001+i) for i, value in enumerate(dict.fromkeys(node.get('val') for node in axis_nodes))}
                for axis_id in axis_nodes:
                    axis_id.set('val', axis_map[axis_id.get('val')])
                chart.has_title = False
                typeface(chart.font)
                chart.font.size = Pt(13)
                chart.font.color.rgb = color('text')
                chart.has_legend = len(content['series']) > 1 or content['type'] == 'pie'
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                series_colors = ['accent', 'title', '52667C', 'B05C32']
                for i, s in enumerate(chart.series):
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = color(series_colors[i])
                    s.format.line.color.rgb = color(series_colors[i])
                if content['type'] != 'pie':
                    typeface(chart.category_axis.tick_labels.font)
                    typeface(chart.value_axis.tick_labels.font)
                    chart.category_axis.tick_labels.font.size = Pt(13)
                    chart.value_axis.tick_labels.font.size = Pt(12)
                    if content['type'] in ('bar', 'column'):
                        values = [v for s in content['series'] for v in s['values']]
                        if min(values) >= 0:
                            chart.value_axis.minimum_scale = 0
                    chart.value_axis.has_major_gridlines = True
                    chart.value_axis.major_gridlines.format.line.color.rgb = color('tint')
                if len(content['categories']) <= 6:
                    plot = chart.plots[0]
                    plot.has_data_labels = True
                    typeface(plot.data_labels.font)
                    plot.data_labels.font.size = Pt(12)
                    plot.data_labels.font.color.rgb = color('text')
                    if content['type'] in ('bar', 'column'):
                        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
                native['charts'] += 1
            elif kind == 'image':
                im = row[kind]
                picture = slide.shapes.add_picture(BytesIO(base64.b64decode(im['data'])), box[0], box[1], height=box[3])
                if picture.width > box[2]:
                    ratio = box[2]/picture.width
                    picture.width, picture.height = int(picture.width*ratio), int(picture.height*ratio)
                picture.left += int((box[2]-picture.width)/2)
                native['images'] += 1
    prs.core_properties.title, prs.core_properties.subject = data['title'], data['subtitle']
    prs.save(str(draft))
    return native
