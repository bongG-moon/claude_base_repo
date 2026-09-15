"""Bounded Korean PPT choices and local template inspection. No downloads/code evaluation."""
from __future__ import annotations

from collections import Counter
from pathlib import Path
import posixpath
import re
import zipfile


DESIGNS = {
    'business': ('깔끔한 업무 보고형', '흰 배경·남색 제목·청록 강조',
                 {'title':'17324D','accent':'087F8C','text':'30465B','muted':'52667C','background':'FFFFFF','tint':'EEF4F7'}),
    'monochrome': ('흑백 간결형', '흰 배경·짙은 회색 글자·절제된 강조',
                   {'title':'222222','accent':'444444','text':'333333','muted':'555555','background':'FFFFFF','tint':'F0F0F0'}),
    'warm': ('따뜻한 설명형', '밝은 크림 배경·갈색 제목·차분한 강조',
             {'title':'45342B','accent':'80452C','text':'41382F','muted':'65584B','background':'FAF7F0','tint':'EEE7DC'}),
}


def choices(spec, template=None, *, for_preview=False):
    from .business_artifacts import ArtifactError, _failure
    try:
        if not isinstance(spec, dict):
            raise ArtifactError('invalid_choice', 'PPT 제작 조건의 형식을 확인해 주세요.')
        for key, allowed in (('creationMode', ('new','reference','saved')),
                             ('referenceMode', ('style','preserve')),
                             ('designPreset', tuple(DESIGNS))):
            if key in spec and spec[key] not in allowed:
                raise ArtifactError('invalid_choice', '지원하는 PPT 제작 방식을 선택해 주세요.')
        if 'slideCount' in spec and (type(spec['slideCount']) is not int or not 1 <= spec['slideCount'] <= 60):
            raise ArtifactError('invalid_choice', '장수는 1~60장 사이로 선택해 주세요.')
        for key in ('purpose','audience'):
            if key in spec and (not isinstance(spec[key], str) or not spec[key].strip() or len(spec[key]) > 200):
                raise ArtifactError('invalid_choice', '목적과 대상을 짧은 문장으로 알려 주세요.')
        known = {k:spec[k] for k in ('creationMode','referenceMode','purpose','audience','slideCount','designPreset','designReview') if k in spec}
        if spec.get('creationMode') == 'new' and (template or 'referenceMode' in spec):
            raise ArtifactError('invalid_choice', '새 디자인 제작과 기존 양식 유지 조건이 함께 지정되었습니다. 제작 방식을 확인해 주세요.')
        result = {'ok':False,'status':'input_required','preservedChoices':known,'waitForUser':True}
        if 'creationMode' not in spec:
            return {**result,'stage':'method','missing':['creationMode'],
                    'question':'어떤 방식으로 PPT를 만들까요?',
                    'options':['내용에 맞춰 새 디자인으로 만들기','기존 PPT·회사 양식을 첨부해서 만들기','이전에 저장한 내 양식으로 만들기']}
        if spec['creationMode'] != 'new':
            if not template:
                return {**result,'stage':'reference_file','missing':['template'],
                        'question':'참고할 PPTX 파일을 첨부하거나 이 PC의 파일 경로를 알려 주세요.'}
            if 'referenceMode' not in spec:
                return {**result,'stage':'reference_scope','missing':['referenceMode'],
                        'question':'기존 양식을 어느 정도 유지할까요?',
                        'options':['색감·폰트·분위기를 참고해서 새로 구성','기존 슬라이드 배치와 개체를 유지하고 내용 교체']}
        missing = [key for key in ('purpose','audience','slideCount') if key not in spec]
        if missing:
            return {**result,'stage':'brief','missing':missing,'question':'아직 알려주지 않은 목적·대상·장수만 알려 주세요.'}
        if spec['creationMode'] == 'new' and 'designPreset' not in spec:
            return {**result,'stage':'design','missing':['designPreset'],
                    'question':'새 PPT의 디자인을 선택해 주세요. 기존 PPT를 참고하는 방식으로 바꿀 수도 있습니다.',
                    'designOptions':[{'id':key,'label':name,'description':description} for key,(name,description,_) in DESIGNS.items()],
                    'referenceOption':'참고 PPT 첨부로 변경'}
        if for_preview:
            return {'ok':True,'status':'preview_ready','stage':'preview_ready','selection':known}
        review = spec.get('designReview')
        if review is None:
            return {**result,'status':'preview_required','stage':'design_preview','missing':['designReview'],'waitForUser':False,
                    'question':'전체 구성과 실제 대표 슬라이드 미리보기를 준비한 뒤 승인을 기다리세요. 최종 PPT는 아직 만들지 않습니다.',
                    'previewCommand':'business ppt-design-preview --spec FULL_JOB.json --output NEW_DRAFT.pptx'}
        if (not isinstance(review,dict) or set(review) != {'specSha256','previewPath','previewSha256','confirmed'}
                or type(review.get('confirmed')) is not bool
                or any(not isinstance(review.get(k),str) or not re.fullmatch('[a-f0-9]{64}',review[k]) for k in ('specSha256','previewSha256'))
                or not isinstance(review.get('previewPath'),str) or not Path(review['previewPath']).is_absolute()
                or len(review['previewPath'])>2048 or Path(review['previewPath']).suffix.lower()!='.pptx'):
            raise ArtifactError('invalid_choice','대표 디자인 확인 정보를 다시 준비해 주세요.')
        if not review['confirmed']:
            return {**result,'stage':'design_confirm','missing':['designReview.confirmed'],
                    'question':'구성과 대표 슬라이드를 확인해 주세요. 이대로 제작 / 수정 요청 / 참고 PPT로 변경 중 선택해 주세요.'}
        return {'ok':True,'status':'choices_ready','stage':'ready','selection':known,
                'message':'대표 디자인 승인 조건이 입력되었습니다. 생성기는 원본 명세와 미리보기 일치 여부를 확인한 뒤 제작합니다.'}
    except Exception as exc:
        return _failure(exc)


def design_digest(spec, template=None):
    import hashlib
    import json
    from .business_artifacts import _source
    reference = hashlib.sha256(_source(Path(template),('.pptx',)).read_bytes()).hexdigest() if template else None
    images=[]
    for row in spec.get('sections',spec.get('slides',[])):
        if isinstance(row,dict) and isinstance(row.get('image'),dict):
            from .business_artifacts import _image
            image=_image(row['image'])
            images.append(hashlib.sha256(image['data'].encode('ascii')).hexdigest())
    data = {'spec':{k:v for k,v in spec.items() if k!='designReview'},'templateSha256':reference,'imageSha256':images}
    return hashlib.sha256(json.dumps(data,ensure_ascii=True,sort_keys=True,separators=(',',':'),allow_nan=False).encode()).hexdigest()


def verify_design_review(spec, template=None):
    import hashlib
    from .business_artifacts import ArtifactError, _source
    review=spec['designReview']
    if review['specSha256']!=design_digest(spec,template):
        raise ArtifactError('design_review_changed','구성·내용·디자인 또는 참고 양식이 달라졌습니다. 새 대표 미리보기를 보여주고 다시 확인해 주세요.')
    preview=_source(Path(review['previewPath']),('.pptx',))
    if hashlib.sha256(preview.read_bytes()).hexdigest()!=review['previewSha256']:
        raise ArtifactError('design_preview_changed','대표 미리보기 파일이 변경되었습니다. 다시 확인해 주세요.')


def installed_fonts():
    """Names only; no font installation or access to other user profiles."""
    result = set()
    if __import__('os').name == 'nt':
        import winreg
        for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                with winreg.OpenKey(hive, r'SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts') as key:
                    for i in range(winreg.QueryInfoKey(key)[1]):
                        name = winreg.EnumValue(key, i)[0]
                        result.add(re.sub(r'\s*\([^)]*\)$', '', name).strip())
            except OSError:
                pass
    return result


def analyze(template):
    from .business_artifacts import inspect_template, _xml, A, P
    import hashlib
    inspected = inspect_template(Path(template))
    if not inspected.get('ok') or inspected.get('hasExternalRelationships'):
        return inspected if not inspected.get('ok') else {'ok':False,'status':'blocked','code':'external_template_links',
                'message':'외부 연결이 포함된 양식은 자동 분석하지 않습니다.'}
    fonts, colors, pages, theme = Counter(), Counter(), [], {}
    with zipfile.ZipFile(template) as package:
        order={}
        if 'ppt/_rels/presentation.xml.rels' in package.namelist():
            rels=_xml(package.read('ppt/_rels/presentation.xml.rels'))
            targets={r.get('Id'):posixpath.normpath(r.get('Target','').lstrip('/') if r.get('Target','').startswith('/')
                                                 else 'ppt/'+r.get('Target','')) for r in rels}
            presentation=_xml(package.read('ppt/presentation.xml'))
            for i,slide in enumerate(presentation.findall(f'.//{{{P}}}sldId')):
                rid=slide.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                order[targets.get(rid,'')]=i
        for name in sorted(package.namelist()):
            if name.startswith('ppt/theme/') and name.endswith('.xml'):
                root = _xml(package.read(name))
                for node in root.findall(f'.//{{{A}}}clrScheme/*'):
                    if len(node):
                        color = node[0].get('val') if node[0].tag.endswith('srgbClr') else node[0].get('lastClr')
                        if color and re.fullmatch('[a-fA-F0-9]{6}',color):
                            theme.setdefault(node.tag.split('}')[-1],color.upper())
                for node in root.findall(f'.//{{{A}}}font'):
                    if node.get('script') == 'Hang' and node.get('typeface'):
                        fonts[node.get('typeface')] += 1
            if not re.fullmatch(r'ppt/slides/slide\d+\.xml', name):
                continue
            root = _xml(package.read(name))
            shapes=[]
            for node in root.findall(f'.//{{{P}}}sp') + root.findall(f'.//{{{P}}}graphicFrame'):
                identity=node.find(f'.//{{{P}}}cNvPr')
                if identity is None:
                    continue
                xfrm=node.find(f'.//{{{A}}}xfrm')
                geometry={}
                if xfrm is not None:
                    for part in xfrm:
                        geometry[part.tag.split('}')[-1]]=dict(part.attrib)
                # Text is only a bounded preview in this request, never a durable profile.
                text=' '.join(t.text or '' for t in node.findall(f'.//{{{A}}}t'))
                shapes.append({'shapeId':int(identity.get('id','0')),'name':identity.get('name','')[:120],
                               'kind':node.tag.split('}')[-1],
                               'textPreview':text[:160],'geometryEmu':geometry})
            for tag in ('latin','ea'):
                for node in root.findall(f'.//{{{A}}}{tag}'):
                    if node.get('typeface') and not node.get('typeface').startswith('+'):
                        fonts[node.get('typeface')]+=1
            for node in root.findall(f'.//{{{A}}}srgbClr'):
                if re.fullmatch('[a-fA-F0-9]{6}',node.get('val','')):
                    colors[node.get('val').upper()]+=1
            pages.append({'part':name,'sourceSlide':order.get(name),'shapes':shapes[:80], 'shapesTruncated':len(shapes)>80})
    if hashlib.sha256(Path(template).read_bytes()).hexdigest() != inspected['sha256']:
        return {'ok':False,'status':'blocked','code':'template_changed','message':'분석 중 양식이 변경되었습니다. 저장을 마친 뒤 다시 확인해 주세요.'}
    pages.sort(key=lambda page: page['sourceSlide'] if page['sourceSlide'] is not None else 9999)
    available=installed_fonts()
    return {'ok':True,'status':'analyzed','sha256':inspected['sha256'],'slideCount':inspected['slideCount'],
            'sizeEmu':inspected['sizeEmu'],'themeColors':theme,'declaredColors':[c for c,_ in colors.most_common(12)],
            'declaredFonts':[{'name':f,'available':f in available} for f,_ in fonts.most_common(12)],
            'slides':pages[:60], 'visualReview':'not_performed', 'savedToMemory':False,
            'message':'양식의 개체·명시된 색상·폰트를 확인했습니다. 상속된 최종 모양과 혼합 서식은 실제 미리보기에서 확인해야 합니다.'}


def quality(path):
    """Objective native-object checks; not an aesthetic or source accuracy claim."""
    import importlib.util
    if importlib.util.find_spec('pptx') is None:
        return {'status':'unavailable','message':'PPT 구조 검사 라이브러리가 준비되지 않았습니다.'}
    from pptx import Presentation
    deck=Presentation(str(path))
    issues=[]
    for index,slide in enumerate(deck.slides,1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left+shape.width > deck.slide_width+12700 or shape.top+shape.height > deck.slide_height+12700:
                issues.append({'slide':index,'shapeId':shape.shape_id,'code':'out_of_bounds','message':'개체가 슬라이드 영역을 벗어납니다.'})
            frames=[shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.size is not None and run.font.size.pt < 10:
                            issues.append({'slide':index,'shapeId':shape.shape_id,'code':'tiny_text','message':'10pt 미만 글자가 있어 확인이 필요합니다.'})
    return {'status':'issues_found' if issues else 'checked','issues':issues[:100],
            'visualReview':'not_performed','fontFallback':'not_fully_verified','sourceAccuracy':'not_verified'}


def optional_archforge(path):
    """Use an already provisioned, reviewed checker version; never install it."""
    import importlib.metadata
    import json
    import subprocess
    import sys
    import tempfile
    try:
        version=importlib.metadata.version('archforge')
    except importlib.metadata.PackageNotFoundError:
        return {'status':'not_installed','automaticInstall':False}
    if version!='0.11.0':
        return {'status':'unreviewed_version','automaticInstall':False}
    try:
        with tempfile.TemporaryDirectory(prefix='company-ppt-lint-') as cwd:
            proc=subprocess.run([sys.executable,'-I','-m','archforge',str(Path(path).resolve()),'--profile','core','--no-config','--json'],
                                cwd=cwd,capture_output=True,timeout=30,
                                creationflags=getattr(__import__('os'),'CREATE_NO_WINDOW',0))
        if len(proc.stdout)>1_000_000:
            raise ValueError('too much output')
        parsed=json.loads(proc.stdout.decode('utf-8-sig'))
        if not isinstance(parsed,dict) or proc.returncode not in (0,1):
            raise ValueError('invalid linter response')
        summary=parsed.get('summary',{})
        if (not isinstance(summary,dict) or type(summary.get('pass')) is not bool
                or any(type(summary.get(key)) is not int or summary[key]<0 for key in ('error_count','warn_count'))):
            raise ValueError('invalid linter summary')
        # Do not relay raw slide text/private paths from an external checker.
        return {'status':'checked' if proc.returncode==0 and summary['pass'] and not summary['error_count'] and not summary['warn_count'] and not summary.get('incomplete') else 'issues_found',
                'errors':summary['error_count'],'warnings':summary['warn_count'],
                'version':version,'visualReview':'not_performed','rawOutputStored':False}
    except (OSError,ValueError,subprocess.TimeoutExpired):
        return {'status':'unavailable','message':'추가 PPT 검사기를 완료하지 못했습니다. 기본 검사 결과와 별개입니다.'}


def fill_template(data, draft, template, mappings):
    """Fill existing native slots; no arbitrary code, cloning, or screenshot slides."""
    from .business_artifacts import ArtifactError
    from .presentation_design import _lines
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    if not isinstance(mappings,list) or len(mappings)!=len(data['sections']):
        raise ArtifactError('template_mapping_required','각 장에서 사용할 원본 슬라이드와 교체할 개체를 먼저 지정해야 합니다.')
    deck=Presentation(str(template))
    selected=[]
    native={'text':0,'tables':0,'charts':0,'images':0}
    def fail():
        raise ArtifactError('template_mapping_invalid','양식의 교체 대상·남겨둘 개체·내용 분량을 다시 확인해야 합니다. 임의 새 배치로 바꾸지 않았습니다.')
    def replace(frame,text):
        # Mixed-run/paragraph styles cannot be flattened while claiming exact fidelity.
        if len(frame.paragraphs)!=1 or len(frame.paragraphs[0].runs)>1:
            raise ArtifactError('mixed_text_style','여러 문단·혼합 글꼴 개체는 양식 유지 자동 교체를 지원하지 않습니다. 분위기 참고 방식을 선택하거나 양식을 단순화해 주세요.')
        paragraph=frame.paragraphs[0]
        run=paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        run.text=text
    for row,mapping in zip(data['sections'],mappings):
        if not isinstance(mapping,dict) or set(mapping)-{'sourceSlide','title','body','table','chart','keepShapeIds'}:
            fail()
        idx=mapping.get('sourceSlide')
        if type(idx) is not int or idx<0 or idx>=len(deck.slides) or idx in selected:
            fail()
        selected.append(idx)
        slide=deck.slides[idx]
        shapes={s.shape_id:s for s in slide.shapes}
        used=set()
        keep=mapping.get('keepShapeIds',[])
        if not isinstance(keep,list) or any(type(v) is not int or v not in shapes for v in keep):
            fail()
        if row.get('bullets') or row.get('kpis') or row.get('image') or row.get('eyebrow') or row.get('takeaway') or row.get('source'):
            raise ArtifactError('template_content_unsupported','양식 유지 모드는 현재 제목·본문·표·일반 차트 교체를 지원합니다. 그 외 내용은 지원되는 칸에 명시적으로 통합하거나 분위기 참고를 선택해 주세요.')
        for field in ('title','body','table','chart'):
            value=row.get(field)
            if not value:
                if field in mapping:
                    fail()
                continue
            sid=mapping.get(field)
            if type(sid) is not int or sid not in shapes or sid in used or sid in keep:
                fail()
            used.add(sid)
            shape=shapes[sid]
            if field in ('title','body'):
                if not shape.has_text_frame:
                    fail()
                frame=shape.text_frame
                sizes=[r.font.size.pt for p in frame.paragraphs for r in p.runs if r.font.size]
                size=max(sizes or [30 if field=='title' else 18])
                if _lines(value,shape.width/12700,size)*size*1.3 > shape.height/12700:
                    raise ArtifactError('template_text_overflow','새 내용이 양식의 글자 칸에 비해 깁니다. 내용을 줄이거나 분위기 참고 방식으로 재구성해 주세요.')
                replace(frame,value)
                native['text']+=1
            elif field=='table':
                if not shape.has_table:
                    fail()
                rows=[value['headers'],*value['rows']]
                table=shape.table
                if len(rows)!=len(table.rows) or any(len(r)!=len(table.columns) for r in rows):
                    raise ArtifactError('template_table_size','기존 표와 새 표의 행·열 수가 다릅니다. 몰래 잘라 넣지 않았습니다.')
                for ri,values in enumerate(rows):
                    for ci,text in enumerate(values):
                        cell=table.cell(ri,ci)
                        if cell.is_merge_origin or cell.is_spanned:
                            fail()
                        if _lines(text,table.columns[ci].width/12700,16)*20 > table.rows[ri].height/12700:
                            raise ArtifactError('template_text_overflow','새 표 내용이 기존 셀보다 깁니다. 내용을 줄이거나 표 구성을 바꿔 주세요.')
                        replace(cell.text_frame,text)
                native['tables']+=1
            else:
                if not shape.has_chart or len(shape.chart.plots)!=1:
                    fail()
                chart=shape.chart
                from pptx.enum.chart import XL_CHART_TYPE
                types={'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'bar':XL_CHART_TYPE.BAR_CLUSTERED,
                       'line':XL_CHART_TYPE.LINE,'pie':XL_CHART_TYPE.PIE}
                if chart.chart_type!=types[value['type']] or len(chart.series)!=len(value['series']):
                    raise ArtifactError('template_chart_type','기존 차트의 유형·계열 수가 달라 양식을 그대로 유지할 수 없습니다.')
                cd=CategoryChartData()
                cd.categories=value['categories']
                for series in value['series']:
                    cd.add_series(series['name'],series['values'])
                chart.replace_data(cd)
                if value.get('title'):
                    chart.has_title=True
                    replace(chart.chart_title.text_frame,value['title'])
                native['charts']+=1
        # Avoid silently leaving old report data in unmapped editable objects.
        for sid,shape in shapes.items():
            if hasattr(shape,'shapes'):
                raise ArtifactError('template_group_unsupported','그룹 개체는 양식 유지 자동 교체를 지원하지 않습니다. 원본은 보존했습니다.')
            meaningful=(shape.has_text_frame and bool(shape.text.strip())) or shape.has_table or shape.has_chart
            if meaningful and sid not in used and sid not in keep:
                raise ArtifactError('template_unmapped_content','이전 내용이 남을 수 있는 개체가 있습니다. 교체할 내용과 유지할 로고·고정 문구를 먼저 구분해 주세요.')
        # Speaker notes can retain prior business content: clear on the new copy only.
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            slide.notes_slide.notes_text_frame.clear()
    original=list(deck.slides._sldIdLst)
    for i,relation in enumerate(original):
        deck.slides._sldIdLst.remove(relation)
        if i not in selected:
            deck.part.drop_rel(relation.rId)
    for i in selected:
        deck.slides._sldIdLst.append(original[i])
    deck.save(str(draft))
    return native
