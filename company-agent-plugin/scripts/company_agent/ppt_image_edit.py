"""Resize selected pictures in an existing PPTX; preserve originals and other shapes."""
from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path
import tempfile

from . import business_artifacts as a
from .business_safety import blocked_input
from .ppt_scene import fit_picture, number


def resize(spec,output,preview_directory=None):
    try:
        if (blocked:=blocked_input(spec)) is not None:
            return blocked
        a._check_content_scope(spec)
        output=a._target(Path(output),'.pptx')
        source=a._source(Path(spec.get('file','')),('.pptx',))
        checked=a.inspect_template(source)
        if not checked.get('ok'): return checked
        if checked['hasExternalRelationships']:
            raise a.ArtifactError('external_template_links','외부 연결이 있는 PPT는 자동 수정하지 않습니다.')
        if importlib.util.find_spec('pptx') is None:
            raise a.ArtifactError('ppt_edit_engine_unavailable','기존 그림 수정에 필요한 python-pptx가 없습니다. 자동 설치하거나 새 그림으로 대체하지 않았습니다.')
        fit=spec.get('fit','contain')
        if fit not in ('cover','contain','stretch'):
            raise a.ArtifactError('invalid_image_fit','fit은 cover(꽉 채움), contain(전체 보기), stretch(비율 변경)입니다.')
        selected=spec.get('slides',list(range(1,checked['slideCount']+1)))
        if not isinstance(selected,list) or not selected or len(selected)!=len(set(selected)) or any(type(i) is not int or not 1<=i<=checked['slideCount'] for i in selected):
            raise a.ArtifactError('invalid_image_selection','수정할 slides는 중복 없는 1부터 시작하는 장 번호입니다.')
        ids=spec.get('shapeIds',{})
        if not isinstance(ids,dict) or set(ids)-{str(n) for n in selected}:
            raise a.ArtifactError('invalid_image_selection','shapeIds는 선택한 장 번호별 개체 번호 목록입니다.')
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.util import Pt
        with tempfile.TemporaryDirectory(prefix='company-ppt-images-',ignore_cleanup_errors=True) as temp:
            work=Path(temp); copy=work/'source.pptx'; copy.write_bytes(source.read_bytes())
            if hashlib.sha256(copy.read_bytes()).hexdigest()!=checked['sha256']:
                raise a.ArtifactError('source_changed','PPT가 변경되었습니다. 다시 확인해 주세요.')
            deck=Presentation(str(copy)); edits=[]
            frame=spec.get('box')
            if frame is None:
                box=(0,0,deck.slide_width,deck.slide_height)
            else:
                if not isinstance(frame,dict) or set(frame)!={'x','y','w','h'}:
                    raise a.ArtifactError('invalid_image_box','box는 포인트 단위 x,y,w,h입니다.')
                box=tuple(Pt(number(frame[k],0 if k in ('x','y') else .01,deck.slide_width.pt if k in ('x','w') else deck.slide_height.pt)) for k in ('x','y','w','h'))
                if box[0]+box[2]>deck.slide_width or box[1]+box[3]>deck.slide_height:
                    raise a.ArtifactError('invalid_image_box','그림 영역이 슬라이드 밖으로 나갑니다.')
            for index in selected:
                images=[s for s in deck.slides[index-1].shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE]
                wanted=ids.get(str(index))
                if wanted is None:
                    if len(images)!=1:
                        raise a.ArtifactError('image_selection_required',f'{index}장에는 그림이 {len(images)}개입니다. ppt-analyze로 확인한 shapeIds를 지정해 주세요. 임의로 고르지 않았습니다.')
                else:
                    if not isinstance(wanted,list) or not wanted or any(type(i) is not int for i in wanted) or len(wanted)!=len(set(wanted)) or not set(wanted)<={s.shape_id for s in images}:
                        raise a.ArtifactError('invalid_image_selection','해당 장의 실제 그림 shapeId만 지정해 주세요.')
                    images=[s for s in images if s.shape_id in wanted]
                for picture in images:
                    if picture.rotation:
                        raise a.ArtifactError('rotated_image_unsupported','회전된 그림은 위치 변경 전에 회전 유지 여부를 확인해야 합니다.')
                    fit_picture(picture,box,fit)
                    edits.append({'slide':index,'shapeId':picture.shape_id,'fit':fit})
            draft=work/'result.pptx'; deck.save(str(draft))
            structural=a.inspect_template(draft)
            if not structural.get('ok') or structural['slideCount']!=checked['slideCount']:
                raise a.ArtifactError('ppt_validation_failed','수정한 PPT 구조 검사에 실패했습니다.')
            from .ppt_workflow import quality
            validation=quality(draft)
            visual=a._office({},draft,None,work,True)
            if hashlib.sha256(source.read_bytes()).hexdigest()!=checked['sha256']:
                raise a.ArtifactError('source_changed','수정 중 원본 PPT가 변경되었습니다. 결과를 전달하지 않았습니다.')
            previews=[]
            if visual.get('ok'):
                target=Path(preview_directory) if preview_directory else Path(tempfile.mkdtemp(prefix='company-ppt-review-'))
                a._reject_reparse(target)
                target.mkdir(exist_ok=preview_directory is None)
                for image in sorted((work/'preview').glob('slide-*.png')):
                    a._publish(image,target/image.name); previews.append(str(target/image.name))
            a._publish(draft,output)
        warnings=['그림의 크기/자르기만 수정했습니다. 그림 안 글자는 편집 가능한 텍스트로 바뀌지 않습니다.']
        if not visual.get('ok'): warnings.append(visual.get('message','이미지 미리보기를 확인하지 못했습니다.'))
        return {'ok':True,'status':'created' if visual.get('ok') and validation['status']=='checked' else 'partial',
                'outputPath':str(output),'sourcePreserved':True,'edits':edits,'slides':checked['slideCount'],
                'previews':previews,'warnings':warnings,'validation':{'structure':'passed','quality':validation,'render':visual,'visualReview':'required'}}
    except Exception as exc:
        return a._failure(exc)
