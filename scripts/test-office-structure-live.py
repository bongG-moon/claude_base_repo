"""Synthetic-only Office structure check, without accepting corporate source paths."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import time

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'company-agent-plugin/scripts'))
from company_agent.office_reader import _invoke, normalize
from office_read_test import run_test


def main():
    if sys.argv[1:] != ['--synthetic-only']:
        raise SystemExit('Use --synthetic-only; no existing source files are accepted.')
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document
    with tempfile.TemporaryDirectory(prefix='company-structure-',ignore_cleanup_errors=True) as folder:
        root=Path(folder)
        deck=Presentation()
        for i in range(4):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            group=slide.shapes.add_group_shape()
            group.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1)).text=f'그룹 테스트 {i+1}'
            table=slide.shapes.add_table(2,2,Inches(1),Inches(3),Inches(4),Inches(1)).table
            table.cell(0,0).text='부서'; table.cell(0,1).text='실적'
            table.cell(1,0).text='가팀'; table.cell(1,1).text='123'
        ppt=root/'four-slides.pptx'; deck.save(ppt)
        doc=Document(); doc.add_heading('실적보고',level=1)
        table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='부서'; table.cell(0,1).text='실적'
        table.cell(1,0).text='가팀'; table.cell(1,1).text='123'
        word=root/'table.docx'; doc.save(word)
        reports=[]
        for path,end in ((ppt,4),(word,20)):
            before=hashlib.sha256(path.read_bytes()).hexdigest()
            started=time.perf_counter()
            result=_invoke(normalize({'file':str(path),'start':1,'end':end,'maxChars':10000}))
            structures=[x.get('structure',{}) for x in result.get('items',[])]
            has_table=any(x.get('type')=='table-cell' and x.get('row')==2 and x.get('column')==2 for x in structures)
            has_groups=path==word or any('/group:' in x['location'] for x in result.get('items',[]))
            four=path==word or {1,2,3,4}=={x.get('slide') for x in structures}
            report={'file':path.name,'ok':result.get('ok') is True and has_table and has_groups and four,
                    'sourceUnchanged':before==hashlib.sha256(path.read_bytes()).hexdigest(),
                    'hasTableCoordinates':has_table,'groupRead':has_groups,'fourSlidesRead':four,
                    'elapsedMs':round((time.perf_counter()-started)*1000),'drmTested':False}
            reports.append(report); print(json.dumps(report),flush=True)
            if not report['ok']: print(json.dumps({'code':result.get('code'),'coverage':result.get('coverage'),'structures':structures},ensure_ascii=True),flush=True)
        independent=run_test(ppt,options={'start':1,'end':4,'max_chars':10000})
        report={'standalone':True,'ok':independent.get('ok') is True and independent.get('coverage',{}).get('end')==4,
                'sourceUnchanged':independent.get('sourceUnchanged'), 'elapsedMs':independent.get('elapsedMs'),'drmTested':False}
        reports.append(report); print(json.dumps(report),flush=True)
        if not reports[1]['ok']:
            from office_read_test import read_office, Preview
            import traceback
            try: read_office(word,'word',Preview(),{})
            except Exception: traceback.print_exc()
        return 0 if all(r['ok'] and r['sourceUnchanged'] for r in reports) else 1


if __name__=='__main__': raise SystemExit(main())
