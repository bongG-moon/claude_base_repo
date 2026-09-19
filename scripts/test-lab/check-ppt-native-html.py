"""Opt-in local roundtrip; actual source is never committed or modified."""
import argparse
from collections import Counter
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path
import sys
import re
import time
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'company-agent-plugin/scripts'))
from company_agent import ppt_html, ppt_html_import, artifact_delivery


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--source',required=True)
    parser.add_argument('--output',required=True)
    parser.add_argument('--diagnose',action='store_true')
    parser.add_argument('--browser',help='Explicit installed browser for this test, never a denial fallback')
    args=parser.parse_args()
    output=Path(args.output).resolve(); output.mkdir(parents=True,exist_ok=True)
    source=Path(args.source).resolve()
    before=hashlib.sha256(source.read_bytes()).hexdigest()
    actual=ppt_html_import.subprocess.run
    def trace(*a,**k):
        result=actual(*a,**k)
        if args.diagnose:
            print({'exit':result.returncode,'stdoutBytes':len(result.stdout),'stderr':result.stderr[-2500:].decode('utf-8','replace')})
        return result
    spec={'creationMode':'reference','title':'HTML 배치 변환 검증','purpose':'변환 검증','audience':'검토자','slideCount':5,
          'htmlSource':{'path':str(source)}}
    start=artifact_delivery.start(output/'state',output/'result.pptx')
    assert start.get('ok'),start
    work=start['workFile']; started=time.perf_counter()
    browser=Path(args.browser) if args.browser else ppt_html_import.browser_path()
    with patch.object(ppt_html_import.subprocess,'run',trace), patch.object(ppt_html_import,'browser_path',return_value=browser):
        preview=artifact_delivery.build(output/'state',work,'ppt-design-preview',spec)
    assert preview.get('ok'),preview
    preview_ms=round((time.perf_counter()-started)*1000)
    spec['designReview']={**preview['designReview'],'confirmed':True}
    # This is test approval of the fixture conversion, not a business sign-off.
    with patch.object(ppt_html_import,'capture',side_effect=AssertionError('HTML must not be rendered twice')):
        result=artifact_delivery.build(output/'state',work,'ppt',spec)
    assert result.get('ok'),result
    plan=ppt_html.prepare(spec)[0]['presentationPlan']
    saved=ppt_html.save_template(spec,output/'state'/'representative.html')
    assert saved.get('ok'),saved
    metadata=ppt_html.read_template(saved['outputPath'])
    assert len(metadata['layouts'])==5
    assert all(e.get('text','예시')=='예시' for p in metadata['layouts'] for e in p['elements'])
    reused={'creationMode':'saved','purpose':'새 자료','audience':'검토자','slideCount':5,
            'slides':[{'title':f'예시 {i+1}','layoutIndex':i+1,'texts':['새']*p['textSlots']} for i,p in enumerate(metadata['layouts'])]}
    reused_plan=ppt_html.prepare(reused,Path(saved['outputPath']))[0]['presentationPlan']
    assert [len(p['elements']) for p in reused_plan['pages']]==[len(p['elements']) for p in plan['pages']]
    final=artifact_delivery.publish(output/'state',work)
    assert final.get('ok'),final
    assert hashlib.sha256(source.read_bytes()).hexdigest()==before
    class Content(HTMLParser):
        def __init__(self): super().__init__(); self.depth=0; self.parts=[]
        def handle_starttag(self,tag,attrs):
            if tag=='section': self.depth+=1
        def handle_endtag(self,tag):
            if tag=='section': self.depth-=1
        def handle_data(self,value):
            if self.depth: self.parts.append(value)
    parser=Content(); parser.feed(source.read_text(encoding='utf-8-sig'))
    from pptx import Presentation
    actual=''.join(s.text for p in Presentation(final['outputPath']).slides for s in p.shapes if s.has_text_frame)
    assert re.sub(r'\s+','',''.join(parser.parts))==re.sub(r'\s+','',actual),'HTML/PPT content differs'
    receipt={'previewMs':preview_ms,'sourcePreserved':True,'textPreservedIgnoringWhitespace':True,'templateGeometryRoundtrip':True,'templateContentRemoved':True,'counts':dict(Counter(e['kind'] for p in plan['pages'] for e in p['elements'])),
             'build':result,'delivery':final}
    (output/'receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(receipt,ensure_ascii=False))


if __name__=='__main__':
    main()
