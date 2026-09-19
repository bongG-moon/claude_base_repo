"""Synthetic acceptance test through the shipped HTML -> native PPT workflow."""
import argparse
import json
from pathlib import Path
import sys
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'company-agent-plugin/scripts'))
from company_agent import business_artifacts as artifacts, ppt_html, artifact_delivery


def check(output, render=False, single_delivery=False):
    output.mkdir(parents=True, exist_ok=False)
    job = {'creationMode':'new','purpose':'월간 실적 보고','audience':'부서장','slideCount':5,
           'designPreset':'business','title':'6~8월 월간 실적 보고',
           'facts':[
               {'id':'actual','label':'누적 실적','unit':'백만원','op':'sum','inputs':[{'chart':[2,1,i]} for i in range(3)]},
               {'id':'target','label':'누적 목표','unit':'백만원','op':'sum','inputs':[{'chart':[2,0,i]} for i in range(3)]},
               {'id':'rate','label':'목표 달성률','unit':'%','op':'ratio','decimals':2,'inputs':[{'fact':'actual'},{'fact':'target'}]}],
           'slides':[
               {'title':'6~8월 월간 실적 보고','eyebrow':'부서장 보고', 'body':'월별 목표와 실적을 비교합니다.','source':'검증용 가상 자료'},
               {'title':'누적 목표 대비 {{fact:rate}}% 달성','kpis':[{'fact':'actual'},{'fact':'target'},{'fact':'rate'}],
                'body':'6월은 목표 미달, 7월과 8월은 목표 초과입니다.\n실적 변화의 원인은 자료에 없어 추가 확인이 필요합니다.','source':'검증용 가상 자료'},
               {'title':'월별 목표와 실적 비교','chart':{'type':'column','title':'금액 단위: 백만원','categories':['6월','7월','8월'],
                 'series':[{'name':'목표','values':[150,150,150]},{'name':'실적','values':[145,175,210]}]},'source':'검증용 가상 자료'},
               {'title':'월별 상세 실적','table':{'headers':['월','목표 (백만원)','실적 (백만원)'],
                 'rows':[['6월',150,145],['7월',150,175],['8월',150,210]]},'source':'검증용 가상 자료'},
               {'title':'확인할 사항','table':{'headers':['항목','확인 내용','상태'],
                 'rows':[['실적 변동','월별 실적 변화의 원인','미확인'],['자료 범위','집계 기준의 일관성','미확인']]},'source':'원인을 임의로 추정하지 않음'}]}
    state = output/'internal-state'
    project = output/'project' if single_delivery else output
    work = None
    if single_delivery:
        project.mkdir()
        workspace = artifact_delivery.start(state,project/'monthly-editable.pptx')
        assert workspace['ok'], workspace
        work = workspace['workFile']
    started = time.perf_counter()
    if work:
        for _ in range(3):
            draft = artifact_delivery.build(state,work,'ppt-design-preview',job)
            assert draft['ok'], draft
            assert list(project.iterdir()) == [], 'Internal drafts leaked into the project'
    else:
        draft = artifacts.create_ppt(job,output/'monthly-draft.html',require_choices=True,preview_only=True)
    draft_ms = round((time.perf_counter()-started)*1000,1)
    assert draft['ok'], draft
    assert artifacts.create_ppt(job,output/'not-approved.pptx',require_choices=True)['stage']=='design_preview'
    assert not (output/'not-approved.pptx').exists()
    # Test fixture only: simulate the human approval boundary, not a live approval.
    job['designReview'] = {**draft['designReview'],'confirmed':True}
    saved = ppt_html.save_template(job,output/'representative-template.html')
    assert saved['ok'], saved
    def build_final():
        return artifact_delivery.build(state,work,'ppt',job) if work else artifacts.create_ppt(job,project/'monthly-editable.pptx',require_choices=True)
    if render:
        final = build_final()
    else:
        with patch.object(artifacts,'_office',return_value={'ok':False,'code':'test_without_office','message':'자동 테스트의 Office 렌더링 생략'}):
            final = build_final()
    assert final['ok'], final
    published = None
    if work:
        assert list(project.iterdir()) == [], 'PPT or review images leaked before publication'
        published = artifact_delivery.publish(state,work)
        assert published['ok'], published
        assert list(project.iterdir()) == [project/'monthly-editable.pptx']
        assert artifact_delivery.publish(state,work)['alreadyDelivered']
        assert artifact_delivery.build(state,work,'ppt',job)['alreadyDelivered']
        assert list(project.iterdir()) == [project/'monthly-editable.pptx']
    from pptx import Presentation
    deck = Presentation(project/'monthly-editable.pptx')
    assert len(deck.slides)==5
    assert next(s.chart for s in deck.slides[2].shapes if s.has_chart).series[1].values==(145.,175.,210.)
    assert next(s.table for s in deck.slides[3].shapes if s.has_table).cell(3,2).text=='210'
    reuse = {'creationMode':'saved','purpose':'재사용 시험','audience':'팀원','slideCount':1,
             'slides':[{'title':'새 업무의 제목','body':'양식만 재사용한 새 내용'}]}
    reused = ppt_html.create_draft(reuse,output/'reused-draft.html',output/'representative-template.html')
    assert reused['ok'], reused
    result = {'draftMs':draft_ms,'draft':draft,'template':saved,'ppt':final,'reused':reused,
              'publication':published,'projectFiles':[p.name for p in project.iterdir()] if work else None,
              'internalDraftAttempts':3 if work else 1,
              'approval':'simulated-in-test-only','nativeEditabilityChecked':True,'visualInspection':'pending'}
    with (output/'validation.json').open('x',encoding='utf-8') as stream:
        json.dump(result,stream,ensure_ascii=False,indent=2)
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--render',action='store_true',help='Also exercise the installed PowerPoint renderer.')
    parser.add_argument('--single-delivery',action='store_true',help='Keep three drafts and native previews out of the deliverable folder.')
    args = parser.parse_args()
    print(json.dumps(check(args.output,args.render,args.single_delivery),ensure_ascii=False,indent=2))
