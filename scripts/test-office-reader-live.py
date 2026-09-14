"""Developer-only synthetic Office smoke check. Never accepts existing source paths.

Calls the fixed helper directly, not the user confirmation UI. Test data is made
here and is not company data. This is not an authorization bypass for real files.
"""
import hashlib
import json
from pathlib import Path
import sys
import tempfile

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'company-agent-plugin/scripts'))
from company_agent.office_reader import _invoke,normalize


def main():
    if sys.argv[1:]!=['--synthetic-only']:
        raise SystemExit('Use --synthetic-only on a test PC with no active Office work.')
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document
    reports=[]
    with tempfile.TemporaryDirectory(prefix='company-office-synthetic-',ignore_cleanup_errors=True) as tmp:
        root=Path(tmp)
        wb=Workbook(); wb.active['A1']='한글 실적'; wb.active['B1']=123; wb.save(root/'sample.xlsx')
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='한글 실적 123'
        deck.save(root/'sample.pptx')
        doc=Document(); doc.add_paragraph('한글 실적 123'); doc.save(root/'sample.docx')
        for name,selection in (('sample.xlsx',{'sheet':1,'range':'A1:B1'}),
                               ('sample.pptx',{'start':1,'end':1}),('sample.docx',{'start':1,'end':1})):
            path=root/name; before=hashlib.sha256(path.read_bytes()).hexdigest()
            try:
                result=_invoke(normalize({'file':str(path),**selection}))
                texts=' '.join(i.get('text','') for i in result.get('items',[]))
                report={'sample':name,'ok':result.get('ok') is True,'code':result.get('code'),
                        'stage':result.get('stage'),'hangulMatches':'한글 실적' in texts,'numberMatches':'123' in texts,
                        'sourceUnchanged':before==hashlib.sha256(path.read_bytes()).hexdigest(),'drmTested':False}
            except Exception as exc:
                report={'sample':name,'ok':False,'exceptionType':type(exc).__name__,'drmTested':False}
            reports.append(report)
            print(json.dumps(report,ensure_ascii=True),flush=True)
    return 0 if all(r.get('ok') and r.get('hangulMatches') and r.get('numberMatches') and r.get('sourceUnchanged') for r in reports) else 1


if __name__=='__main__': raise SystemExit(main())
