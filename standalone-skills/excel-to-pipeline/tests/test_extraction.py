"""Extractor invariants. In-memory workbooks only; fixtures authored separately."""
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from excel_pipeline import (InputError, discover, extract_table, regions_from_points,
                            validate_plan, preflight, column_names, publish)
from formula_engine import FormulaEngine


class ExtractionTests(unittest.TestCase):
    def book(self):
        wb = openpyxl.Workbook()
        wb.active.title = "한글 시트"
        return wb, wb.active

    def test_regions_split_side_by_side_and_stacked(self):
        points = {(r, c) for r in (3, 4, 5) for c in (2, 3, 7, 8)}
        points.update((r, c) for r in (9, 10) for c in (2, 3))
        self.assertEqual(regions_from_points(points), [(2, 3, 3, 5), (7, 3, 8, 5), (2, 9, 3, 10)])

    def test_explicit_table_retains_internal_blank_row_boundary(self):
        wb, ws = self.book()
        ws.append(["사번", "값"])
        ws.append(["001", 0])
        ws.append([None, None])
        ws.append(["002", 2])
        ws.add_table(Table(displayName="Data", ref="A1:B4"))
        regions, _, merges = discover(ws)
        self.assertEqual(len(regions), 1)
        regions[0]["id"] = "t"
        table = extract_table(ws, regions[0], merges, FormulaEngine(wb))
        self.assertEqual(table["rows"], [["001", 0], ["002", 2]])
        self.assertEqual(table["source_rows"], [2, 4])
        self.assertEqual(table["blank_rows"], [3])

    def test_raw_strings_and_case_distinct_ids_are_preserved(self):
        wb, ws = self.book()
        ws.append(["식별자", "문자열"])
        values = [["c03", " aBc "], ["C03", "AbC"], ["001aB", "straße"]]
        for row in values:
            ws.append(row)
        ws.add_table(Table(displayName="CaseSensitiveData", ref="A1:B4"))
        regions, _, merges = discover(ws)
        regions[0]["id"] = "t"
        result = extract_table(ws, regions[0], merges, FormulaEngine(wb))
        self.assertEqual(result["rows"], values)
        self.assertEqual(ws["B2"].value, " aBc ")

    def test_merged_multirow_header(self):
        wb, ws = self.book()
        ws["B3"], ws["C3"], ws["D3"] = "부서", "실적", None
        ws["C4"], ws["D4"] = "건수", "비율"
        ws.merge_cells("B3:B4")
        ws.merge_cells("C3:D3")
        ws["B5"], ws["C5"], ws["D5"] = "기획", 10, 0.5
        regions, _, merges = discover(ws)
        self.assertEqual(regions[0]["header_rows"], 2)
        self.assertEqual(column_names(ws, regions[0], merges), ["부서", "실적 / 건수", "실적 / 비율"])

    def test_no_cache_formulas_remain_occupied_and_recompute(self):
        wb, ws = self.book()
        ws.append(["수량", "결과"])
        ws.append([4, "=A2*3"])
        ws.append([5, "=A3*3"])
        regions, _, merges = discover(ws)
        self.assertEqual(regions[0]["range"], "A1:B3")
        regions[0]["id"] = "t"
        table = extract_table(ws, regions[0], merges, FormulaEngine(wb))
        self.assertEqual(table["rows"], [[4, 12], [5, 15]])

    def test_unsupported_is_not_cached_or_zero(self):
        wb, ws = self.book()
        ws.append(["사번", "결과"])
        ws.append(["001", '=INDIRECT("A2")'])
        regions, _, merges = discover(ws)
        regions[0]["id"] = "t"
        table = extract_table(ws, regions[0], merges, FormulaEngine(wb))
        self.assertEqual(table["rows"], [["001", None]])
        self.assertTrue(table["formula_errors"])

    def test_merged_values_fill_only_actual_merge(self):
        wb, ws = self.book()
        ws.append(["부서", "사번", "값"])
        ws.append(["기획", "001", 1])
        ws.append([None, "002", None])
        ws.append([None, "003", 0])
        ws.merge_cells("A2:A3")
        for cell in ws[1]: cell.font = Font(bold=True)
        regions, _, merges = discover(ws)
        regions[0]["id"] = "t"
        table = extract_table(ws, regions[0], merges, FormulaEngine(wb))
        self.assertEqual(table["rows"], [["기획", "001", 1], ["기획", "002", None], [None, "003", 0]])

    def test_total_rows_not_silently_deleted(self):
        wb, ws = self.book()
        ws.append(["구분", "값"])
        ws.append(["A", 10])
        ws.append(["합계", "=SUM(B2:B2)"])
        regions, _, merges = discover(ws)
        regions[0]["id"] = "t"
        table = extract_table(ws, regions[0], merges, FormulaEngine(wb))
        self.assertEqual(len(table["rows"]), 2)
        self.assertEqual(table["row_kinds"], ["data", "summary_candidate"])

    def test_duplicate_header_dedup_is_unique(self):
        wb, ws = self.book()
        ws.append(["값", "값", "값__2"])
        ws.append([1, 2, 3])
        regions, _, merges = discover(ws)
        self.assertEqual(len(set(column_names(ws, regions[0], merges))), 3)

    def test_all_text_header_ambiguity_exposed(self):
        wb, ws = self.book()
        ws.append(["서울", "홍길동"])
        ws.append(["부산", "김하나"])
        regions, _, _ = discover(ws)
        self.assertIn("header_ambiguous", regions[0]["flags"])

    def test_plan_rejects_traversal_overlap_unknown_sheet(self):
        wb, _ = self.book()
        base = {"id": "t", "sheet": "한글 시트", "range": "A1:B5", "header_rows": 1}
        for change in [{"id": "../bad"}, {"sheet": "missing"}, {"header_rows": 9}, {"range": "A1:XFD1048576"}, {"exclude_rows": [1]}]:
            with self.subTest(change=change), self.assertRaises(InputError):
                validate_plan({"tables": [{**base, **change}]}, wb)
        with self.assertRaises(InputError):
            validate_plan({"tables": [base, {**base, "id": "t2"}]}, wb)

    def test_preflight_rejects_fake_extension_and_does_not_modify(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
            source = Path(folder) / "wrong.xlsx"
            source.write_bytes(b"not an Excel archive")
            with self.assertRaises(InputError): preflight(source)
            self.assertEqual(source.read_bytes(), b"not an Excel archive")

    def test_publish_will_not_overwrite_existing_output(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as folder:
            with self.assertRaises(InputError): publish({}, Path(folder))


if __name__ == "__main__":
    unittest.main()
