"""Replay audit failures through formulas and the generated DATA/PROCESS project."""
import datetime as dt
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import openpyxl

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))
from excel_pipeline import analyze, publish, safe_value
from formula_engine import FormulaEngine
from project_builder import make_rules
from project_runtime import _decode, calculate_data, read_excel_data
from test_fixture_paths import fixture_root


class NumericConsistencyTests(unittest.TestCase):
    def test_sum_functions_share_accumulation_before_rounding(self):
        for sign in (1, -1):
            with self.subTest(sign=sign):
                wb = openpyxl.Workbook()
                ws = wb.active
                for number in [0.1 * sign] * 68 + [0.05 * sign]:
                    ws.append(["포함", number])
                formulas = ['SUM(B1:B69)', 'SUMIF(A1:A69,"포함",B1:B69)',
                            'SUMIFS(B1:B69,A1:A69,"포함")', 'AVERAGE(B1:B69)*69']
                for row, formula in enumerate(formulas, 1):
                    ws.cell(row, 4, "=" + formula)
                    ws.cell(row, 5, "=ROUND(" + formula + ",1)")
                engine = FormulaEngine(wb)
                values = [engine.evaluate(ws.title, f"D{r}") for r in range(1, 4)]
                self.assertTrue(all(v["status"] == "computed" for v in values), values)
                self.assertEqual(1, len({v["value"] for v in values}))
                for row in range(1, 5):
                    self.assertEqual({"status": "computed", "value": 6.9 * sign},
                                     engine.evaluate(ws.title, f"E{row}"))
                wb.close()

    def test_cancellation_and_integer_precision_use_same_policy(self):
        for numbers, expected in [([1e16, 1.0, -1e16], 1.0), ([2**54, 1, -2**54], 1)]:
            with self.subTest(numbers=numbers):
                wb = openpyxl.Workbook()
                ws = wb.active
                for value in numbers:
                    ws.append(["선택", value])
                for row, formula in enumerate(['=SUM(B1:B3)', '=SUMIF(A1:A3,"선택",B1:B3)',
                                                '=SUMIFS(B1:B3,A1:A3,"선택")'], 1):
                    ws.cell(row, 4, formula)
                engine = FormulaEngine(wb)
                self.assertEqual([expected] * 3, [engine.evaluate(ws.title, f"D{r}")["value"] for r in range(1, 4)])
                wb.close()

    def test_sum_overflow_reports_formula_error(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["선택", 1e308])
        ws.append(["선택", 1e308])
        for row, formula in enumerate(['=SUM(B1:B2)', '=SUMIF(A1:A2,"선택",B1:B2)',
                                        '=SUMIFS(B1:B2,A1:A2,"선택")'], 1):
            ws.cell(row, 4, formula)
        engine = FormulaEngine(wb)
        for row in range(1, 4):
            result = engine.evaluate(ws.title, f"D{row}")
            self.assertEqual(("error", "#NUM!"), (result["status"], result["code"]))
        wb.close()


class DurationProjectTests(unittest.TestCase):
    def test_duration_roundtrip_preserves_sign_days_and_fraction(self):
        cases = [(dt.timedelta(hours=26), "26:00:00"), (dt.timedelta(0), "0:00:00"),
                 (dt.timedelta(microseconds=-1), "-0:00:00.000001"),
                 (dt.timedelta(hours=-49, minutes=-3, seconds=-4, microseconds=-125000), "-49:03:04.125")]
        for value, text in cases:
            with self.subTest(value=value):
                self.assertEqual(text, safe_value(value))
                self.assertEqual(value, _decode(text, {"kind": "duration"}))
                self.assertEqual(value, _decode(value, {"kind": "duration"}))

    def test_duration_rejects_dates_ambiguous_numbers_and_invalid_clock(self):
        for value in ("2026-10-29", dt.datetime(2026, 10, 29), dt.time(2), 26, "26:60:00", "26:00:60", "nan"):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "기간"):
                _decode(value, {"kind": "duration"})

    def test_workbook_keeps_duration_separate_from_date_time_and_text(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "누적시간.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "시간"
            ws.append(["누적시간", "기준일", "시각", "원문", "계산시간"])
            ws.append([dt.timedelta(hours=26), dt.datetime(2026, 10, 29), dt.time(2), "26:00:00", "=A2*24"])
            ws["A2"].number_format = "[h]:mm:ss"
            wb.save(source)
            wb.close()
            result = analyze(source)
            rules = make_rules(result)
            self.assertEqual(["duration", "date", "date", "text"], [rules["inputs"][f"시간!{c}2"]["kind"] for c in "ABCD"])
            data = read_excel_data(source, rules)
            self.assertEqual("26:00:00", data["cells"]["시간!A2"])
            self.assertEqual(result["tables"], calculate_data(data, rules)["tables"])
            self.assertEqual(26, calculate_data(data, rules)["formulas"][0]["value"])

    def test_complex_project_reruns_and_data_only_replacement_changes_result(self):
        # Use the established 12-sheet / 6,798-formula fixture, adding duration
        # and the reported rounding boundary without running its VBA semantics.
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "복잡한_기간_계산.xlsx"
            wb = openpyxl.load_workbook(fixture_root() / "고급_수식매크로_혼합.xlsx")
            ws = wb.create_sheet("누적시간")
            ws.append(["포함여부", "금액", "시간", "시간합계", "금액합계", "조건합계", "다중조건합계"])
            for value in [0.1] * 68 + [0.05]:
                ws.append(["포함", value])
            ws["C2"] = dt.timedelta(hours=26)
            ws["C2"].number_format = "[h]:mm:ss"
            ws["D2"] = "=C2*24"
            ws["E2"] = "=ROUND(SUM(B2:B70),1)"
            ws["F2"] = '=ROUND(SUMIF(A2:A70,"포함",B2:B70),1)'
            ws["G2"] = '=ROUND(SUMIFS(B2:B70,A2:A70,"포함"),1)'
            wb.save(source)
            wb.close()
            source_digest = hashlib.sha256(source.read_bytes()).hexdigest()
            result = analyze(source)
            self.assertEqual(6802, result["summary"]["formula_count"])
            self.assertEqual(13, result["summary"]["sheet_count"])
            project = root / "기간_프로젝트"
            publish(result, project, context_chars=4000)
            before = {f.name: hashlib.sha256(f.read_bytes()).hexdigest() for f in (project / "PROCESS").glob("*.py")}
            first = root / "첫번째"
            command = [sys.executable, "-X", "utf8", "-B", str(project / "main.py"), "--diagnostics", "--output", str(first)]
            run = subprocess.run(command, cwd=root, capture_output=True, text=True, encoding="utf-8", timeout=90)
            self.assertEqual(0, run.returncode, run.stdout + run.stderr)
            values = json.loads((first / "_details" / "formulas.json").read_text(encoding="utf-8"))
            self.assertEqual({"D2": 26.0, "E2": 6.9, "F2": 6.9, "G2": 6.9},
                             {f["cell"]: f["value"] for f in values if f["sheet"] == "누적시간"})
            rules = json.loads((project / "_support" / "contract.json").read_text(encoding="utf-8"))
            changed = read_excel_data(source, rules)
            changed["cells"]["누적시간!C2"] = "51:30:00"
            changed.pop("sha256", None)
            changed.pop("input_evidence", None)
            (project / "DATA" / "records.json").write_text(json.dumps(changed, ensure_ascii=False), encoding="utf-8")
            (project / "DATA" / "load.py").write_text(
                'import json\nfrom pathlib import Path\ndef load_data():\n'
                '    return json.loads(Path(__file__).with_name("records.json").read_text(encoding="utf-8"))\n', encoding="utf-8")
            second = root / "교체후"
            command[-1] = str(second)
            run = subprocess.run(command, cwd=root, capture_output=True, text=True, encoding="utf-8", timeout=90)
            self.assertEqual(0, run.returncode, run.stdout + run.stderr)
            values = json.loads((second / "_details" / "formulas.json").read_text(encoding="utf-8"))
            self.assertAlmostEqual(51.5, next(f["value"] for f in values if f["sheet"] == "누적시간" and f["cell"] == "D2"), places=12)
            self.assertEqual(before, {f.name: hashlib.sha256(f.read_bytes()).hexdigest() for f in (project / "PROCESS").glob("*.py")})
            self.assertEqual(source_digest, hashlib.sha256(source.read_bytes()).hexdigest())
            self.assertLessEqual(len((second / "_details" / "model_context.json").read_text(encoding="utf-8")), 4000)


if __name__ == "__main__":
    unittest.main()
