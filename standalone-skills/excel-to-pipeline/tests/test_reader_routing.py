"""Independent reader-boundary tests. No test imports xlwings or starts Excel."""
import builtins
import contextlib
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import MagicMock, Mock, patch
import zipfile
from test_fixture_paths import fixture_root

import openpyxl

SKILL = Path(__file__).resolve().parents[1]
ROOT = SKILL.parents[1]
FIXTURES = fixture_root()
sys.path.insert(0, str(SKILL / "scripts"))
import excel_pipeline as pipeline
import workbook_input
from model_context import build_model_context, serialize_model_context
from vba_extract import outcome


class FakeNativeReadError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


class ReaderRoutingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="reader-route-", dir=SKILL / "tests")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.fixture = FIXTURES / "complex_workbook.xlsx"
        self.native = types.ModuleType("xlwings_reader")
        self.native.NativeReadError = FakeNativeReadError
        self.native.load_via_excel = Mock(side_effect=self.normalized_books)
        self.module_patch = patch.dict(sys.modules, {"xlwings_reader": self.native})
        self.module_patch.start()
        self.addCleanup(self.module_patch.stop)

    def normalized_books(self, source, timeout=90, vba=None, incomplete=None):
        """Reuse an existing fixture as an in-memory native-reader response."""
        wb = openpyxl.load_workbook(self.fixture, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(self.fixture, data_only=True, keep_links=False)
        evidence = {"reader": "xlwings", "excel_started": True,
                    "drm_status": "not_determined", "values_origin": "excel_observed_value",
                    "incomplete_metadata": incomplete or [], "macro_executed": False}
        return wb, cached, ["macros"] if vba else [], vba or outcome("no_vba"), evidence

    def binary(self, name="opaque.xlsx"):
        source = self.folder / name
        source.write_bytes(b"\x80\x00" + b"synthetic vendor envelope, not a real DRM document")
        return source

    def ole_binary(self):
        source = self.folder / "ole-envelope.xlsx"
        source.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"mocked OLE directory, not a workbook")
        return source

    def ole_module(self, streams):
        module = types.ModuleType("olefile")
        container = MagicMock()
        container.__enter__.return_value = container
        container.listdir.return_value = streams
        module.OleFileIO = Mock(return_value=container)
        return module, container

    def assert_rejected_without_native(self, source, reader="auto", expected=Exception):
        with self.assertRaises(expected):
            pipeline.analyze(source, reader=reader)
        self.native.load_via_excel.assert_not_called()

    def test_auto_ooxml_never_imports_native_reader_or_xlwings(self):
        real_import = builtins.__import__

        def guarded_import(name, *args, **kwargs):
            if name in {"xlwings", "xlwings_reader", "olefile"}:
                self.fail("Ordinary OOXML must not import native Excel support")
            return real_import(name, *args, **kwargs)

        before = hashlib.sha256(self.fixture.read_bytes()).hexdigest()
        with patch("builtins.__import__", side_effect=guarded_import):
            result = pipeline.analyze(self.fixture)
        self.assertEqual(result["summary"]["reader"], "openpyxl")
        self.assertFalse(result["input_evidence"]["excel_started"])
        self.assertEqual(result["summary"]["formula_status"], {"computed": 82})
        self.assertEqual(hashlib.sha256(self.fixture.read_bytes()).hexdigest(), before)
        self.native.load_via_excel.assert_not_called()

    def test_binary_candidate_uses_native_once_without_claiming_drm(self):
        source = self.binary()
        result = pipeline.analyze(source, excel_timeout=45)
        self.native.load_via_excel.assert_called_once_with(source.resolve(), timeout=45)
        self.assertEqual(result["summary"]["sheet_count"], 5)
        self.assertEqual(result["summary"]["table_count"], 7)
        self.assertEqual(result["input_evidence"]["selection_reason"], "non_ooxml_candidate_not_drm_proof")
        self.assertEqual(result["input_evidence"]["drm_status"], "not_determined")

    def test_explicit_openpyxl_keeps_ordinary_reader(self):
        result = pipeline.analyze(self.fixture, reader="openpyxl")
        self.assertEqual(result["summary"]["reader"], "openpyxl")
        self.native.load_via_excel.assert_not_called()

    def test_explicit_openpyxl_does_not_retry_binary(self):
        self.assert_rejected_without_native(self.binary(), reader="openpyxl", expected=pipeline.InputError)

    def test_ole_root_legacy_workbook_stream_blocks_native_before_open(self):
        source = self.ole_binary()
        for root_name in ("Workbook", "Book", "wOrKbOoK", "bOoK"):
            for reader in ("auto", "xlwings"):
                with self.subTest(root_name=root_name, reader=reader):
                    module, container = self.ole_module([[root_name], ["Metadata"]])
                    with patch.dict(sys.modules, {"olefile": module}):
                        self.assert_rejected_without_native(source, reader=reader, expected=pipeline.InputError)
                    container.listdir.assert_called_once_with(streams=True, storages=False)

    def test_ole_malformed_container_never_retries_native(self):
        source = self.ole_binary()
        module, _ = self.ole_module([])
        module.OleFileIO.side_effect = OSError("malformed OLE header")
        with patch.dict(sys.modules, {"olefile": module}):
            self.assert_rejected_without_native(source, expected=pipeline.InputError)
        module.OleFileIO.assert_called_once()

    def test_ole_directory_failure_never_retries_native(self):
        source = self.ole_binary()
        module, container = self.ole_module([])
        container.listdir.side_effect = ValueError("invalid stream directory")
        with patch.dict(sys.modules, {"olefile": module}):
            self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_ole_missing_dependency_is_actionable_without_native_start(self):
        source = self.ole_binary()
        with patch.dict(sys.modules, {"olefile": None}):
            with self.assertRaisesRegex(pipeline.InputError, "olefile|requirements"):
                pipeline.analyze(source)
        self.native.load_via_excel.assert_not_called()

    def test_encrypted_ooxml_ole_can_use_native_without_drm_claim(self):
        source = self.ole_binary()
        module, container = self.ole_module([["EncryptedPackage"], ["EncryptionInfo"]])
        with patch.dict(sys.modules, {"olefile": module}):
            result = pipeline.analyze(source)
        self.native.load_via_excel.assert_called_once()
        self.assertEqual(result["input_evidence"]["drm_status"], "not_determined")
        self.assertEqual(result["input_evidence"]["selection_reason"], "non_ooxml_candidate_not_drm_proof")
        container.listdir.assert_called_once_with(streams=True, storages=False)
        # The optional OLE inspection must never request read/write mode.
        self.assertIsNot(module.OleFileIO.call_args.kwargs.get("write_mode", False), True)
        container.write_stream.assert_not_called()

    def test_nested_workbook_stream_is_not_mistaken_for_legacy_root(self):
        source = self.ole_binary()
        module, _ = self.ole_module([["EncryptedPackage"], ["EncryptionInfo"], ["Metadata", "Workbook"]])
        with patch.dict(sys.modules, {"olefile": module}):
            result = pipeline.analyze(source)
        self.native.load_via_excel.assert_called_once()
        self.assertEqual(result["summary"]["reader"], "xlwings")

    def test_explicit_xlwings_uses_native_for_valid_ooxml(self):
        result = pipeline.analyze(self.fixture, reader="xlwings")
        self.native.load_via_excel.assert_called_once()
        self.assertEqual(result["input_evidence"]["selection_reason"], "explicit_excel")
        self.assertEqual(result["formulas"][0]["cached_value_origin"], "excel_observed_value")

    def test_unknown_reader_is_rejected(self):
        self.assert_rejected_without_native(self.fixture, reader="guessed", expected=pipeline.InputError)

    def test_native_failure_is_not_retried_or_relabelled_success(self):
        source = self.binary()
        self.native.load_via_excel.side_effect = FakeNativeReadError("access_denied", "Excel refused the read")
        with self.assertRaisesRegex(pipeline.InputError, "access_denied"):
            pipeline.analyze(source)
        self.native.load_via_excel.assert_called_once()

    def test_openpyxl_loader_failure_does_not_start_excel(self):
        with patch.object(openpyxl, "load_workbook", side_effect=ValueError("broken workbook")):
            self.assert_rejected_without_native(self.fixture, expected=ValueError)

    def test_broken_zip_is_not_treated_as_protection(self):
        source = self.folder / "broken.xlsx"
        source.write_bytes(b"PK\x03\x04broken zip")
        self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_text_renamed_xlsx_is_not_sent_to_excel_importer(self):
        source = self.folder / "csv.xlsx"
        source.write_bytes(b"id,value\n001,12\n")
        self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_korean_csv_renamed_xlsx_is_not_sent_to_excel_importer(self):
        source = self.folder / "korean-csv.xlsx"
        source.write_bytes("사번,신청과정\n001,데이터분석\n".encode("utf-8"))
        self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_bom_html_renamed_xlsx_is_not_sent_to_excel_importer(self):
        source = self.folder / "html.xlsx"
        source.write_bytes("<html><table><tr><td>자료</td></tr></table></html>".encode("utf-8-sig"))
        self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_utf16_csv_renamed_xlsx_is_not_sent_to_excel_importer(self):
        source = self.folder / "utf16.xlsx"
        source.write_bytes("사번,신청과정\n001,데이터분석\n".encode("utf-16"))
        self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_empty_file_is_rejected(self):
        source = self.folder / "empty.xlsx"
        source.write_bytes(b"")
        self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_missing_file_is_rejected(self):
        self.assert_rejected_without_native(self.folder / "missing.xlsx", expected=pipeline.InputError)

    def test_wrong_suffix_is_rejected(self):
        source = self.binary("legacy.xls")
        self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_oversized_file_is_rejected_before_native(self):
        source = self.binary()
        original = Path.stat
        values = list(original(source))
        values[6] = 65 * 1024 * 1024

        def oversized(path, *args, **kwargs):
            return os.stat_result(values) if path == source else original(path, *args, **kwargs)

        with patch.object(Path, "stat", oversized):
            self.assert_rejected_without_native(source, expected=pipeline.InputError)

    def test_permission_denial_is_not_treated_as_drm(self):
        source = self.binary()
        original = Path.open

        def denied(path, *args, **kwargs):
            if path == source:
                raise PermissionError("not authorized")
            return original(path, *args, **kwargs)

        with patch.object(Path, "open", denied):
            with self.assertRaises(PermissionError):
                workbook_input.load_workbooks(source)
        self.native.load_via_excel.assert_not_called()

    def make_container(self, worksheet=None, extra=None):
        source = self.folder / "container.xlsx"
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("xl/workbook.xml", f'<workbook xmlns="{namespace}"/>')
            archive.writestr("xl/worksheets/sheet1.xml", worksheet or f'<worksheet xmlns="{namespace}"/>')
            for key, value in (extra or {}).items():
                archive.writestr(key, value)
        return source

    def test_range_safety_failure_cannot_be_bypassed_by_forced_native(self):
        source = self.make_container('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1:XFD1048576"/></worksheet>')
        for reader in ("auto", "xlwings"):
            with self.subTest(reader=reader):
                self.assert_rejected_without_native(source, reader=reader, expected=pipeline.InputError)

    def test_malformed_xml_does_not_fall_back(self):
        source = self.make_container("<worksheet><unclosed>")
        self.assert_rejected_without_native(source)

    def test_external_refresh_connections_block_explicit_native(self):
        source = self.make_container(extra={"xl/connections.xml": "<connections/>"})
        with self.assertRaises(pipeline.InputError):
            workbook_input.choose_reader(source, "xlwings")
        self.native.load_via_excel.assert_not_called()

    def test_excel4_macro_sheets_block_explicit_native(self):
        source = self.make_container(extra={"xl/macrosheets/sheet1.xml": "<worksheet/>"})
        with self.assertRaises(pipeline.InputError):
            workbook_input.choose_reader(source, "xlwings")
        self.native.load_via_excel.assert_not_called()

    def denied_analysis(self):
        source = self.binary("opaque.xlsm")
        denied = outcome("access_denied", "VBA project source access denied")
        self.native.load_via_excel.side_effect = lambda path, timeout: self.normalized_books(path, timeout, denied)
        with patch("vba_extract.extract_vba", side_effect=AssertionError("Must not retry denied VBA via another extractor")):
            return pipeline.analyze(source)

    def test_vba_denial_preserves_cell_and_formula_analysis(self):
        result = self.denied_analysis()
        self.assertEqual(result["summary"]["sheet_count"], 5)
        self.assertEqual(result["summary"]["table_count"], 7)
        self.assertEqual(result["summary"]["formula_status"], {"computed": 82})
        self.assertEqual(result["vba"]["status"], "access_denied")
        self.assertIsNone(result["vba_index"])
        self.assertIn("vba_source_unavailable", result["summary"]["review_flags"])
        self.assertFalse(result["summary"]["vba"]["macro_executed"])
        self.assertFalse(result["summary"]["vba"]["translation_generated"])
        self.assertEqual(result["migration_plan"]["translation_status"], "blocked_source_unavailable")

    def test_incomplete_native_metadata_is_visible(self):
        self.native.load_via_excel.side_effect = lambda path, timeout: self.normalized_books(path, timeout, incomplete=["defined_names"])
        result = pipeline.analyze(self.binary())
        self.assertIn("native_metadata_incomplete", result["summary"]["review_flags"])
        self.assertEqual(result["summary"]["status"], "review_required")

    def test_native_partial_context_preserves_budget_and_vba_status(self):
        result = self.denied_analysis()
        before = json.dumps(result, ensure_ascii=False, sort_keys=True)
        for rows, chars in ((0, 1024), (5, 8000), (5, 12000), (1000, 12000), (1001, 50000)):
            with self.subTest(rows=rows, chars=chars):
                context = build_model_context(result, sample_rows=rows, max_chars=chars)
                serialized = serialize_model_context(context)
                self.assertLessEqual(len(serialized), chars)
                self.assertEqual(json.loads(serialized), context)
                self.assertLessEqual(context["limits"]["effective_rows"], 1000)
                self.assertEqual(context["vba"]["status"], "access_denied")
                self.assertEqual(context["totals"]["tables"], 7)
        self.assertEqual(json.dumps(result, ensure_ascii=False, sort_keys=True), before)

    def test_publish_packages_native_runtime_and_optional_dependencies(self):
        result = pipeline.analyze(self.fixture)
        output = self.folder / "project"
        pipeline.publish(result, output, sample_rows=5, context_chars=8000, layout="detailed")
        for filename in ("workbook_input.py", "xlwings_reader.py"):
            self.assertEqual((output / "pipeline_lib" / filename).read_bytes(), (SKILL / "scripts" / filename).read_bytes())
        self.assertIn("xlwings", (output / "requirements-excel.txt").read_text(encoding="utf-8"))
        self.assertIn("olefile", (output / "requirements-excel.txt").read_text(encoding="utf-8"))
        self.assertNotIn("xlwings", (output / "requirements.txt").read_text(encoding="utf-8"))
        self.assertLessEqual(len((output / "model_context.json").read_text(encoding="utf-8")), 8000)
        # A fresh interpreter proves portable imports do not depend on this test's sys.path.
        run = subprocess.run([sys.executable, "-X", "utf8", str(output / "transform.py"),
                              str(FIXTURES / "variant_next_month.xlsx"), "--context-only", "--context-chars", "8000"],
                             capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertLessEqual(len(run.stdout), 8000)
        self.assertEqual(json.loads(run.stdout)["totals"]["formulas"], 89)

    def test_cli_context_only_native_path_stays_bounded(self):
        source = self.binary()
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            code = pipeline.main([str(source), "--reader", "xlwings", "--context-only", "--context-chars", "8000"])
        self.assertEqual(code, 0)
        self.assertLessEqual(len(output.getvalue()), 8000)
        self.assertEqual(json.loads(output.getvalue())["totals"]["formulas"], 82)
        self.native.load_via_excel.assert_called_once()


if __name__ == "__main__":
    unittest.main()
