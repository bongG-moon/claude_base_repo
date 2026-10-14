"""In-memory regression tests; never launches Excel or reads network/workbooks."""

import datetime
import importlib.util
import json
from pathlib import Path
import sys
import unittest

import openpyxl
from openpyxl.workbook.defined_name import DefinedName


MODULE = Path(__file__).resolve().parents[1] / "scripts" / "formula_engine.py"
spec = importlib.util.spec_from_file_location("excel_formula_engine_test", MODULE)
module = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = module
spec.loader.exec_module(module)
FormulaEngine = module.FormulaEngine


class FormulaTests(unittest.TestCase):
    def setUp(self):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = "매출 자료"

    def result(self, formula, **options):
        self.ws["Z1"] = formula
        return FormulaEngine(self.wb, **options).evaluate(self.ws.title, "Z1")

    def computed(self, formula, value, **options):
        result = self.result(formula, **options)
        self.assertEqual(result["status"], "computed", result)
        self.assertEqual(result["value"], value, result)
        json.dumps(result, allow_nan=False)

    def unsupported(self, formula, code=None, **options):
        result = self.result(formula, **options)
        self.assertEqual(result["status"], "unsupported", result)
        if code:
            self.assertEqual(result["code"], code, result)
        self.assertIsNone(result["value"])

    def test_arithmetic_precedence_percent_and_comparison(self):
        for expression, value in [("=2+3*4", 14), ("=(2+3)*4", 20), ("=-2^2", 4),
                                  ("=2^3^2", 64), ("=100*(1+10%)", 110),
                                  ("=2>=1", True), ("=2<>2", False), ("=\"ABC\"=\"abc\"", True)]:
            with self.subTest(expression=expression):
                self.computed(expression, value)

    def test_korean_quoted_sheet_absolute_references_and_escape(self):
        other = self.wb.create_sheet("팀 '서울'")
        other["B3"] = 1200
        self.computed("='팀 ''서울'''!$B$3/2", 600)
        self.ws["A1"] = 3
        self.computed("='매출 자료'!$A$1+1", 4)

    def test_unicode_unquoted_sheet_and_numeric_sheet_name(self):
        self.wb.create_sheet("기준")["A1"] = 4
        self.wb.create_sheet("2026")["A1"] = 5
        self.computed("=기준!A1+2026!A1", 9)

    def test_aggregates_ignore_referenced_text_boolean_and_blank(self):
        for row, value in enumerate([10, "20", True, None, 30, ""], 1):
            self.ws.cell(row, 1, value)
        self.computed("=SUM(A1:A6)", 40)
        self.computed("=AVERAGE(A1:A6)", 20)
        self.computed("=MIN(A1:A6)", 10)
        self.computed("=MAX(A1:A6)", 30)
        self.computed("=COUNT(A1:A6)", 2)
        self.computed("=COUNTA(A1:A6)", 5)
        self.computed('=SUM("20",TRUE)', 21)
        self.computed("=SUM(A2,A3)", 0)

    def test_literal_none_zero_id_and_date_are_distinct(self):
        self.ws["A1"] = "00123"
        self.ws["A2"] = 0
        self.ws["A4"] = datetime.date(2026, 10, 14)
        engine = FormulaEngine(self.wb)
        self.assertEqual(engine.evaluate(self.ws.title, "A1"), {"status": "literal", "value": "00123"})
        self.assertEqual(engine.evaluate(self.ws.title, "A2"), {"status": "literal", "value": 0})
        self.assertEqual(engine.evaluate(self.ws.title, "A3"), {"status": "literal", "value": None})
        self.assertEqual(engine.evaluate(self.ws.title, "A4")["value"], "2026-10-14")
        self.computed('=A1&"-"&A2', "00123-0")
        self.computed("=A3", 0)
        self.computed('=IF(TRUE,")","-")', ")")
        self.computed('="+"&"%"&"("', "+%(")

    def test_if_is_lazy_and_iferror_only_catches_real_excel_errors(self):
        self.computed("=IF(FALSE,1/0,12)", 12)
        self.computed("=IF(TRUE,12,UNKNOWN(1))", 12)
        self.computed("=IFERROR(1/0,99)", 99)
        self.computed("=IFERROR(#N/A,\"미등록\")", "미등록")
        self.unsupported("=IFERROR(UNKNOWN(1),99)", "unsupported_function:UNKNOWN")
        self.unsupported("=IFERROR(SUM(A1:A20000),99)", "range_cell_limit")

    def test_round_halves_away_from_zero(self):
        self.computed("=ROUND(2.675,2)", 2.68)
        self.computed("=ROUND(-2.675,2)", -2.68)
        self.computed("=ROUND(125,-1)", 130)
        self.computed("=ROUND(1e30,0)", 1e30)

    def test_decimal_arithmetic_does_not_lose_half_boundary(self):
        self.computed("=ROUND(2.5*2*(1-0.9),0)", 1)
        self.computed("=ROUND(-2.5*2*(1-0.9),0)", -1)
        self.computed("=ROUND(2.5*2*(1-0.90001),0)", 0)
        self.computed("=0.1+0.2", 0.3)
        self.computed("=ROUND(1/3,3)", 0.333)

    def test_summary_first_rechecks_budget_once_after_dependencies(self):
        self.ws["A1"] = "=SUM(B1:B16)"
        self.ws["A2"] = "=SUM(D1:D1000)"
        for row in range(1, 17):
            self.ws.cell(row, 2, "=SUM(C1:C16)")
            self.ws.cell(row, 3, 1)
        self.assertEqual(FormulaEngine(self.wb, max_operations=90).evaluate(self.ws.title, "A1")["code"], "operation_limit")
        records = FormulaEngine(self.wb, max_operations=90).analyze()
        self.assertEqual(records[0]["status"], "computed")
        self.assertEqual(records[0]["value"], 256)
        self.assertNotIn("code", records[0])
        self.assertEqual(next(r for r in records if r["cell"] == "A2")["code"], "operation_limit")

    def test_sumif_sumifs_countif_countifs(self):
        rows = [("서울", "A", 100), ("서울", "B", 200), ("부산", "A", 400), ("서울", "A", 300)]
        for row in rows:
            self.ws.append(row)
        self.computed('=SUMIF(A1:A4,"서울",C1:C4)', 600)
        self.computed('=SUMIFS(C1:C4,A1:A4,"서울",B1:B4,"A")', 400)
        self.computed('=COUNTIF(C1:C4,">=200")', 3)
        self.computed('=COUNTIFS(A1:A4,"서울",C1:C4,">100")', 2)
        self.unsupported('=SUMIF(A1:A4,"서울",C1:C3)', "criteria_range_shape_mismatch")

    def test_criteria_wildcards_escape_and_blank(self):
        for row, value in enumerate(["서울1", "서울2", "서울*", "부산", None, ""], 1):
            self.ws.cell(row, 1, value)
        self.computed('=COUNTIF(A1:A6,"서울?")', 3)
        self.computed('=COUNTIF(A1:A6,"서울~*")', 1)
        self.computed('=COUNTIF(A1:A6,"")', 2)
        self.computed('=COUNTIF(A1:A6,"<>")', 4)

    def test_sumif_does_not_touch_unmatched_error(self):
        self.ws.append(["선택", 10])
        self.ws.append(["제외", "=1/0"])
        self.computed('=SUMIF(A1:A2,"선택",B1:B2)', 10)

    def test_vlookup_exact_preserves_string_ids_and_missing_error(self):
        self.ws.append(["001", "Alice"])
        self.ws.append([1, "Bob"])
        self.computed('=VLOOKUP("001",A1:B2,2,FALSE)', "Alice")
        self.computed('=VLOOKUP(1,A1:B2,2,0)', "Bob")
        self.computed('=IFERROR(VLOOKUP("002",A1:B2,2,FALSE),"미등록")', "미등록")
        self.unsupported('=VLOOKUP("001",A1:B2,2)', "approximate_lookup")
        self.unsupported('=VLOOKUP("001",A1:B2,2,TRUE)', "approximate_lookup")

    def test_named_ranges_local_scope_preferred_and_direct_only(self):
        other = self.wb.create_sheet("기준표")
        other["A1"], other["A2"] = 2, 3
        self.ws["A1"] = 100
        self.wb.defined_names.add(DefinedName("비율", attr_text="'기준표'!$A$1:$A$2"))
        self.ws.defined_names.add(DefinedName("비율", attr_text="'매출 자료'!$A$1"))
        self.computed("=SUM(비율)", 100)
        other["C1"] = "=SUM(비율)"
        self.assertEqual(FormulaEngine(self.wb).evaluate("기준표", "C1")["value"], 5)
        self.wb.defined_names.add(DefinedName("계산이름", attr_text="SUM('기준표'!A1:A2)"))
        self.unsupported("=계산이름", "named_formula_or_union")
        self.wb.defined_names.add(DefinedName("상대이름", attr_text="'기준표'!A1"))
        self.unsupported("=상대이름", "relative_named_range")

    def test_cycles_and_transitive_formula_dependencies(self):
        self.ws["A1"] = "=B1+1"
        self.ws["B1"] = "=A1+1"
        result = self.result("=IFERROR(A1,0)")
        self.assertEqual(result, {"status": "error", "value": None, "code": "circular_reference"})
        self.ws["B1"] = 10
        self.computed("=A1*2", 22)

    def test_unknown_external_dynamic_array_and_3d_are_explicit(self):
        for formula in ['=XLOOKUP(1,A1:A3,B1:B3)', "='[outside.xlsx]Sheet1'!A1",
                        "=[outside.xlsx]Sheet1!A1", "=SUM(Table1[Amount])", "=A1#",
                        "=SUM(Sheet1:Sheet3!A1)", "=SUM(A:A)", "=SUM(1:1)",
                        "={1,2}", "=INDIRECT(\"A1\")", "=LET(x,1,x+1)"]:
            with self.subTest(formula=formula):
                self.unsupported(formula)

    def test_range_bounds_and_limits(self):
        for formula in ["=A0", "=XFE1", "=A1048577", "=SUM(B2:A1)"]:
            with self.subTest(formula=formula):
                self.unsupported(formula)
        self.unsupported("=SUM(A1:A101)", "range_cell_limit", max_range_cells=100)
        self.unsupported("=SUM(A1:A10)", "operation_limit", max_operations=3)
        self.unsupported("=1+2+3", "formula_length_limit", max_formula_chars=4)
        self.unsupported("=1+2+3", "formula_token_limit", max_tokens=3)
        self.unsupported("=((((1))))", "formula_depth_limit", max_depth=2)
        self.unsupported("=" + "9" * 5000, "numeric_limit")

    def test_no_missing_cell_mutation_and_no_formula_mutation(self):
        self.ws["A1"] = "=SUM(C1:C3)"
        before = dict(self.ws._cells)
        engine = FormulaEngine(self.wb)
        self.assertEqual(engine.evaluate(self.ws.title, "A1")["value"], 0)
        self.assertEqual(engine.evaluate(self.ws.title, "ZZ100")["value"], None)
        self.assertEqual(before, self.ws._cells)
        self.assertEqual(self.ws["A1"].value, "=SUM(C1:C3)")

    def test_analyze_includes_every_formula_without_cache(self):
        self.ws["A1"], self.ws["A2"] = 10, 20
        self.ws["B1"] = "=SUM(A1:A2)"
        self.ws["B2"] = "=IF(A1>0,B1,0)"
        self.ws["B3"] = "=UNKNOWN(A1)"
        self.ws["B4"] = "=1/0"
        records = FormulaEngine(self.wb).analyze()
        self.assertEqual(len(records), 4)
        self.assertEqual(records[0]["dependencies"], [{"sheet": "매출 자료", "range": "A1:A2"}])
        self.assertEqual(records[1]["functions"], ["IF"])
        self.assertEqual(records[1]["value"], 30)
        self.assertEqual(records[2]["status"], "unsupported")
        self.assertEqual(records[2]["functions"], ["UNKNOWN"])
        self.assertEqual(records[3]["code"], "#DIV/0!")
        json.dumps(records, ensure_ascii=False, allow_nan=False)

    def test_named_dependency_metadata_expands_direct_range(self):
        self.ws["A1"] = 100
        self.wb.defined_names.add(DefinedName("원가", attr_text="'매출 자료'!$A$1"))
        self.ws["B1"] = "=원가*2"
        self.assertEqual(FormulaEngine(self.wb).analyze()[0]["dependencies"], [{"sheet": "매출 자료", "range": "A1"}])

    def test_error_literal_arity_and_boolean_functions(self):
        self.ws["A1"] = "#REF!"
        self.computed("=IFERROR(A1,7)", 7)
        self.computed("=AND(1<2,TRUE)", True)
        self.computed("=OR(FALSE,3>2)", True)
        for formula in ["=IF(TRUE)", "=ROUND(2)", "=COUNTIF(A1)", "=COUNTIFS(A1,1,A2)"]:
            with self.subTest(formula=formula):
                self.unsupported(formula, "function_argument_count")

    def test_malicious_text_is_not_executed(self):
        self.computed('="__import__(\"\"os\"\").system(\"\"whoami\"\")"', '__import__("os").system("whoami")')
        self.unsupported('=__import__("os")', "unsupported_function:__IMPORT__")

    def test_aggregate_referenced_errors_counted_or_ignored_only_when_supported(self):
        self.ws["A1"] = "#N/A"
        self.ws["A2"] = 10
        self.computed("=COUNTA(A1:A2)", 2)
        self.computed("=COUNT(A1:A2)", 1)
        self.assertEqual(self.result("=SUM(A1:A2)")["code"], "#N/A")
        self.ws["A1"] = "=UNKNOWN(1)"
        self.unsupported("=COUNT(A1:A2)", "unsupported_function:UNKNOWN")

    def test_dates_compare_and_criteria_use_excel_serials(self):
        self.ws["A1"] = datetime.date(2026, 10, 13)
        self.ws["A2"] = datetime.date(2026, 10, 14)
        self.computed("=A2>A1", True)
        self.computed('=COUNTIF(A1:A2,">="&A2)', 1)
        self.computed("=COUNTIF(A1:A2,A1)", 1)

    def test_numeric_and_string_growth_are_bounded(self):
        self.ws["A1"] = 10 ** 200
        self.assertEqual(self.result("=A1*A1")["status"], "error")
        self.computed("=IFERROR(1e308*1e308,7)", 7)
        self.ws["A1"] = "x" * 20000
        self.unsupported("=A1&A1", "string_length_limit")


if __name__ == "__main__":
    unittest.main()
