"""Developer-only real-XLSM change tests; never opens Office or runs VBA.

The negative fixtures are deliberately edited copies, not reusable user files.
Only the targeted OOXML part or embedded VBA project is changed. Formula caches
are not recalculated; the reviewed Python migration must ignore those caches.
"""
from __future__ import annotations

import argparse
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
import os
from pathlib import Path
import posixpath
import shutil
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
import zipfile

sys.dont_write_bytecode = True
SKILL = Path(__file__).resolve().parents[1]
ROOT = SKILL.parents[1]
TRIAL = ROOT / "outputs/excel-xlsm-migration-test"
SOURCE = TRIAL / "fixtures/고급_내장VBA_혼합.xlsm"
EXAMPLE = SKILL / "examples/macro-migration"
CONTRACT = EXAMPLE / "embedded-source-contract.json"
RUNTIMES = [TRIAL / ".runtime", TRIAL / "native-validation/.runtime"]
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def snapshot(folder):
    return {str(p.relative_to(folder)): sha(p) for p in folder.rglob("*") if p.is_file()}


def write_json(path, value):
    with path.open("x", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)


def sheet_part(source, sheet_name):
    with zipfile.ZipFile(source) as archive:
        sheets = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rid = next(s.attrib[f"{{{REL}}}id"] for s in sheets.findall(f"{{{NS}}}sheets/{{{NS}}}sheet")
               if s.attrib["name"] == sheet_name)
    target = next(r.attrib["Target"] for r in relationships.findall(f"{{{PKG}}}Relationship")
                  if r.attrib["Id"] == rid)
    return target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))


def patch_zip(source, target, part, edit):
    """Copy every ZIP member, replacing only the selected, inspected XML part."""
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(target, "x") as outgoing:
        for info in incoming.infolist():
            content = incoming.read(info.filename)
            outgoing.writestr(info, edit(content) if info.filename == part else content)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target) as changed:
        assert original.namelist() == changed.namelist()
        differences = [name for name in original.namelist() if original.read(name) != changed.read(name)]
    assert differences == [part], differences
    return differences


def cell_edit(coordinate, value, formula=False):
    def edit(raw):
        tree = ET.fromstring(raw)
        cell = tree.find(f".//{{{NS}}}c[@r='{coordinate}']")
        assert cell is not None, coordinate
        item = cell.find(f"{{{NS}}}{'f' if formula else 'v'}")
        assert item is not None, coordinate
        if not formula:
            assert cell.find(f"{{{NS}}}f") is None, "data-only test cannot change a formula"
        item.text = str(value)
        return ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    return edit


def run(output):
    assert SOURCE.is_file() and CONTRACT.is_file(), "Stable fixture and reviewed contract must already exist"
    if output.exists() or output.is_symlink():
        raise ValueError("Use a new output folder; existing evidence is never overwritten")
    for path in [SKILL / "scripts", EXAMPLE, *RUNTIMES]:
        sys.path.insert(0, str(path))
    # The isolated extraction subprocess must see the already prepared oletools.
    previous_pythonpath = os.environ.get("PYTHONPATH")
    os.environ["PYTHONPATH"] = os.pathsep.join([str(p) for p in RUNTIMES] +
                                              ([previous_pythonpath] if previous_pythonpath else []))
    try:
        from openpyxl import load_workbook
        from pyopenvba import ExcelFile
        from excel_pipeline import InputError
        from migration_contract import verify_contract
        from migrate_embedded_training import migrate
        from vba_extract import extract_vba

        source_before, contract_before = sha(SOURCE), sha(CONTRACT)
        original_extraction = extract_vba(SOURCE)
        assert original_extraction["status"] == "extracted", original_extraction
        original_modules = {m["name"]: m["source_sha256"] for m in original_extraction["modules"]}
        output.mkdir(parents=True)
        files = output / "fixtures"
        files.mkdir()
        cases = []

        def checked_gate(path, expected_difference=None):
            tree_before, file_before = snapshot(output), sha(path)
            try:
                evidence = verify_contract(path, CONTRACT)
            except InputError as exc:
                assert expected_difference is not None, str(exc)
                assert expected_difference in str(exc), str(exc)
                evidence = {"status": "rejected", "message": str(exc)}
            else:
                assert expected_difference is None, "Stale source was accepted"
                assert evidence["status"] == "matched"
            assert tree_before == snapshot(output), "Contract verification wrote a file in the test area"
            assert file_before == sha(path) and source_before == sha(SOURCE)
            assert contract_before == sha(CONTRACT), "Contract was regenerated"
            evidence["verification_wrote_no_test_files"] = True
            return evidence

        baseline_gate = checked_gate(SOURCE)
        baseline = migrate(SOURCE, "mixed")
        row = next(r for r in baseline["records"] if r["source_sheet"] == "신청원본" and
                   r["manual_amount"] is None and r["unit_price"] > 0 and r["discount"] < 1)
        quantity = int(row["quantity"]) + 1
        coordinate = f"G{row['source_row']}"
        workbook = load_workbook(SOURCE, data_only=False, keep_links=False)
        try:
            assert workbook["신청원본"][coordinate].value == row["quantity"]
            digits = next(int(workbook["정책"].cell(r, 3).value)
                          for r in range(6, workbook["정책"].max_row + 1)
                          if workbook["정책"].cell(r, 2).value == "금액소수자리")
            formula_cell = next(c for sheet in workbook for cells in sheet for c in cells
                                if c.data_type == "f" and isinstance(c.value, str))
            formula_location = (formula_cell.parent.title, formula_cell.coordinate, formula_cell.value)
        finally:
            workbook.close()
        data_file = files / "data-only.xlsm"
        changed_parts = patch_zip(SOURCE, data_file, sheet_part(SOURCE, "신청원본"), cell_edit(coordinate, quantity))
        data_gate = checked_gate(data_file)
        expected_amount = (Decimal(str(row["unit_price"])) * quantity *
                           (Decimal(1) - Decimal(str(row["discount"]))))
        expected_amount = expected_amount.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)
        delta = expected_amount - Decimal(str(row["amount"]))
        mode_evidence = []
        for mode in ("macro-only", "mixed"):
            changed_result = migrate(data_file, mode)
            changed_rows = {r["application_id"]: r for r in changed_result["records"]}
            assert len(changed_rows) == len(baseline["records"])
            for expected in baseline["records"]:
                actual = changed_rows[expected["application_id"]]
                if expected["application_id"] == row["application_id"]:
                    assert actual["quantity"] == quantity and Decimal(str(actual["amount"])) == expected_amount
                    assert {k: v for k, v in actual.items() if k not in {"quantity", "amount"}} == {
                        k: v for k, v in expected.items() if k not in {"quantity", "amount"}}
                else:
                    assert actual == expected, expected["application_id"]
            assert Decimal(str(changed_result["summary"]["total_amount"])) == (
                Decimal(str(baseline["summary"]["total_amount"])) + delta)
            mode_evidence.append({"mode": mode, "amount": changed_rows[row["application_id"]]["amount"],
                                  "total_amount": changed_result["summary"]["total_amount"],
                                  "record_count": len(changed_rows)})
        cases.append({"case": "data-only", "gate": data_gate, "changed_parts": changed_parts,
                      "cell": "신청원본!" + coordinate, "old_quantity": row["quantity"], "new_quantity": quantity,
                      "expected_amount": str(expected_amount), "expected_total_delta": str(delta),
                      "mode_results": mode_evidence, "source_sha256": sha(data_file)})
        print(json.dumps({"case": "data-only", "status": "passed", "modes": 2}), flush=True)

        sheet, cell, formula = formula_location
        formula_file = files / "formula-changed.xlsm"
        patch_zip(SOURCE, formula_file, sheet_part(SOURCE, sheet), cell_edit(cell, formula[1:] + "+1", formula=True))

        # Exercise sheet-scoped defined-name identity on a real OOXML file too.
        names_file = files / "sheet-name-added.xlsm"
        def add_local_name(raw):
            tree = ET.fromstring(raw)
            names = tree.find(f"{{{NS}}}definedNames")
            assert names is not None, "Fixture should already have defined names"
            name = ET.SubElement(names, f"{{{NS}}}definedName", {"name": "QAReviewProbe", "localSheetId": "0"})
            name.text = "'신청원본'!$G$6"
            return ET.tostring(tree, encoding="utf-8", xml_declaration=True)
        patch_zip(SOURCE, names_file, "xl/workbook.xml", add_local_name)

        negative = [("formula-changed", formula_file, "formula_sha256"),
                    ("sheet-name-added", names_file, "defined_names")]
        for label in ("vba-line-changed", "vba-module-added", "vba-module-deleted"):
            target = files / (label + ".xlsm")
            shutil.copyfile(SOURCE, target)
            with ExcelFile(str(target)) as book:
                project = book.vba_project()
                if label == "vba-line-changed":
                    source = book.get_module("TrainingCalculation")
                    needle = 'item("unit_price") * item("quantity") * (1 - item("discount"))'
                    assert source.count(needle) == 1
                    book.set_module("TrainingCalculation", source.replace(needle, needle + " + 1", 1))
                elif label == "vba-module-added":
                    project.add_module("AddedReviewProbe", 'Attribute VB_Name = "AddedReviewProbe"\r\n'
                                       'Option Explicit\r\nPublic Function ReviewProbe() As Long\r\n'
                                       '    ReviewProbe = 1\r\nEnd Function\r\n')
                else:
                    project.delete_module("EventWrapper")
                book.save()
            extracted = extract_vba(target)
            assert extracted["status"] == "extracted", extracted
            modules = {m["name"]: m["source_sha256"] for m in extracted["modules"]}
            if label == "vba-line-changed":
                assert set(modules) == set(original_modules)
                assert [n for n in modules if modules[n] != original_modules[n]] == ["TrainingCalculation"]
            elif label == "vba-module-added":
                assert set(modules) - set(original_modules) == {"AddedReviewProbe"}
                assert all(modules[n] == value for n, value in original_modules.items())
            else:
                assert set(original_modules) - set(modules) == {"EventWrapper"}
                assert all(original_modules[n] == value for n, value in modules.items())
            negative.append((label, target, "modules"))

        for label, path, difference in negative:
            start = time.perf_counter()
            gate = checked_gate(path, difference)
            attempted_output = output / (label + "-migration-output")
            tree_before, input_before = snapshot(output), sha(path)
            # Real CLI entrypoint, not a mocked processing function: no output is published.
            command = [sys.executable, "-X", "utf8", "-B", str(EXAMPLE / "migrate_embedded_training.py"),
                       str(path), "--output", str(attempted_output)]
            done = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", timeout=90)
            assert done.returncode == 2, done.stdout + done.stderr
            failure = json.loads(done.stderr)
            assert failure["status"] == "not_migrated" and difference in failure["message"]
            assert not attempted_output.exists() and tree_before == snapshot(output)
            assert input_before == sha(path) and contract_before == sha(CONTRACT)
            cases.append({"case": label, "gate": gate, "cli_exit": done.returncode,
                          "cli_status": failure["status"], "no_output_published": True,
                          "source_sha256": input_before, "seconds": round(time.perf_counter() - start, 3)})
            print(json.dumps({"case": label, "status": "passed", "cli_exit": done.returncode}), flush=True)

        assert source_before == sha(SOURCE) and contract_before == sha(CONTRACT)
        report = {"status": "passed", "case_count": len(cases), "baseline_gate": baseline_gate,
                  "source_sha256": source_before, "contract_sha256": contract_before,
                  "original_unchanged": True, "contract_unchanged": True,
                  "office_launched": False, "vba_executed": False, "security_settings_changed": False,
                  "network_or_model_calls": 0, "cases": cases,
                  "scope": "Reviewed example change detection only; VBA is statically extracted, not compiled or executed."}
        write_json(output / "verification.json", report)
        print(json.dumps({"status": "passed", "cases": len(cases), "source_unchanged": True}), flush=True)
        return report
    finally:
        if previous_pythonpath is None:
            os.environ.pop("PYTHONPATH", None)
        else:
            os.environ["PYTHONPATH"] = previous_pythonpath


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="실제 내장 VBA 사본의 변경 감지 독립 검사")
    parser.add_argument("--output", type=Path, default=TRIAL / "change-validation")
    run(parser.parse_args().output.resolve())
