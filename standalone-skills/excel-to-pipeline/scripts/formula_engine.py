"""Small, bounded Excel formula interpreter for *analysis*, not Excel emulation.

No eval/exec, subprocess, file access, network, macros, cached-value fallback,
external-link refresh, or workbook mutation. Unknown syntax/functions are reported
as unsupported. The caller supplies an already opened formula-mode workbook.

Supported: scalar arithmetic/comparisons/concatenation; bounded A1 references;
direct named ranges; SUM, AVERAGE, MIN, MAX, COUNT, COUNTA, IF, IFERROR, ROUND,
SUMIF(S), COUNTIF(S), exact VLOOKUP, AND, OR. Array/dynamic/structured/3D refs,
approximate lookups, VBA and Power Query are deliberately outside this subset.
Results are candidate recomputations, not evidence of Excel compatibility.
"""

from __future__ import annotations

import datetime as dt
import math
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from typing import Any

from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import to_excel


class FormulaFailure(Exception):
    def __init__(self, code: str, status: str = "unsupported"):
        super().__init__(code)
        self.code, self.status = code, status


@dataclass(frozen=True)
class Token:
    kind: str
    text: str


CELL = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$")
CELLISH = re.compile(r"^\$?[A-Za-z]{1,3}\$?[0-9]+$")
NUMBER = re.compile(r"(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][+-]?\d+)?")
ERRORS = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#NULL!")
FUNCTIONS = frozenset({"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "IF",
                       "IFERROR", "ROUND", "SUMIF", "SUMIFS", "COUNTIF", "COUNTIFS",
                       "VLOOKUP", "AND", "OR"})


def _tokens(text: str, max_tokens: int) -> list[Token]:
    result, i = [], 0
    while i < len(text):
        c = text[i]
        if c.isspace():
            i += 1
            continue
        if len(result) >= max_tokens:
            raise FormulaFailure("formula_token_limit")
        if c in ('"', "'"):
            quote, parts = c, []
            i += 1
            while i < len(text):
                if text[i] == quote:
                    if i + 1 < len(text) and text[i + 1] == quote:
                        parts.append(quote)
                        i += 2
                    else:
                        i += 1
                        break
                else:
                    parts.append(text[i])
                    i += 1
            else:
                raise FormulaFailure("unterminated_string")
            value = "".join(parts)
            if quote == "'" and ("[" in value or "]" in value):
                raise FormulaFailure("external_reference")
            result.append(Token("STRING" if quote == '"' else "SHEET", value))
            continue
        if c == "[":
            raise FormulaFailure("external_or_structured_reference")
        if c in "{}":
            raise FormulaFailure("array_constant")
        if c == "@":
            raise FormulaFailure("implicit_intersection")
        if c == "#":
            error = next((e for e in ERRORS if text[i:i + len(e)].upper() == e), None)
            if not error:
                raise FormulaFailure("dynamic_or_unknown_error_reference")
            result.append(Token("ERROR", error))
            i += len(error)
            continue
        match = NUMBER.match(text, i)
        if match:
            result.append(Token("NUMBER", match.group(0)))
            i = match.end()
            continue
        if c.isalpha() or c in "_$\\":
            j = i + 1
            while j < len(text) and (text[j].isalnum() or text[j] in "_.$\\"):
                j += 1
            result.append(Token("ID", text[i:j]))
            i = j
            continue
        if text[i:i + 2] in ("<=", ">=", "<>"):
            result.append(Token("OP", text[i:i + 2]))
            i += 2
            continue
        if c in "+-*/^&=<>%":
            result.append(Token("OP", c))
        elif c in "(),:!":
            result.append(Token(c, c))
        else:
            raise FormulaFailure("unsupported_syntax")
        i += 1
    result.append(Token("EOF", ""))
    return result


class Parser:
    # Excel evaluates exponentiation left to right; unary minus binds before ^.
    PRECEDENCE = {"=": 10, "<>": 10, "<": 10, ">": 10, "<=": 10, ">=": 10,
                  "&": 20, "+": 30, "-": 30, "*": 40, "/": 40, "^": 50}

    def __init__(self, text: str, max_tokens: int, max_depth: int):
        self.tokens = _tokens(text, max_tokens)
        self.i, self.max_depth = 0, max_depth

    def take(self) -> Token:
        token = self.tokens[self.i]
        self.i += 1
        return token

    def peek(self, text=None) -> bool | Token:
        token = self.tokens[self.i]
        return token.text == text and token.kind in ("OP", text) if text is not None else token

    def expect(self, text: str):
        if not self.peek(text):
            raise FormulaFailure("unsupported_syntax")
        self.take()

    def parse(self):
        node = self.expression()
        if self.peek().kind != "EOF":
            raise FormulaFailure("unsupported_syntax")
        return node

    def expression(self, minimum=0, depth=0):
        if depth > self.max_depth:
            raise FormulaFailure("formula_depth_limit")
        token = self.take()
        if token.kind == "OP" and token.text in ("+", "-"):
            left = ("unary", token.text, self.expression(60, depth + 1))
        elif token.kind == "NUMBER":
            if self.peek("!"):
                self.take()
                left = self.reference(token.text)
            else:
                if len(token.text) > 308:
                    raise FormulaFailure("numeric_limit")
                value = float(token.text) if any(c in token.text.lower() for c in ".e") else int(token.text)
                if not math.isfinite(value):
                    raise FormulaFailure("#NUM!", "error")
                left = ("literal", value)
        elif token.kind == "STRING":
            left = ("literal", token.text)
        elif token.kind == "ERROR":
            left = ("error", token.text)
        elif token.text == "(":
            left = self.expression(depth=depth + 1)
            self.expect(")")
        elif token.kind in ("ID", "SHEET"):
            if self.peek("!"):
                self.take()
                left = self.reference(token.text)
            elif token.kind == "SHEET":
                raise FormulaFailure("unsupported_syntax")
            elif self.peek("("):
                self.take()
                args = []
                if not self.peek(")"):
                    while True:
                        args.append(self.expression(depth=depth + 1))
                        if not self.peek(","):
                            break
                        self.take()
                self.expect(")")
                left = ("call", token.text.upper(), args)
            elif token.text.upper() in ("TRUE", "FALSE"):
                left = ("literal", token.text.upper() == "TRUE")
            else:
                left = self.reference(None, token)
        else:
            raise FormulaFailure("unsupported_syntax")
        while self.peek("%"):
            self.take()
            left = ("unary", "%", left)
        while self.peek().kind == "OP" and self.peek().text in self.PRECEDENCE and self.PRECEDENCE[self.peek().text] >= minimum:
            operator = self.take().text
            right = self.expression(self.PRECEDENCE[operator] + 1, depth + 1)
            left = ("binary", operator, left, right)
        return left

    def reference(self, sheet, token=None):
        token = token or self.take()
        if token.kind != "ID":
            raise FormulaFailure("whole_axis_or_invalid_reference")
        if not CELL.fullmatch(token.text):
            if CELLISH.fullmatch(token.text):
                raise FormulaFailure("invalid_reference")
            if self.peek(":"):
                raise FormulaFailure("whole_axis_or_3d_reference")
            return ("name", sheet, token.text)
        end = token.text
        if self.peek(":"):
            self.take()
            end_token = self.take()
            if end_token.kind != "ID" or not CELL.fullmatch(end_token.text):
                raise FormulaFailure("invalid_reference")
            end = end_token.text
        return ("ref", sheet, token.text, end)


@dataclass
class Context:
    remaining: int
    stack: list[tuple[str, str]]

    def tick(self):
        self.remaining -= 1
        if self.remaining < 0:
            raise FormulaFailure("operation_limit")


@dataclass
class CellRange:
    engine: "FormulaEngine"
    sheet: str
    r1: int
    c1: int
    r2: int
    c2: int
    context: Context

    @property
    def shape(self):
        return self.r2 - self.r1 + 1, self.c2 - self.c1 + 1

    def at(self, row, col):
        return self.engine._cell(self.sheet, f"{get_column_letter(self.c1 + col)}{self.r1 + row}", self.context)

    def values(self):
        for row in range(self.shape[0]):
            for col in range(self.shape[1]):
                yield self.at(row, col)


class FormulaEngine:
    def __init__(self, workbook, *, max_range_cells=10000, max_operations=100000,
                 max_depth=64, max_formula_chars=8192, max_tokens=2048, compiled_nodes=None):
        self.workbook = workbook
        self.max_range_cells = max_range_cells
        self.max_operations = max_operations
        self.max_depth = max_depth
        self.max_formula_chars = max_formula_chars
        self.max_tokens = max_tokens
        self._parsed = {}
        self._cache = {}
        self._sheets = {name.casefold(): name for name in workbook.sheetnames}
        self.compiled_nodes = compiled_nodes

    @staticmethod
    def _json_value(value):
        if isinstance(value, (dt.datetime, dt.date, dt.time)):
            return value.isoformat()
        if isinstance(value, dt.timedelta):
            return value.total_seconds() / 86400
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            try:
                if not math.isfinite(value):
                    raise FormulaFailure("#NUM!", "error")
            except OverflowError:
                raise FormulaFailure("#NUM!", "error")
        if isinstance(value, str) and len(value) > 32767:
            raise FormulaFailure("string_length_limit")
        if value is None or isinstance(value, (str, bool, int, float)):
            return value
        raise FormulaFailure("unsupported_literal_type")

    def _sheet(self, sheet):
        name = self._sheets.get(sheet.casefold()) if isinstance(sheet, str) else None
        if name is None:
            raise FormulaFailure("#REF!", "error")
        return name

    def _coordinate(self, coordinate):
        match = CELL.fullmatch(coordinate) if isinstance(coordinate, str) else None
        if not match:
            raise FormulaFailure("invalid_reference")
        col = column_index_from_string(match.group(1).upper())
        row = int(match.group(2))
        if col > 16384 or row > 1048576:
            raise FormulaFailure("invalid_reference")
        return row, col, f"{get_column_letter(col)}{row}"

    def _parse(self, formula):
        if not isinstance(formula, str) or not formula.startswith("="):
            raise FormulaFailure("array_or_nontext_formula")
        if len(formula) > self.max_formula_chars:
            raise FormulaFailure("formula_length_limit")
        if formula not in self._parsed:
            self._parsed[formula] = Parser(formula[1:], self.max_tokens, self.max_depth).parse()
        return self._parsed[formula]

    def evaluate(self, sheet: str, coordinate: str) -> dict[str, Any]:
        """Recompute one cell, never substituting saved Excel cached results."""
        try:
            name = self._sheet(sheet)
            row, col, coordinate = self._coordinate(coordinate)
            cell = self.workbook[name]._cells.get((row, col))
            formula = cell is not None and cell.data_type == "f"
            value = self._cell(name, coordinate, Context(self.max_operations, []))
            return {"status": "computed" if formula else "literal", "value": self._json_value(value)}
        except FormulaFailure as exc:
            return {"status": exc.status, "value": None, "code": exc.code}
        except (OverflowError, ZeroDivisionError, InvalidOperation):
            return {"status": "error", "value": None, "code": "#NUM!"}
        except RecursionError:
            return {"status": "unsupported", "value": None, "code": "formula_depth_limit"}

    def _cell(self, sheet, coordinate, context):
        context.tick()
        row, col, coordinate = self._coordinate(coordinate)
        key = sheet, coordinate
        if key in context.stack:
            raise FormulaFailure("circular_reference", "error")
        if key in self._cache:
            return self._cache[key]
        # ws.cell() / ws[coord] would create missing cells: avoid any mutation.
        cell = self.workbook[sheet]._cells.get((row, col))
        if cell is None:
            return None
        if cell.data_type == "e":
            raise FormulaFailure(str(cell.value), "error")
        if cell.data_type != "f":
            return cell.value
        if len(context.stack) >= self.max_depth:
            raise FormulaFailure("formula_depth_limit")
        context.stack.append(key)
        try:
            ast = (self.compiled_nodes[key] if self.compiled_nodes is not None
                   else self._parse(cell.value))
            value = self._eval(ast, sheet, context)
            if isinstance(value, CellRange):
                if value.shape != (1, 1):
                    raise FormulaFailure("array_result")
                value = value.at(0, 0)
            # Excel's scalar reference to a physically empty cell returns zero.
            value = 0 if value is None else value
            self._json_value(value)
            self._cache[key] = value
            return value
        finally:
            context.stack.pop()

    def _range(self, node, sheet, context):
        name = self._sheet(node[1] or sheet)
        r1, c1, _ = self._coordinate(node[2])
        r2, c2, _ = self._coordinate(node[3])
        if r2 < r1 or c2 < c1:
            raise FormulaFailure("reversed_reference")
        if (r2 - r1 + 1) * (c2 - c1 + 1) > self.max_range_cells:
            raise FormulaFailure("range_cell_limit")
        return CellRange(self, name, r1, c1, r2, c2, context)

    def _name_node(self, name, sheet):
        definition = None
        local = getattr(self.workbook[sheet], "defined_names", {})
        for key, item in local.items():
            if key.casefold() == name.casefold():
                definition = item
                break
        if definition is None:
            for key, item in self.workbook.defined_names.items():
                if key.casefold() == name.casefold():
                    definition = item
                    break
        if definition is None:
            raise FormulaFailure("unknown_name")
        text = definition.attr_text
        if not isinstance(text, str):
            raise FormulaFailure("unsupported_named_range")
        node = self._parse(text if text.startswith("=") else "=" + text)
        if node[0] != "ref":
            raise FormulaFailure("named_formula_or_union")
        absolute = re.compile(r"^\$[A-Za-z]{1,3}\$[1-9][0-9]*$")
        if node[1] is None or not absolute.fullmatch(node[2]) or not absolute.fullmatch(node[3]):
            # Relative named references depend on active-cell context, which is
            # unavailable from the workbook value graph alone.
            raise FormulaFailure("relative_named_range")
        return node

    def _number(self, value):
        if isinstance(value, CellRange):
            if value.shape != (1, 1):
                raise FormulaFailure("array_expression")
            value = value.at(0, 0)
        if value is None:
            return 0
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, (dt.datetime, dt.date, dt.time, dt.timedelta)):
            return to_excel(value, epoch=self.workbook.epoch)
        if isinstance(value, (int, float)):
            if not math.isfinite(value):
                raise FormulaFailure("#NUM!", "error")
            return value
        if isinstance(value, str):
            try:
                number = float(value.strip())
                if math.isfinite(number):
                    return number
            except ValueError:
                pass
        raise FormulaFailure("#VALUE!", "error")

    @staticmethod
    def _finite(value):
        try:
            if isinstance(value, complex) or not math.isfinite(value):
                raise FormulaFailure("#NUM!", "error")
        except OverflowError:
            raise FormulaFailure("#NUM!", "error")
        return value

    def _sum_numbers(self, values):
        # SUM 계열은 같은 누적 방식을 사용합니다. +=의 이진 오차 때문에
        # 조건 없는 합과 조건부 합의 ROUND 결과가 달라지지 않게 합니다.
        # 정수만 있으면 정수 정밀도를 유지하고, 소수는 보상 합산합니다.
        try:
            total = sum(values) if all(isinstance(v, int) for v in values) else math.fsum(values)
        except (ValueError, OverflowError):
            raise FormulaFailure("#NUM!", "error")
        return self._finite(total)

    def _text(self, value):
        if value is None:
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, (dt.datetime, dt.date, dt.time, dt.timedelta)):
            value = self._number(value)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    def _truth(self, value):
        if isinstance(value, CellRange):
            if value.shape != (1, 1):
                raise FormulaFailure("array_expression")
            value = value.at(0, 0)
        if value is None:
            return False
        if isinstance(value, str):
            if value.upper() in ("TRUE", "FALSE"):
                return value.upper() == "TRUE"
            raise FormulaFailure("#VALUE!", "error")
        return bool(value)

    def _compare(self, left, right, op):
        # Excel scalar comparisons order numbers < strings < booleans.
        if isinstance(left, (dt.datetime, dt.date, dt.time, dt.timedelta)):
            left = self._number(left)
        if isinstance(right, (dt.datetime, dt.date, dt.time, dt.timedelta)):
            right = self._number(right)
        if left is None:
            left = "" if isinstance(right, str) else 0
        if right is None:
            right = "" if isinstance(left, str) else 0
        def key(value):
            if isinstance(value, bool):
                return 2, int(value)
            if isinstance(value, str):
                return 1, value.casefold()
            if isinstance(value, (int, float)):
                return 0, value
            raise FormulaFailure("comparison_type")
        a, b = key(left), key(right)
        return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
                "<=": a <= b, ">=": a >= b}[op]

    def _eval(self, node, sheet, context):
        context.tick()
        kind = node[0]
        if kind == "unsupported":
            raise FormulaFailure(node[1])
        if kind == "literal":
            return node[1]
        if kind == "error":
            raise FormulaFailure(node[1], "error")
        if kind == "ref":
            region = self._range(node, sheet, context)
            return region.at(0, 0) if region.shape == (1, 1) else region
        if kind == "name":
            scope = self._sheet(node[1] or sheet)
            return self._eval(self._name_node(node[2], scope), scope, context)
        if kind == "unary":
            value = self._number(self._eval(node[2], sheet, context))
            return -value if node[1] == "-" else value / 100 if node[1] == "%" else value
        if kind == "binary":
            left, right = self._eval(node[2], sheet, context), self._eval(node[3], sheet, context)
            if isinstance(left, CellRange) or isinstance(right, CellRange):
                raise FormulaFailure("array_expression")
            op = node[1]
            if op == "&":
                a, b = self._text(left), self._text(right)
                if len(a) + len(b) > 32767:
                    raise FormulaFailure("string_length_limit")
                return a + b
            if op in ("=", "<>", "<", ">", "<=", ">="):
                return self._compare(left, right, op)
            a, b = self._number(left), self._number(right)
            if op in ("+", "-", "*", "/"):
                if op == "/" and b == 0:
                    raise FormulaFailure("#DIV/0!", "error")
                # Preserve decimal input semantics across ordinary operations.
                # Binary float 5*(1-.9) otherwise reaches ROUND as .4999999999999999.
                # This bounded subset is not Excel's full floating-point model.
                decimal = isinstance(a, float) or isinstance(b, float) or op == "/"
                x, y = (Decimal(str(a)), Decimal(str(b))) if decimal else (a, b)
                value = x + y if op == "+" else x - y if op == "-" else x * y if op == "*" else x / y
                return self._finite(float(value) if decimal else value)
            if op == "^":
                if abs(b) > 1000 or abs(a) > 1e100:
                    raise FormulaFailure("numeric_limit")
                try:
                    value = a ** b
                except (ValueError, OverflowError, ZeroDivisionError):
                    raise FormulaFailure("#NUM!", "error")
                if isinstance(value, complex) or not math.isfinite(value):
                    raise FormulaFailure("#NUM!", "error")
                return value
        if kind == "call":
            return self._call(node[1], node[2], sheet, context)
        raise FormulaFailure("unsupported_syntax")

    @staticmethod
    def _arity(args, minimum, maximum=None):
        if len(args) < minimum or (maximum is not None and len(args) > maximum):
            raise FormulaFailure("function_argument_count")

    def _reference_arg(self, node, sheet, context):
        if node[0] == "name":
            scope = self._sheet(node[1] or sheet)
            return self._reference_arg(self._name_node(node[2], scope), scope, context)
        if node[0] != "ref":
            raise FormulaFailure("range_argument_required")
        return self._range(node, sheet, context)

    def _aggregate(self, function, nodes, sheet, context):
        numbers, nonempty = [], 0
        for node in nodes:
            # A single referenced cell follows the same text/boolean exclusion
            # rules as a range. Literal SUM("2", TRUE) is different from SUM(A1).
            referenced = node[0] in ("ref", "name")
            region = self._reference_arg(node, sheet, context) if referenced else None
            value = region if region else self._eval(node, sheet, context)
            coordinates = ((row, col) for row in range(region.shape[0]) for col in range(region.shape[1])) if region else [(0, 0)]
            for row, col in coordinates:
                context.tick()
                try:
                    item = region.at(row, col) if region else value
                except FormulaFailure as exc:
                    if referenced and function in ("COUNT", "COUNTA") and exc.status == "error" and exc.code in ERRORS:
                        nonempty += 1
                        continue
                    raise
                if item is not None:
                    nonempty += 1
                if item is None:
                    continue
                if isinstance(item, (dt.date, dt.datetime, dt.time, dt.timedelta)):
                    numbers.append(self._number(item))
                elif isinstance(item, (int, float)) and not isinstance(item, bool):
                    numbers.append(item)
                elif not referenced and not isinstance(value, CellRange):
                    if function not in ("COUNTA",):
                        try:
                            numbers.append(self._number(item))
                        except FormulaFailure:
                            if function != "COUNT":
                                raise
        if function == "COUNTA":
            return nonempty
        if function == "COUNT":
            return len(numbers)
        if function == "SUM":
            return self._sum_numbers(numbers)
        if function == "AVERAGE":
            if not numbers:
                raise FormulaFailure("#DIV/0!", "error")
            return self._finite(self._sum_numbers(numbers) / len(numbers))
        if function == "MIN":
            return min(numbers) if numbers else 0
        return max(numbers) if numbers else 0

    def _criteria(self, criterion):
        if isinstance(criterion, CellRange):
            raise FormulaFailure("array_criterion")
        operator, target = "=", criterion
        if isinstance(target, (dt.datetime, dt.date, dt.time, dt.timedelta)):
            target = self._number(target)
        if isinstance(criterion, str):
            match = re.match(r"^(<=|>=|<>|=|<|>)(.*)$", criterion, flags=re.DOTALL)
            if match:
                operator, target = match.groups()
            if isinstance(target, str) and re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][+-]?\d+)?", target):
                target = float(target)
        if isinstance(target, str) and operator in ("=", "<>") and any(c in target for c in "*?~"):
            pieces, i = [], 0
            while i < len(target):
                if target[i] == "~" and i + 1 < len(target):
                    i += 1
                    pieces.append(re.escape(target[i]))
                else:
                    pieces.append(".*" if target[i] == "*" else "." if target[i] == "?" else re.escape(target[i]))
                i += 1
            expression = re.compile("^" + "".join(pieces) + "$", flags=re.IGNORECASE | re.DOTALL)
            return lambda value: (isinstance(value, str) and bool(expression.fullmatch(value))) == (operator == "=")
        def check(value):
            if isinstance(value, (dt.datetime, dt.date, dt.time, dt.timedelta)):
                value = self._number(value)
            if target == "":
                equal = value is None or value == ""
                return equal if operator == "=" else not equal if operator == "<>" else self._compare(value, target, operator)
            if isinstance(target, (int, float)) and not isinstance(target, bool):
                if value is None:
                    value = 0
                if isinstance(value, str):
                    try:
                        value = float(value)
                    except ValueError:
                        return operator == "<>"
                if not isinstance(value, (int, float)) or isinstance(value, bool):
                    return operator == "<>"
            return self._compare(value, target, operator)
        return check

    def _conditional(self, function, args, sheet, context):
        if function in ("SUMIF", "COUNTIF"):
            self._arity(args, 2, 3 if function == "SUMIF" else 2)
            conditions = [(self._reference_arg(args[0], sheet, context), self._criteria(self._eval(args[1], sheet, context)))]
            target = self._reference_arg(args[2] if len(args) == 3 else args[0], sheet, context) if function == "SUMIF" else None
        else:
            offset = 1 if function == "SUMIFS" else 0
            self._arity(args, offset + 2)
            if (len(args) - offset) % 2:
                raise FormulaFailure("function_argument_count")
            target = self._reference_arg(args[0], sheet, context) if offset else None
            conditions = [(self._reference_arg(args[i], sheet, context), self._criteria(self._eval(args[i + 1], sheet, context)))
                          for i in range(offset, len(args), 2)]
        shape = conditions[0][0].shape
        if any(region.shape != shape for region, _ in conditions) or (target and target.shape != shape):
            raise FormulaFailure("criteria_range_shape_mismatch")
        total, numbers = 0, []
        for row in range(shape[0]):
            for col in range(shape[1]):
                context.tick()
                if all(check(region.at(row, col)) for region, check in conditions):
                    if target is None:
                        total += 1
                    else:
                        value = target.at(row, col)
                        if isinstance(value, (int, float)) and not isinstance(value, bool):
                            numbers.append(value)
                        elif isinstance(value, (dt.date, dt.datetime, dt.time, dt.timedelta)):
                            numbers.append(self._number(value))
        return total if target is None else self._sum_numbers(numbers)

    def _call(self, function, args, sheet, context):
        if function not in FUNCTIONS:
            raise FormulaFailure("unsupported_function:" + function)
        if function == "IF":
            self._arity(args, 2, 3)
            condition = self._truth(self._eval(args[0], sheet, context))
            return self._eval(args[1], sheet, context) if condition else self._eval(args[2], sheet, context) if len(args) == 3 else False
        if function == "IFERROR":
            self._arity(args, 2, 2)
            try:
                return self._eval(args[0], sheet, context)
            except FormulaFailure as exc:
                if exc.status != "error" or exc.code not in ERRORS:
                    raise
                return self._eval(args[1], sheet, context)
        if function in ("SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA"):
            return self._aggregate(function, args, sheet, context)
        if function == "ROUND":
            self._arity(args, 2, 2)
            number = self._number(self._eval(args[0], sheet, context))
            raw_digits = self._number(self._eval(args[1], sheet, context))
            digits = int(raw_digits)
            if abs(digits) > 100 or abs(number) > 1e100:
                raise FormulaFailure("numeric_limit")
            try:
                with localcontext() as decimal_context:
                    decimal_context.prec = 220  # bounded number/digit limits above
                    return float(Decimal(str(number)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))
            except InvalidOperation:
                raise FormulaFailure("numeric_limit")
        if function in ("SUMIF", "SUMIFS", "COUNTIF", "COUNTIFS"):
            return self._conditional(function, args, sheet, context)
        if function == "VLOOKUP":
            self._arity(args, 3, 4)
            if len(args) != 4 or self._truth(self._eval(args[3], sheet, context)):
                raise FormulaFailure("approximate_lookup")
            lookup = self._eval(args[0], sheet, context)
            if isinstance(lookup, CellRange):
                raise FormulaFailure("array_lookup")
            region = self._reference_arg(args[1], sheet, context)
            column_value = self._number(self._eval(args[2], sheet, context))
            column = int(column_value)
            if column < 1 or column > region.shape[1]:
                raise FormulaFailure("#REF!", "error")
            # Exact lookup preserves string IDs versus numeric IDs; no coercion.
            wildcard = self._criteria(lookup) if isinstance(lookup, str) and any(c in lookup for c in "*?~") else None
            for row in range(region.shape[0]):
                key = region.at(row, 0)
                if wildcard(key) if wildcard else self._compare(key, lookup, "="):
                    value = region.at(row, column - 1)
                    return 0 if value is None else value
            raise FormulaFailure("#N/A", "error")
        if function in ("AND", "OR"):
            self._arity(args, 1)
            booleans = []
            for node in args:
                value = self._eval(node, sheet, context)
                items = value.values() if isinstance(value, CellRange) else [value]
                referenced = node[0] in ("ref", "name")
                for item in items:
                    if referenced and (item is None or isinstance(item, str)):
                        continue
                    booleans.append(self._truth(item))
            if not booleans:
                raise FormulaFailure("#VALUE!", "error")
            return all(booleans) if function == "AND" else any(booleans)
        raise FormulaFailure("unsupported_function:" + function)

    def _metadata(self, node, sheet, functions, dependencies):
        kind = node[0]
        if kind == "call":
            functions.add(node[1])
            for arg in node[2]:
                self._metadata(arg, sheet, functions, dependencies)
        elif kind in ("ref", "name"):
            if kind == "name":
                try:
                    scope = self._sheet(node[1] or sheet)
                    node = self._name_node(node[2], scope)
                except FormulaFailure:
                    return
            try:
                region = self._range(node, sheet, Context(self.max_operations, []))
                start = f"{get_column_letter(region.c1)}{region.r1}"
                end = f"{get_column_letter(region.c2)}{region.r2}"
                reference = start if start == end else start + ":" + end
                dependencies.add((region.sheet, reference))
            except FormulaFailure:
                return
        elif kind == "binary":
            self._metadata(node[2], sheet, functions, dependencies)
            self._metadata(node[3], sheet, functions, dependencies)
        elif kind == "unary":
            self._metadata(node[2], sheet, functions, dependencies)

    def analyze(self) -> list[dict[str, Any]]:
        """Return one bounded analysis record per formula, including no-cache cells."""
        records = []
        for sheet in self.workbook:
            for (_, _), cell in sorted(sheet._cells.items()):
                if cell.data_type != "f":
                    continue
                functions, dependencies = set(), set()
                try:
                    ast = (self.compiled_nodes[(sheet.title, cell.coordinate)] if self.compiled_nodes is not None
                           else self._parse(cell.value))
                    self._metadata(ast, sheet.title, functions, dependencies)
                except (FormulaFailure, RecursionError, OverflowError, ValueError):
                    pass
                records.append({"sheet": sheet.title, "cell": cell.coordinate,
                                "formula": cell.value if isinstance(cell.value, str) else str(type(cell.value).__name__),
                                "functions": sorted(functions),
                                "dependencies": [{"sheet": name, "range": reference} for name, reference in sorted(dependencies)],
                                **self.evaluate(sheet.title, cell.coordinate)})
        # A summary sheet can exhaust a call's budget while filling dependency
        # caches. Revisit only operation-limit results once after all formulas
        # have had a bounded evaluation. Never raise budgets, retry cycles, or
        # substitute saved Excel values. This also keeps table/formula reports
        # consistent after downstream formulas have populated the cache.
        for record in records:
            if record.get("code") == "operation_limit":
                result = self.evaluate(record["sheet"], record["cell"])
                record.pop("code", None)
                record.update(result)
        return records
