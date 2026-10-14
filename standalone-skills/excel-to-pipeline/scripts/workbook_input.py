"""Choose one input path; container errors are not evidence of DRM."""
from pathlib import Path
import codecs
import zipfile


def looks_like_text(head):
    encoding = ("utf-32" if head.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)) else
                "utf-16" if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)) else "utf-8-sig")
    try:
        text = codecs.getincrementaldecoder(encoding)(errors="strict").decode(head, final=False)
    except UnicodeError:
        return False
    return bool(text) and all(char.isprintable() or char in "\t\r\n" for char in text)


def reject_legacy_ole(source, head):
    """Do not launch legacy BIFF/XLM files renamed as .xlsx/.xlsm."""
    from excel_pipeline import InputError
    if not head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return
    try:
        import olefile
    except ImportError as error:
        raise InputError("OLE 파일의 실제 형식 확인에 olefile이 필요합니다. 스킬 의존성을 준비한 뒤 다시 요청하세요.") from error
    try:
        with olefile.OleFileIO(str(source)) as container:
            streams = container.listdir(streams=True, storages=False)
    except Exception as error:
        raise InputError("OLE 파일 구조를 확인하지 못해 Excel로 열지 않았습니다.") from error
    if any(len(path) == 1 and path[0].casefold() in {"workbook", "book"} for path in streams):
        raise InputError("구형 Excel 바이너리 형식은 지원하지 않습니다. 확장자만 바꾼 파일을 Excel로 열지 않습니다.")


def choose_reader(source, requested="auto"):
    from excel_pipeline import InputError, preflight
    source = Path(source)
    if requested not in {"auto", "openpyxl", "xlwings"}:
        raise InputError("읽기 방식은 auto/openpyxl/xlwings 중 하나여야 합니다.")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise InputError("xlsx/xlsm만 지원합니다. 확장자를 바꾸어 처리하지 마세요.")
    if not source.is_file():
        raise InputError("입력 파일을 찾을 수 없습니다.")
    if source.stat().st_size > 64 * 1024 * 1024:
        raise InputError("시험판 입력 한도는 64 MiB입니다.")
    # Permission errors are surfaced, not relabelled as DRM or retried.
    with source.open("rb") as stream:
        head = stream.read(512)
    if not head:
        raise InputError("입력 파일이 비어 있습니다.")
    if zipfile.is_zipfile(source):
        # Even a forced Excel read cannot evade XML/size/range validation.
        features = preflight(source)
        selected = "xlwings" if requested == "xlwings" else "openpyxl"
        if selected == "xlwings":
            with zipfile.ZipFile(source) as archive:
                names = archive.namelist()
            if any(name.startswith(("xl/macrosheets/", "xl/dialogsheets/")) for name in names):
                raise InputError("Excel 4.0 매크로·대화 시트가 있는 파일은 Excel 자동 읽기로 열지 않습니다.")
            if any(feature in features for feature in ("connections", "query_tables")):
                raise InputError("열 때 외부 연결이 갱신될 수 있어 이 파일의 Excel 자동 읽기는 중단했습니다. 일반 분석은 가능합니다.")
        return selected, features, "explicit_excel" if selected == "xlwings" else "readable_ooxml"
    if head.startswith(b"PK"):
        raise InputError("ZIP 구조가 손상되었습니다. DRM으로 단정하거나 Excel로 재시도하지 않습니다.")
    if requested == "openpyxl":
        raise InputError("일반 OOXML 파일이 아닙니다. 암호화·DRM·실제 파일 형식을 확인하세요.")
    # Do not ask Excel to import CSV/HTML renamed to .xlsx. Vendor protection
    # wrappers are not standardized: a binary candidate is not a DRM diagnosis.
    if looks_like_text(head):
        raise InputError("Excel 통합문서가 아닌 텍스트 형식으로 보입니다. 파일 형식을 확인하세요.")
    reject_legacy_ole(source, head)
    return "xlwings", [], "non_ooxml_candidate_not_drm_proof"


def load_workbooks(source, reader="auto", excel_timeout=90):
    """Return normalized books, features, optional VBA result, and evidence."""
    from excel_pipeline import InputError
    selected, features, reason = choose_reader(source, reader)
    if selected == "openpyxl":
        import openpyxl
        wb = openpyxl.load_workbook(source, data_only=False, keep_links=False)
        try:
            cached = openpyxl.load_workbook(source, data_only=True, keep_links=False)
        except Exception:
            wb.close()
            raise
        return wb, cached, features, None, {
            "reader": "openpyxl", "selection_reason": reason,
            "drm_status": "not_determined", "excel_started": False,
            "values_origin": "saved_ooxml_cache", "incomplete_metadata": [],
        }
    from xlwings_reader import NativeReadError, load_via_excel
    try:
        wb, cached, native_features, extraction, evidence = load_via_excel(source, timeout=excel_timeout)
    except NativeReadError as error:
        raise InputError(f"Excel 읽기를 마치지 못했습니다 [{error.code}]. {error}") from error
    evidence = {**evidence, "reader": "xlwings", "selection_reason": reason}
    return wb, cached, sorted(set(features + native_features)), extraction, evidence
