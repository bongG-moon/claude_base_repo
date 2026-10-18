"""Local, bounded workbook discovery. Never modifies the source workbook."""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import re
import shutil
import sys
import time
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from formula_engine import FormulaEngine

MAX_CELLS = 250_000
MAX_RECTANGLE = 1_000_000
MAX_UNCOMPRESSED = 128 * 1024 * 1024
MAX_REGIONS = 500
VERSION = "0.6.0"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
FLAG_LABELS = {
    "header_ambiguous": "헤더 여부 확인", "blank_boundary_review": "빈 줄·열의 표 경계 확인",
    "summary_rows_retained": "합계 후보 행 보존", "internal_blank_rows": "표 내부 빈 행 위치 보존",
    "unresolved_formula_values": "계산하지 못한 수식 있음", "repeated_header_rows": "표 안에 반복 헤더 후보",
}


class InputError(ValueError):
    pass


def safe_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def has_value(cell):
    return cell.value is not None and cell.value != ""


def bounded_range(reference, label, max_area=MAX_RECTANGLE, single_cell=False):
    """Validate XML ranges before a loader can expand implicit/merged cells."""
    try:
        if not isinstance(reference, str) or not re.fullmatch(
                r"\$?[A-Za-z]+\$?[1-9][0-9]*(?::\$?[A-Za-z]+\$?[1-9][0-9]*)?", reference):
            raise ValueError()
        c1, r1, c2, r2 = range_boundaries(reference.upper())
        if c2 < c1 or r2 < r1 or (single_cell and (c1 != c2 or r1 != r2)):
            raise ValueError()
    except (TypeError, ValueError) as error:
        raise InputError(f"잘못된 {label} 범위입니다.") from error
    if c2 > 512 or r2 > 100_000:
        raise InputError(f"{label} 범위가 시험판 최대 512열·100,000행을 넘습니다.")
    if (c2 - c1 + 1) * (r2 - r1 + 1) > max_area:
        raise InputError(f"{label} 범위가 시험판 셀 면적 한도를 넘습니다.")
    return c1, r1, c2, r2


def preflight(source):
    """Inspect container limits before openpyxl allocates worksheet objects."""
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise InputError("xlsx/xlsm만 지원합니다. 확장자를 바꾸어 처리하지 마세요.")
    if not source.is_file():
        raise InputError("입력 파일을 찾을 수 없습니다.")
    if source.stat().st_size > 64 * 1024 * 1024:
        raise InputError("시험판 입력 한도는 64 MiB입니다.")
    if not zipfile.is_zipfile(source):
        raise InputError("일반 OOXML 파일이 아닙니다. 암호화·DRM·실제 파일 형식을 확인하세요.")
    with zipfile.ZipFile(source) as archive:
        infos = archive.infolist()
        if len(infos) > 10_000 or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
            raise InputError("압축 해제 크기 또는 항목 수가 시험판 한도를 넘습니다.")
        names = {i.filename for i in infos}
        if "xl/workbook.xml" not in names:
            raise InputError("Excel 워크북 구조가 없는 ZIP 파일입니다.")
        features = []
        for prefix, label in [("xl/externalLinks/", "external_links"),
                              ("xl/pivotTables/", "pivot_tables"),
                              ("xl/queryTables/", "query_tables"),
                              ("xl/vbaProject", "macros"),
                              ("xl/connections", "connections")]:
            if any(n.startswith(prefix) for n in names):
                features.append(label)
        count, merged_area, table_area, table_count = 0, 0, 0, 0
        for info in infos:
            if not info.filename.endswith((".xml", ".rels")):
                continue
            raw = archive.read(info)
            # UTF-16/32 XML has NUL-padded ASCII markup. Check before parsing it.
            markup = raw.replace(b"\x00", b"").upper()
            if b"<!DOCTYPE" in markup or b"<!ENTITY" in markup:
                raise InputError("DTD/entity가 있는 XML은 처리하지 않습니다.")
            if re.fullmatch(r"xl/tables/[^/]+\.xml", info.filename):
                root = ET.fromstring(raw)
                c1, r1, c2, r2 = bounded_range(root.attrib.get("ref"), "Excel Table")
                table_count += 1
                table_area += (c2 - c1 + 1) * (r2 - r1 + 1)
                if table_count > MAX_REGIONS or table_area > MAX_RECTANGLE:
                    raise InputError("Excel Table 누적 면적 또는 개수가 시험판 한도를 넘습니다.")
                continue
            if not re.fullmatch(r"xl/worksheets/[^/]+\.xml", info.filename):
                continue
            root = ET.fromstring(raw)
            dimension = root.find(NS + "dimension")
            if dimension is not None:
                bounded_range(dimension.attrib.get("ref"), "시트")
            for merged in root.findall(".//" + NS + "mergeCell"):
                c1, r1, c2, r2 = bounded_range(merged.attrib.get("ref"), "병합", MAX_CELLS)
                merged_area += (c2 - c1 + 1) * (r2 - r1 + 1)
                if merged_area > MAX_CELLS:
                    raise InputError("병합 범위 누적 셀 수가 시험판 한도 250,000개를 넘습니다.")
            cells = root.findall(".//" + NS + "c")
            count += len(cells)
            if count > MAX_CELLS:
                raise InputError("셀 수가 시험판 한도 250,000개를 넘습니다.")
            for cell in cells:
                bounded_range(cell.attrib.get("r", "A1"), "셀", single_cell=True)
        return features


def bounds_text(bounds):
    c1, r1, c2, r2 = bounds
    return f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"


def inside(row, col, bounds):
    c1, r1, c2, r2 = bounds
    return r1 <= row <= r2 and c1 <= col <= c2


def regions_from_points(points):
    """XY-cut: repeated blank row/column separators; never use cached formula values."""
    if not points:
        return []
    todo, result = [points], []
    while todo:
        cells = todo.pop()
        rows, cols = sorted({r for r, _ in cells}), sorted({c for _, c in cells})
        split = False
        for axis, values in [(0, rows), (1, cols)]:
            gaps = [right for left, right in zip(values, values[1:]) if right > left + 1]
            if gaps:
                buckets = []
                start = values[0]
                for end in gaps + [values[-1] + 1]:
                    buckets.append({p for p in cells if start <= p[axis] < end})
                    start = end
                todo.extend(buckets)
                split = True
                break
        if not split:
            result.append((cols[0], rows[0], cols[-1], rows[-1]))
        if len(todo) + len(result) > MAX_REGIONS:
            raise InputError("표 후보가 500개를 넘습니다. 분석 범위를 나누어 주세요.")
    return sorted(result, key=lambda b: (b[1], b[0]))


def merged_lookup(ws):
    result = {}
    for merged in ws.merged_cells.ranges:
        if (merged.max_row - merged.min_row + 1) * (merged.max_col - merged.min_col + 1) > MAX_CELLS:
            raise InputError("병합 범위가 시험판 한도를 넘습니다.")
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                result[(row, col)] = (merged.min_row, merged.min_col)
        if len(result) > MAX_CELLS:
            raise InputError("병합 셀 수가 시험판 한도를 넘습니다.")
    return result


def row_labels(ws, row, left, right, merges):
    return [ws.cell(*merges.get((row, col), (row, col))).value for col in range(left, right + 1)]


def is_label(value):
    return isinstance(value, str) and bool(value.strip()) and not value.startswith("=")


def infer_header(ws, bounds, merges):
    c1, r1, c2, r2 = bounds
    width = c2 - c1 + 1
    flags = []
    values = row_labels(ws, r1, c1, c2, merges)
    if sum(is_label(v) for v in values) < max(1, width * 0.5):
        return 0, ["header_ambiguous"]
    depth = 1
    for merged in ws.merged_cells.ranges:
        if merged.min_row == r1 and merged.max_row > r1 and merged.min_col >= c1 and merged.max_col <= c2:
            depth = max(depth, merged.max_row - r1 + 1)
        if (merged.min_row == r1 and merged.max_col > merged.min_col and
                merged.min_col >= c1 and merged.max_col <= c2 and r1 < r2):
            below = row_labels(ws, r1 + 1, c1, c2, merges)
            if sum(is_label(v) for v in below) >= width * 0.5:
                depth = max(depth, 2)
    if depth > 3 or r1 + depth > r2:
        return 0, ["header_ambiguous"]
    first_data = row_labels(ws, r1 + depth, c1, c2, merges)
    header_styled = any(ws.cell(r1, c).font.bold or ws.cell(r1, c).fill.patternType for c in range(c1, c2 + 1))
    if all(v is None or is_label(v) for v in first_data) and not header_styled:
        flags.append("header_ambiguous")
    return depth, flags


def discover(ws):
    merges = merged_lookup(ws)
    explicit = []
    for table in ws.tables.values():
        explicit.append({"range": table.ref, "bounds": range_boundaries(table.ref),
                         "header_rows": 1 if table.headerRowCount else 0,
                         "origin": "excel_table", "name": table.name, "flags": []})
    points = set()
    byrow = {}
    for row in ws.iter_rows():
        for cell in row:
            if has_value(cell) and not any(inside(cell.row, cell.column, t["bounds"]) for t in explicit):
                points.add((cell.row, cell.column))
                byrow.setdefault(cell.row, []).append(cell)
    notes = []
    for merged in ws.merged_cells.ranges:
        if (merged.min_row == merged.max_row and merged.max_col > merged.min_col and
                len(byrow.get(merged.min_row, [])) == 1):
            cell = ws.cell(merged.min_row, merged.min_col)
            if is_label(cell.value):
                points.discard((cell.row, cell.column))
                notes.append({"range": str(merged), "value": cell.value, "reason": "single_merged_heading"})
    for point, anchor in merges.items():
        if anchor in points and not any(inside(*point, t["bounds"]) for t in explicit):
            points.add(point)
    detected = []
    for bounds in regions_from_points(points):
        c1, r1, c2, r2 = bounds
        if r1 == r2 or c1 == c2:
            notes.append({"range": bounds_text(bounds), "reason": "non_tabular_or_single_dimension"})
            continue
        depth, flags = infer_header(ws, bounds, merges)
        detected.append({"range": bounds_text(bounds), "bounds": bounds,
                         "header_rows": depth, "origin": "detected", "flags": flags})
    if len(detected) > 1:
        for table in detected:
            table["flags"].append("blank_boundary_review")
    return sorted(explicit + detected, key=lambda t: (t["bounds"][1], t["bounds"][0])), notes, merges


def validate_plan(plan, wb):
    if not isinstance(plan, dict) or not isinstance(plan.get("tables"), list) or len(plan["tables"]) > MAX_REGIONS:
        raise InputError("plan.json의 tables 배열을 확인하세요.")
    checked, ids = {}, set()
    for table in plan["tables"]:
        sheet, table_id = table.get("sheet"), table.get("id")
        if sheet not in wb.sheetnames or not isinstance(table_id, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,64}", table_id):
            raise InputError("명세 시트명 또는 표 ID가 잘못되었습니다.")
        if table_id in ids:
            raise InputError("표 ID가 중복됩니다.")
        ids.add(table_id)
        try:
            bounds = range_boundaries(table["range"])
            c1, r1, c2, r2 = bounds
            depth = table["header_rows"]
            if (not all(isinstance(x, int) and x > 0 for x in bounds) or c2 < c1 or r2 < r1 or
                    c2 > 512 or r2 > 100_000 or (r2-r1+1)*(c2-c1+1) > MAX_RECTANGLE or
                    type(depth) is not int or not 0 <= depth <= min(10, r2-r1)):
                raise ValueError()
            excluded = table.get("exclude_rows", [])
            if not isinstance(excluded, list) or any(type(r) is not int or not r1+depth <= r <= r2 for r in excluded):
                raise ValueError()
        except (KeyError, TypeError, ValueError) as error:
            raise InputError("명세의 범위·헤더 행 수·제외 행이 잘못되었습니다.") from error
        for old in checked.get(sheet, []):
            a,b,c,d = old["bounds"]
            if not (c2 < a or c1 > c or r2 < b or r1 > d):
                raise InputError("명세에 겹치는 표 범위가 있습니다.")
        checked.setdefault(sheet, []).append({**table, "bounds": bounds, "origin": "plan", "flags": []})
    return checked


def column_names(ws, table, merges):
    c1, r1, c2, _ = table["bounds"]
    depth = table["header_rows"]
    names, counts = [], Counter()
    for col in range(c1, c2 + 1):
        parts = []
        for row in range(r1, r1 + depth):
            val = ws.cell(*merges.get((row, col), (row, col))).value
            if val is not None and str(val).strip() and str(val).strip() not in parts:
                parts.append(str(val).strip())
        name = " / ".join(parts) or f"column_{get_column_letter(col)}"
        counts[name] += 1
        names.append(name if counts[name] == 1 else f"{name}__{counts[name]}")
    # Names created by suffixing must also not collide with original labels.
    used = set()
    for index, name in enumerate(names):
        base, suffix = name, 2
        while name in used:
            name = f"{base}__{suffix}"
            suffix += 1
        used.add(name)
        names[index] = name
    return names


def extract_table(ws, table, merges, engine):
    c1, r1, c2, r2 = table["bounds"]
    columns = column_names(ws, table, merges)
    rows, source_rows, row_kinds, blank_rows, repeated_headers = [], [], [], [], []
    errors = []
    for row in range(r1 + table["header_rows"], r2 + 1):
        if row in table.get("exclude_rows", []):
            continue
        cells = [ws.cell(*merges.get((row, col), (row, col))) for col in range(c1, c2 + 1)]
        if not any(has_value(cell) for cell in cells):
            blank_rows.append(row)
            continue
        values = []
        for cell in cells:
            if cell.data_type == "f":
                answer = engine.evaluate(ws.title, cell.coordinate)
                if answer["status"] not in {"computed", "literal"}:
                    errors.append({"cell": cell.coordinate, "code": answer.get("code", answer["status"])})
                    values.append(None)
                else:
                    values.append(safe_value(answer.get("value")))
            else:
                values.append(safe_value(cell.value))
        is_summary = any(isinstance(v, str) and re.fullmatch(r"(총\s*합계|합계|소계|총계|total|subtotal|grand total)", v.strip(), re.I)
                         for v in values[:min(2, len(values))])
        repeated_header = table["header_rows"] == 1 and values == columns
        if repeated_header:
            repeated_headers.append(row)
        rows.append(values)
        source_rows.append(row)
        row_kinds.append("header_candidate" if repeated_header else "summary_candidate" if is_summary else "data")
    flags = list(table["flags"])
    if "summary_candidate" in row_kinds:
        flags.append("summary_rows_retained")
    if errors:
        flags.append("unresolved_formula_values")
    if blank_rows:
        flags.append("internal_blank_rows")
    if repeated_headers:
        flags.append("repeated_header_rows")
    return {"id": table["id"], "sheet": ws.title, "range": table["range"],
            "origin": table["origin"], "header_rows": table["header_rows"], "columns": columns,
            "rows": rows, "source_rows": source_rows, "row_kinds": row_kinds,
            "blank_rows": blank_rows, "repeated_header_rows": repeated_headers,
            "excluded_rows": table.get("exclude_rows", []),
            "flags": flags, "formula_errors": errors,
            "merged_ranges": [str(m) for m in ws.merged_cells.ranges
                              if inside(m.min_row, m.min_col, table["bounds"])]}


def equivalent(a, b):
    if isinstance(a, (int, float)) and not isinstance(a, bool) and isinstance(b, (int, float)) and not isinstance(b, bool):
        return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-9)
    return safe_value(a) == safe_value(b)


def processing_plan(tables, formulas):
    """Map cell lineage to table-level processing stages without claiming semantic ETL."""
    stages = []
    for table in tables:
        bounds = range_boundaries(table["range"])
        own = [f for f in formulas if f["sheet"] == table["sheet"] and inside(
            range_boundaries(f["cell"])[1], range_boundaries(f["cell"])[0], bounds)]
        dependencies, column_rules, unmatched = set(), {}, []
        for formula in own:
            col = range_boundaries(formula["cell"])[0] - bounds[0]
            name = table["columns"][col]
            rule = column_rules.setdefault(name, {"column": name, "formula_count": 0,
                                                  "functions": set(), "examples": [], "unresolved": 0})
            rule["formula_count"] += 1
            rule["functions"].update(formula["functions"])
            rule["unresolved"] += formula["status"] != "computed"
            if len(rule["examples"]) < 2:
                rule["examples"].append({"cell": formula["cell"], "formula": formula["formula"]})
            for dep in formula["dependencies"]:
                try:
                    a, b, c, d = range_boundaries(dep["range"])
                    matches = []
                    for other in tables:
                        if other["sheet"] != dep["sheet"]:
                            continue
                        x, y, z, w = range_boundaries(other["range"])
                        if not (c < x or a > z or d < y or b > w):
                            matches.append(other["id"])
                    dependencies.update(m for m in matches if m != table["id"])
                    if not matches and dep not in unmatched:
                        unmatched.append(dep)
                except (ValueError, TypeError, KeyError):
                    unmatched.append(dep)
        for rule in column_rules.values():
            rule["functions"] = sorted(rule["functions"])
        stages.append({"table": table["id"], "sheet": table["sheet"], "range": table["range"],
                       "role": "formula_processing" if own else "input_table",
                       "depends_on_tables": sorted(dependencies), "formula_columns": list(column_rules.values()),
                       "references_outside_detected_tables": unmatched})
    remaining = {stage["table"]: set(stage["depends_on_tables"]) for stage in stages}
    order = []
    while remaining:
        ready = sorted(key for key, deps in remaining.items() if not deps)
        if not ready:
            break
        order.extend(ready)
        remaining = {key: deps - set(ready) for key, deps in remaining.items() if key not in ready}
    return {"stages": stages, "suggested_table_order": order,
            "nonlinear_table_dependencies": sorted(remaining),
            "note": "표 수준 처리 순서입니다. 실제 계산은 셀 의존성을 따르며 업무 의미의 pandas 변환은 별도입니다."}


def analyze(source, plan=None, reader="auto", excel_timeout=90):
    from workbook_input import load_workbooks
    source = Path(source).resolve()
    started = time.perf_counter()
    # Cheap file checks precede hashing; no Excel imports on ordinary OOXML.
    if not source.is_file():
        raise InputError("입력 파일을 찾을 수 없습니다.")
    if source.stat().st_size > 64 * 1024 * 1024:
        raise InputError("시험판 입력 한도는 64 MiB입니다.")
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    wb, cached, features, vba_result, input_evidence = load_workbooks(source, reader, excel_timeout)
    try:
        engine = FormulaEngine(wb)
        formulas = engine.analyze()
        for formula in formulas:
            cache = cached[formula["sheet"]][formula["cell"]].value
            formula["cached_value"] = safe_value(cache)
            formula["cached_value_origin"] = input_evidence["values_origin"]
            formula["value"] = safe_value(formula.get("value"))
            formula["cache_comparison"] = ("unavailable" if cache is None else
                                           "not_computed" if formula["status"] != "computed" else
                                           "match" if equivalent(cache, formula["value"]) else "mismatch")
        planned = validate_plan(plan, wb) if plan is not None else None
        tables, sheets, observations = [], [], []
        for sheet_index, ws in enumerate(wb.worksheets, 1):
            if ws.max_row * ws.max_column > MAX_RECTANGLE:
                raise InputError("시트 사용 범위가 시험판 한도를 넘습니다.")
            detected, notes, merges = discover(ws)
            specs = detected if planned is None else planned.get(ws.title, [])
            for index, spec in enumerate(specs, 1):
                spec.setdefault("id", f"s{sheet_index:02d}_t{index:02d}")
                tables.append(extract_table(ws, spec, merges, engine))
            sheets.append({"name": ws.title, "state": ws.sheet_state,
                           "range": ws.calculate_dimension(), "table_count": len(specs), "notes": notes,
                           "unselected_detected_regions": [] if planned is None else [s["range"] for s in detected if not any(s["range"] == t["range"] for t in specs)]})
            for row in ws.iter_rows():
                for cell in row:
                    if not has_value(cell):
                        continue
                    observations.append({"sheet": ws.title, "cell": cell.coordinate,
                                         "value": safe_value(cell.value) if cell.data_type != "f" else None,
                                         "formula": safe_value(cell.value) if cell.data_type == "f" else None,
                                         "type": cell.data_type, "number_format": cell.number_format,
                                         "cached_value": safe_value(cached[ws.title][cell.coordinate].value) if cell.data_type == "f" else None})
        if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
            raise InputError("분석 중 원본 파일이 바뀌었습니다. 저장이 끝난 뒤 다시 실행해 주세요.")
        flags = sorted({flag for table in tables for flag in table["flags"]})
        if features:
            flags.append("non_cell_features_not_executed")
        if any(f["cache_comparison"] == "mismatch" for f in formulas):
            flags.append("cached_result_mismatch")
        if any(f["status"] != "computed" for f in formulas):
            flags.append("formula_review_required")
        if input_evidence.get("incomplete_metadata"):
            flags.append("native_metadata_incomplete")
        summary = {"version": VERSION, "source_name": source.name, "sha256": digest,
                   "sheet_count": len(sheets), "table_count": len(tables),
                   "data_row_count": sum(len(t["rows"]) for t in tables),
                   "formula_count": len(formulas), "formula_status": dict(Counter(f["status"] for f in formulas)),
                   "cache_comparison": dict(Counter(f["cache_comparison"] for f in formulas)),
                   "review_flags": flags, "status": "review_required" if flags else "extracted",
                   "source_unchanged": True, "elapsed_seconds": round(time.perf_counter()-started, 4),
                   "non_cell_features": features, "model_calls": 0,
                   "reader": input_evidence["reader"],
                   "values_origin": input_evidence["values_origin"],
                   "scope": "cell_structure_and_supported_formula_subset_not_full_excel_compatibility"}
        result = {"summary": summary, "sheets": sheets, "tables": tables,
                  "pipeline_plan": processing_plan(tables, formulas),
                  "formulas": formulas, "source_cells": observations, "input_evidence": input_evidence}
        # Reusable projects keep calculation rules apart from replaceable input values.
        result["project_source"] = str(source)
        result["workbook_structure"] = {
            "epoch": wb.epoch.isoformat(),
            "names": [{"name": n.name, "text": n.attr_text, "sheet": scope}
                      for scope, names in [(None, wb.defined_names)] +
                      [(ws.title, getattr(ws, "defined_names", {})) for ws in wb.worksheets]
                      for n in names.values()],
            "merges": {ws.title: [str(m) for m in ws.merged_cells.ranges] for ws in wb.worksheets},
        }
        # Ordinary xlsx files never import the VBA dependency or spawn its worker.
        if source.suffix.lower() == ".xlsm" or "macros" in features or (vba_result and vba_result["status"] != "no_vba"):
            attach_vba_analysis(source, result, vba_result)
        if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
            raise InputError("분석 중 원본 파일이 바뀌었습니다. 결과를 발행하지 않았습니다.")
        summary["elapsed_seconds"] = round(time.perf_counter() - started, 4)
        return result
    finally:
        wb.close()
        cached.close()


def attach_vba_analysis(source, result, extraction=None):
    """Gather source evidence, never interpret workbook VBA as Python code."""
    from vba_extract import extract_vba
    from vba_inspect import inspect_module_texts
    # A denied Excel VBA-project read stays denied. Never retry via a different
    # extractor or turn missing source into invented macro business rules.
    if extraction is None:
        extraction = extract_vba(source)
    index = None
    if extraction["status"] == "extracted" and extraction.get("modules"):
        index = inspect_module_texts(extraction["modules"])
    procedures = [{"module": m["file"], **p} for m in (index or {}).get("modules", [])
                  for p in m["procedures"]]
    functions = {p["name"].casefold(): p["name"] for p in procedures if p["kind"].lower() == "function"}
    references = []
    for formula in result["formulas"]:
        expression = re.sub(r'"(?:[^"\n]|"")*"', '""', formula["formula"])
        names = sorted({functions[n.casefold()] for n in re.findall(r"\b([\w]+)\s*\(", expression)
                        if n.casefold() in functions})
        if names:
            references.append({"sheet": formula["sheet"], "cell": formula["cell"], "functions": names})
    needs = extraction["status"] != "no_vba"
    result["vba"] = extraction
    result["vba_index"] = index
    result["migration_plan"] = {
        "kind": "source_evidence_not_automatic_translation", "vba_status": extraction["status"],
        "translation_status": "requires_semantic_translation" if extraction["status"] == "extracted" else
                              "not_required" if not needs else "blocked_source_unavailable",
        "macro_executed": False, "procedures": procedures,
        "formula_udf_candidates": references,
        "calculation_order_candidates": [{"module": p["module"], "procedure": p["name"],
                                           "locations": [h["line"] for h in p["review_hints"]
                                                         if h["kind"] == "calculation_order"]}
                                          for p in procedures if any(h["kind"] == "calculation_order" for h in p["review_hints"])],
        "table_processing_plan": "pipeline_plan.json", "formula_details": "formulas.json",
        "source_directory": "vba/", "reviewed_python_entrypoint": None,
        "limits": ["호출·UDF·순서 목록은 원문 위치를 찾는 후보이며 완전한 자료 흐름 분석이 아닙니다.",
                   "현재 셀 계산은 매크로 실행 전 관찰값입니다. VBA가 셀을 바꾸는 순서는 별도로 이전해야 합니다.",
                   "VBA 실행 코드와 원문 일치, 외부 라이브러리, Excel 이벤트 동작은 자동 검증하지 않습니다."]}
    result["summary"]["vba"] = {"status": extraction["status"], "module_count": len(extraction.get("modules", [])),
                                "procedure_count": len(procedures), "macro_executed": False,
                                "translation_generated": False}
    if needs:
        result["summary"]["review_flags"].append("vba_semantic_migration_required" if index else "vba_source_unavailable")
        result["summary"]["status"] = "review_required"


def write_json(path, data):
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")


def markdown_report(result):
    s = result["summary"]
    lines = ["# 엑셀 구조 분석 결과", "", f"파일: {s['source_name']}", "",
             f"{s['sheet_count']}개 시트 · {s['table_count']}개 표 후보 · {s['formula_count']}개 수식", "",
             "원본은 수정하지 않았습니다. 표 경계는 업무 의미까지 확정한 것이 아닙니다.", "",
             "모델에는 model_context.json의 제한된 구조·표본부터 전달하세요. 전체 데이터는 로컬 결과 파일에 있습니다.",
             "표본 행 수는 Python의 전체 처리 범위를 줄이지 않습니다. 표본만 보고 전체 검증을 주장하지 마세요.", "",
             "| 표 | 시트 | 범위 | 행 수 | 확인할 점 |", "|---|---|---|---:|---|"]
    for table in result["tables"]:
        sheet = table['sheet'].replace('|', '\\|').replace('\n', ' ')
        lines.append(f"| {table['id']} | {sheet} | {table['range']} | {len(table['rows'])} | {', '.join(FLAG_LABELS.get(f, f) for f in table['flags']) or '없음'} |")
    lines += ["", "## 계산 결과", "", f"Python 계산 상태: {s['formula_status']}",
              f"Excel 관찰값 비교: {s['cache_comparison']}", "",
              "직접 읽기는 파일의 저장 캐시, xlwings는 Excel에서 관찰한 값입니다. 어느 쪽도 새로 재계산한 검증은 아닙니다.",
              "미지원·오류 수식은 표 값에 null로 남기고 formulas.json에 원인과 캐시를 따로 보존합니다.",
              "합계 후보 행은 삭제하지 않았습니다. 데이터 집계에 포함할지 확인하세요.", "",
              "## 처리 구조", "",
              "표별 원천 데이터·수식 처리 단계와 열별 대표 계산은 pipeline_plan.json에 있습니다.",
              "표 수준 의존성이 서로 얽혀도 셀 수준 순환과 같지는 않습니다. 실제 순환 오류는 formulas.json을 확인하세요.", "",
              "## 다시 사용하기", "",
              '`python -X utf8 main.py` — 생성된 프로젝트 폴더에서 실행합니다.', "",
              "DATA/sources.py에서 입력 파일을 바꾸거나 DATA/load.py에서 다른 자료를 불러옵니다.",
              "DATA는 입력, PROCESS는 고정한 수식·업무 처리를 담당합니다. 한국어 수정 안내는 README.md에 있습니다.",
              "행·열·수식 구조가 같을 때 PROCESS를 그대로 재사용합니다. 구조가 달라졌다면 변경 부분부터 재검토하세요.",
              "이전 방식의 구조 재분석이 필요한 경우에만 다음 호환 명령을 사용합니다.",
              '`python -X utf8 transform.py "다음파일.xlsx" --output "다음결과"`',
              "이 명령은 새 표 경계를 다시 찾는 분석이며, 고정 PROCESS 재사용 검증과는 다릅니다.", "",
              "## 분석 범위", "", "이미지·차트의 의미, VBA·피벗·Power Query 실행, 외부 파일 읽기는 포함하지 않습니다.",
              "생성 JSON에도 원본 업무 데이터가 있으므로 원본과 같은 보안 범위에서 보관하세요.", ""]
    evidence = result.get("input_evidence", {})
    lines += ["## 파일 읽기 방식", "", f"사용한 방식: {s.get('reader', 'openpyxl')}",
              "읽기 방식과 DRM 판정은 다릅니다. 직접 해석할 수 없는 파일이라고 DRM이 확정되는 것은 아닙니다.",
              "세부 범위·제한: input_evidence.json", ""]
    if evidence.get("incomplete_metadata"):
        lines += ["확인하지 못한 메타데이터: " + ", ".join(evidence["incomplete_metadata"]), ""]
    if "vba" in result:
        vba = result["vba"]
        lines += ["## 내장 VBA와 수식 함께 보기", "", f"내장 원문 추출: {vba['status']}",
                  "vba/의 원문, vba-index.json의 절차 위치, migration_plan.json의 수식·재계산 후보를 함께 확인하세요.",
                  "추출 성공은 업무 코드 이전 완료가 아닙니다. transform.py는 구조 분석용이며 VBA 동작을 실행하지 않습니다.",
                  "매크로가 값을 변경하는 순서를 원문으로 확인한 후 업무별 Python 진입점과 대응표를 작성합니다.",
                  "이전한 Python은 검토한 VBA·수식 계약에 연결하고, 원문이 바뀌면 해당 부분을 다시 검토합니다.",
                  vba.get("message") or "", ""]
    return "\n".join(lines)


TRANSFORM = '''"""독립 재실행 진입점. 원본을 수정하지 않고 새 결과 폴더를 생성합니다."""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent / "pipeline_lib"))
from excel_pipeline import main


def customize_tables(tables):
    """필요한 업무별 정제/결합을 여기에 추가하세요.

    각 표는 columns, rows, source_rows, row_kinds를 포함합니다.
    summary_candidate는 삭제되지 않은 합계 후보입니다.
    null은 빈 값일 수도, 미지원 수식일 수도 있으므로 formula_errors를 확인하세요.
    """
    return tables


if __name__ == "__main__":
    raise SystemExit(main(customize=customize_tables, emit_project=False))
'''


def publish(result, output, emit_project=True, sample_rows=5, context_chars=12000):
    from model_context import build_model_context, serialize_model_context
    context = build_model_context(result, sample_rows=sample_rows, max_chars=context_chars)
    output = Path(output).absolute()
    if output.exists() or output.is_symlink():
        raise InputError("결과 폴더가 이미 있습니다. 기존 자료 보호를 위해 새 폴더를 지정해 주세요.")
    output.mkdir(parents=True, exist_ok=False)
    (output / "tables").mkdir()
    # Keep compact serialization: pretty-printing would exceed the chosen budget.
    (output / "model_context.json").write_text(serialize_model_context(context), encoding="utf-8")
    write_json(output / "summary.json", result["summary"])
    write_json(output / "input_evidence.json", result.get("input_evidence", {}))
    write_json(output / "workbook.json", {"summary": result["summary"], "sheets": result["sheets"],
                                         "tables": [{k: v for k, v in t.items() if k not in {"rows"}} for t in result["tables"]]})
    write_json(output / "formulas.json", result["formulas"])
    write_json(output / "pipeline_plan.json", result["pipeline_plan"])
    write_json(output / "source_cells.json", result["source_cells"])
    for table in result["tables"]:
        write_json(output / "tables" / (table["id"] + ".json"), table)
    write_json(output / "plan.json", {"version": 1, "tables": [
        {"id": t["id"], "sheet": t["sheet"], "range": t["range"], "header_rows": t["header_rows"],
         "exclude_rows": t["excluded_rows"]} for t in result["tables"]]})
    (output / "report.md").write_text(markdown_report(result), encoding="utf-8")
    if "vba" in result:
        from vba_extract import publish_sources
        publish_sources(result["vba"], output / "vba")
        write_json(output / "vba-index.json", result["vba_index"])
        write_json(output / "migration_plan.json", result["migration_plan"])
    if emit_project:
        (output / "pipeline_lib").mkdir()
        for filename in ("excel_pipeline.py", "formula_engine.py", "vba_extract.py", "vba_inspect.py", "migration_contract.py", "model_context.py", "workbook_input.py", "xlwings_reader.py", "project_builder.py", "project_runtime.py"):
            shutil.copyfile(Path(__file__).resolve().parent / filename, output / "pipeline_lib" / filename)
        (output / "transform.py").write_text(TRANSFORM, encoding="utf-8")
        (output / "requirements.txt").write_text("openpyxl>=3.1,<4\n", encoding="utf-8")
        (output / "requirements-vba.txt").write_text("-r requirements.txt\noletools==0.60.2\n", encoding="utf-8")
        (output / "requirements-excel.txt").write_text("-r requirements.txt\nxlwings>=0.33,<0.34; sys_platform == 'win32'\nolefile>=0.47,<0.48\n", encoding="utf-8")
        from project_builder import write_project
        write_project(result, output, sample_rows, context_chars)


def main(argv=None, customize=None, emit_project=True):
    parser = argparse.ArgumentParser(description="Excel 구조를 표와 Python 재실행 프로젝트로 분해합니다.")
    parser.add_argument("source", help="원본 xlsx/xlsm 경로")
    parser.add_argument("--output", help="새 결과 폴더")
    parser.add_argument("--inspect", action="store_true", help="요약만 확인, 파일 생성 안 함")
    parser.add_argument("--plan", help="확인한 표 범위/헤더 JSON (선택)")
    parser.add_argument("--reader", choices=("auto", "openpyxl", "xlwings"), default="auto",
                        help="기본 auto: 일반 OOXML 직접 읽기, 그 외 보호 파일 후보는 Excel 읽기. xlwings는 경로 시험용")
    parser.add_argument("--excel-timeout", type=int, default=90,
                        help="Excel 읽기 대기 상한 10~300초 (기본 90), 직접 읽기에는 적용 안 함")
    parser.add_argument("--sample-rows", type=int, default=5,
                        help="모델용 표별 앞부분 표본 행 수, 0~1000 (기본 5, 전체 Python 처리에는 영향 없음)")
    parser.add_argument("--context-chars", type=int, default=12000,
                        help="모델용 요약의 최대 글자 수, 1024~50000 (토큰 수가 아님)")
    parser.add_argument("--context-only", action="store_true",
                        help="파일 생성 없이 제한된 모델용 요약만 출력")
    args = parser.parse_args(argv)
    if not 10 <= args.excel_timeout <= 300:
        parser.error("--excel-timeout은 10~300초 범위입니다.")
    if not 0 <= args.sample_rows <= 1000:
        parser.error("--sample-rows는 0~1000 범위입니다.")
    if not 1024 <= args.context_chars <= 50000:
        parser.error("--context-chars는 1024~50000 범위입니다.")
    if args.context_only and (args.output or args.inspect):
        parser.error("--context-only는 --output/--inspect와 함께 사용하지 않습니다.")
    if not args.inspect and not args.context_only and not args.output:
        parser.error("--output 또는 --inspect가 필요합니다.")
    try:
        if args.output and (Path(args.output).exists() or Path(args.output).is_symlink()):
            raise InputError("결과 폴더가 이미 있습니다. 새 폴더를 지정해 주세요.")
        plan = json.loads(Path(args.plan).read_text(encoding="utf-8")) if args.plan else None
        result = analyze(args.source, plan, args.reader, args.excel_timeout)
        if customize is not None:
            result["tables"] = customize(result["tables"])
        if args.context_only:
            from model_context import build_model_context, serialize_model_context
            sys.stdout.write(serialize_model_context(build_model_context(result, sample_rows=args.sample_rows,
                                                                         max_chars=args.context_chars)))
            return 0
        if not args.inspect:
            publish(result, args.output, emit_project, args.sample_rows, args.context_chars)
        print(json.dumps(result["summary"], ensure_ascii=False, indent=2))
        return 0
    except (InputError, OSError, ValueError, KeyError, zipfile.BadZipFile, ET.ParseError) as error:
        print(json.dumps({"status": "not_processed", "message": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
