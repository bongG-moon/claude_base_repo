"""User-adjusted training migration, not a VBA compiler or exact VBA clone.

Reads an XLSX/XLSM without executing Excel/VBA or trusting formula caches. Produces
typed JSON in a new directory. No network, external process or automatic installer.
Unlike the preserved VBA source's CodeOf/UCase rule, identifiers keep their case.
"""
from __future__ import annotations

import argparse
import datetime as dt
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
import hashlib
import json
from pathlib import Path
import sys
import time

import openpyxl

SKILL_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from excel_pipeline import InputError, preflight
from formula_engine import FormulaEngine

SCHEMAS = {
    "신청원본": ["신청ID", "사번", "부서", "과정코드", "교육일", "수량", "상태", "할인율", "수동금액", "갱신순번", "메모", "회차"],
    "출결원본": ["신청ID", "차시", "참석", "시간"],
    "평가원본": ["신청ID", "점수", "응시"],
    "사원기준": ["사번", "이름", "부서", "직급", "활성"],
    "교육기준": ["과정코드", "과정명", "단가", "교육시간", "수료점수"],
    "정책": ["항목", "값"],
}


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def text(value):
    return "" if value is None else str(value).strip()


def code(value):
    """Trim approved surrounding whitespace without changing identifier case."""
    return text(value)


def number(value, name, *, low=None, high=None, integer=False):
    if blank(value) or isinstance(value, bool):
        raise ValueError(f"{name}: 숫자가 필요합니다")
    try:
        result = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError(f"{name}: 숫자가 아닙니다") from None
    if not result.is_finite():
        raise ValueError(f"{name}: 유한한 숫자가 필요합니다")
    if low is not None and result < low or high is not None and result > high:
        raise ValueError(f"{name}: 허용 범위를 벗어났습니다")
    if integer and result != result.to_integral_value():
        raise ValueError(f"{name}: 정수가 필요합니다")
    return result


def flag(value, name):
    token = code(value)
    # Recognized Boolean spellings are parsed, not rewritten into identifiers.
    if token in {"Y", "y", "TRUE", "True", "true", "1", "참석", "응시"}:
        return True
    if token in {"N", "n", "FALSE", "False", "false", "0", "불참", "미응시", ""}:
        return False
    raise ValueError(f"{name}: 알 수 없는 참/거짓 값입니다")


def training_date(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    try:
        return dt.date.fromisoformat(text(value)).isoformat()
    except ValueError:
        raise ValueError("교육일: 날짜 또는 YYYY-MM-DD가 필요합니다") from None


def json_number(value):
    return int(value) if value == value.to_integral_value() else float(value)


def excel_round(value, digits=0):
    """Excel ROUND semantics; VBA Round/Python round use a different tie rule."""
    with localcontext() as context:
        context.prec = 60  # Inputs are bounded to 1e15; a product can reach 1e30.
        return Decimal(str(value)).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)


def source_rows(workbook, sheet, schema, engine):
    if sheet not in workbook.sheetnames:
        raise InputError(f"필수 시트가 없습니다: {sheet}")
    ws = workbook[sheet]
    actual = [text(ws.cell(5, c).value) for c in range(2, len(schema) + 2)]
    if actual != schema:
        raise InputError(f"{sheet}!B5 헤더가 이 업무 예제의 스키마와 다릅니다")
    rows = []
    for r in range(6, ws.max_row + 1):
        cells = [ws.cell(r, c) for c in range(2, len(schema) + 2)]
        if all(blank(cell.value) for cell in cells):
            continue
        values = []
        for cell in cells:
            if cell.data_type == "f":
                result = engine.evaluate(sheet, cell.coordinate)
                if result["status"] != "computed":
                    raise InputError(f"원본 입력 수식 계산 불가: {sheet}!{cell.coordinate} ({result.get('code')})")
                values.append(result["value"])
            elif cell.data_type == "e":
                raise InputError(f"원본 입력에 Excel 오류가 있습니다: {sheet}!{cell.coordinate}")
            else:
                values.append(cell.value)
        rows.append(dict(zip(schema, values), source_sheet=sheet, source_row=r))
    return rows


def unique_lookup(rows, key):
    result = {}
    for row in rows:
        item = code(row[key])
        if not item or item in result:
            raise InputError(f"{row['source_sheet']}: 기준표 키가 비어 있거나 중복됩니다 ({key})")
        result[item] = row
    return result


def reason_row(row, reason):
    return {"application_id": code(row.get("신청ID")), "source_sheet": row["source_sheet"],
            "source_row": row["source_row"], "reason": reason}


def deduplicate(rows):
    """Highest numeric revision wins, later physical input wins a tie."""
    latest, rejected, superseded = {}, [], []
    for row in rows:
        key = code(row["신청ID"])
        try:
            if not key:
                raise ValueError("신청ID가 비어 있습니다")
            revision = number(row["갱신순번"], "갱신순번", low=0, high=10**15, integer=True)
        except ValueError as exc:
            rejected.append(reason_row(row, str(exc)))
            continue
        previous = latest.get(key)
        if previous is None or revision >= previous[0]:
            if previous is not None:
                superseded.append(reason_row(previous[1], "더 높은 갱신순번 또는 동률 뒤 행으로 교체"))
            latest[key] = revision, row
        else:
            superseded.append(reason_row(row, "최신 갱신순번보다 오래된 행"))
    return [entry[1] for entry in latest.values()], rejected, superseded


def evidence(rows, *, attendance):
    values, issues, lineage = {}, {}, {}
    for row in rows:
        key = code(row["신청ID"])
        lineage.setdefault(key, []).append({"sheet": row["source_sheet"], "row": row["source_row"]})
        try:
            if not key:
                raise ValueError("신청ID가 비어 있습니다")
            if attendance:
                session = int(number(row["차시"], "차시", low=1, high=10**9, integer=True))
                attended = flag(row["참석"], "참석")
                values.setdefault(key, set())
                if attended:
                    values[key].add(session)
            else:
                taken = flag(row["응시"], "응시")
                score = None if blank(row["점수"]) else number(row["점수"], "점수", low=0, high=100)
                if taken and score is None:
                    raise ValueError("응시했지만 점수가 없습니다")
                # The original example uses the last assessment row, not max score.
                values[key] = taken, score
        except ValueError as exc:
            issues.setdefault(key, []).append(str(exc))
    return values, issues, lineage


def preprocess(inputs):
    staff = unique_lookup(inputs["사원기준"], "사번")
    courses = unique_lookup(inputs["교육기준"], "과정코드")
    policy_rows = unique_lookup(inputs["정책"], "항목")
    required = {"필수참석차시", "기본수료점수", "금액소수자리"}
    if not required <= set(policy_rows):
        raise InputError("정책 시트에 필수참석차시·기본수료점수·금액소수자리가 필요합니다")
    try:
        policy = {
            "required_sessions": int(number(policy_rows["필수참석차시"]["값"], "필수참석차시", low=1, high=100, integer=True)),
            "default_score": number(policy_rows["기본수료점수"]["값"], "기본수료점수", low=0, high=100),
            "digits": int(number(policy_rows["금액소수자리"]["값"], "금액소수자리", low=0, high=4, integer=True)),
        }
    except ValueError as exc:
        raise InputError(str(exc)) from None
    attendance, attendance_issues, attendance_lines = evidence(inputs["출결원본"], attendance=True)
    assessments, assessment_issues, assessment_lines = evidence(inputs["평가원본"], attendance=False)
    latest, rejected, superseded = deduplicate(inputs["신청원본"])
    accepted, excluded, lineage = [], [], []
    for row in latest:
        key = code(row["신청ID"])
        status = text(row["상태"])
        if status in {"대기", "취소"}:
            excluded.append(reason_row(row, status))
            continue
        try:
            if status != "승인":
                raise ValueError("알 수 없는 신청 상태입니다")
            # IDs are stored as text; lost leading zeros are never guessed back.
            if not isinstance(row["사번"], str):
                raise ValueError("사번은 앞자리 0을 보존한 텍스트여야 합니다")
            employee_id, course_code = code(row["사번"]), code(row["과정코드"])
            if employee_id not in staff:
                raise ValueError("사원기준에 없는 사번입니다")
            if course_code not in courses:
                raise ValueError("교육기준에 없는 과정코드입니다")
            employee, course = staff[employee_id], courses[course_code]
            department = text(row["부서"])
            if department != text(employee["부서"]):
                raise ValueError("신청 부서와 사원기준 부서가 다릅니다. 자동 교정하지 않았습니다")
            day = training_date(row["교육일"])
            quantity = int(number(row["수량"], "수량", low=0, high=10**15, integer=True))
            discount = number(row["할인율"], "할인율", low=0, high=1)
            manual = None if blank(row["수동금액"]) else number(row["수동금액"], "수동금액", low=0, high=10**15)
            unit_price = number(course["단가"], "단가", low=0, high=10**15)
            pass_score = policy["default_score"] if blank(course["수료점수"]) else number(course["수료점수"], "수료점수", low=0, high=100)
            problems = attendance_issues.get(key, []) + assessment_issues.get(key, [])
            if problems:
                raise ValueError("출결·평가 확인 필요: " + "; ".join(problems))
            taken, score = assessments.get(key, (False, None))
            accepted.append({
                "application_id": key, "employee_id": employee_id, "department": department,
                "course_code": course_code, "training_date": day, "quantity": quantity,
                "discount": discount, "manual_amount": manual, "unit_price": unit_price,
                "attended_sessions": len(attendance.get(key, set())), "exam_taken": taken,
                "score": score, "pass_score": pass_score, "source_sheet": row["source_sheet"],
                "source_row": row["source_row"], "round": text(row["회차"]),
            })
            lineage.append({"application_id": key, "application": {"sheet": row["source_sheet"], "row": row["source_row"]},
                            "employee": {"sheet": "사원기준", "row": employee["source_row"]},
                            "course": {"sheet": "교육기준", "row": course["source_row"]},
                            "attendance": attendance_lines.get(key, []), "assessment": assessment_lines.get(key, [])})
        except ValueError as exc:
            rejected.append(reason_row(row, str(exc)))
    return accepted, excluded, rejected, superseded, lineage, policy, len(latest)


def calculate_macro_only(records, policy):
    for row in records:
        # Preserve the VBA branch: an override must not evaluate an unused formula.
        if row["manual_amount"] is not None:
            row["amount"] = row["manual_amount"]
        else:
            with localcontext() as context:
                context.prec = 60
                row["amount"] = excel_round(row["unit_price"] * row["quantity"] * (1 - row["discount"]), policy["digits"])
        row["completed"] = (row["attended_sessions"] >= policy["required_sessions"] and row["exam_taken"]
                            and row["score"] is not None and row["score"] >= row["pass_score"])
    return []


def calculate_mixed(records, policy):
    """In-memory calculation graph only; never write or execute a workbook."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "계산"
    for r, row in enumerate(records, 2):
        values = [row["unit_price"], row["quantity"], row["discount"], row["manual_amount"],
                  row["attended_sessions"], row["exam_taken"], row["score"], row["pass_score"]]
        for c, value in enumerate(values, 1):
            ws.cell(r, c, json_number(value) if isinstance(value, Decimal) else value)
        ws.cell(r, 9, f'=IF(COUNTA(D{r})=0,ROUND(A{r}*B{r}*(1-C{r}),{policy["digits"]}),D{r})')
        ws.cell(r, 10, f'=AND(E{r}>={policy["required_sessions"]},F{r}=TRUE,COUNTA(G{r})>0,G{r}>=H{r})')
    engine, trace = FormulaEngine(book), []
    for r, row in enumerate(records, 2):
        amount, complete = engine.evaluate("계산", f"I{r}"), engine.evaluate("계산", f"J{r}")
        if amount["status"] != "computed" or complete["status"] != "computed":
            raise InputError(f"이전한 계산 수식 실패: 신청 {row['application_id']}")
        row["amount"], row["completed"] = Decimal(str(amount["value"])), bool(complete["value"])
        trace.append({"application_id": row["application_id"], "amount_formula": ws.cell(r, 9).value,
                      "completion_formula": ws.cell(r, 10).value, "amount": amount["value"],
                      "completed": complete["value"], "cached_values_used": False})
    book.close()
    return trace


def aggregate(records):
    departments = {}
    for row in records:
        entry = departments.setdefault(row["department"], {"department": row["department"],
            "approved_count": 0, "total_amount": Decimal(0), "completed_count": 0})
        entry["approved_count"] += 1
        with localcontext() as context:
            context.prec = 60
            entry["total_amount"] += row["amount"]
        entry["completed_count"] += int(row["completed"])
    return [departments[key] for key in sorted(departments)]


def serializable(value):
    if isinstance(value, Decimal):
        return json_number(value)
    if isinstance(value, dict):
        return {k: serializable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [serializable(v) for v in value]
    return value


def process_inputs(inputs, mode="macro-only"):
    if mode not in {"macro-only", "mixed"}:
        raise InputError("지원하지 않는 처리 모드입니다")
    records, excluded, rejected, superseded, lineage, policy, deduplicated = preprocess(inputs)
    trace = calculate_macro_only(records, policy) if mode == "macro-only" else calculate_mixed(records, policy)
    with localcontext() as context:
        context.prec = 60
        total = sum((row["amount"] for row in records), Decimal(0))
    summary = {"status": "review_required" if rejected else "processed", "mode": mode,
               "input_rows": len(inputs["신청원본"]), "deduplicated_rows": deduplicated,
               "approved_count": len(records), "excluded_count": len(excluded), "rejected_count": len(rejected),
               "superseded_count": len(superseded), "total_amount": total,
               "completed_count": sum(row["completed"] for row in records), "vba_executed": False,
               "formula_caches_used": False, "model_calls": 0,
               "migration_type": "user_adjusted_example_not_exact_vba_equivalence",
               "identifier_policy": "trim_whitespace_preserve_case_case_sensitive_match",
               "source_rule_changes": ["CodeOf: original VBA UCase removed in Python at user request"],
               "original_vba_equivalent": False}
    return serializable({"summary": summary, "records": records, "departments": aggregate(records),
                         "rejected": rejected, "excluded": excluded, "superseded": superseded,
                         "lineage": lineage, "formula_steps": trace})


def process_workbook(source, mode="macro-only", include_exceptions=False):
    start = time.perf_counter()
    source = Path(source).resolve()
    features = preflight(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(source, data_only=False, keep_links=False)
    try:
        engine = FormulaEngine(workbook)
        inputs = {sheet: source_rows(workbook, sheet, schema, engine) for sheet, schema in SCHEMAS.items()}
        if include_exceptions:
            inputs["신청원본"] += source_rows(workbook, "예외신청", SCHEMAS["신청원본"], engine)
        result = process_inputs(inputs, mode)
    finally:
        workbook.close()
    after = hashlib.sha256(source.read_bytes()).hexdigest()
    if before != after:
        raise InputError("입력 파일이 처리 중 변경되어 결과를 발행하지 않았습니다")
    result["summary"].update(source_sha256=before, source_unchanged=True,
                             container_features=features, elapsed_seconds=round(time.perf_counter() - start, 4))
    return result


def publish(result, output):
    output = Path(output)
    if output.exists() or output.is_symlink():
        raise InputError("기존 결과 폴더는 덮어쓰지 않습니다. 새 폴더를 지정하세요")
    output.mkdir(parents=True, exist_ok=False)
    for key, value in result.items():
        (output / f"{key}.json").write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    summary = result["summary"]
    lines = ["# 매크로 업무의 Python 이전 결과", "", f"- 실행 방식: {summary['mode']}",
             f"- 입력 {summary['input_rows']}행, 최신 신청 {summary['deduplicated_rows']}건",
             f"- 승인 처리 {summary['approved_count']}건, 대기·취소 제외 {summary['excluded_count']}건, 확인 필요 {summary['rejected_count']}건",
             f"- 금액 합계 {summary['total_amount']:,}, 수료 {summary['completed_count']}건", "",
             "이 결과는 함께 제공한 VBA 업무 규칙에 사용자 요청을 반영한 Python 예제입니다.",
             "식별자는 앞뒤 공백만 제거하고 대소문자를 보존합니다. 조인·중복 판정도 대소문자를 구분합니다.",
             "원본 VBA의 CodeOf/UCase 규칙과 의도적으로 다르며 원본 VBA 동등 결과로 보고하지 않습니다.",
             "실제 VBA 실행, 범용 자동 변환, 모든 Excel 함수 지원을 검증한 결과는 아닙니다.",
             "수식의 저장 계산값은 사용하지 않았습니다. 원본 Excel 파일은 저장하지 않았습니다.",
             "결과는 식별자 앞자리 0과 null을 보존하는 JSON입니다. 불명확한 입력은 rejected.json에 있습니다."]
    (output / "결과.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description="검토한 교육 정산 VBA 예제를 Python으로 실행합니다 (매크로 실행 없음)")
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--mode", choices=["macro-only", "mixed"], default="macro-only")
    parser.add_argument("--include-exceptions", action="store_true")
    args = parser.parse_args()
    try:
        if args.output.exists() or args.output.is_symlink():
            raise InputError("기존 결과 폴더는 덮어쓰지 않습니다")
        result = process_workbook(args.source, args.mode, args.include_exceptions)
        publish(result, args.output)
        print(json.dumps(result["summary"], ensure_ascii=False, indent=2))
    except (InputError, OSError, ValueError) as exc:
        print(json.dumps({"ok": False, "message": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
